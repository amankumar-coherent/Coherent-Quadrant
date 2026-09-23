#!/usr/bin/env python3
"""Recompute quadrants for an already-exported market, ignoring unscored rows.

Why this exists
---------------
Quadrant thresholds are the cohort medians. Unscored companies were being fed
into that median as zeros, which dragged both thresholds down until
">= median" was true for nearly every real company: CPaaS shipped 276 of 287
companies in two quadrants with Trailblazers completely empty, and Dry Eye did
the same. Markets that scored every row were unaffected, which is exactly the
fingerprint of this bug.

The scoring fix is in expand_quadrant_score.py; this script applies the same
correction to reports that were already written, without re-scoring anything.

    .\\.venv\\Scripts\\python.exe scripts\\fix_quadrants.py --market "Global CPaaS Market"
    .\\.venv\\Scripts\\python.exe scripts\\fix_quadrants.py --all
"""
from __future__ import annotations

import argparse
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "src") not in sys.path:
    sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook  # noqa: E402

from vendor_intel.quadrant.rating_map import (  # noqa: E402
    assign_quadrants_absolute_median,
)

SHEET = "Company Details"


def _scored(value: object) -> bool:
    return str(value or "").strip() not in ("", "0")


def fix_workbook(xlsx: Path) -> tuple[int, Counter, Counter] | None:
    """Rewrite the Quadrant column from the SCORED rows' medians."""
    wb = load_workbook(xlsx)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb[wb.sheetnames[0]]
    header = [c.value for c in ws[1]]
    try:
        qi = header.index("Quadrant") + 1
        xi = header.index("X") + 1
        yi = header.index("Y") + 1
    except ValueError:
        wb.close()
        return None

    before = Counter()
    scored: list[tuple[int, int, int]] = []  # (excel_row, x, y)
    for r in range(2, ws.max_row + 1):
        before[str(ws.cell(r, qi).value or "blank")] += 1
        x, y = ws.cell(r, xi).value, ws.cell(r, yi).value
        if _scored(x) and _scored(y):
            scored.append((r, int(x), int(y)))

    if not scored:
        wb.close()
        return None

    quads, _, _ = assign_quadrants_absolute_median(
        [x for _, x, _ in scored], [y for _, _, y in scored]
    )
    after = Counter()
    for (row, _, _), quad in zip(scored, quads):
        ws.cell(row, qi).value = quad
        after[quad] += 1

    wb.save(xlsx)
    wb.close()
    return len(scored), before, after


def fix_checkpoint(ckpt_path: Path) -> tuple[int, Counter] | None:
    """Same correction applied to the checkpoint's `detail_rows`.

    The HTML report is regenerated from the checkpoint, not from the Excel
    file, so fixing only the workbook leaves the chart showing the old,
    collapsed quadrants.
    """
    import json

    if not ckpt_path.exists():
        return None
    state = json.loads(ckpt_path.read_text(encoding="utf-8"))
    rows = (state.get("data") or {}).get("detail_rows") or []
    scored = [r for r in rows if _scored(r.get("X")) and _scored(r.get("Y"))]
    if not scored:
        return None

    quads, _, _ = assign_quadrants_absolute_median(
        [int(r["X"]) for r in scored], [int(r["Y"]) for r in scored]
    )
    after = Counter()
    for row, quad in zip(scored, quads):
        row["Quadrant"] = quad
        after[quad] += 1
    ckpt_path.write_text(
        json.dumps(state, ensure_ascii=False), encoding="utf-8"
    )
    return len(scored), after


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--market", help="market name, e.g. 'Global CPaaS Market'")
    ap.add_argument("--all", action="store_true", help="every exported market")
    args = ap.parse_args()

    base = ROOT / "output" / "chatgpt_expand"
    if args.all:
        targets = sorted(
            p
            for p in base.glob("*/*_FINAL.xlsx")
            # Skip archived runs: _old_, _pre300_, _failed_lock_ and friends.
            if not any(
                tag in p.parent.name
                for tag in ("_old", "_pre", "_failed", "_blocked", "_soured", "_quota")
            )
        )
    elif args.market:
        from vendor_intel.pipeline.web_expand import default_output_dir

        out = default_output_dir(args.market, "global")
        targets = sorted(Path(out).glob("*_FINAL.xlsx"))
    else:
        ap.print_help()
        return 2

    if not targets:
        print("no exported markets found")
        return 1

    for xlsx in targets:
        result = fix_workbook(xlsx)
        name = xlsx.parent.name[:44]
        if result is None:
            print(f"{name:46} skipped (no scored rows / no Quadrant column)")
            continue
        n, before, after = result
        moved = sum((after - before).values())
        print(f"{name:46} {n} scored | {moved} reassigned")
        print(
            f"{'':46} before  L={before.get('Leaders', 0):3} "
            f"C={before.get('Challengers', 0):3} "
            f"T={before.get('Trailblazers', 0):3} "
            f"E={before.get('Emerging Players', 0):3}"
        )
        print(
            f"{'':46} after   L={after.get('Leaders', 0):3} "
            f"C={after.get('Challengers', 0):3} "
            f"T={after.get('Trailblazers', 0):3} "
            f"E={after.get('Emerging Players', 0):3}"
        )
        # The HTML chart is rendered from the checkpoint, so it needs the
        # same correction or it keeps showing the collapsed quadrants.
        ck = fix_checkpoint(xlsx.parent / "chatgpt_checkpoint_batch_all.json")
        if ck:
            n_ck, after_ck = ck
            print(
                f"{'':46} ckpt    {n_ck} rows -> "
                f"T={after_ck.get('Trailblazers', 0)}"
            )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
