#!/usr/bin/env python3
"""Gap-fill missing Website/HQ/Summary for LNG Landscape via DeepSeek."""
from __future__ import annotations

import asyncio
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.chatgpt_env import apply_chatgpt_expand_env, deepseek_chat_config
from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx

SLUG = "global_liquefied_natural_gas_market_global"
QUERY = "Global Liquefied Natural Gas Market"
OUT = ROOT / "output" / "chatgpt_expand" / SLUG
XLSX = OUT / f"{SLUG}_FINAL.xlsx"


def _norm(s: str) -> str:
    s = str(s or "").strip().lower()
    s = re.sub(r"\([^)]*\)", " ", s)
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _needs(r: dict) -> bool:
    return not (
        str(r.get("Website") or "").strip()
        and str(r.get("Headquarters") or "").strip()
        and str(r.get("Summary") or "").strip()
    )


async def main() -> int:
    apply_chatgpt_expand_env(root=ROOT)
    key, base, model = deepseek_chat_config()
    from openai import AsyncOpenAI

    client = AsyncOpenAI(api_key=key, base_url=base)

    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Landscape"]
    rows_raw = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows_raw[0]]
    landscape = []
    for r in rows_raw[1:]:
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        if str(d.get("Company") or "").strip():
            landscape.append(d)

    need = [r for r in landscape if _needs(r)]
    print(f"gap-fill {len(need)}/{len(landscape)}", flush=True)

    for i in range(0, len(need), 10):
        batch = need[i : i + 10]
        names = [str(r.get("Company") or "") for r in batch]
        prompt = (
            "Fill Global LNG market landscape fields. Return JSON array with keys:\n"
            "company, website, founded, headquarters (City, Country), continent_geography,\n"
            "operational_presence, ownership, employees, specialty_focus, summary,\n"
            "country_code (ISO2). Facts only; empty string if unknown.\n"
            f"Companies: {json.dumps(names, ensure_ascii=False)}"
        )
        try:
            resp = await client.chat.completions.create(
                model=model,
                temperature=0.1,
                messages=[
                    {"role": "system", "content": "Return only valid JSON array."},
                    {"role": "user", "content": prompt},
                ],
            )
            text = (resp.choices[0].message.content or "").strip()
            if text.startswith("```"):
                text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.I | re.S)
            data = json.loads(text)
            by = {_norm(x.get("company") or ""): x for x in data if isinstance(x, dict)}
        except Exception as e:
            print(f"batch fail {i}: {e}", flush=True)
            by = {}
        for r in batch:
            x = by.get(_norm(r.get("Company") or ""), {})
            if not x:
                continue
            for dst, src in (
                ("Website", "website"),
                ("Founded", "founded"),
                ("Headquarters", "headquarters"),
                ("Continent / Geography", "continent_geography"),
                ("Operational Presence", "operational_presence"),
                ("Ownership", "ownership"),
                ("Employees", "employees"),
                ("Specialty Focus", "specialty_focus"),
                ("Summary", "summary"),
                ("Country Code", "country_code"),
            ):
                val = str(x.get(src) or "").strip()
                if val and not str(r.get(dst) or "").strip():
                    r[dst] = val
        print(f"  filled {min(i+10, len(need))}/{len(need)}", flush=True)

    audit_path = OUT / "chatgpt_expand_batch_all_audit.json"
    audit = {}
    if audit_path.exists():
        audit = json.loads(audit_path.read_text(encoding="utf-8")).get("xy_scoring") or {}
    detail_rows = to_company_detail_rows(landscape, QUERY, audit)
    by_co = {
        _norm(str(r.get("Company") or "")): str(r.get("Role") or "Brand") for r in landscape
    }
    for d in detail_rows:
        role = by_co.get(_norm(d.get("Brand") or ""))
        if role in ("Brand", "Marketer"):
            d["Role"] = role

    write_final_xlsx(
        XLSX,
        landscape,
        "Companies",
        {"query": QUERY, "gap_fill": "2026-08-14", "xy_scoring": audit},
        detail_rows=detail_rows,
    )
    export_expand_quadrant_outputs(
        OUT, detail_rows, QUERY, country="global", audit=audit, chart_n=20
    )
    still = sum(1 for r in landscape if _needs(r))
    print(f"done still_missing={still}/{len(landscape)}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
