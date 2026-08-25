#!/usr/bin/env python3
"""Reclassify Livongo from Marketer -> Brand (digital health + connected devices)."""
from __future__ import annotations

import json
import re
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx
from vendor_intel.quadrant.brand_meta import format_acquired_suffix
from vendor_intel.quadrant.rating_map import assign_quadrants_half_median

OUT = ROOT / "output" / "chatgpt_expand"
SLUG = "global_wearable_medical_devices_market_global"
QUERY = "Global Wearable Medical Devices Market"
FOLDER = OUT / SLUG


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def main() -> int:
    xlsx = FOLDER / f"{SLUG}_FINAL.xlsx"
    wb = load_workbook(xlsx, data_only=True)
    ws = wb["Landscape"]
    rows_raw = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows_raw[0]]

    details_by: dict[str, dict] = {}
    if "Company Details" in wb.sheetnames:
        dws = wb["Company Details"]
        drows = list(dws.iter_rows(values_only=True))
        dhdr = [str(h) for h in drows[0]]
        for r in drows[1:]:
            if not r:
                continue
            d = {dhdr[i]: ("" if r[i] is None else r[i]) for i in range(len(dhdr)) if i < len(r)}
            key = _norm(d.get("Brand") or d.get("Company") or "")
            if key:
                details_by[key] = d

    kept: list[dict] = []
    for r in rows_raw[1:]:
        if not r:
            continue
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        brand = str(d.get("Company") or d.get("Brand") or "").strip()
        if not brand:
            continue
        key = _norm(brand)
        det = details_by.get(key)
        if det:
            if det.get("Brand"):
                brand = str(det["Brand"]).strip()
                key = _norm(brand)
                d["Brand"] = brand
            if det.get("Company") and str(det["Company"]).startswith("("):
                d["_display_company"] = det["Company"]
            if det.get("X") not in (None, ""):
                d["X Score"] = det["X"]
                d["Y Score"] = det["Y"]
                d["Overall Score"] = det["Overall"]
                d["Quadrant"] = det.get("Quadrant") or ""
            if det.get("Found in"):
                d["Headquarters"] = det["Found in"]
                d["Found in"] = det["Found in"]
            if det.get("Role") in {"Brand", "Marketer"}:
                d["Role"] = det["Role"]
                d["Distribution Type"] = det["Role"]

        if key == "livongo":
            d["Role"] = "Brand"
            d["Distribution Type"] = "Brand"
            d["Specialty Focus"] = (
                "Digital health Brand: connected glucose/BP devices + AI coaching "
                "for chronic care (diabetes, hypertension)"
            )
            d["Core Categories"] = "Connected diabetes RPM; Digital health platform"
            d["Summary"] = (
                "Livongo — Brand. Digital health company with connected monitoring devices, "
                "AI analytics, and coaching; acquired by Teladoc Health in 2020."
            )
            d["Ownership"] = "acquired by Teladoc Health, 2020"
            d["_display_company"] = format_acquired_suffix(
                "Teladoc Health", relation="acquired_by", year="2020"
            )
            print("Fixed Livongo -> Brand", flush=True)

        d["Brand"] = brand
        d["Company"] = brand
        d["Role"] = d.get("Role") if d.get("Role") in {"Brand", "Marketer"} else "Brand"
        d["Distribution Type"] = d["Role"]
        kept.append(d)

    xs, ys = [], []
    for r in kept:
        try:
            xs.append(float(r.get("X Score") or 50))
        except (TypeError, ValueError):
            xs.append(50.0)
        try:
            ys.append(float(r.get("Y Score") or 50))
        except (TypeError, ValueError):
            ys.append(50.0)
    quads, _, _ = assign_quadrants_half_median(xs, ys)
    for r, q in zip(kept, quads):
        r["Quadrant"] = q
        try:
            r["Overall Score"] = round(
                (float(r.get("X Score") or 50) + float(r.get("Y Score") or 50)) / 2
            )
        except (TypeError, ValueError):
            pass

    audit: dict = {}
    audit_path = FOLDER / "chatgpt_expand_batch_all_audit.json"
    if audit_path.exists():
        audit = json.loads(audit_path.read_text(encoding="utf-8")).get("xy_scoring") or {}
    qjson = FOLDER / f"{SLUG}_quadrant.json"
    if qjson.exists():
        try:
            payload = json.loads(qjson.read_text(encoding="utf-8"))
            crit = payload.get("criteria") or {}
            if crit.get("parameter_definitions"):
                audit["parameter_definitions"] = crit["parameter_definitions"]
            labels = crit.get("axis_labels") or {}
            if labels.get("x"):
                audit["axis_x"] = labels["x"]
            if labels.get("y"):
                audit["axis_y"] = labels["y"]
            if crit.get("x_axis"):
                audit["x_features"] = crit["x_axis"]
            if crit.get("y_axis"):
                audit["y_features"] = crit["y_axis"]
        except Exception:
            pass

    detail_rows = to_company_detail_rows(kept, QUERY, audit)
    for det in detail_rows:
        key = _norm(det.get("Brand") or "")
        src = next((r for r in kept if _norm(r.get("Brand") or "") == key), None)
        if not src:
            continue
        det["Role"] = src.get("Role") or "Brand"
        if src.get("_display_company"):
            det["Company"] = src["_display_company"]
        elif not str(det.get("Company") or "").startswith("("):
            det["Company"] = det.get("Brand") or det.get("Company")

    write_final_xlsx(
        xlsx,
        kept,
        "Companies",
        {"query": QUERY, "xy_scoring": audit},
        detail_rows=detail_rows,
    )
    extras = export_expand_quadrant_outputs(
        FOLDER, detail_rows, QUERY, country="global", audit=audit, chart_n=20
    )
    print(
        f"roles={dict(Counter(r.get('Role') for r in detail_rows))} n={len(detail_rows)}",
        flush=True,
    )
    print(f"html={extras.get('html')}", flush=True)

    # Prevent future Marketer reclass
    path = ROOT / "scripts" / "wearable_add_marketers.py"
    txt = path.read_text(encoding="utf-8")
    old = '    "livongo": "Markets connected diabetes RPM programs/devices under Teladoc Health",'
    new = (
        "    # livongo kept as Brand (connected devices + coaching platform), "
        "not Marketer — do not reclass\n"
    )
    if old in txt:
        path.write_text(txt.replace(old, new), encoding="utf-8")
        print("Patched wearable_add_marketers.py", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
