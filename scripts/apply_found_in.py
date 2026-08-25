"""Apply verified Found in locations to FINAL Excel + rebuild HTML.

Only writes locations from resolve JSON + WEB_OVERRIDES (web-verified).
Leaves blank when unverified — never invents.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx

ROOT = Path(r"D:/Coherent-Quadrant/output/chatgpt_expand")
AUDIT = ROOT / "_audit"

# Extra web-verified HQ (own web search). Format: City, Country — never invent.
WEB_OVERRIDES: dict[str, str] = {
    "alucoat": "Linares, Jaén, Spain",
    "al ghurair packaging": "Dubai, United Arab Emirates",
    "al watania plastics": "Riyadh, Saudi Arabia",
    "qatar plastic products co.": "Mesaieed, Qatar",
    "flex middle east": "Dubai, United Arab Emirates",
    "flex america": "Altamira, Tamaulipas, Mexico",
    "napco national": "Dammam, Saudi Arabia",
    "white lotus industries": "Surat, Gujarat, India",
    "parkside flexibles": "Normanton, West Yorkshire, United Kingdom",
    "gulf packaging industries": "Dubai, United Arab Emirates",
    # Round-2 blanks (name/domain cleaned + web verified)
    "flex-pack engineering": "Uniontown, Ohio, USA",
    "al khaleej polypropylene": "Sohar, Oman",
    "flexible packaging solutions (fps)": "Amstelveen, Netherlands",
    "flexible packaging solutions": "Amstelveen, Netherlands",
    "pouch partners": "Brunello, Varese, Italy",
    "flexo films (india)": "Noida, Uttar Pradesh, India",
    "biobeat technologies ltd.": "Petah Tikva, Israel",
    "vitalthings": "Trondheim, Norway",
    "zensorium": "Singapore, Singapore",
    "everist health": "Ann Arbor, Michigan, USA",
    "vivalink": "Campbell, California, USA",
    "physiq": "Chicago, Illinois, USA",
    "cardiac insight": "Bellevue, Washington, USA",
    "nexxto": "Sao Paulo, Brazil",
    "vitalerter": "Airport City, Israel",
    "swasth": "Mumbai, Maharashtra, India",
    "quantified ag": "Lincoln, Nebraska, USA",
    "quro medical": "Sandton, Gauteng, South Africa",
    "aviro health": "Cape Town, South Africa",
    "cure bionics": "Sousse, Tunisia",
    "biologix sistemas": "Sao Paulo, Brazil",
    "vitaliti": "Kitchener, Ontario, Canada",
    "vitaltracer": "Montreal, Quebec, Canada",
    "sens4care": "Cornellà de Llobregat, Barcelona, Spain",
    "sense4care": "Cornellà de Llobregat, Barcelona, Spain",
    "medlevensohn": "Serra, Espírito Santo, Brazil",
    "avertus": "Toronto, Ontario, Canada",
    "vitaliberty": "Mannheim, Germany",
    "healthwatch technologies": "Kfar Saba, Israel",
    "cardiodiagnostics": "Campbell, California, USA",
    "biosign technologies": "Mississauga, Ontario, Canada",
    "philips (middle east)": "Dubai, United Arab Emirates",
    "philips (canada)": "Markham, Ontario, Canada",
    "baxter (middle east)": "Dubai, United Arab Emirates",
    "roche (middle east)": "Dubai, United Arab Emirates",
    "omron healthcare (canada)": "Toronto, Ontario, Canada",
    "apple (canada)": "Toronto, Ontario, Canada",
    # Round-3 blanks (cleaned name / domain + web verified)
    "flex-pack": "Itasca, Illinois, USA",
    "flexoprint": "Shah Alam, Selangor, Malaysia",
    "thai packaging co": "Samut Prakan, Thailand",
    "thai packaging co., ltd.": "Samut Prakan, Thailand",
    "zhejiang yamei new materials": "Haining, Zhejiang, China",
    "dongil industries": "Pohang, South Korea",
    "vital beats": "Copenhagen, Denmark",
    "sonofit": "Trondheim, Norway",
    "imedtrix": "Milpitas, California, USA",
    "healthq technologies": "Stellenbosch, South Africa",
    "medpass": "Sao Paulo, Brazil",
    "medihelp": "Pretoria, South Africa",
    "pluri sistemas": "Juiz de Fora, Minas Gerais, Brazil",
    # Clear bad wiki match — Biomedical Systems do Brasil ≠ US lab
    "biomedical systems do brasil": "",
}

# LLM placeholder / non-findable names — drop from market (cannot invent HQ)
DROP_FAKES: set[str] = {
    "embalagens flexíveis ltda.",
    "embalagens flexíveis (brasil converters)",
    "embalagens flexíveis do brasil (efb)",
    "embalagens flexíveis do sul (efs)",
    "embalagens flexíveis nordeste",
    "embalagens flexíveis nordeste (efn)",
    "embalagens flexíveis rio de janeiro (efrj)",
    "embalagens flexíveis rio grande do sul",
    "embalagens flexíveis sul",
    "embalagens flexíveis são paulo",
    "embalagens flexíveis são paulo (efsp)",
    "flexibras embalagens ltda.",
    "flexo embalagens",
    "flexpack embalagens",
    "plastrel embalagens",
    "polipropileno do brasil ltda.",
    "tricorder",
    "aerobics",
    "afya papyrus",
    "vigilância em saúde (vigilsaúde)",
    "tecnologia em saúde (tecsaude)",
    "tecnologia em saúde (tecsaude)",
    "mediwrist (by mediwrist ltd)",
    "mediwear",
    "sensify medical",
    "biometrix",
    "cardioMaps".lower(),
    "cardiomaps",
    "nora health",
    "vitaltrak",
    "vitalis tecnologia",
    "viva medical",
    "delfi medical",
    "biocare",
    "vitalsigns",
    "sana health",
    "biomedical systems do brasil",
    # Remaining flex blanks — no reliable unique HQ after web search
    "al tajir packaging",
    "dubai packaging",
    "plastibras embalagens ltda.",
    "saudi industrial packaging co.",
    "advance packaging (india)",
    "fleming packaging",
    "zhongshan hongsu industrial co., ltd.",
}

# Wiki/manual results that are known wrong — blank them
CLEAR_BAD: set[str] = {
    "biomedical systems do brasil",
}


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def _lookup(brand: str, table: dict[str, str]) -> str | None:
    b = _norm(brand)
    if b in table:
        return table[b]
    # Prefer longer keys; require brand to contain key as prefix/whole — not short
    # brand inheriting a longer company's HQ (e.g. Flex-Pack ≠ Flex-Pack Engineering).
    best: tuple[int, str] | None = None
    for k, v in table.items():
        if not k or len(k) < 6:
            continue
        if b == k or b.startswith(k + " ") or b.startswith(k + ",") or b.startswith(k + "("):
            cand = (len(k), v if v is not None else "")
            if best is None or cand[0] > best[0]:
                best = cand
        elif k.startswith(b + " ") or k.startswith(b + ",") or k.startswith(b + "("):
            # brand is shorter parent of key — do NOT inherit
            continue
        elif k in b and len(k) >= 12:
            cand = (len(k), v if v is not None else "")
            if best is None or cand[0] > best[0]:
                best = cand
    if best is None:
        return None
    return best[1]


def _load_resolved(slug: str) -> dict[str, dict]:
    path = AUDIT / f"{slug}_found_in_resolved.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    out: dict[str, dict] = {}
    for row in data:
        out[_norm(row["brand"])] = row
    return out


def _final_location(brand: str, resolved: dict[str, dict]) -> tuple[str, str]:
    b = _norm(brand)
    if b in CLEAR_BAD:
        return "", "cleared_bad"
    ov = _lookup(brand, WEB_OVERRIDES)
    if ov is not None:
        return ov, "web_override" if ov else "web_override_blank"
    row = resolved.get(b)
    if not row:
        for k, r in resolved.items():
            if k in b or b in k:
                row = r
                break
    if not row:
        return "", "missing"
    loc = str(row.get("found_in") or "").strip()
    src = str(row.get("source") or "")
    # Drop country-only leftovers if somehow present
    if loc and "," not in loc and loc.lower() in {
        "usa",
        "us",
        "india",
        "japan",
        "china",
        "germany",
        "brazil",
        "canada",
        "australia",
        "switzerland",
        "united kingdom",
        "united states",
        "uae",
        "singapore",
        "france",
        "israel",
        "ireland",
        "finland",
        "spain",
        "italy",
        "south korea",
        "south africa",
        "malaysia",
        "thailand",
        "mexico",
        "chile",
        "indonesia",
        "saudi arabia",
        "qatar",
        "belgium",
        "netherlands",
        "austria",
        "new zealand",
        "hong kong",
        "kenya",
        "seed_file",
        "baden",
        "tait",
        "lengerich",
    }:
        return "", "rejected_country_only"
    return loc, src


def _apply_market(slug: str, query: str) -> None:
    folder = ROOT / slug
    xlsx = folder / f"{slug}_FINAL.xlsx"
    resolved = _load_resolved(slug)

    wb = load_workbook(xlsx, data_only=True)
    # Landscape
    land_ws = wb["Landscape"] if "Landscape" in wb.sheetnames else wb[wb.sheetnames[0]]
    land_rows = list(land_ws.iter_rows(values_only=True))
    land_hdr = [str(h) for h in land_rows[0]]
    landscape: list[dict] = []
    for r in land_rows[1:]:
        if not r:
            continue
        d = {
            land_hdr[i]: ("" if r[i] is None else r[i])
            for i in range(len(land_hdr))
            if i < len(r)
        }
        if not str(d.get("Company") or "").strip():
            continue
        landscape.append(d)

    dropped = []
    kept = []
    for row in landscape:
        brand = str(row.get("Company") or "").strip()
        bn = _norm(brand)
        if bn in DROP_FAKES or any(bn.startswith(f) for f in DROP_FAKES if len(f) > 20):
            dropped.append(brand)
            continue
        if "tecsaude" in bn.replace("ú", "u").replace("á", "a") or "tecnologia em sa" in bn:
            dropped.append(brand)
            continue
        # also drop by contains for embalagens flexíveis regional placeholders
        if "embalagens flex" in bn and (
            "nordeste" in bn
            or "são paulo" in bn
            or "sao paulo" in bn
            or "rio de janeiro" in bn
            or "rio grande" in bn
            or "do brasil" in bn
            or "do sul" in bn
            or "brasil converters" in bn
            or bn.endswith("sul")
            or bn == "embalagens flexíveis ltda."
        ):
            dropped.append(brand)
            continue
        kept.append(row)
    landscape = kept
    if dropped:
        print(f"  dropped fakes ({len(dropped)}): {', '.join(dropped[:6])}{'...' if len(dropped)>6 else ''}")

    details_by: dict[str, dict] = {}
    if "Company Details" in wb.sheetnames:
        ws = wb["Company Details"]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(h) for h in rows[0]]
        for r in rows[1:]:
            if not r:
                continue
            d = {
                hdr[i]: ("" if r[i] is None else str(r[i]))
                for i in range(len(hdr))
                if i < len(r)
            }
            key = _norm(d.get("Brand") or d.get("Company") or "")
            if key:
                details_by[key] = d

    report = []
    filled = 0
    blank = 0
    for row in landscape:
        brand = str(row.get("Company") or "").strip()
        # Prefer matching Company Details brand if present
        det = details_by.get(_norm(brand))
        brand_for_lookup = str((det or {}).get("Brand") or brand)
        loc, src = _final_location(brand_for_lookup, resolved)
        if not loc:
            # Prefer exact / longer brand only — never short name borrowing longer HQ
            bn = _norm(brand)
            for dk, dv in details_by.items():
                if dk == bn or dk.startswith(bn + " ") or dk.startswith(bn + "("):
                    # longer detail key — don't use for short brand
                    continue
                if bn == dk or bn.startswith(dk + " ") or bn.startswith(dk + "("):
                    loc, src = _final_location(dv.get("Brand") or brand, resolved)
                    if loc:
                        break
        old = str((det or {}).get("Found in") or row.get("Headquarters") or "")
        if loc:
            row["Headquarters"] = loc
            filled += 1
        else:
            row["Headquarters"] = ""
            blank += 1
        if det:
            for col in ("Quadrant", "X", "Y", "Overall", "Role"):
                if det.get(col) not in (None, ""):
                    row[col] = det[col]
                    if col == "X":
                        row["X Score"] = det[col]
                    elif col == "Y":
                        row["Y Score"] = det[col]
                    elif col == "Overall":
                        row["Overall Score"] = det[col]
            if det.get("Role"):
                row["Distribution Type"] = det["Role"]
        report.append(
            {
                "brand": brand_for_lookup,
                "old": old,
                "found_in": loc,
                "source": src,
            }
        )

    audit_path = folder / "chatgpt_expand_batch_all_audit.json"
    audit: dict = {}
    if audit_path.exists():
        audit = json.loads(audit_path.read_text(encoding="utf-8")).get("xy_scoring") or {}

    detail_rows = to_company_detail_rows(landscape, query, audit)
    # Force Found in from our verified map
    by_brand = {_norm(r["brand"]): r["found_in"] for r in report}
    for d in detail_rows:
        key = _norm(d.get("Brand") or d.get("Company") or "")
        loc = by_brand.get(key)
        if loc is None:
            # exact / longer-name contains only — never short name inherits longer
            for k, v in by_brand.items():
                if not k:
                    continue
                if key == k or key.startswith(k + " ") or key.startswith(k + "("):
                    loc = v
                    break
        d["Found in"] = loc or ""

    write_final_xlsx(
        xlsx,
        landscape,
        "Companies",
        {"query": query, "xy_scoring": audit, "found_in_fill": True},
        detail_rows=detail_rows,
    )
    extras = export_expand_quadrant_outputs(
        folder, detail_rows, query, country="global", audit=audit, chart_n=20
    )
    out_rep = AUDIT / f"{slug}_found_in_applied.json"
    out_rep.write_text(
        json.dumps(
            {
                "filled": filled,
                "blank_unverified": blank,
                "html": str(extras.get("html") or ""),
                "rows": report,
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    print(f"{slug}: filled={filled} blank_unverified={blank}")
    print(f"  html -> {extras.get('html')}")
    # show a few samples
    for r in report[:8]:
        print(f"  {r['brand'][:40]!r}: {r['old']!r} -> {r['found_in']!r}")


def main() -> None:
    _apply_market(
        "global_flexible_packaging_market_global",
        "Global Flexible Packaging Market",
    )
    _apply_market(
        "global_wearable_medical_devices_market_global",
        "Global Wearable Medical Devices Market",
    )


if __name__ == "__main__":
    main()
