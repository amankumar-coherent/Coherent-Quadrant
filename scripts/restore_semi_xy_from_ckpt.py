#!/usr/bin/env python3
"""Restore semiconductor X/Y/Overall from rescore checkpoint by company name (fix scramble)."""
from __future__ import annotations

import csv
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx
from vendor_intel.quadrant.rating_map import assign_quadrants_half_median

OUT = ROOT / "output" / "chatgpt_expand"
SLUG = "global_semiconductor_market_global"
QUERY = "Global Semiconductor Market"
FOLDER = OUT / SLUG
CKPT = OUT / "_audit" / "semi_xy_rescore_checkpoint.json"
INDUSTRY = "ICT, Automation, Semiconductor / Semiconductors"


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def _plain(name: str) -> str:
    return re.sub(
        r"\s*\((?:acquired by|subsidiary of|merged into)[^)]*\)\s*$",
        "",
        str(name or ""),
        flags=re.I,
    ).strip()


def main() -> None:
    done = json.loads(CKPT.read_text(encoding="utf-8")).get("done") or {}
    audit = json.loads(CKPT.read_text(encoding="utf-8")).get("audit") or {}

    xlsx = FOLDER / f"{SLUG}_FINAL.xlsx"
    wb = load_workbook(xlsx, data_only=True)
    ws = wb["Landscape"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows[0]]
    landscape: list[dict] = []
    missing: list[str] = []
    for r in rows[1:]:
        if not r:
            continue
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        name = _plain(str(d.get("Company") or d.get("Brand") or ""))
        if not name:
            continue
        key = _norm(name)
        sc = done.get(key)
        if not sc:
            # try brand
            sc = done.get(_norm(_plain(str(d.get("Brand") or ""))))
        if not sc:
            missing.append(name)
            landscape.append(d)
            continue
        d["Company"] = name
        d["Brand"] = str(d.get("Brand") or name)
        if str(d.get("Brand") or "").lower().startswith("("):
            d["Brand"] = name
        x = int(str(sc.get("X Score") or 0) or 0)
        y = int(str(sc.get("Y Score") or 0) or 0)
        o = int(str(sc.get("Overall Score") or 0) or 0)
        if not o and (x or y):
            o = int(round((x + y) / 2.0))
        d["X Score"] = str(x)
        d["Y Score"] = str(y)
        d["Overall Score"] = str(o)
        d["X"] = str(x)
        d["Y"] = str(y)
        d["Overall"] = str(o)
        d["Role"] = "Solution Provider"
        d["Distribution Type"] = "Solution Provider"
        d["Industry Category"] = sc.get("Industry Category") or INDUSTRY
        landscape.append(d)

    if missing:
        print(f"WARNING missing checkpoint scores for {len(missing)}: {missing[:10]}")

    # Preserve ownership display from Company Details
    details_by: dict[str, dict] = {}
    if "Company Details" in wb.sheetnames:
        ws = wb["Company Details"]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(h) for h in rows[0]]
        for r in rows[1:]:
            if not r:
                continue
            d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
            key = _norm(_plain(d.get("Brand") or d.get("Company") or ""))
            if key:
                details_by[key] = d
    for row in landscape:
        key = _norm(_plain(row.get("Company") or ""))
        det = details_by.get(key)
        if not det:
            continue
        brand = str(det.get("Brand") or "").strip()
        company_col = str(det.get("Company") or "").strip()
        if brand and not brand.lower().startswith("("):
            row["Brand"] = brand
        if company_col and company_col != brand and (
            "acquired by" in company_col.lower()
            or "subsidiary of" in company_col.lower()
            or company_col != brand
        ):
            row["_display_company"] = company_col
        if det.get("Found in"):
            row["Headquarters"] = det["Found in"]
            row["Found in"] = det["Found in"]

    xs = [float(r.get("X Score") or 0) for r in landscape]
    ys = [float(r.get("Y Score") or 0) for r in landscape]
    quads, mid_x, mid_y = assign_quadrants_half_median(xs, ys)
    for row, q in zip(landscape, quads):
        row["Quadrant"] = q

    audit = dict(audit or {})
    audit["score_midpoints"] = {"x": mid_x, "y": mid_y}
    audit["quadrant_mid_x"] = mid_x
    audit["quadrant_mid_y"] = mid_y
    audit["scored"] = len(landscape)
    audit["overall_formula"] = "(X+Y)/2"
    audit["rows"] = [
        {
            "company": r.get("Company"),
            "x": int(str(r.get("X Score") or 0) or 0),
            "y": int(str(r.get("Y Score") or 0) or 0),
            "overall": int(str(r.get("Overall Score") or 0) or 0),
            "quadrant": r.get("Quadrant"),
            "crawled": True,
        }
        for r in landscape
    ]

    # Build details from landscape scores (to_company_detail_rows already copies X/Y from row)
    detail_rows = to_company_detail_rows(landscape, QUERY, audit)
    # Map by brand/company — NEVER zip by position (details are sorted by Overall)
    by_key = {_norm(_plain(r.get("Company") or "")): r for r in landscape}
    for det in detail_rows:
        key = _norm(_plain(det.get("Brand") or det.get("Company") or ""))
        src = by_key.get(key)
        if not src:
            continue
        det["X"] = int(str(src.get("X Score") or 0) or 0)
        det["Y"] = int(str(src.get("Y Score") or 0) or 0)
        det["Overall"] = int(str(src.get("Overall Score") or 0) or 0)
        det["Quadrant"] = src.get("Quadrant")
        det["Role"] = "Solution Provider"
        if src.get("_display_company"):
            det["Company"] = src["_display_company"]
        if src.get("Brand"):
            det["Brand"] = src["Brand"]

    write_final_xlsx(
        xlsx,
        landscape,
        "Companies",
        {"query": QUERY, "xy_scoring": audit, "scores_restored_from_checkpoint": True},
        detail_rows=detail_rows,
    )
    extras = export_expand_quadrant_outputs(
        FOLDER, detail_rows, QUERY, country="global", audit=audit, chart_n=20
    )
    (FOLDER / "chatgpt_xy_scores_batch_all.json").write_text(
        json.dumps(audit, indent=2), encoding="utf-8"
    )
    csv_path = FOLDER / f"{SLUG}_companies.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(
            f,
            fieldnames=["Brand", "Company", "Role", "Quadrant", "X", "Y", "Overall", "Found in"],
            extrasaction="ignore",
        )
        w.writeheader()
        for d in detail_rows:
            w.writerow(d)

    # Verify TSMC etc.
    checks = ["TSMC", "MediaTek", "NVIDIA", "Phison Electronics", "Aura Semiconductor", "Intel Corporation"]
    print(f"Restored {len(landscape)} rows; mid_x={mid_x} mid_y={mid_y}")
    for name in checks:
        src = by_key.get(_norm(name))
        if src:
            print(
                f"  {name}: X={src.get('X Score')} Y={src.get('Y Score')} "
                f"O={src.get('Overall Score')} Q={src.get('Quadrant')}"
            )
    print(f"html -> {extras.get('html')}")


if __name__ == "__main__":
    main()
