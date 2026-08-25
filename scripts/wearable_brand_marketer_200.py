#!/usr/bin/env python3
"""Wearable medical: keep Brand/Marketer only, purge off-market, add 70+ brands, rescore."""
from __future__ import annotations

import asyncio
import json
import os
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

from openpyxl import load_workbook

from vendor_intel.placeholders.load_keys import apply_env_overrides
from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    score_expand_rows,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx
from vendor_intel.quadrant.brand_meta import format_acquired_suffix

OUT = ROOT / "output" / "chatgpt_expand"
SLUG = "global_wearable_medical_devices_market_global"
QUERY = "Global Wearable Medical Devices Market"
FOLDER = OUT / SLUG
INDUSTRY = "Healthcare / Wearable Medical Devices"

OEM_REMOVE = {
    "bluespark technologies",
    "sensirion",
    "sensitec",
    "valencell",
    "leman micro devices",
    "tdk",
}
WRONG_VERTICAL = {"vivonic", "nymi"}
CONSUMER_GADGET = {
    "garmin",
    "whoop",
    "polar electro",
    "suunto",
    "oura",
    "fitbit",
    "apple",
    "xiaomi",
    "huawei",
    "sony",
    "lg electronics",
    "seiko",
    "seiko epson",
    "samsung electronics",
    "panasonic",
    "fujitsu",
    "nec",
    "toshiba",
    "ttp",
    "samsung medison",
}
REGIONAL_RE = re.compile(
    r"\((canada|middle east|brazil|latin america)\)|"
    r"healthcare middle east|healthcare \(canada\)|"
    r"philips brazil|nihon kohden latin",
    re.I,
)

# Ownership to re-apply after rebuild (Brand stays plain; Company = suffix)
OWNERSHIP = {
    "fitbit": ("Google", "acquired_by", "2021"),  # removed as consumer
    "hillrom": ("Baxter", "acquired_by", "2021"),
    "biotelemetry": ("Philips", "acquired_by", "2021"),
    "preventice solutions": ("Boston Scientific", "acquired_by", "2021"),
    "smiths medical": ("ICU Medical", "acquired_by", "2022"),
    "ectosense": ("ResMed", "acquired_by", "2021"),
    "everion": ("Biofourmis", "acquired_by", "2019"),
    "zephyr technology": ("Medtronic", "acquired_by", "2011"),
    "zoll medical": ("Asahi Kasei", "acquired_by", "2012"),
    "smi (sensomotoric instruments)": ("Apple", "acquired_by", "2017"),
    "cardiac insight": ("Dreamtech", "acquired_by", "2022"),
    "cosinuss": ("corpuls", "acquired_by", "2025"),
    "neurometrix": ("electroCore", "acquired_by", "2025"),
    "masimo": ("Danaher", "acquired_by", "2026"),
    "physiq": ("Prolaio", "acquired_by", ""),
    "beijing choice (choicemmed)": ("Tianjin Chase Sun Pharmaceutical", "acquired_by", "2015"),
    "medisana": ("Ogawa Smart HealthCare (Xiamen Comfort Science)", "acquired_by", "2016"),
    "intelesens": ("Renew Health", "acquired_by", "2017"),
    "hearx group": ("LXE Hearing", "merged_into", "2025"),
    "earlysense": None,  # special plain
    "bosch healthcare solutions": ("Robert Bosch GmbH", "subsidiary_of", ""),
    "omron healthcare": ("Omron", "subsidiary_of", ""),
    "a&d": ("A&D Holon Holdings", "subsidiary_of", ""),
    "sensium healthcare": ("The Surgical Company Group", "subsidiary_of", ""),
    "vitaliberty": ("vitagroup AG", "subsidiary_of", ""),
    "bardy diagnostics": ("Baxter", "acquired_by", "2021"),
    "biotelemetry heart": ("Philips", "acquired_by", "2021"),
    "preventice bodyguardian": ("Boston Scientific", "acquired_by", "2021"),
    "bigfoot biomedical": ("Abbott", "acquired_by", "2023"),
    "livongo": ("Teladoc Health", "acquired_by", "2020"),
    "eargo": ("LXE Hearing", "merged_into", "2025"),
    "masimo w1": ("Danaher", "acquired_by", "2026"),
}


def _norm(s: str) -> str:
    s = str(s or "").strip().lower()
    s = re.sub(r"[®™]", "", s)
    s = re.sub(
        r"\b(ltd\.?|limited|inc\.?|incorporated|llc|gmbh|ag|corp\.?|corporation|co\.|company|plc)\b\.?",
        "",
        s,
    )
    s = re.sub(r"\s+", " ", s).strip(" ,.")
    return s


def _remove_reason(brand: str) -> str | None:
    k = _norm(brand)
    if k in OEM_REMOVE:
        return "oem"
    if k in WRONG_VERTICAL:
        return "wrong_vertical"
    if k in CONSUMER_GADGET:
        return "consumer_gadget"
    if REGIONAL_RE.search(brand or ""):
        return "regional_parent"
    if k in {"omron healthcare (canada)", "omron healthcare middle east", "philips brazil", "nihon kohden latin america", "apple (canada)"}:
        return "regional_parent"
    return None


def _row_from_candidate(c: dict) -> dict:
    name = c["company"]
    hq = c.get("hq") or ""
    return {
        "Company": name,
        "Website": c.get("website") or "",
        "Founded": c.get("founded") or "",
        "Headquarters": hq,
        "Continent / Geography": c.get("continent") or "",
        "Operational Presence": c.get("presence") or "",
        "Ownership": "Independent",
        "Employees": "",
        "Core Categories": c.get("categories") or "",
        "Specialty Focus": c.get("specialty") or "",
        "Key Brands Represented": name,
        "Retail / E-commerce / Both": "No",
        "Distribution Type": "Brand",
        "Contact Person": "",
        "Role": "Brand",
        "Email": "",
        "LinkedIn": "",
        "Office No.": "",
        "Country Code": c.get("country") or "",
        "Region Code": c.get("region") or "",
        "Summary": f"{name} — wearable medical device brand. {c.get('specialty') or ''}".strip(),
        "X Score": "",
        "Y Score": "",
        "Overall Score": "",
        "Quadrant": "",
        "Industry Category": INDUSTRY,
        "Brand": name,
        "Found in": hq,
    }


def _parent_stem(name: str) -> str:
    k = _norm(name)
    for stem in (
        "abbott",
        "insulet",
        "masimo",
        "philips",
        "medtronic",
        "biofourmis",
        "current health",
        "tandem",
        "preventice",
        "biotelemetry",
        "biobeat",
        "bio beat",
    ):
        if k == stem or k.startswith(stem + " ") or stem in k:
            return stem
    return k


async def main() -> int:
    apply_env_overrides()
    # Fill landscape from research; score via DeepSeek without full crawl for speed
    os.environ["EXPAND_XY_SKIP_CRAWL"] = "1"
    os.environ["EXPAND_MARKET_AXIS_LLM"] = "1"
    os.environ.setdefault("EXPAND_XY_CONCURRENT", "6")

    cands = json.loads(
        (OUT / "_audit" / "wearable_new_brands_70plus.json").read_text(encoding="utf-8")
    )
    xlsx = FOLDER / f"{SLUG}_FINAL.xlsx"
    wb = load_workbook(xlsx, data_only=True)
    ws = wb["Landscape"]
    rows_raw = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows_raw[0]]

    details_by: dict[str, dict] = {}
    if "Company Details" in wb.sheetnames:
        dws = wb["Company Details"]
        drows = list(dws.iter_rows(values_only=True))
        dhdr = [str(h) for h in drows[0]]
        for r in drows[1:]:
            if not r:
                continue
            d = {dhdr[i]: ("" if r[i] is None else r[i]) for i in range(len(dhdr)) if i < len(r)}
            key = _norm(d.get("Brand") or d.get("Company") or "")
            if key:
                details_by[key] = d

    kept: list[dict] = []
    removed: list[dict] = []
    existing = set()
    parent_stems = set()
    for r in rows_raw[1:]:
        if not r:
            continue
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        brand = str(d.get("Company") or d.get("Brand") or "").strip()
        if not brand:
            continue
        # Prefer Details brand/role/scores
        key = _norm(brand)
        det = details_by.get(key)
        if det:
            if det.get("Brand"):
                brand = str(det["Brand"]).strip()
                d["Brand"] = brand
                d["Company"] = brand
            if det.get("Role"):
                d["Role"] = det["Role"]
                d["Distribution Type"] = det["Role"]
            if det.get("X") not in (None, ""):
                d["X Score"] = det["X"]
                d["Y Score"] = det["Y"]
                d["Overall Score"] = det["Overall"]
                d["Quadrant"] = det.get("Quadrant") or ""
            if det.get("Found in"):
                d["Headquarters"] = det["Found in"]
                d["Found in"] = det["Found in"]
            # restore owned company display later
            if det.get("Company") and str(det["Company"]).startswith("("):
                d["_display_company"] = det["Company"]

        reason = _remove_reason(brand)
        role = str(d.get("Role") or d.get("Distribution Type") or "Brand").strip()
        if role.upper() == "OEM" or _norm(role) == "oem":
            reason = reason or "oem"
        if reason:
            removed.append({"brand": brand, "reason": reason})
            continue
        # Force Brand (Marketer unused unless already Marketer)
        if role not in {"Brand", "Marketer"}:
            d["Role"] = "Brand"
            d["Distribution Type"] = "Brand"
        else:
            d["Role"] = role
            d["Distribution Type"] = role
        d["Brand"] = brand
        d["Company"] = brand
        d["Industry Category"] = INDUSTRY
        kept.append(d)
        existing.add(_norm(brand))
        parent_stems.add(_parent_stem(brand))

    added: list[dict] = []
    skipped_dup: list[str] = []
    for c in cands:
        name = c["company"]
        key = _norm(name)
        stem = _parent_stem(name)
        # Skip if already present or product-line of existing parent
        if key in existing:
            skipped_dup.append(name)
            continue
        if stem in parent_stems and stem != key:
            # e.g. Abbott FreeStyle Libre when Abbott exists; Masimo W1 when Masimo exists
            skipped_dup.append(f"{name} (parent {stem})")
            continue
        # Biobeat vs BioBeat
        if key.replace(" ", "") in {e.replace(" ", "") for e in existing}:
            skipped_dup.append(name)
            continue
        if "biobeat" in key and any("biobeat" in e.replace(" ", "") for e in existing):
            skipped_dup.append(name)
            continue
        if "biofourmis" in key and any("biofourmis" in e for e in existing):
            skipped_dup.append(name)
            continue
        if "current health" in key and any("current health" in e for e in existing):
            skipped_dup.append(name)
            continue
        row = _row_from_candidate(c)
        kept.append(row)
        added.append(row)
        existing.add(key)
        parent_stems.add(stem)

    # EarlySense special company text
    for row in kept:
        if _norm(row.get("Brand") or "") == "earlysense":
            row["_display_company"] = (
                "(ceased; hospital IP to Hillrom, 2021; remaining assets to TytoCare, 2022)"
            )

    # Apply ownership suffixes for known owned brands
    for row in kept:
        key = _norm(row.get("Brand") or "")
        info = OWNERSHIP.get(key)
        if info is None and key == "earlysense":
            continue
        if not info:
            continue
        owner, relation, year = info
        row["_display_company"] = format_acquired_suffix(owner, relation=relation, year=year)
        row["Ownership"] = f"{relation.replace('_', ' ')} {owner}" + (f", {year}" if year else "")
        row["parent_owner"] = owner
        row["ownership_relation"] = relation
        if year:
            row["ownership_year"] = year

    target = 200
    print(f"Kept {len(kept) - len(added)} + added {len(added)} = {len(kept)} (removed {len(removed)})", flush=True)
    if len(kept) < target:
        print(f"WARNING: only {len(kept)} < {target}; need more candidates", flush=True)

    # Score newly added (and any missing scores) via DeepSeek
    need_score = [r for r in kept if r.get("X Score") in (None, "")]
    print(f"Scoring {len(need_score)} rows via DeepSeek (skip crawl)...", flush=True)
    if need_score:
        scored, _score_audit = await score_expand_rows(need_score, QUERY, country="global")
        by_name = {_norm(r.get("Company") or r.get("Brand") or ""): r for r in scored}
        for r in kept:
            key = _norm(r.get("Company") or r.get("Brand") or "")
            s = by_name.get(key)
            if not s:
                continue
            for col in ("X Score", "Y Score", "Overall Score", "Quadrant", "Summary"):
                if s.get(col) not in (None, ""):
                    r[col] = s[col]

    # Half-median quadrant reassignment across full set
    from vendor_intel.quadrant.rating_map import assign_quadrants_half_median

    xs: list[float] = []
    ys: list[float] = []
    for r in kept:
        try:
            xs.append(float(r.get("X Score") or r.get("X") or 50))
        except (TypeError, ValueError):
            xs.append(50.0)
        try:
            ys.append(float(r.get("Y Score") or r.get("Y") or 50))
        except (TypeError, ValueError):
            ys.append(50.0)
    quads, _, _ = assign_quadrants_half_median(xs, ys)
    for r, q in zip(kept, quads):
        r["Quadrant"] = q
        try:
            x = float(r.get("X Score") or 50)
            y = float(r.get("Y Score") or 50)
            r["Overall Score"] = round((x + y) / 2)
        except (TypeError, ValueError):
            pass

    audit: dict = {}
    audit_path = FOLDER / "chatgpt_expand_batch_all_audit.json"
    if audit_path.exists():
        audit = json.loads(audit_path.read_text(encoding="utf-8")).get("xy_scoring") or {}
    qjson = FOLDER / f"{SLUG}_quadrant.json"
    if qjson.exists():
        try:
            payload = json.loads(qjson.read_text(encoding="utf-8"))
            crit = payload.get("criteria") or {}
            if crit.get("parameter_definitions"):
                audit["parameter_definitions"] = crit["parameter_definitions"]
            labels = crit.get("axis_labels") or {}
            if labels.get("x"):
                audit["axis_x"] = labels["x"]
            if labels.get("y"):
                audit["axis_y"] = labels["y"]
            if crit.get("x_axis"):
                audit["x_features"] = crit["x_axis"]
            if crit.get("y_axis"):
                audit["y_features"] = crit["y_axis"]
        except Exception:
            pass

    detail_rows = to_company_detail_rows(kept, QUERY, audit)
    for det in detail_rows:
        key = _norm(det.get("Brand") or "")
        src = next((r for r in kept if _norm(r.get("Brand") or r.get("Company") or "") == key), None)
        det["Role"] = (src.get("Role") if src else None) or "Brand"
        if src and src.get("_display_company"):
            det["Company"] = src["_display_company"]
        elif key == "earlysense":
            det["Company"] = (
                "(ceased; hospital IP to Hillrom, 2021; remaining assets to TytoCare, 2022)"
            )
        else:
            # Brand == Company for independents
            if not str(det.get("Company") or "").startswith("("):
                det["Company"] = det.get("Brand") or det.get("Company")

    write_final_xlsx(
        xlsx,
        kept,
        "Companies",
        {
            "query": QUERY,
            "xy_scoring": audit,
            "removed": removed,
            "added": [r.get("Company") for r in added],
            "skipped_dup": skipped_dup,
        },
        detail_rows=detail_rows,
    )
    extras = export_expand_quadrant_outputs(
        FOLDER, detail_rows, QUERY, country="global", audit=audit, chart_n=20
    )

    # Re-apply ownership script lightly already done; ensure roles Brand-only
    out = {
        "before_approx": 167,
        "removed": removed,
        "added_count": len(added),
        "added": [r.get("Company") for r in added],
        "skipped_dup": skipped_dup,
        "final_count": len(detail_rows),
        "roles": {},
        "html": str(extras.get("html") or ""),
    }
    from collections import Counter

    out["roles"] = dict(Counter(r.get("Role") for r in detail_rows))
    audit_out = OUT / "_audit" / "wearable_brand_marketer_200.json"
    audit_out.write_text(json.dumps(out, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Final {out['final_count']} roles={out['roles']}", flush=True)
    print(f"Removed {len(removed)}; added {len(added)}; skipped dups {len(skipped_dup)}", flush=True)
    print(f"html -> {extras.get('html')}", flush=True)
    print(f"audit -> {audit_out}", flush=True)
    return 0 if out["final_count"] >= 200 else 2


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
