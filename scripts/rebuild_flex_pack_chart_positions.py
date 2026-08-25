#!/usr/bin/env python3
"""Re-place flex-pack chart dots after cell-bound fix (no X/Y/Q changes)."""
from __future__ import annotations

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import export_expand_quadrant_outputs

SLUG = "global_flexible_packaging_market_global"
QUERY = "Global Flexible Packaging Market"
OUT = ROOT / "output" / "chatgpt_expand" / SLUG
XLSX = OUT / f"{SLUG}_FINAL.xlsx"
JSONP = OUT / f"{SLUG}_quadrant.json"


def main() -> None:
    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Company Details"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows[0]]
    details: list[dict] = []
    for r in rows[1:]:
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        if not str(d.get("Brand") or "").strip():
            continue
        for k in ("X", "Y", "Overall"):
            try:
                d[k] = int(round(float(d.get(k) or 0)))
            except (TypeError, ValueError):
                d[k] = 0
        details.append(d)

    audit: dict = {}
    if JSONP.exists():
        prior = json.loads(JSONP.read_text(encoding="utf-8"))
        crit = prior.get("criteria") or {}
        for k in (
            "parameter_definitions",
            "x_features",
            "y_features",
            "axis_x",
            "axis_y",
            "industry_group",
            "industry_category",
        ):
            if crit.get(k) is not None:
                audit[k] = crit[k]
            elif prior.get(k) is not None:
                audit[k] = prior[k]
        # Prefer nested criteria axes labels
        if crit.get("x_axis_label"):
            audit["axis_x"] = crit["x_axis_label"]
        if crit.get("y_axis_label"):
            audit["axis_y"] = crit["y_axis_label"]

    extras = export_expand_quadrant_outputs(
        OUT, details, QUERY, country="global", audit=audit, chart_n=20
    )
    d = json.loads(Path(extras["json"]).read_text(encoding="utf-8"))
    for b in d.get("brands") or []:
        if not b.get("on_chart"):
            continue
        name = str(b.get("brand") or "")
        left = int(b.get("left_pct") or 0)
        top = int(b.get("top_pct") or 0)
        flag = "  <-- near mid" if (40 < left < 60) or (40 < top < 60) else ""
        print(f"{name[:40]:40} Q={b.get('quadrant'):18} L={left:3} T={top:3}{flag}")
    creative = next(
        (b for b in d["brands"] if "creative" in str(b.get("brand") or "").lower()),
        None,
    )
    if creative:
        print(
            "Creative Polypack:",
            f"left={creative.get('left_pct')}",
            f"top={creative.get('top_pct')}",
            f"Q={creative.get('quadrant')}",
        )


if __name__ == "__main__":
    main()
