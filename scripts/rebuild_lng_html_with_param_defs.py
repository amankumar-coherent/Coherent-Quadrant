#!/usr/bin/env python3
"""Rebuild LNG HTML with Parameter Definitions at the end (from xy audit)."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)

OUT = ROOT / "output" / "chatgpt_expand"
SLUG = "global_liquefied_natural_gas_market_global"
QUERY = "Global Liquefied Natural Gas Market"
FOLDER = OUT / SLUG


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def main() -> int:
    xlsx = FOLDER / f"{SLUG}_FINAL.xlsx"
    wb = load_workbook(xlsx, data_only=True)

    dws = wb["Company Details"]
    drows = list(dws.iter_rows(values_only=True))
    dhdr = [str(h) for h in drows[0]]
    details: list[dict] = []
    for r in drows[1:]:
        if not r:
            continue
        d = {dhdr[i]: ("" if r[i] is None else r[i]) for i in range(len(dhdr)) if i < len(r)}
        if d.get("Brand"):
            details.append(d)

    ws = wb["Landscape"]
    lrows = list(ws.iter_rows(values_only=True))
    lhdr = [str(h) for h in lrows[0]]
    land_by: dict[str, dict] = {}
    for r in lrows[1:]:
        if not r:
            continue
        d = {lhdr[i]: ("" if r[i] is None else r[i]) for i in range(len(lhdr)) if i < len(r)}
        key = _norm(d.get("Company") or d.get("Brand") or "")
        if key:
            land_by[key] = d

    landscape: list[dict] = []
    for det in details:
        brand = str(det.get("Brand") or "").strip()
        key = _norm(brand)
        src = land_by.get(key) or {}
        row = dict(src)
        row["Brand"] = brand
        row["Company"] = brand
        row["Role"] = det.get("Role") or "Brand"
        row["Distribution Type"] = row["Role"]
        row["X Score"] = det.get("X")
        row["Y Score"] = det.get("Y")
        row["Overall Score"] = det.get("Overall")
        row["Quadrant"] = det.get("Quadrant")
        row["Found in"] = det.get("Found in") or ""
        row["Headquarters"] = det.get("Found in") or row.get("Headquarters") or ""
        if det.get("Company") and (
            str(det["Company"]).startswith("(") or _norm(det.get("Company")) != key
        ):
            row["_display_company"] = det["Company"]
        landscape.append(row)

    xy = json.loads((FOLDER / "chatgpt_xy_scores_batch_all.json").read_text(encoding="utf-8"))
    param_defs = xy.get("parameter_definitions") or {"x": {}, "y": {}}
    # Prefer deep-rescore axes if xy audit is empty/wrong
    if not any((param_defs.get("x") or {}).values()):
        from vendor_intel.placeholders.load_keys import apply_env_overrides
        from vendor_intel.quadrant.axis_define import explain_market_parameters

        apply_env_overrides()
        param_defs = explain_market_parameters(
            QUERY,
            list(xy.get("x_features") or []),
            list(xy.get("y_features") or []),
            axis_x=str(xy.get("axis_x") or ""),
            axis_y=str(xy.get("axis_y") or ""),
            geography="global",
        )
        xy["parameter_definitions"] = param_defs
        (FOLDER / "chatgpt_xy_scores_batch_all.json").write_text(
            json.dumps(xy, indent=2, ensure_ascii=False), encoding="utf-8"
        )

    audit = {
        "industry_group": xy.get("industry_group") or "Others",
        "industry_category": xy.get("industry_category") or "Energy",
        "axis_x": xy.get("axis_x") or "Asset & Operating Capability",
        "axis_y": xy.get("axis_y") or "Portfolio & Commercial Strategy",
        "x_features": xy.get("x_features")
        or [
            "Asset / Project Portfolio Breadth",
            "Operational Efficiency & Throughput",
            "Reliability & Asset Performance",
            "Safety & Regulatory Compliance",
            "Sustainability & Energy Transition",
        ],
        "y_features": xy.get("y_features")
        or [
            "Geographic & Corridor Reach",
            "Offtake / Customer Contracts",
            "Financial Performance",
            "Growth & Project Pipeline",
            "Partnerships & Business Expansion",
        ],
        "parameter_definitions": param_defs,
        "quadrant_mid_x": xy.get("quadrant_mid_x") or xy.get("mid_x"),
        "quadrant_mid_y": xy.get("quadrant_mid_y") or xy.get("mid_y"),
    }

    detail_rows = to_company_detail_rows(landscape, QUERY, audit)
    for det in detail_rows:
        key = _norm(det.get("Brand") or "")
        src = next((r for r in landscape if _norm(r.get("Brand") or "") == key), None)
        if not src:
            continue
        det["Role"] = src.get("Role") or "Brand"
        if src.get("_display_company"):
            det["Company"] = src["_display_company"]
        det["Found in"] = src.get("Found in") or det.get("Found in") or ""

    paths = export_expand_quadrant_outputs(
        FOLDER, detail_rows, QUERY, country="global", audit=audit, chart_n=20
    )
    html = Path(paths.get("html") or (FOLDER / f"{SLUG}_report.html"))
    text = html.read_text(encoding="utf-8")
    has = "Parameter Definitions" in text
    pending = text.count("Definition pending for this market parameter.")
    print(f"html={html}")
    print(f"parameter_definitions_section={has}")
    print(f"pending_placeholders={pending}")
    print(f"x_defs={len((param_defs.get('x') or {}))} y_defs={len((param_defs.get('y') or {}))}")
    print(f"axis_x={audit['axis_x']}")
    print(f"axis_y={audit['axis_y']}")
    return 0 if has and pending == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
