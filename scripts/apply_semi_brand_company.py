"""Apply web-verified Brand/Company ownership for semiconductor market.

Company column uses ``(acquired by Parent)`` / ``(subsidiary of Parent)`` /
``(merged into Parent)`` when ownership is confirmed; otherwise Brand == Company.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx
from vendor_intel.quadrant.brand_meta import brand_display_fields, format_acquired_suffix

ROOT = Path(r"D:\Coherent-Quadrant\output\chatgpt_expand")
SLUG = "global_semiconductor_market_global"
QUERY = "Global Semiconductor Market"

# brand_key (normalized) -> ownership. Only entries verified via public sources.
# relation: acquired_by | subsidiary_of | merged_into
# brand_rename: optional new Brand label (plain name, no suffix)
OWNERSHIP: dict[str, dict[str, str]] = {
    # Acquisitions
    "graphcore": {
        "owner": "SoftBank Group",
        "relation": "acquired_by",
        "year": "2024",
        "source": "SoftBank / Graphcore July 2024; Reuters",
    },
    "gan systems": {
        "owner": "Infineon Technologies",
        "relation": "acquired_by",
        "year": "2023",
        "source": "Infineon closed GaN Systems Oct 2023",
    },
    "autotalks": {
        "owner": "Qualcomm",
        "relation": "acquired_by",
        "year": "2025",
        "source": "Qualcomm completed Autotalks Jun 2025",
    },
    "redlen technologies": {
        "owner": "Canon",
        "relation": "acquired_by",
        "year": "2021",
        "source": "Canon press 29 Sep 2021",
    },
    "ampere computing llc": {
        "owner": "SoftBank Group",
        "relation": "acquired_by",
        "year": "2025",
        "source": "SoftBank completed Ampere 25 Nov 2025",
        "brand_rename": "Ampere Computing",
    },
    "ampere computing": {
        "owner": "SoftBank Group",
        "relation": "acquired_by",
        "year": "2025",
        "source": "SoftBank completed Ampere 25 Nov 2025 (SoftBank Group press 26 Nov 2025)",
        "brand_rename": "Ampere Computing",
    },
    "plessey semiconductors": {
        "owner": "Haylo Labs",
        "relation": "acquired_by",
        "year": "2025",
        "source": "Plessey / Haylo Labs Aug 2025",
    },
    "solantro semiconductor": {
        "owner": "Huada Semiconductor",
        "relation": "acquired_by",
        "year": "2018",
        "source": "Huada (HDSC) acquired Solantro 2018",
    },
    "saankhya labs": {
        "owner": "Tejas Networks",
        "relation": "acquired_by",
        "year": "2022",
        "source": "Tejas Networks acquisition announcement",
    },
    "stats chippac": {
        "owner": "JCET Group",
        "relation": "acquired_by",
        "year": "2015",
        "source": "JCET acquired STATS ChipPAC from Temasek",
    },
    "siliconware precision industries (spil)": {
        "owner": "ASE Technology Holding",
        "relation": "acquired_by",
        "year": "2018",
        "source": "ASE Holding / SPIL combination",
        "brand_rename": "SPIL (Siliconware Precision Industries)",
    },
    "key foundry co., ltd.": {
        "owner": "SK Hynix",
        "relation": "acquired_by",
        "year": "2022",
        "source": "SK hynix completed Key Foundry Aug 2022",
        "brand_rename": "Key Foundry",
    },
    "teledyne dalsa": {
        "owner": "Teledyne Technologies",
        "relation": "acquired_by",
        "year": "2011",
        "source": "Teledyne acquired DALSA 2011",
    },
    "nexperia": {
        "owner": "Wingtech Technology",
        "relation": "acquired_by",
        "year": "2019",
        "source": "Wingtech acquired Nexperia (legal parent; Dutch governance disputes ongoing)",
    },
    "silead india": {
        "owner": "GigaDevice",
        "relation": "acquired_by",
        "year": "2019",
        "source": "GigaDevice completed Silead Inc. acquisition 2019",
        "brand_rename": "Silead",
    },
    "silead": {
        "owner": "GigaDevice",
        "relation": "acquired_by",
        "year": "2019",
        "source": "GigaDevice completed Silead Inc. acquisition 2019",
        "brand_rename": "Silead",
    },
    "arm holdings": {
        "owner": "SoftBank Group",
        "relation": "subsidiary_of",
        "source": "SoftBank acquired Arm 2016; still ~90% owner after 2023 IPO (Wikipedia / Arm filings)",
    },
    # Subsidiaries / divisions
    "bosch sensortec": {
        "owner": "Robert Bosch GmbH",
        "relation": "subsidiary_of",
        "source": "Bosch Sensortec wholly owned Bosch MEMS/sensor unit",
    },
    "mobileye": {
        "owner": "Intel Corporation",
        "relation": "subsidiary_of",
        "source": "Intel majority / controlling stake in Mobileye Global",
    },
    "sony semiconductor solutions": {
        "owner": "Sony Group",
        "relation": "subsidiary_of",
        "source": "Sony Semiconductor Solutions Corp. is Sony Group company",
    },
    "toshiba electronic devices & storage corporation": {
        "owner": "Toshiba",
        "relation": "subsidiary_of",
        "source": "Toshiba Electronic Devices & Storage is Toshiba group company",
        "brand_rename": "Toshiba Electronic Devices & Storage",
    },
    "toshiba electronic devices & storage": {
        "owner": "Toshiba",
        "relation": "subsidiary_of",
        "source": "Toshiba Electronic Devices & Storage is Toshiba group company",
        "brand_rename": "Toshiba Electronic Devices & Storage",
    },
    "carsem": {
        "owner": "Malaysian Pacific Industries",
        "relation": "subsidiary_of",
        "source": "Carsem is MPI / Hong Leong Group OSAT operating company",
    },
    "samsung electro-mechanics": {
        "owner": "Samsung",
        "relation": "subsidiary_of",
        "source": "Samsung Electro-Mechanics is Samsung Group affiliate",
    },
    "lg innotek": {
        "owner": "LG",
        "relation": "subsidiary_of",
        "source": "LG Innotek is LG Group affiliate",
    },
    "siemens eda (formerly mentor graphics)": {
        "owner": "Siemens",
        "relation": "subsidiary_of",
        "source": "Siemens acquired Mentor Graphics 2017; operates as Siemens EDA division",
        "brand_rename": "Siemens EDA",
    },
    "siemens eda": {
        "owner": "Siemens",
        "relation": "subsidiary_of",
        "source": "Siemens acquired Mentor Graphics 2017; operates as Siemens EDA division",
        "brand_rename": "Siemens EDA",
    },
    # Absorbed / renamed
    "lapis technology co., ltd.": {
        "owner": "Rohm Semiconductor",
        "relation": "merged_into",
        "year": "2024",
        "source": "ROHM absorption-type merger of LAPIS Technology Apr 2024",
        "brand_rename": "LAPIS Technology",
    },
    "lapis technology": {
        "owner": "Rohm Semiconductor",
        "relation": "merged_into",
        "year": "2024",
        "source": "ROHM absorption-type merger of LAPIS Technology Apr 2024",
        "brand_rename": "LAPIS Technology",
    },
    "ams ag": {
        "owner": "ams-OSRAM AG",
        "relation": "merged_into",
        "year": "2022",
        "source": "ams AG renamed ams-OSRAM AG after OSRAM integration",
        "brand_rename": "ams OSRAM",
        "company_plain": "ams-OSRAM AG",
    },
    "ams osram": {
        "owner": "ams-OSRAM AG",
        "relation": "merged_into",
        "year": "2022",
        "source": "ams AG renamed ams-OSRAM AG after OSRAM integration",
        "brand_rename": "ams OSRAM",
        "company_plain": "ams-OSRAM AG",
    },
    # JV note (not acquisition)
    "ssmc (systems on silicon manufacturing company)": {
        "owner": "NXP Semiconductors / TSMC",
        "relation": "subsidiary_of",
        "source": "SSMC is JV of NXP and TSMC",
        "brand_rename": "SSMC",
        "company_plain": "NXP Semiconductors / TSMC (JV)",
    },
    "ssmc": {
        "owner": "NXP Semiconductors / TSMC",
        "relation": "subsidiary_of",
        "source": "SSMC is JV of NXP and TSMC",
        "brand_rename": "SSMC",
        "company_plain": "NXP Semiconductors / TSMC (JV)",
    },
    "silterra malaysia": {
        "owner": "Dagang NeXchange (DNeX)",
        "relation": "subsidiary_of",
        "source": "DNeX holds ~60% of SilTerra Malaysia",
        "brand_rename": "SilTerra Malaysia",
    },
    "semikron danfoss": {
        "owner": "Danfoss",
        "relation": "subsidiary_of",
        "year": "2026",
        "source": "Danfoss took 100% ownership of Semikron Danfoss (Mar 2026, Danfoss press)",
    },
    "on semiconductor": {
        "owner": "onsemi",
        "relation": "merged_into",
        "year": "2021",
        "source": "ON Semiconductor rebranded as onsemi (Aug 2021)",
        "brand_rename": "onsemi",
        "company_plain": "onsemi",
    },
    "onsemi": {
        "owner": "onsemi",
        "relation": "merged_into",
        "year": "2021",
        "source": "ON Semiconductor rebranded as onsemi (Aug 2021)",
        "brand_rename": "onsemi",
        "company_plain": "onsemi",
    },
    # Additional web-verified (Aug 2026 pass)
    "ansys": {
        "owner": "Synopsys",
        "relation": "acquired_by",
        "year": "2025",
        "source": "Synopsys completed Ansys acquisition 17 Jul 2025 (Synopsys / SEC 8-K)",
    },
    "omnivision technologies": {
        "owner": "OmniVision Group",
        "relation": "acquired_by",
        "year": "2019",
        "source": "Will Semiconductor (now OmniVision Group) acquired OmniVision May 2019; parent renamed 2025",
        "brand_rename": "OmniVision Technologies",
    },
    "will semiconductor": {
        "owner": "OmniVision Group",
        "relation": "merged_into",
        "year": "2025",
        "source": "Will Semiconductor Co., Ltd. Shanghai renamed OmniVision Integrated Circuits Group / OmniVision Group (Jun 2025)",
        "brand_rename": "OmniVision Group",
        "company_plain": "OmniVision Group",
    },
    "hitachi high-tech": {
        "owner": "Hitachi, Ltd.",
        "relation": "subsidiary_of",
        "year": "2020",
        "source": "Hitachi tender offer; Hitachi High-Tech wholly owned subsidiary since May 2020",
        "brand_rename": "Hitachi High-Tech",
    },
    "imagination technologies": {
        "owner": "Canyon Bridge Capital Partners",
        "relation": "acquired_by",
        "year": "2017",
        "source": "Canyon Bridge took Imagination private Nov 2017 (still owner; sale process reported 2025)",
    },
    "lx semicon": {
        "owner": "LX Holdings",
        "relation": "subsidiary_of",
        "source": "LX Semicon (ex-Silicon Works) is LX Holdings affiliate / group company",
        "brand_rename": "LX Semicon",
    },
    "shinko electric industries": {
        "owner": "JIC Capital (JICC-04)",
        "relation": "acquired_by",
        "year": "2025",
        "source": "JICC-04 tender offer completed Mar 2025; Fujitsu exited; Shinko taken private under JIC Capital",
    },
}

# DO NOT invent Cyient Semiconductors rename — original row is Cyient (parent).
# Ineda Systems: DO NOT label acquired by Intel (acqui-hire only; IP stayed with Ineda).

# Plain Brand cleanups (no ownership change) — drop legal suffixes for display
BRAND_CLEAN: dict[str, str] = {
    "cadence design systems, inc.": "Cadence Design Systems",
    "synopsys, inc.": "Synopsys",
    "power integrations, inc.": "Power Integrations",
    "magnachip semiconductor corporation": "Magnachip",
    "lx semicon co., ltd.": "LX Semicon",
    "key foundry": "Key Foundry",
    "ampere computing": "Ampere Computing",
    "silead": "Silead",
    "ams osram": "ams OSRAM",
    "siemens eda": "Siemens EDA",
    "toshiba electronic devices & storage": "Toshiba Electronic Devices & Storage",
    "spil (siliconware precision industries)": "SPIL (Siliconware Precision Industries)",
    "cyient semiconductors": "Cyient",  # revert invented subsidiary brand name
    "will semiconductor": "OmniVision Group",
    "omnivision group": "OmniVision Group",
    "hitachi high-tech": "Hitachi High-Tech",
    "global unichip corporation (guc)": "Global Unichip Corporation (GUC)",
}



def _norm(s: str) -> str:
    s = str(s or "").strip().lower()
    s = re.sub(r"\s+", " ", s)
    return s


def _identity(row: dict) -> str:
    # Prefer Brand; Landscape often has plain Company while Details may already
    # have "(acquired by …)" — never match on the ownership suffix alone.
    brand = str(row.get("Brand") or "").strip()
    company = str(row.get("Company") or "").strip()
    if brand and not brand.lower().startswith("("):
        return brand
    if company and not company.lower().startswith("("):
        return company
    return brand or company


def _load(xlsx: Path) -> tuple[list[dict], dict]:
    wb = load_workbook(xlsx, data_only=True)
    landscape: list[dict] = []
    ws = wb["Landscape"] if "Landscape" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows[0]]
    for r in rows[1:]:
        if not r:
            continue
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        if str(d.get("Company") or "").strip():
            landscape.append(d)

    details_by: dict[str, dict] = {}
    if "Company Details" in wb.sheetnames:
        ws = wb["Company Details"]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(h) for h in rows[0]]
        for r in rows[1:]:
            if not r:
                continue
            d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
            key = _norm(d.get("Brand") or d.get("Company") or "")
            if key:
                details_by[key] = d

    for row in landscape:
        key = _norm(row.get("Company") or "")
        det = details_by.get(key)
        if not det:
            for dk, dv in details_by.items():
                if dk in key or key in dk:
                    det = dv
                    break
        if det:
            for col in ("Quadrant", "X", "Y", "Overall", "Role"):
                # Prefer Landscape X/Y/Overall when already present (avoid Details sort scramble)
                if col in ("X", "Y", "Overall") and (
                    row.get("X Score") not in (None, "")
                    or row.get("Y Score") not in (None, "")
                ):
                    continue
                if det.get(col) not in (None, ""):
                    row[col] = det[col]
                    if col == "X":
                        row["X Score"] = det[col]
                    elif col == "Y":
                        row["Y Score"] = det[col]
                    elif col == "Overall":
                        row["Overall Score"] = det[col]
            if det.get("Role"):
                row["Distribution Type"] = det["Role"]
            if det.get("Found in"):
                row["Headquarters"] = det["Found in"]
                row["Found in"] = det["Found in"]
            brand = str(det.get("Brand") or "").strip()
            company_col = str(det.get("Company") or "").strip()
            if brand and not brand.lower().startswith("("):
                row["Brand"] = brand
            if company_col and company_col != brand:
                row["_display_company"] = company_col
    return landscape, details_by


def _match_ownership(name: str) -> tuple[str, dict] | None:
    n = _norm(name)
    if n in OWNERSHIP:
        return n, OWNERSHIP[n]
    for key, info in OWNERSHIP.items():
        if key in n or n in key:
            return key, info
        rename = _norm(info.get("brand_rename") or "")
        if rename and (rename == n or rename in n or n in rename):
            return key, info
    return None


def main() -> None:
    folder = ROOT / SLUG
    xlsx = folder / f"{SLUG}_FINAL.xlsx"
    landscape, _ = _load(xlsx)

    applied: list[dict] = []
    for row in landscape:
        ident = _identity(row)
        # Always use semiconductor commercial role
        row["Distribution Type"] = "Solution Provider"
        row["Role"] = "Solution Provider"
        # Drop any prior Ownership / parent fields — only re-apply verified map
        for k in (
            "Ownership",
            "parent",
            "parent_owner",
            "ownership_relation",
            "ownership_year",
            "parent_or_independent",
        ):
            if k in row:
                row[k] = ""

        hit = _match_ownership(ident)
        if not hit:
            # Independent: Brand == Company (plain). Never invent parent labels.
            cleaned = BRAND_CLEAN.get(_norm(ident)) or ident
            # Strip leftover ownership suffixes from prior exports
            if cleaned.lower().startswith("(") or "acquired by" in cleaned.lower():
                cleaned = ident if not ident.lower().startswith("(") else cleaned
            row["Company"] = cleaned
            row["Brand"] = cleaned
            continue
        key, info = hit
        brand = str(info.get("brand_rename") or ident).strip()
        owner = info["owner"]
        relation = info.get("relation") or "acquired_by"
        year = info.get("year") or ""
        if info.get("company_plain"):
            company_col = info["company_plain"]
            ownership_label = company_col
        else:
            company_col = format_acquired_suffix(owner, relation=relation, year=year)
            verb = {
                "acquired_by": "acquired by",
                "merged_into": "merged into",
                "subsidiary_of": "subsidiary of",
            }.get(relation, "acquired by")
            ownership_label = f"{verb} {owner}" + (f", {year}" if year else "")

        row["Company"] = brand
        row["Brand"] = brand
        row["Ownership"] = ownership_label
        row["parent_owner"] = owner
        row["ownership_relation"] = relation
        if year:
            row["ownership_year"] = year
        applied.append(
            {
                "brand": brand,
                "company": company_col,
                "owner": owner,
                "relation": relation,
                "year": year,
                "source": info.get("source") or "",
                "matched_key": key,
                "from": ident,
            }
        )

    audit_path = folder / "chatgpt_expand_batch_all_audit.json"
    audit: dict = {}
    if audit_path.exists():
        audit = json.loads(audit_path.read_text(encoding="utf-8")).get("xy_scoring") or {}

    # Recompute quadrants from Landscape X/Y (do not trust stale Details labels)
    from vendor_intel.quadrant.rating_map import assign_quadrants_half_median

    xs: list[float] = []
    ys: list[float] = []
    for row in landscape:
        try:
            x = float(row.get("X Score") or row.get("X") or 0)
            y = float(row.get("Y Score") or row.get("Y") or 0)
        except (TypeError, ValueError):
            x, y = 0.0, 0.0
        overall = int(round((x + y) / 2.0))
        row["X Score"] = str(int(round(x)))
        row["Y Score"] = str(int(round(y)))
        row["Overall Score"] = str(overall)
        row["X"] = row["X Score"]
        row["Y"] = row["Y Score"]
        row["Overall"] = row["Overall Score"]
        xs.append(x)
        ys.append(y)
    quads, mid_x, mid_y = assign_quadrants_half_median(xs, ys)
    for row, q in zip(landscape, quads):
        row["Quadrant"] = q
    audit["score_midpoints"] = {"x": mid_x, "y": mid_y}
    audit["quadrant_mid_x"] = mid_x
    audit["quadrant_mid_y"] = mid_y

    # Build details with consumer_brand mode so Company shows (acquired by …)
    detail_rows = to_company_detail_rows(landscape, QUERY, audit)
    by_brand = {_norm(r.get("Brand") or r.get("Company") or ""): r for r in detail_rows}
    by_land = {_norm(r.get("Brand") or r.get("Company") or ""): r for r in landscape}
    # Sync scores/quadrant onto details BY NAME (never zip — details are Overall-sorted)
    for det in detail_rows:
        key = _norm(det.get("Brand") or det.get("Company") or "")
        src = by_land.get(key)
        if not src:
            continue
        det["X"] = int(str(src.get("X Score") or 0) or 0)
        det["Y"] = int(str(src.get("Y Score") or 0) or 0)
        det["Overall"] = int(str(src.get("Overall Score") or 0) or 0)
        det["Quadrant"] = src.get("Quadrant")
        det["Role"] = "Solution Provider"
    for item in applied:
        brand = item["brand"]
        key = _norm(brand)
        det = by_brand.get(key)
        if not det:
            for dk, dv in by_brand.items():
                if dk in key or key in dk:
                    det = dv
                    break
        if not det:
            continue
        # Force display fields with consumer_brand mode
        meta = {
            "company": brand,
            "brand": brand,
            "parent_owner": item["owner"],
            "ownership_relation": item["relation"],
            "ownership_year": item.get("year") or "",
            "parent": item.get("company") if item["company"].startswith("(") else f"acquired by {item['owner']}",
        }
        b, c, _ = brand_display_fields(
            meta,
            market=QUERY,
            industry_category="food",  # force consumer_brand path
            mode="consumer_brand",
        )
        if item.get("company") and not item["company"].startswith("("):
            # Plain legal / JV label
            det["Brand"] = brand
            det["Company"] = item["company"]
        else:
            det["Brand"] = b or brand
            det["Company"] = c or item["company"]

    # Ensure every independent row has Brand == Company (plain name)
    owned_brands = {_norm(a["brand"]) for a in applied}
    for det in detail_rows:
        brand = str(det.get("Brand") or "").strip()
        company = str(det.get("Company") or "").strip()
        key = _norm(brand) or _norm(company)
        cleaned = BRAND_CLEAN.get(_norm(brand)) or BRAND_CLEAN.get(_norm(company))
        if key not in owned_brands:
            # No verified ownership — never leave parent / acquired-by text
            plain = cleaned or brand or company
            if plain.lower().startswith("(") or "acquired by" in plain.lower() or "subsidiary of" in plain.lower():
                plain = brand if brand and not brand.lower().startswith("(") else company
            det["Brand"] = plain
            det["Company"] = plain
        elif cleaned and brand == company:
            det["Brand"] = cleaned
            det["Company"] = cleaned
        if not company or company.lower() in {"independent", "n/a"}:
            det["Company"] = brand
        if not str(det.get("Company") or "").strip():
            det["Company"] = brand
        # Semiconductor market role
        det["Role"] = "Solution Provider"

    # Prefer parameter definitions + axes from existing quadrant JSON if present
    qjson = folder / f"{SLUG}_quadrant.json"
    if qjson.exists():
        try:
            payload = json.loads(qjson.read_text(encoding="utf-8"))
            crit = payload.get("criteria") or {}
            if crit.get("parameter_definitions"):
                audit["parameter_definitions"] = crit["parameter_definitions"]
            labels = crit.get("axis_labels") or {}
            if labels.get("x"):
                audit["axis_x"] = labels["x"]
            if labels.get("y"):
                audit["axis_y"] = labels["y"]
            if crit.get("x_axis"):
                audit["x_features"] = crit["x_axis"]
            if crit.get("y_axis"):
                audit["y_features"] = crit["y_axis"]
        except Exception:
            pass

    write_final_xlsx(
        xlsx,
        landscape,
        "Companies",
        {
            "query": QUERY,
            "xy_scoring": audit,
            "brand_company_ownership": applied,
        },
        detail_rows=detail_rows,
    )
    extras = export_expand_quadrant_outputs(
        folder, detail_rows, QUERY, country="global", audit=audit, chart_n=20
    )

    audit_out = ROOT / "_audit" / "semiconductor_brand_company_ownership.json"
    audit_out.parent.mkdir(exist_ok=True)
    audit_out.write_text(
        json.dumps(
            {
                "total_companies": len(detail_rows),
                "ownership_applied": len(applied),
                "applied": applied,
                "html": str(extras.get("html") or ""),
            },
            indent=2,
        ),
        encoding="utf-8",
    )

    md = ROOT / "_audit" / "semiconductor_brand_company_ownership.md"
    lines = [
        "# Semiconductor Brand / Company ownership verification",
        "",
        f"Total companies: **{len(detail_rows)}**",
        f"Ownership labels applied: **{len(applied)}** (web-verified only)",
        "",
        "| Brand | Company | Source |",
        "|---|---|---|",
    ]
    for a in sorted(applied, key=lambda x: x["brand"].lower()):
        lines.append(f"| {a['brand']} | {a['company']} | {a['source']} |")
    lines.extend(
        [
            "",
            "All other rows keep Brand = Company (independent / no confirmed acquirer).",
            "",
        ]
    )
    md.write_text("\n".join(lines), encoding="utf-8")

    print(f"Applied ownership to {len(applied)} / {len(detail_rows)}")
    for a in sorted(applied, key=lambda x: x["brand"].lower()):
        print(f"  {a['brand']} -> {a['company']}")
    print(f"html -> {extras.get('html')}")
    print(f"audit -> {md}")


if __name__ == "__main__":
    main()
