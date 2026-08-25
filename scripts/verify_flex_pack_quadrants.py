#!/usr/bin/env python3
"""Verify Quadrant column for Global Flexible Packaging vs Coherent definitions.

X = Solution Capability (product/tech platforms, converting/film depth)
Y = Business Strategy (scale, supply chain, GTM, partnerships, market reach)

  Leaders:      High X + High Y — mature scalable brands, strong supply chains
  Challengers:  Lower X + High Y — broad reach/partnerships; deepening platforms
  Trailblazers: High X + Lower Y — advanced capability; limited global scale / GTM
  Emerging:     Lower X + Lower Y — niche/newer; focused offerings, limited presence

Invalid legacy labels (Visionaries, Niche Players) are remapped via X/Y.
"""
from __future__ import annotations

import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx
from vendor_intel.quadrant.rating_map import assign_quadrants_half_median

SLUG = "global_flexible_packaging_market_global"
QUERY = "Global Flexible Packaging Market"
OUT = ROOT / "output" / "chatgpt_expand" / SLUG
XLSX = OUT / f"{SLUG}_FINAL.xlsx"
AUDIT = (
    ROOT
    / "output"
    / "chatgpt_expand"
    / "_audit"
    / f"{SLUG}_quadrant_verify.json"
)

VALID = {"Leaders", "Challengers", "Trailblazers", "Emerging Players"}

# Stem → (X, Y, note). Scores are definition-aligned priors (0–100).
SEED_XY: dict[str, tuple[int, int, str]] = {
    # —— Leaders: high capability + high strategy / scale ——
    "amcor": (88, 90, "Global #1 flexible packaging; Berry integration; proven scale"),
    "mondi": (84, 86, "Major European flexible packaging; strong supply chain + GTM"),
    "huhtamaki": (82, 85, "Global food flexible packaging; scalable platforms"),
    "constantia flexibles": (80, 84, "Large European flexible converter; broad CPG reach"),
    "sealed air": (83, 85, "Cryovac/food flexible franchise; global channels"),
    "berry global": (78, 88, "Broad flexible portfolio; commercial scale (now Amcor)"),
    "coveris": (78, 82, "Large European flexible packaging group"),
    "proampac": (80, 83, "Major NA flexible converter; PE-backed scale"),
    "printpack": (79, 81, "Large US flexible converter; proven CPG adoption"),
    "sonoco": (77, 82, "Diversified packaging scale incl. flexible"),
    "uflex": (78, 80, "Large Indian flexible packaging; global film+convert"),
    "sudpack verpackungen": (81, 80, "Major German flexible packaging; strong EU supply chain"),
    "winpak": (80, 78, "High-barrier flexible; solid NA/EU commercial franchise"),
    "wipak": (79, 77, "High-barrier medical/food flexible; Wihuri scale"),
    "transcontinental": (76, 80, "Large NA flexible packaging operations"),
    "novolex": (74, 82, "Broad packaging portfolio + Pactiv flexible scale"),
    "kluckner pentaplast": (77, 79, "Large rigid+flexible films; global reach"),
    "klockner pentaplast": (77, 79, "Large rigid+flexible films; global reach"),
    "polyplex corporation": (78, 76, "Large BOPET film producer; global markets"),
    "taghleef": (77, 78, "Major BOPP films; Al Ghurair scale"),
    "toray industries": (84, 80, "Advanced polyester/specialty films + global GTM"),
    "jindal poly films": (76, 77, "Large BOPP/BOPET films; group scale"),
    "jindal films": (75, 76, "Specialty films under B.C. Jindal scale"),
    "epl limited": (74, 78, "Global laminated tubes; Blackstone-backed scale"),
    "essel": (74, 78, "Global laminated tubes franchise"),
    "schur flexibles": (75, 76, "European flexible packaging group"),
    "bischof + klein": (74, 75, "Major German industrial/flexible packer"),
    "bischof klein": (74, 75, "Major German industrial/flexible packer"),
    "scg packaging": (73, 77, "ASEAN packaging scale under SCG"),
    "scientex": (72, 74, "Large Malaysian flexible packaging"),
    "reynolds consumer": (70, 80, "Broad consumer wrap brand reach + supply"),
    "tcpl packaging": (71, 73, "Large Indian flexible converter"),
    "clondalkin": (72, 74, "European flexible packaging group"),
    "pactiv evergreen": (70, 76, "Scale food packaging under Novolex"),
    # —— Challengers: high strategy/reach, lower own solution depth ——
    "innovia": (58, 76, "CCL film commercial reach; platform deepening"),
    "cosmo films": (52, 74, "BOPP commercial brand under Cosmo First — reach > own tech depth"),
    "cosmo first": (54, 75, "Film marketer scale; commercial reach ahead of specialty depth"),
    "al ghurair packaging": (58, 73, "Group packaging commercial reach"),
    "arabian flexible packaging": (55, 70, "Regional group commercial channel"),
    "smurfit westrock bag": (58, 74, "Parent scale; flexible is one platform"),
    "sigma stretch film": (56, 72, "Stretch film commercial footprint"),
    "sigma plastics": (55, 73, "Large PE film commercial group"),
    "charter next generation": (57, 74, "PE film commercial scale"),
    "srf limited packaging": (60, 72, "Films under SRF commercial group"),
    "nan ya plastics films": (58, 71, "Formosa group commercial film reach"),
    "far eastern new century films": (57, 70, "Group film commercial footprint"),
    "shinkong synthetic": (56, 69, "Taiwan group film commercial reach"),
    "formosa idemitsu": (55, 68, "JV film commercial channel"),
    "hyosung chemical films": (58, 70, "Chaebol film commercial reach"),
    "kolon industries films": (57, 69, "Group film commercial footprint"),
    "indorama ventures packaging": (56, 71, "Indorama commercial scale; films unit"),
    "manjushree technopack flexible": (54, 68, "Parent commercial reach; flexible niche"),
    "mpact flexible": (55, 67, "Parent listed scale; flexible segment"),
    "saica flex": (54, 66, "SAICA group commercial channel"),
    "gascogne flexible": (53, 65, "Group commercial; focused flexible arm"),
    "al watania plastics": (52, 66, "Regional group commercial packaging"),
    "ampac holdings": (58, 72, "Legacy brand under ProAmpac commercial"),
    "aep industries": (55, 70, "Legacy brand; Amcor/Berry commercial"),
    "rpc group": (54, 71, "Legacy packaging brand under Amcor/Berry"),
    "rpc bpi": (53, 70, "Legacy brand; commercial via Amcor/Berry"),
    "bemis": (56, 74, "Legacy Amcor brand; commercial residual"),
    "nordfolien": (52, 68, "Legacy site brand; Amcor commercial"),
    "aluflexpack": (58, 72, "Under Constantia commercial umbrella"),
    "walki": (57, 71, "Under Oji commercial reach"),
    "selig": (54, 69, "Closure/flexible niche under CC Industries"),
    # —— Trailblazers: high capability, lower global GTM maturity ——
    # Keep Y clearly below cohort strategy median so half-median lands bottom-right.
    "kuraray": (82, 48, "EVAL EVOH barrier tech leader; focused specialty GTM"),
    "honeywell": (80, 46, "Aclar PCTFE barrier tech; specialty medical packaging"),
    "3m scotchpak": (78, 45, "Scotchpak barrier films; specialty channel"),
    "dupont tyvek": (81, 47, "Tyvek medical packaging material; specialty GTM"),
    "soarus": (79, 44, "SoarnoL EVOH specialist; focused markets"),
    "nippon gohsei": (78, 43, "EVOH specialty; Mitsubishi Chemical channel niche"),
    "mitsubishi polyester film": (80, 48, "Advanced PET films; specialty industrial GTM"),
    "mitsubishi gas chemical": (77, 45, "Specialty barrier materials; focused markets"),
    "mitsui chemicals": (76, 46, "Specialty film resins; selective GTM"),
    "mylar specialty films": (78, 42, "Advanced PET film tech; rebuilding commercial scale"),
    "treofan": (74, 44, "Technical BOPP (Polyopt); scale reset post Al Ghurair"),
    "toppan speciality films": (75, 47, "Specialty barrier films; Toppan but focused unit"),
    "max speciality films": (73, 45, "Specialty films; focused India/Asia GTM"),
    "ahlstrom": (74, 48, "Fiber/specialty packaging materials; selective flex GTM"),
    "vacmet": (72, 44, "Metallized/specialty films; regional tech strength"),
    "ester filmtech": (71, 42, "Specialty PET films; early commercial scale"),
    "ester industries": (70, 48, "Polyester films tech; mid-scale GTM"),
    "chiripal poly films": (69, 43, "Specialty films; regional capability focus"),
    "polyplex thailand": (68, 45, "Regional film tech arm of Polyplex"),
    "toray advanced film": (76, 46, "Advanced Toray films Japan; focused unit GTM"),
    "asahi kasei packaging": (73, 44, "Specialty materials; packaging niche GTM"),
    "wipf": (72, 47, "High-barrier Swiss tech; limited global scale"),
    "futamura": (74, 46, "NatureFlex compostable films; specialty GTM"),
    "kureha": (75, 45, "Specialty PVDC/barrier; focused markets"),
    "unitika": (73, 44, "Specialty nylon/films; Japan-centric GTM"),
    "toyobo": (74, 47, "Specialty films; selective global reach"),
    "okura industrial": (71, 43, "Specialty films Japan; limited global GTM"),
    "fujimori": (70, 44, "Specialty packaging Japan; focused markets"),
    "wentus": (69, 42, "Specialty EU films; mid-scale"),
    "bilcare": (68, 41, "Pharma barrier films; niche GTM"),
    "alucoat": (70, 43, "Specialty alu-coat tech; focused markets"),
    "al khaleej": (65, 40, "PP film tech site under Taghleef; limited brand GTM"),
    # mid-size converters with capable ops but not Leader-scale GTM
    "interflex": (64, 50, "Capable US converter; not global Leader scale"),
    "huangshan novel": (58, 48, "China regional flexible; capability ahead of global GTM"),
    "wipf ag": (72, 47, "High-barrier Swiss tech; limited global scale"),
    # —— Emerging: lower capability + lower strategy / niche regional ——
    "sudpack iberica": (52, 48, "Regional Südpack arm; limited independent presence"),
    "epac": (50, 52, "Digital flexible packager; growing niche — not yet Leader/Challenger scale"),
    "bryce": (55, 51, "US mid-size converter; focused accounts"),
    "flexopack": (53, 49, "Regional EU flexible; limited global presence"),
    "parkside flexibles": (52, 48, "UK niche flexible converter"),
    "polifilm": (54, 50, "Mid-size EU films; focused markets"),
    "cedo": (48, 46, "Niche consumer wrap; limited presence"),
    "oerlemans": (50, 47, "Regional NL packaging"),
    "sharpak": (49, 45, "Niche flexible converter"),
    "pt trias sentosa": (51, 48, "Indonesia film; regional"),
    "pt argha karya": (52, 49, "Indonesia BOPP; regional"),
    "thai film industries": (53, 50, "Thailand films; regional"),
    "thong guan": (52, 49, "Malaysia stretch film; regional"),
    "bp plastics": (51, 48, "Malaysia films; regional"),
    "takween": (50, 47, "MENA packaging; regional"),
    "qatar plastic": (49, 46, "Regional Gulf plastics"),
    "c-p flexible": (51, 48, "US mid converter; focused"),
    "flair flexible": (52, 49, "US mid converter"),
    "ppc flexible": (51, 48, "US mid converter"),
    "paardekooper": (50, 47, "Regional Benelux packaging"),
    "advanced film company thailand": (52, 49, "Regional Thai films"),
    "jiangsu shuangxing": (53, 50, "China films; regional"),
    "shanghai zijiang": (52, 49, "China packaging; regional"),
    "zhejiang yamei": (51, 48, "China materials; regional"),
    "gulf packaging industries": (50, 47, "Regional Gulf packaging"),
    "thai packaging": (51, 48, "Regional Thai converter"),
    "hosokawa yoko": (54, 50, "Japan pouch niche"),
    "rengo": (58, 55, "Japan packaging; limited global flex presence"),
    "dai nippon": (60, 56, "Japan packaging print; flex niche globally"),
    "toyo seikan": (59, 55, "Japan packaging; flexible segment niche globally"),
    "toppan printing": (61, 57, "Japan packaging giant; flex is one niche globally"),
}


def _norm(s: str) -> str:
    s = str(s or "").strip().lower()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = s.replace("ö", "o").replace("ü", "u").replace("ä", "a")
    s = re.sub(r"[+/|&,_.:\-()]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _f(v, default=50.0) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def _match_seed(name: str) -> tuple[int, int, str] | None:
    n = _norm(name)
    best = None
    best_len = -1
    for stem, trip in SEED_XY.items():
        if stem in n and len(stem) > best_len:
            best = trip
            best_len = len(stem)
    return best


def _is_compressed(x: float, y: float) -> bool:
    """Scores stuck near cohort floor / mid band with little signal."""
    if (x == 52 and y == 52) or (x == 57 and y == 60):
        return True
    if 48 <= x <= 58 and 48 <= y <= 61 and abs(x - y) <= 4:
        return True
    return False


def _heuristic_xy(row: dict) -> tuple[int, int]:
    """Definition-aligned spread when no seed and scores are compressed."""
    role = str(row.get("Role") or row.get("Distribution Type") or "Brand")
    blob = " ".join(
        [
            str(row.get("Specialty Focus") or ""),
            str(row.get("Core Categories") or ""),
            str(row.get("Summary") or ""),
            str(row.get("Company") or ""),
            str(row.get("Operational Presence") or ""),
            str(row.get("Continent / Geography") or ""),
        ]
    ).lower()
    emp = str(row.get("Employees") or "")
    # Brand converters default slightly higher capability; marketers slightly higher strategy
    x, y = (58, 54) if role == "Brand" else (52, 60)

    if any(k in blob for k in ("barrier", "evoh", "metalliz", "bopet", "bopp", "cpp", "mono-material", "recycl")):
        x += 6
    if any(k in blob for k in ("global", "europe", "north america", "multi-continent", "worldwide")):
        y += 8
    if any(k in blob for k in ("regional", "niche", "startup", "local", "emerging")):
        y -= 6
        x -= 2
    if any(k in blob for k in ("acquired by", "subsidiary of", "legacy")):
        # commercial parent reach but weaker independent brand capability signal
        y += 4
        x -= 3
    # rough employee bands
    digits = re.sub(r"[^\d]", "", emp.split("-")[-1] if emp else "")
    try:
        n = int(digits) if digits else 0
    except ValueError:
        n = 0
    if n >= 10000:
        y += 8
        x += 4
    elif n >= 1000:
        y += 4
        x += 2
    elif 0 < n < 200:
        y -= 4
        x -= 2

    return max(38, min(90, x)), max(38, min(90, y))


def _load_landscape() -> tuple[list[dict], dict[str, dict]]:
    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Landscape"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows[0]]
    landscape = []
    for r in rows[1:]:
        d = {hdr[i]: ("" if r[i] is None else str(r[i])) for i in range(len(hdr)) if i < len(r)}
        if d.get("Company"):
            landscape.append(d)

    details_by: dict[str, dict] = {}
    if "Company Details" in wb.sheetnames:
        ws = wb["Company Details"]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(h) for h in rows[0]]
        for r in rows[1:]:
            d = {hdr[i]: ("" if r[i] is None else str(r[i])) for i in range(len(hdr)) if i < len(r)}
            key = _norm(d.get("Brand") or "")
            if key:
                details_by[key] = d
    return landscape, details_by


def main() -> None:
    landscape, details_by = _load_landscape()
    before = {
        str(r.get("Company") or ""): {
            "quadrant": str(r.get("Quadrant") or ""),
            "x": _f(r.get("X Score") or r.get("X")),
            "y": _f(r.get("Y Score") or r.get("Y")),
        }
        for r in landscape
    }

    seeded: list[dict] = []
    heur: list[dict] = []
    invalid_fixed: list[dict] = []

    for row in landscape:
        name = str(row.get("Company") or "").strip()
        det = details_by.get(_norm(name)) or {}
        # Prefer detail Brand for identity; keep ownership-driven Company display later
        brand = str(det.get("Brand") or name).strip()
        row["Brand"] = brand
        row["Company"] = brand  # landscape key = brand identity
        if det.get("Company") and str(det["Company"]).startswith("("):
            row["_display_company"] = det["Company"]
        elif det.get("Company") and det["Company"] != brand:
            row["_display_company"] = det["Company"]

        role = str(row.get("Role") or det.get("Role") or "Brand")
        if role not in ("Brand", "Marketer"):
            role = "Brand"
        row["Role"] = role
        row["Distribution Type"] = role

        x = _f(row.get("X Score") or row.get("X") or det.get("X"), 50)
        y = _f(row.get("Y Score") or row.get("Y") or det.get("Y"), 50)
        old_q = str(row.get("Quadrant") or det.get("Quadrant") or "")

        seed = _match_seed(brand) or _match_seed(name)
        if seed:
            sx, sy, note = seed
            row["X Score"] = sx
            row["Y Score"] = sy
            row["Overall Score"] = int(round((sx + sy) / 2))
            seeded.append({"brand": brand, "x": sx, "y": sy, "note": note, "old_q": old_q})
        elif _is_compressed(x, y) or old_q not in VALID:
            hx, hy = _heuristic_xy(row)
            # If invalid label but scores look differentiated, keep scores and only fix label later
            if old_q not in VALID and not _is_compressed(x, y) and (x >= 62 or y >= 62 or x <= 45 or y <= 45):
                row["X Score"] = int(round(x))
                row["Y Score"] = int(round(y))
                row["Overall Score"] = int(round((x + y) / 2))
                invalid_fixed.append({"brand": brand, "old_q": old_q, "x": x, "y": y, "mode": "keep_xy"})
            else:
                row["X Score"] = hx
                row["Y Score"] = hy
                row["Overall Score"] = int(round((hx + hy) / 2))
                heur.append({"brand": brand, "x": hx, "y": hy, "old_q": old_q})
                if old_q not in VALID:
                    invalid_fixed.append({"brand": brand, "old_q": old_q, "x": hx, "y": hy, "mode": "heuristic"})
        else:
            row["X Score"] = int(round(x))
            row["Y Score"] = int(round(y))
            row["Overall Score"] = int(round((x + y) / 2))

    xs = [_f(r.get("X Score")) for r in landscape]
    ys = [_f(r.get("Y Score")) for r in landscape]
    quads, mid_x, mid_y = assign_quadrants_half_median(xs, ys)

    changes: list[dict] = []
    for row, q in zip(landscape, quads):
        old = before.get(str(row.get("Brand") or row.get("Company") or ""), {}).get("quadrant") or ""
        # also try original company key
        if not old:
            old = str(row.get("Quadrant") or "")
        row["Quadrant"] = q
        row["X"] = row["X Score"]
        row["Y"] = row["Y Score"]
        row["Overall"] = row["Overall Score"]
        if old and old != q:
            changes.append(
                {
                    "brand": row.get("Brand") or row.get("Company"),
                    "from": old,
                    "to": q,
                    "x": _f(row.get("X Score")),
                    "y": _f(row.get("Y Score")),
                    "role": row.get("Role"),
                }
            )

    # Rebuild company details (preserves ownership Company via brand_display_fields + override)
    audit_path = OUT / "chatgpt_expand_batch_all_audit.json"
    audit: dict = {}
    if audit_path.exists():
        try:
            audit = json.loads(audit_path.read_text(encoding="utf-8")).get("xy_scoring") or {}
        except Exception:
            audit = {}
    audit["axis_x"] = "Solution Capability"
    audit["axis_y"] = "Business Strategy"
    audit["quadrant_method"] = "half_median_rank_split"
    audit["quadrant_definitions"] = {
        "Leaders": "High Solution Capability + High Business Strategy",
        "Challengers": "Lower Solution Capability + High Business Strategy",
        "Trailblazers": "High Solution Capability + Lower Business Strategy",
        "Emerging Players": "Lower Solution Capability + Lower Business Strategy",
    }
    audit["mid_x"] = mid_x
    audit["mid_y"] = mid_y

    details = to_company_detail_rows(landscape, QUERY, audit)
    by_brand = {_norm(r.get("Brand") or r.get("Company") or ""): r for r in landscape}
    for d in details:
        src = by_brand.get(_norm(d.get("Brand") or ""))
        if not src:
            continue
        d["Role"] = src.get("Role") or d.get("Role") or "Brand"
        d["Quadrant"] = src.get("Quadrant")
        d["X"] = src.get("X Score")
        d["Y"] = src.get("Y Score")
        d["Overall"] = src.get("Overall Score")
        if src.get("_display_company"):
            d["Company"] = src["_display_company"]

    write_final_xlsx(
        XLSX,
        landscape,
        "Companies",
        {
            "query": QUERY,
            "quadrant_verify": "2026-08-14",
            "seeded": len(seeded),
            "changes": len(changes),
        },
        detail_rows=details,
    )
    extras = export_expand_quadrant_outputs(
        OUT, details, QUERY, country="global", audit=audit, chart_n=20
    )

    report = {
        "n": len(landscape),
        "mid_x": mid_x,
        "mid_y": mid_y,
        "counts_after": dict(Counter(quads)),
        "counts_before": dict(Counter(v["quadrant"] for v in before.values())),
        "seeded": len(seeded),
        "heuristic": len(heur),
        "invalid_label_fixed": invalid_fixed,
        "changes": changes,
        "seeded_sample": seeded[:40],
        "by_quadrant_top": {},
        "html": str(extras.get("html") or ""),
    }
    by_q: dict[str, list] = {q: [] for q in VALID}
    for row in landscape:
        by_q.setdefault(str(row["Quadrant"]), []).append(row)
    for q, items in by_q.items():
        items.sort(key=lambda r: -_f(r.get("Overall Score")))
        report["by_quadrant_top"][q] = [
            {
                "brand": r.get("Brand") or r.get("Company"),
                "x": _f(r.get("X Score")),
                "y": _f(r.get("Y Score")),
                "overall": _f(r.get("Overall Score")),
                "role": r.get("Role"),
            }
            for r in items[:15]
        ]

    AUDIT.parent.mkdir(parents=True, exist_ok=True)
    AUDIT.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")

    print(
        f"DONE n={len(landscape)} seeded={len(seeded)} heur={len(heur)} "
        f"changes={len(changes)} mid=({mid_x:.1f},{mid_y:.1f})"
    )
    print("counts_before", report["counts_before"])
    print("counts_after ", report["counts_after"])
    if invalid_fixed:
        print("invalid labels remapped:")
        for x in invalid_fixed:
            print(f"  {x['brand']}: {x['old_q']} ({x['mode']})")
    print("Top per quadrant:")
    for q in ("Leaders", "Challengers", "Trailblazers", "Emerging Players"):
        print(f"  [{q}]")
        for t in report["by_quadrant_top"].get(q, [])[:8]:
            print(
                f"    O={t['overall']:.0f} X={t['x']:.0f} Y={t['y']:.0f} "
                f"[{t['role']}] {t['brand']}"
            )
    print("audit ->", AUDIT)


if __name__ == "__main__":
    main()
