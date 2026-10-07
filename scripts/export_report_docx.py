#!/usr/bin/env python3
"""Word (.docx) version of a market's Coherent Quadrant report, matching the PDF.

The content is read from the same rendered page the PDF is printed from
(export_report_pdf's EXPAND/COVER/TOC scripts run first), so every name,
score, bullet and table row is the text the PDF shows -- nothing is
re-derived from the JSON. The quadrant chart is the same chart, captured as an
image. Layout follows the PDF: cover with key figures, contents, each section
on a new page, one company per page in Parameter Detail, back cover last.
A branded front cover (_pipeline/front_cover.pdf, else config/front_cover.pdf)
goes first, its "xxxxxx" title slot filled with the market name.

When Microsoft Word is installed, it refreshes the contents page numbers
(a TOC field) before the file is saved.

    .venv\\Scripts\\python.exe scripts\\export_report_docx.py --market "..." --country global
    .venv\\Scripts\\python.exe scripts\\export_report_docx.py --html path\\to\\x_report.html
"""
from __future__ import annotations

import argparse
import asyncio
import io
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "scripts"))

from export_report_pdf import (  # noqa: E402
    COVER_JS, EXPAND_JS, PRINT_CSS, TOC_JS, back_cover_for, filled_front_cover, front_cover_for)

NAVY, ACCENT, TEAL = "002857", "1A759F", "52B69A"
SLATE9, SLATE7, SLATE5 = "0F172A", "334155", "64748B"
QUAD_COLORS = {"Leaders": "7CFC00", "Challengers": "7EB6FF", "Trailblazers": "C4B5FD",
               "Evolving Players": "A7F3D0"}

# Reads the rendered report into plain structures, section by section.
EXTRACT_JS = r"""
() => {
  // Only what the printed page shows: hidden elements (screen-only hints) are skipped.
  const shown = el => el && (el.nodeType !== 1 || getComputedStyle(el).display !== 'none');
  const txt = el => (shown(el) ? el.textContent.replace(/\s+/g, ' ').trim() : '');
  const out = {cover: {}, tiles: [], sections: []};
  const cv = document.querySelector('.hero.cover');
  out.cover = {brand: txt(cv.querySelector('.cv-brand')), market: txt(cv.querySelector('.cv-market')),
               geo: txt(cv.querySelector('.cv-meta')), scope: txt(cv.querySelector('.cv-scope'))};
  out.tiles = [...document.querySelectorAll('.cv-tile')].map(t => [txt(t.querySelector('b')), txt(t.querySelector('span'))]);
  const cell = td => {
    const dots = td.querySelector('.dots');
    if (dots) return {dots: dots.querySelectorAll('i.on').length, of: dots.querySelectorAll('i').length};
    const pill = td.querySelector('.pill');
    if (pill) return {pill: txt(pill)};
    const disc = td.querySelector('.disc');
    if (disc) return {disc: txt(disc)};
    return {t: txt(td)};
  };
  for (const sec of document.querySelectorAll('section.card')) {
    if (sec.classList.contains('cv-toc')) continue;
    const h = sec.querySelector('h2');
    const s = {title: txt(h && h.childNodes[0]), sub: txt(h && h.querySelector('.sub'))};
    if (sec.querySelector('.plot')) {
      s.kind = 'chart';
      s.mixTitle = txt(sec.querySelector('.mix .t'));
      const mix = sec.querySelector('.mix');
      s.mix = mix ? txt(mix).replace(s.mixTitle, '').trim() : '';
    } else if (sec.querySelector('.qdef')) {
      s.kind = 'qdefs';
      s.items = [...sec.querySelectorAll('.qdef')].map(q => ({tag: txt(q.querySelector('.tag')),
        title: txt(q.querySelector('.qt')), sub: txt(q.querySelector('.qp')), text: txt(q.querySelector('.qd'))}));
    } else if (sec.querySelector('.cblock')) {
      s.kind = 'detail';
      s.companies = [...sec.querySelectorAll('.cblock')].map(b => ({
        name: txt(b.querySelector('.chead .nm')), brand: txt(b.querySelector('.chead .br')),
        quadrant: txt(b.querySelector('.chead .pill')),
        scores: [...b.querySelectorAll('.chead .sc > span')].map(x => [txt(x.childNodes[0]), txt(x.querySelector('strong'))]),
        axes: [...b.querySelectorAll('.axwrap')].map(a => ({title: txt(a.querySelector('.axttl')),
          params: [...a.querySelectorAll('.pdetail')].map(p => ({h: txt(p.querySelector('.h')),
            basis: txt(p.querySelector('.basis')), sh: txt(p.querySelector('.sh')),
            items: [...p.querySelectorAll('li')].map(li => txt(li)),
            none: txt(p.querySelector('.none'))}))}))}));
    } else if (sec.querySelector('table')) {
      s.kind = 'table';
      const t = sec.querySelector('table');
      s.head = [...t.querySelectorAll('thead th')].map(th => txt(th));
      s.rows = [...t.querySelectorAll('tbody tr')].map(tr => [...tr.querySelectorAll('td')].map(cell));
    } else if (sec.querySelector('.steps') || sec.querySelector('.mintro')) {
      s.kind = 'method';
      s.intro = txt(sec.querySelector('.mintro'));
      s.steps = [...sec.querySelectorAll('.step')].map(st => ({n: txt(st.querySelector('.n')),
        t: txt(st.querySelector('.t')), d: txt(st.querySelector('.d'))}));
      s.note = txt(sec.querySelector('.mnote'));
    } else {
      s.kind = 'panels';
      s.lead = txt(sec.querySelector('.lead'));
      s.panels = [...sec.querySelectorAll('.panel')].map(p => {
        const parts = [];
        for (const el of p.children) {
          if (el.classList.contains('lbl')) parts.push({lbl: txt(el)});
          else if (el.classList.contains('val')) parts.push({val: txt(el)});
          else if (el.classList.contains('muted')) parts.push({muted: txt(el)});
          else if (el.tagName === 'OL') parts.push({ol: [...el.children].map(li => ({
            name: txt(li.querySelector('.pname')) || txt(li), def: txt(li.querySelector('.pdef'))}))});
        }
        return parts;
      });
    }
    out.sections.push(s);
  }
  return out;
}
"""


async def read_report(html: Path) -> tuple[dict, bytes]:
    from patchright.async_api import async_playwright

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        page = await browser.new_page(viewport={"width": 1280, "height": 900}, device_scale_factor=2)
        await page.goto(html.resolve().as_uri())
        await page.wait_for_timeout(1500)
        await page.evaluate(EXPAND_JS)
        await page.evaluate(COVER_JS)
        await page.add_style_tag(content=PRINT_CSS)
        await page.emulate_media(media="screen")
        await page.evaluate(TOC_JS, None)
        await page.wait_for_timeout(400)
        data = await page.evaluate(EXTRACT_JS)
        box = await page.evaluate("""() => {
          const a = document.querySelector('.chartrow').getBoundingClientRect();
          const b = document.querySelector('.xaxis').getBoundingClientRect();
          return {x: a.left + scrollX, y: a.top + scrollY, width: a.width, height: b.bottom - a.top};
        }""")
        # A busy machine (10 browsers) can take longer than the 30 s default
        # to capture a large chart; allow 2 minutes and one retry.
        try:
            chart = await page.screenshot(clip=box, full_page=True, timeout=120_000)
        except Exception:  # noqa: BLE001
            await page.wait_for_timeout(3000)
            chart = await page.screenshot(clip=box, full_page=True, timeout=120_000)
        await browser.close()
    return data, chart


# -- docx helpers ----------------------------------------------------------------

def _shade(el, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    pr = el.get_or_add_tcPr() if el.tag.endswith("}tc") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def _border(par, side: str, color: str, size: int = 12, space: int = 4) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    ppr = par._p.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        ppr.append(bdr)
    e = OxmlElement(f"w:{side}")
    for k, v in (("val", "single"), ("sz", str(size)), ("space", str(space)), ("color", color)):
        e.set(qn(f"w:{k}"), v)
    bdr.append(e)


def _cell_margins(cell, twips: int) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(twips))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tcpr.append(mar)


def _field(run, code: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for kind, text in (("begin", None), (None, code), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            if kind == "begin":
                fc.set(qn("w:dirty"), "true")
            run._r.append(fc)
        elif text == code:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {code} "
            run._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)


def _table_borders(t, data: bool) -> None:
    """Layout tables (cover, tiles, panels) get no borders; data tables get
    light horizontal rules only, like the report's tables."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tblpr = t._tbl.tblPr
    old = tblpr.find(qn("w:tblBorders"))
    if old is not None:
        tblpr.remove(old)
    bd = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        rule = data and side in ("bottom", "insideH")
        e.set(qn("w:val"), "single" if rule else "nil")
        if rule:
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "E2E8F0")
        bd.append(e)
    tblpr.append(bd)


def _ptab(run, relative_to: str = "margin") -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    e = OxmlElement("w:ptab")
    e.set(qn("w:relativeTo"), relative_to)
    e.set(qn("w:alignment"), "right")
    e.set(qn("w:leader"), "none")
    run._r.append(e)


def _repeat_header(row) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    trpr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trpr.append(e)
    c = OxmlElement("w:cantSplit")
    trpr.append(c)


def _no_split(row) -> None:
    from docx.oxml import OxmlElement

    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


class Builder:
    def __init__(self, market: str, front_png: bytes | None = None):
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.shared import Mm, Pt, RGBColor

        self.Pt, self.RGB, self.Mm, self.WD_SECTION = Pt, RGBColor, Mm, WD_SECTION
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        if front_png:
            # Full-bleed branded cover, then the report starts on a new section.
            sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Mm(0)
            sec.header_distance = sec.footer_distance = Mm(0)
            fp = self.doc.paragraphs[0] if self.doc.paragraphs else self.doc.add_paragraph()
            fp.paragraph_format.space_after = Pt(0)
            fp.paragraph_format.line_spacing = 1.0
            fp.add_run().add_picture(io.BytesIO(front_png), width=Mm(210), height=Mm(296))
            sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
            sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.top_margin = sec.bottom_margin = Mm(16)
        sec.left_margin = sec.right_margin = Mm(16)
        self.width = sec.page_width - sec.left_margin - sec.right_margin
        st = self.doc.styles["Normal"]
        st.font.name = "Arial"
        st.font.size = Pt(9.5)
        st.font.color.rgb = RGBColor.from_string(SLATE7)
        st.element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "Arial")
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.15
        for name, size, color in (("Heading 1", 20, NAVY), ("Heading 2", 14, "FFFFFF")):
            h = self.doc.styles[name]
            h.font.name, h.font.size, h.font.bold = "Arial", Pt(size), True
            h.font.color.rgb = RGBColor.from_string(color)
            h.element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}asciiTheme", "")
            h.paragraph_format.space_before = Pt(0)
            h.paragraph_format.space_after = Pt(8)
            h.paragraph_format.keep_with_next = True
        self.market = market
        self._footer(sec, market)
        self._front = bool(front_png)

    def _footer(self, sec, market: str, indent=None) -> None:
        sec.footer.is_linked_to_previous = False
        p = sec.footer.paragraphs[0]
        for e in list(p._p):
            if not e.tag.endswith("}pPr"):
                p._p.remove(e)
        if indent is not None:
            p.paragraph_format.left_indent = p.paragraph_format.right_indent = indent
        def r(text=""):
            run = p.add_run(text)
            run.font.size = self.Pt(7)
            run.font.color.rgb = self.RGB.from_string(SLATE5)
            return run
        r(f"Coherent Quadrant \u00b7 {market}")
        _ptab(r(), "indent" if indent is not None else "margin")
        r("Page ")
        _field(r(), "PAGE")
        r(" of ")
        _field(r(), "NUMPAGES")

    # paragraphs
    def para(self, text="", *, size=None, bold=False, color=None, after=None, before=None, italic=False,
             caps=False, spacing=None):
        p = self.doc.add_paragraph()
        if text:
            self.run(p, text, size=size, bold=bold, color=color, italic=italic, caps=caps, spacing=spacing)
        if after is not None:
            p.paragraph_format.space_after = self.Pt(after)
        if before is not None:
            p.paragraph_format.space_before = self.Pt(before)
        return p

    def run(self, p, text, *, size=None, bold=False, color=None, italic=False, caps=False, spacing=None):
        run = p.add_run(text)
        if size:
            run.font.size = self.Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.all_caps = caps
        if color:
            run.font.color.rgb = self.RGB.from_string(color)
        if spacing:
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn

            sp = OxmlElement("w:spacing")
            sp.set(qn("w:val"), str(spacing))
            run._r.get_or_add_rPr().append(sp)
        return run

    def heading(self, title: str, sub: str = "", *, new_page=True, toc=True):
        if toc:
            h = self.doc.add_heading(title, level=1)
        else:  # same look, kept out of the contents (as in the PDF)
            h = self.para(title, size=20, bold=True, color=NAVY, after=8)
            h.paragraph_format.keep_with_next = True
        h.paragraph_format.page_break_before = new_page
        _border(h, "bottom", TEAL, size=24, space=6)
        if sub:
            self.para(sub, size=9.5, color=SLATE5, after=10)
        return h

    def table(self, ncols: int, widths: list[float], *, data: bool = False):
        t = self.doc.add_table(rows=0, cols=ncols)
        _table_borders(t, data)
        t.autofit = False
        t.alignment = 1
        total = sum(widths)
        self._widths = [int(self.width * w / total) for w in widths]
        return t

    def row(self, t, values, *, header=False, fills=None):
        r = t.add_row()
        (_repeat_header if header else _no_split)(r)
        for i, c in enumerate(r.cells):
            c.width = self._widths[i]
            if fills and fills[i]:
                _shade(c._tc, fills[i])
        # A one-colour layout block (cover, tiles) is also shaded at table
        # level, so Word's PDF output shows no hairline seams between lines.
        if fills and not header and len(set(fills)) == 1 and fills[0] == NAVY:
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn

            tblpr = t._tbl.tblPr
            if tblpr.find(qn("w:shd")) is None:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), NAVY)
                tblpr.append(shd)
        return r


def render_front_cover(path: Path, title: str) -> bytes:
    """The filled cover page as an image (Word shows it full-bleed)."""
    doc = filled_front_cover(path, title)
    try:
        return doc[0].get_pixmap(dpi=200).tobytes("png")
    finally:
        doc.close()


def build(data: dict, chart_png: bytes, back_cover: Path | None, out: Path,
          front_cover: Path | None = None) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_BREAK
    from docx.shared import Mm, Pt

    cover = data["cover"]
    front_png = render_front_cover(front_cover, cover["market"]) if front_cover else None
    b = Builder(cover["market"], front_png)
    doc = b.doc

    # Cover: navy block with the market; Market Classification; key figures.
    t = b.table(1, [1])
    cell = b.row(t, [None], fills=[NAVY]).cells[0]
    _cell_margins(cell, 600)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(120)
    b.run(cp, cover["brand"], size=10.5, bold=True, color="E2E8F0", spacing=60)
    for text, size, bold, color, after in ((cover["market"], 26, True, "FFFFFF", 6),
                                           (cover["geo"], 11, True, "FFFFFF", 2),
                                           (cover["scope"], 10, False, "CBD5E1", 18)):
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        b.run(p, text, size=size, bold=bold, color=color)
        if text == cover["market"]:
            _border(p, "bottom", TEAL, size=24, space=10)
    b.para("", after=6)

    sections = list(data["sections"])
    first = sections.pop(0)
    _panels(b, first, new_page=False, toc=False)
    if data["tiles"]:
        t = b.table(len(data["tiles"]), [1] * len(data["tiles"]))
        r = b.row(t, [None] * len(data["tiles"]), fills=[NAVY] * len(data["tiles"]))
        for c, (num, label) in zip(r.cells, data["tiles"]):
            _cell_margins(c, 220)
            p = c.paragraphs[0]
            b.run(p, num, size=24, bold=True, color="FFFFFF")
            p2 = c.add_paragraph()
            b.run(p2, label, size=9, color="E2E8F0")

    # Contents: a Word TOC field (Word fills in the page numbers).
    p = b.para("Contents", size=20, bold=True, color=NAVY, after=8)
    p.paragraph_format.page_break_before = True
    _border(p, "bottom", TEAL, size=24, space=6)
    _field(b.para().add_run(), 'TOC \\o "1-2" \\h \\z \\u')

    for s in sections:
        kind = s.get("kind")
        if kind == "chart":
            b.heading(s["title"], s["sub"])
            pic = b.para(after=6)
            pic.alignment = 1
            pic.add_run().add_picture(io.BytesIO(chart_png), width=b.width)
            if s.get("mixTitle"):
                b.para(s["mixTitle"], size=9, bold=True, color=SLATE7, after=1)
                b.para(s["mix"], size=9, color=SLATE7)
        elif kind == "qdefs":
            b.heading(s["title"], s["sub"])
            for q in s["items"]:
                t = b.table(1, [1])
                c = b.row(t, [None], fills=["F8FAFC"]).cells[0]
                _cell_margins(c, 160)
                p = c.paragraphs[0]
                b.run(p, q["tag"], size=8, bold=True, color=SLATE9)
                b.run(p, "   " + q["title"], size=11, bold=True, color=SLATE9)
                b.run(c.add_paragraph(), q["sub"], size=8.5, color=SLATE5)
                b.run(c.add_paragraph(), q["text"], size=9.5, color=SLATE7)
                b.para("", after=4)
        elif kind == "table":
            _table(b, s)
        elif kind == "detail":
            _detail(b, s)
        elif kind == "method":
            b.heading(s["title"], s["sub"])
            if s["intro"]:
                b.para(s["intro"], size=9.5, color=SLATE5, after=10)
            t = b.table(2, [1, 9])
            for st in s["steps"]:
                r = b.row(t, [None, None], fills=[ACCENT, "F8FAFC"])
                for c in r.cells:
                    _cell_margins(c, 140)
                r.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = r.cells[0].paragraphs[0]
                p.alignment = 1
                b.run(p, st["n"], size=14, bold=True, color="FFFFFF")
                b.run(r.cells[1].paragraphs[0], st["t"], size=10, bold=True, color=SLATE9)
                if st["d"]:
                    b.run(r.cells[1].add_paragraph(), st["d"], size=9, color=SLATE7)
            if s["note"]:
                b.para(s["note"], size=8.5, color=SLATE5, before=10)
        else:
            _panels(b, s, new_page=True)

    # Table cells: tight paragraph spacing (the Normal style's spacing is for body text).
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for cp in c.paragraphs:
                    pf = cp.paragraph_format
                    if pf.space_after is None:
                        pf.space_after = Pt(1)
                    pf.line_spacing = 1.0
    if back_cover:
        _back_cover(b, back_cover)
    if b._front:
        first = doc.sections[0]
        first.footer.is_linked_to_previous = False
        for fp in first.footer.paragraphs:
            for e in list(fp._p):
                if not e.tag.endswith("}pPr"):
                    fp._p.remove(e)
    doc.save(str(out))


def _panels(b: "Builder", s: dict, *, new_page: bool, toc: bool = True) -> None:
    from docx.shared import Mm

    b.heading(s["title"], s["sub"], new_page=new_page, toc=toc)
    if s.get("lead"):
        b.para(s["lead"], size=9.5, color=SLATE5, after=8)
    panels = s.get("panels") or []
    if len(panels) == 2 and not any("ol" in part and any(i.get("def") for i in part["ol"])
                                    for pnl in panels for part in pnl):
        t = b.table(2, [1, 1])
        r = b.row(t, [None, None], fills=["F8FAFC", "F8FAFC"])
        targets = r.cells
    else:
        targets = [None] * len(panels)
    for pnl, cell in zip(panels, targets):
        if cell is None:
            t = b.table(1, [1])
            cell = b.row(t, [None], fills=["F8FAFC"]).cells[0]
        _cell_margins(cell, 160)
        first = True
        for part in pnl:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            if "lbl" in part:
                b.run(p, part["lbl"], size=10, bold=True, color=SLATE9)
            elif "val" in part:
                b.run(p, part["val"], size=9.5, color=SLATE7)
            elif "muted" in part:
                b.run(p, part["muted"], size=8.5, color=SLATE5)
            elif "ol" in part:
                for i, item in enumerate(part["ol"], 1):
                    if i > 1:
                        p = cell.add_paragraph()
                    b.run(p, f"{i}. ", size=9.5, bold=True, color=SLATE9)
                    b.run(p, item["name"], size=9.5, bold=bool(item.get("def")), color=SLATE9)
                    if item.get("def"):
                        dp = cell.add_paragraph()
                        dp.paragraph_format.left_indent = Mm(5)
                        b.run(dp, item["def"], size=9, color="475569")
        if len(targets) and targets[0] is None:
            b.para("", after=4)


def _table(b: "Builder", s: dict) -> None:
    b.heading(s["title"], s["sub"])
    head = s["head"]
    tall = {"Brand": 2.0, "Company": 2.2, "HQ": 1.7, "Role": 1.2, "Quadrant": 1.75,
            "Relative Positioning": 1.75}
    widths = [tall.get(h, 1.35) for h in head]
    t = b.table(len(head), widths, data=True)
    r = b.row(t, head, header=True, fills=[NAVY] * len(head))
    for c, h in zip(r.cells, head):
        _cell_margins(c, 60)
        b.run(c.paragraphs[0], h.upper(), size=6.5, bold=True, color="FFFFFF")
    for i, cells in enumerate(s["rows"]):
        r = b.row(t, cells, fills=["F8FAFC" if i % 2 else "FFFFFF"] * len(cells))
        for c, v in zip(r.cells, cells):
            _cell_margins(c, 45)
            p = c.paragraphs[0]
            if "dots" in v:
                for k in range(v["of"]):
                    b.run(p, "●" if k < v["dots"] else "○", size=8 + k, color=NAVY)
            elif "pill" in v:
                b.run(p, "● ", size=8, color=QUAD_COLORS.get(v["pill"], "CBD5E1"))
                b.run(p, v["pill"], size=8, bold=True, color=SLATE9)
            elif "disc" in v:
                b.run(p, v["disc"], size=9, bold=True, color="1E3A8A")
            else:
                b.run(p, v.get("t", ""), size=8, color=SLATE7)


def _detail(b: "Builder", s: dict) -> None:
    from docx.shared import Mm

    b.heading(s["title"], s["sub"])
    for n, co in enumerate(s["companies"]):
        h = b.doc.add_heading(co["name"], level=2)
        h.paragraph_format.page_break_before = n > 0
        h.paragraph_format.space_after = 0
        _shade(h._p, NAVY)
        meta = b.para(after=10)
        _shade(meta._p, NAVY)
        if co["brand"]:
            b.run(meta, co["brand"] + "   ", size=9, color="CBD5E1")
        b.run(meta, "● ", size=9, color=QUAD_COLORS.get(co["quadrant"], "CBD5E1"))
        b.run(meta, co["quadrant"], size=9, bold=True, color="FFFFFF")
        for label, val in co["scores"]:
            b.run(meta, f"     {label} ", size=9, color="CBD5E1")
            b.run(meta, val, size=10, bold=True, color="FFFFFF")
        for ax in co["axes"]:
            p = b.para(ax["title"], size=9.5, bold=True, color=ACCENT, before=6, after=6, caps=True)
            _border(p, "bottom", "E2E8F0", size=6, space=3)
            p.paragraph_format.keep_with_next = True
            for prm in ax["params"]:
                hp = b.para(prm["h"], size=10, bold=True, color=SLATE9, before=6, after=3)
                hp.paragraph_format.keep_with_next = True
                hp.paragraph_format.left_indent = Mm(3)
                _border(hp, "left", TEAL, size=24, space=6)
                if prm["basis"]:
                    bp = b.para(prm["basis"], size=9.5, color=SLATE7, after=4)
                    bp.paragraph_format.left_indent = Mm(3)
                    _border(bp, "left", TEAL, size=24, space=6)
                    bp.paragraph_format.keep_with_next = bool(prm["items"])
                if prm["items"]:
                    sp = b.para(prm["sh"], size=7.5, bold=True, color=SLATE5, after=2, caps=True)
                    sp.paragraph_format.left_indent = Mm(3)
                    sp.paragraph_format.keep_with_next = True
                    for it in prm["items"]:
                        lp = b.doc.add_paragraph(style="List Bullet")
                        lp.paragraph_format.left_indent = Mm(9)
                        lp.paragraph_format.space_after = b.Pt(3)
                        b.run(lp, it, size=9.5, color=SLATE7)
                if prm["none"]:
                    b.para(prm["none"], size=9, color=SLATE5)


def _back_cover(b: "Builder", path: Path) -> None:
    import pymupdf

    with pymupdf.open(str(path)) as d:
        png = d[0].get_pixmap(dpi=200, clip=pymupdf.Rect(0, 0, d[0].rect.width, d[0].rect.height - 40)).tobytes("png")
        w_pt = d[0].rect.width
    sec = b.doc.add_section(b.WD_SECTION.NEW_PAGE)
    sec.top_margin = sec.left_margin = sec.right_margin = b.Mm(0)
    sec.bottom_margin = b.Mm(12)
    sec.footer_distance = b.Mm(6)
    b._footer(sec, b.market, indent=b.Mm(16))
    p = b.doc.add_paragraph()
    p.paragraph_format.space_after = 0
    p.add_run().add_picture(io.BytesIO(png), width=b.Mm(210))


def word_finish(path: Path) -> bool:
    """Refresh the TOC and page fields in Word (if installed), then save."""
    try:
        import win32com.client  # type: ignore
    except ImportError:
        import subprocess

        ps = ("$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
              f"$d = $w.Documents.Open('{path}'); "
              "foreach ($t in $d.TablesOfContents) { $t.Update() }; $d.Fields.Update() | Out-Null; "
              "$d.Save(); $d.Close(); $w.Quit()")
        return subprocess.call(["powershell", "-NoProfile", "-Command", ps]) == 0
    w = win32com.client.Dispatch("Word.Application")
    w.Visible = False
    d = w.Documents.Open(str(path))
    for t in d.TablesOfContents:
        t.Update()
    d.Fields.Update()
    d.Save()
    d.Close()
    w.Quit()
    return True


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--market")
    ap.add_argument("--country", default="global")
    ap.add_argument("--html")
    ap.add_argument("--out")
    ap.add_argument("--no-word", action="store_true", help="skip the Word field refresh")
    args = ap.parse_args()
    if args.html:
        html = Path(args.html)
    elif args.market:
        from vendor_intel.pipeline.web_expand import default_output_dir

        out_dir = Path(default_output_dir(args.market, args.country))
        html = out_dir / f"{out_dir.name}_report.html"
    else:
        ap.error("give --market or --html")
    out = Path(args.out) if args.out else html.with_suffix(".docx")
    data, chart = asyncio.run(read_report(html))
    build(data, chart, back_cover_for(html.parent), out, front_cover_for(html.parent))
    refreshed = False if args.no_word else word_finish(out.resolve())
    print(f"wrote {out} ({len(data['sections'])} sections"
          + (", contents refreshed in Word" if refreshed else ", open in Word and update fields for page numbers")
          + ")")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
