#!/usr/bin/env python3
"""Apply web-verified Found in (HQ City, Country) fixes for semiconductor market."""
from __future__ import annotations

import csv
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.enrichment.hq_city_country import normalize_city_country
from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx

OUT = ROOT / "output" / "chatgpt_expand"
SLUG = "global_semiconductor_market_global"
QUERY = "Global Semiconductor Market"
FOLDER = OUT / SLUG

# Only changes verified via public company / filings / official sites.
# Do not invent HQs — omit anything not confirmed.
FOUND_IN_FIXES: dict[str, str] = {
    # siliconmotion.com Taiwan HQ
    "Silicon Motion": "Zhubei, Taiwan",
    # SEC principal executive office (not Luxembourg registry-only)
    "Magnachip": "Cheongju, South Korea",
    # LinkedIn / company HQ Bengaluru (Mumbai is Tata Group corporate)
    "Tata Electronics": "Bengaluru, Karnataka, India",
    # SEC 10-K principal executive offices (moved from Chicago)
    "Littelfuse": "Rosemont, Illinois, USA",
    # sensirion.com contact — Laubisruetistrasse 50, 8712 Stäfa (encoding repair)
    "Sensirion": "Stäfa, Switzerland",
    # ams-osram.com headquarters Premstätten / Premstaetten (encoding repair)
    "ams OSRAM": "Premstätten, Austria",
    # chipus.com.br / chipus-ip.com design center HQ (encoding repair)
    "Chipus Microelectronics": "Florianópolis, Brazil",
}


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def main() -> int:
    xlsx = FOLDER / f"{SLUG}_FINAL.xlsx"
    wb = load_workbook(xlsx, data_only=True)
    ws = wb["Landscape"]
    rows_raw = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows_raw[0]]
    landscape: list[dict] = []
    for r in rows_raw[1:]:
        if not r:
            continue
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        if str(d.get("Company") or "").strip():
            landscape.append(d)

    details_by: dict[str, dict] = {}
    if "Company Details" in wb.sheetnames:
        ws = wb["Company Details"]
        rows_raw = list(ws.iter_rows(values_only=True))
        dhdr = [str(h) for h in rows_raw[0]]
        for r in rows_raw[1:]:
            if not r:
                continue
            d = {dhdr[i]: ("" if r[i] is None else r[i]) for i in range(len(dhdr)) if i < len(r)}
            key = _norm(d.get("Brand") or d.get("Company") or "")
            if key:
                details_by[key] = d

    applied: list[dict] = []
    for name, raw_hq in FOUND_IN_FIXES.items():
        hq = normalize_city_country(raw_hq) or raw_hq
        key = _norm(name)
        # Landscape
        for row in landscape:
            ident = _norm(row.get("Company") or row.get("Brand") or "")
            if ident == key or key in ident or ident in key:
                old = str(row.get("Headquarters") or row.get("Found in") or "").strip()
                row["Headquarters"] = hq
                row["Found in"] = hq
                if old != hq:
                    applied.append({"brand": name, "from": old or "(empty)", "to": hq, "sheet": "Landscape"})
        # Details overlay identity
        det = details_by.get(key)
        if det:
            old = str(det.get("Found in") or "").strip()
            det["Found in"] = hq
            # keep scores/role
            for row in landscape:
                if _norm(row.get("Company") or "") == key or _norm(row.get("Brand") or "") == key:
                    row["Brand"] = str(det.get("Brand") or name)
                    if det.get("Role"):
                        row["Distribution Type"] = det["Role"]
                        row["Role"] = det["Role"]
                    if det.get("X") not in (None, ""):
                        row["X Score"] = det["X"]
                        row["Y Score"] = det["Y"]
                        row["Overall Score"] = det["Overall"]
                        row["Quadrant"] = det.get("Quadrant") or ""
                    company_col = str(det.get("Company") or "").strip()
                    brand = str(det.get("Brand") or "").strip()
                    if company_col and company_col != brand:
                        row["_display_company"] = company_col
                    row["Headquarters"] = hq
                    row["Found in"] = hq
            if old != hq:
                applied.append({"brand": name, "from": old or "(empty)", "to": hq, "sheet": "Details"})

    # Merge details onto all landscape for export
    for row in landscape:
        key = _norm(row.get("Company") or row.get("Brand") or "")
        det = details_by.get(key)
        if not det:
            continue
        brand = str(det.get("Brand") or "").strip()
        if brand:
            row["Brand"] = brand
            row["Company"] = brand
        if det.get("Role"):
            row["Distribution Type"] = det["Role"]
            row["Role"] = det["Role"]
        if det.get("Found in"):
            # Prefer already-fixed Found in on row
            if not str(row.get("Found in") or "").strip():
                row["Found in"] = det["Found in"]
                row["Headquarters"] = det["Found in"]
        if det.get("X") not in (None, ""):
            row["X Score"] = det["X"]
            row["Y Score"] = det["Y"]
            row["Overall Score"] = det["Overall"]
            row["Quadrant"] = det.get("Quadrant") or ""
        company_col = str(det.get("Company") or "").strip()
        if company_col and company_col != brand:
            row["_display_company"] = company_col

    # Pull axes from quadrant json
    audit: dict = {}
    qjson = FOLDER / f"{SLUG}_quadrant.json"
    if qjson.exists():
        payload = json.loads(qjson.read_text(encoding="utf-8"))
        crit = payload.get("criteria") or {}
        if crit.get("parameter_definitions"):
            audit["parameter_definitions"] = crit["parameter_definitions"]
        labels = crit.get("axis_labels") or {}
        if labels.get("x"):
            audit["axis_x"] = labels["x"]
        if labels.get("y"):
            audit["axis_y"] = labels["y"]
        if crit.get("x_axis"):
            audit["x_features"] = crit["x_axis"]
        if crit.get("y_axis"):
            audit["y_features"] = crit["y_axis"]
        if payload.get("industry_group"):
            audit["industry_group"] = payload["industry_group"]
        if payload.get("industry_category"):
            audit["industry_category"] = payload["industry_category"]

    detail_rows = to_company_detail_rows(landscape, QUERY, audit)
    for det in detail_rows:
        key = _norm(det.get("Brand") or "")
        # Force verified Found in
        for name, raw_hq in FOUND_IN_FIXES.items():
            if _norm(name) == key:
                det["Found in"] = normalize_city_country(raw_hq) or raw_hq
        src = next(
            (r for r in landscape if _norm(r.get("Company") or r.get("Brand") or "") == key),
            None,
        )
        if src and src.get("_display_company"):
            det["Company"] = src["_display_company"]
        if src and src.get("Found in") and not det.get("Found in"):
            det["Found in"] = src["Found in"]

    write_final_xlsx(
        xlsx,
        landscape,
        "Companies",
        {"query": QUERY, "xy_scoring": audit, "found_in_fixes": applied},
        detail_rows=detail_rows,
    )
    extras = export_expand_quadrant_outputs(
        FOLDER, detail_rows, QUERY, country="global", audit=audit, chart_n=20
    )

    audit_out = {
        "verified_fixes": [
            {
                "brand": k,
                "found_in": normalize_city_country(v) or v,
                "source": {
                    "Silicon Motion": "siliconmotion.com Taiwan HQ Zhubei",
                    "Magnachip": "SEC principal executive office Cheongju, Korea",
                    "Tata Electronics": "company HQ Bengaluru (Mumbai is Tata corporate)",
                    "Littelfuse": "SEC 10-K: 6133 North River Road, Suite 500, Rosemont, Illinois 60018",
                    "Sensirion": "sensirion.com contact: 8712 Stäfa, Switzerland",
                    "ams OSRAM": "ams-osram.com HQ Tobelbader Strasse 30, 8141 Premstätten, Austria",
                    "Chipus Microelectronics": "chipus.com.br HQ Florianópolis, SC, Brazil",
                }.get(k, "web verified"),
            }
            for k, v in FOUND_IN_FIXES.items()
        ],
        "empty_after": [
            r.get("Brand")
            for r in detail_rows
            if not str(r.get("Found in") or "").strip()
        ],
        "n": len(detail_rows),
        "html": extras.get("html"),
    }
    # Re-read CSV after ownership apply
    csv_path = FOLDER / f"{SLUG}_companies.csv"
    if csv_path.exists():
        rows = list(csv.DictReader(csv_path.open(encoding="utf-8")))
        audit_out["empty_after"] = [
            r.get("Brand") for r in rows if not str(r.get("Found in") or "").strip()
        ]
        audit_out["n"] = len(rows)
        for k in FOUND_IN_FIXES:
            hit = next((r for r in rows if r.get("Brand") == k), None)
            if hit:
                audit_out.setdefault("csv_check", {})[k] = hit.get("Found in")

    out = OUT / "_audit" / "semi_found_in_verification.json"
    out.write_text(json.dumps(audit_out, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Applied {len(FOUND_IN_FIXES)} Found in fixes", flush=True)
    for k, v in FOUND_IN_FIXES.items():
        print(f"  {k} -> {normalize_city_country(v) or v}", flush=True)
    print(f"empty remaining: {audit_out.get('empty_after')}", flush=True)
    print(f"audit -> {out}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
