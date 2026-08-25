"""Audit Found in column quality for both markets."""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"D:/Coherent-Quadrant/output/chatgpt_expand")
OUT = ROOT / "_audit"
OUT.mkdir(exist_ok=True)

COUNTRY_ONLY = re.compile(
    r"^(united states|usa|us|united kingdom|uk|india|japan|china|germany|france|"
    r"italy|spain|brazil|canada|australia|switzerland|austria|finland|sweden|"
    r"norway|denmark|netherlands|belgium|ireland|israel|uae|united arab emirates|"
    r"saudi arabia|qatar|malaysia|thailand|singapore|south korea|korea|"
    r"south africa|mexico|chile|indonesia|hong kong|taiwan|poland|greece|"
    r"turkey|russia|new zealand|portugal|luxembourg|seed_file)$",
    re.I,
)


def score(v: str) -> str:
    v = (v or "").strip()
    if not v:
        return "empty"
    if v.lower() == "seed_file":
        return "seed_file"
    if COUNTRY_ONLY.match(v):
        return "country_only"
    if re.search(r"\b(New|Kuala|Hong|Sao|Saint|Los|San),\s+", v):
        return "broken_city"
    if v.count(",") >= 3 and re.search(r"United,\s*States|Macfarlane|Kl", v, re.I):
        return "messy"
    if "," in v:
        return "okish"
    return "single_token"


def main() -> None:
    for slug in [
        "global_flexible_packaging_market_global",
        "global_wearable_medical_devices_market_global",
    ]:
        x = ROOT / slug / f"{slug}_FINAL.xlsx"
        wb = load_workbook(x, data_only=True)
        ws = wb["Company Details"]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(h) for h in rows[0]]
        items = []
        buckets: dict[str, int] = {}
        for r in rows[1:]:
            d = {
                hdr[i]: ("" if r[i] is None else str(r[i]))
                for i in range(len(hdr))
                if i < len(r)
            }
            brand = d.get("Brand") or d.get("Company") or ""
            found = d.get("Found in") or ""
            b = score(found)
            buckets[b] = buckets.get(b, 0) + 1
            items.append(
                {
                    "brand": brand,
                    "company": d.get("Company", ""),
                    "found_in": found,
                    "status": b,
                }
            )
        (OUT / f"{slug}_found_in_audit.json").write_text(
            json.dumps(items, indent=2), encoding="utf-8"
        )
        need = [it for it in items if it["status"] != "okish"]
        print(slug, "n=", len(items), buckets, "need_fix=", len(need))
        for it in need:
            print(f"  [{it['status']}] {it['brand']!r} => {it['found_in']!r}")


if __name__ == "__main__":
    main()
