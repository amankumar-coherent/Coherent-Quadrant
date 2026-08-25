#!/usr/bin/env python3
"""Verify LNG Brand vs Marketer — deterministic only (no hallucinated flips).

Rules (Global Liquefied Natural Gas Market):
  Brand    = owns/operates liquefaction, LNG terminals/FSRUs, upstream gas for LNG,
             integrated IOC/NOC/utility whose identity is producing or supplying LNG
             under its operating brand.
  Marketer = primary LNG activity is cargo trading / offtake marketing / portfolio
             trading — independent commodity traders, dedicated LNG trading arms,
             and Japanese sogo shosha whose LNG role is marketing/trading
             (even if they hold minority project equity).

Only companies matching verified MARKETER_STEMS are set to Marketer.
Everyone else stays Brand. No LLM inventing roles.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx

SLUG = "global_liquefied_natural_gas_market_global"
QUERY = "Global Liquefied Natural Gas Market"
OUT = ROOT / "output" / "chatgpt_expand" / SLUG
XLSX = OUT / f"{SLUG}_FINAL.xlsx"
AUDIT = ROOT / "output" / "chatgpt_expand" / "_audit"

# Verified LNG marketers / traders / trading-house marketers (web-confirmed roles).
# Sources: WoodMac (Trafigura/Vitol/Gunvor/Glencore LNG traders); Shell press
# (Pavilion LNG trading); Mitsui/Marubeni LNG marketing units; sogo shosha literature.
MARKETER_STEMS: dict[str, str] = {
    # Independent commodity traders
    "vitol": "independent LNG/commodity trader",
    "gunvor": "independent LNG/commodity trader",
    "trafigura": "independent LNG/commodity trader",
    "mercuria": "independent LNG/commodity trader",
    "glencore": "independent LNG/commodity trader",
    "castleton commodities": "US gas/LNG merchant trader",
    "hartree partners": "energy/LNG merchant trader",
    "freepoint commodities": "US gas/LNG merchant trader",
    "bb energy": "independent oil & LNG trader",
    # Dedicated LNG trading / marketing arms
    "pavilion energy": "LNG trading portfolio (acquired by Shell Apr 2025)",
    "qatarenergy trading": "QatarEnergy LNG marketing/trading arm",
    "lng japan": "Japanese LNG procurement/marketing company",
    "sk gas trading": "SK Group LNG/gas trading arm",
    "ptt global lng": "PTT LNG marketing subsidiary",
    # Japanese sogo shosha — LNG marketing/trading organizers (not plant operators)
    "mitsubishi corporation": "sogo shosha LNG marketing/trading",
    "mitsui co": "sogo shosha LNG marketing/trading",
    "mitsui & co": "sogo shosha LNG marketing/trading",
    "marubeni": "sogo shosha LNG marketing/trading",
    "itochu": "sogo shosha LNG marketing/trading",
    "sumitomo corporation": "sogo shosha LNG marketing/trading",
    "sojitz": "sogo shosha LNG marketing/trading",
}


def _norm(s: str) -> str:
    s = str(s or "").strip().lower()
    s = s.replace("&", " ")
    s = re.sub(r"\([^)]*\)", " ", s)
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def classify(name: str) -> tuple[str, str]:
    """Return (Brand|Marketer, reason). Marketer only on verified stems."""
    n = _norm(name)
    best = None
    best_len = -1
    for stem, reason in MARKETER_STEMS.items():
        sn = _norm(stem)
        if sn in n and len(sn) > best_len:
            best = (stem, reason)
            best_len = len(sn)
    if best:
        return "Marketer", f"{best[0]}: {best[1]}"
    return "Brand", "default_producer_operator_or_utility"


def main() -> None:
    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Landscape"]
    rows_raw = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows_raw[0]]
    landscape: list[dict] = []
    changed: list[dict] = []
    for r in rows_raw[1:]:
        if not r:
            continue
        d = {
            hdr[i]: ("" if r[i] is None else r[i])
            for i in range(len(hdr))
            if i < len(r)
        }
        name = str(d.get("Company") or "").strip()
        if not name:
            continue
        old = str(d.get("Role") or d.get("Distribution Type") or "Brand").strip()
        new, reason = classify(name)
        if old != new:
            changed.append(
                {"company": name, "from": old, "to": new, "reason": reason}
            )
        d["Role"] = new
        d["Distribution Type"] = new
        landscape.append(d)

    brand_n = sum(1 for r in landscape if r["Role"] == "Brand")
    marketer_n = sum(1 for r in landscape if r["Role"] == "Marketer")

    audit_path = OUT / "chatgpt_expand_batch_all_audit.json"
    audit: dict = {}
    if audit_path.exists():
        audit = json.loads(audit_path.read_text(encoding="utf-8")).get("xy_scoring") or {}

    detail_rows = to_company_detail_rows(landscape, QUERY, audit)
    by_co = {_norm(str(r.get("Company") or "")): str(r.get("Role") or "Brand") for r in landscape}
    for d in detail_rows:
        role = by_co.get(_norm(d.get("Brand") or ""))
        if role in ("Brand", "Marketer"):
            d["Role"] = role

    write_final_xlsx(
        XLSX,
        landscape,
        "Companies",
        {
            "query": QUERY,
            "role_verify": "2026-08-14",
            "brand": brand_n,
            "marketer": marketer_n,
            "xy_scoring": audit,
        },
        detail_rows=detail_rows,
    )
    extras = export_expand_quadrant_outputs(
        OUT, detail_rows, QUERY, country="global", audit=audit, chart_n=20
    )

    report = {
        "total": len(landscape),
        "brand": brand_n,
        "marketer": marketer_n,
        "changed": changed,
        "marketer_stems_used": sorted(MARKETER_STEMS.keys()),
        "rule": (
            "Brand = LNG producer/terminal/FSRU/IOC-NOC-utility operator; "
            "Marketer = verified commodity trader, LNG trading arm, or sogo shosha LNG marketer only"
        ),
        "html": str(extras.get("html") or ""),
    }
    AUDIT.mkdir(parents=True, exist_ok=True)
    (AUDIT / f"{SLUG}_role_verify.json").write_text(
        json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    lines = ["Brand\tCompany\tRole"]
    for d in detail_rows:
        lines.append(f"{d.get('Brand')}\t{d.get('Company')}\t{d.get('Role')}")
    (AUDIT / f"{SLUG}_roles_after.tsv").write_text("\n".join(lines), encoding="utf-8")

    print(f"total={len(landscape)} Brand={brand_n} Marketer={marketer_n} changed={len(changed)}")
    for c in changed:
        print(f"  {c['from']}->{c['to']}: {c['company']} ({c['reason']})")


if __name__ == "__main__":
    main()
