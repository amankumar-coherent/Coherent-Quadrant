#!/usr/bin/env python3
"""Restore Terns Pharmaceuticals and regenerate GLP-1 report (chart label fix)."""
from __future__ import annotations

import csv
import json
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx
from vendor_intel.quadrant.rating_map import assign_quadrants_absolute_median

SLUG = "global_glp_1_receptor_agonist_market_global"
QUERY = "Global GLP-1 Receptor Agonist Market"
OUT = ROOT / "output" / "chatgpt_expand" / SLUG
XLSX = OUT / f"{SLUG}_FINAL.xlsx"

# Restored from pre-drop FINAL + deep-rescore audit (evidence-only scores)
TERNS = {
    "Company": "Terns Pharmaceuticals",
    "Brand": "Terns Pharmaceuticals",
    "Website": "https://www.ternspharma.com",
    "Founded": "2018",
    "Headquarters": "Foster City, USA",
    "Found in": "Foster City, USA",
    "Continent / Geography": "North America",
    "Operational Presence": "United States, China",
    "Ownership": "Independent",
    "Employees": "",
    "Core Categories": "Healthcare / Pharmaceuticals",
    "Specialty Focus": "GLP-1 / incretin therapies",
    "Key Brands Represented": "TERN-601",
    "Retail / E-commerce / Both": "No",
    "Distribution Type": "Brand",
    "Role": "Brand",
    "Country Code": "US",
    "Region Code": "NA",
    "Summary": "Terns Pharmaceuticals — deep-crawl scored X=37 Y=34 (ok; evid=37504; crawled=True; q=thin)",
    "X Score": 37,
    "Y Score": 34,
    "Overall Score": 36,
    "Industry Category": "Healthcare / Pharmaceutical",
}


def _norm(s: str) -> str:
    return " ".join(str(s or "").lower().split())


def main() -> int:
    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Landscape"]
    rows_raw = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows_raw[0]]

    details_by = {}
    if "Company Details" in wb.sheetnames:
        dws = wb["Company Details"]
        drows = list(dws.iter_rows(values_only=True))
        dhdr = [str(h) for h in drows[0]]
        for r in drows[1:]:
            if not r or not r[0]:
                continue
            d = {dhdr[i]: ("" if r[i] is None else r[i]) for i in range(len(dhdr))}
            details_by[_norm(d.get("Brand") or "")] = d

    kept = []
    for r in rows_raw[1:]:
        if not r or not r[0]:
            continue
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(min(len(hdr), len(r)))}
        brand = str(d.get("Company") or d.get("Brand") or "").strip()
        if not brand:
            continue
        det = details_by.get(_norm(brand)) or {}
        if det.get("Brand"):
            brand = str(det["Brand"]).strip()
            d["Brand"] = brand
        if det.get("Company") and str(det["Company"]).startswith("("):
            d["_display_company"] = det["Company"]
        if det.get("Found in"):
            d["Headquarters"] = det["Found in"]
            d["Found in"] = det["Found in"]
        if det.get("Role") in {"Brand", "Marketer"}:
            d["Role"] = det["Role"]
            d["Distribution Type"] = det["Role"]
        d["Company"] = brand
        d["Brand"] = brand
        d["X Score"] = det.get("X") if det.get("X") not in (None, "") else d.get("X Score")
        d["Y Score"] = det.get("Y") if det.get("Y") not in (None, "") else d.get("Y Score")
        d["Overall Score"] = (
            det.get("Overall") if det.get("Overall") not in (None, "") else d.get("Overall Score")
        )
        kept.append(d)

    if not any(_norm(r.get("Brand") or "") == "terns pharmaceuticals" for r in kept):
        kept.append(dict(TERNS))
        print("restored Terns Pharmaceuticals")
    else:
        print("Terns already present")

    xs = [float(r.get("X Score") or 50) for r in kept]
    ys = [float(r.get("Y Score") or 50) for r in kept]
    quads, mid_x, mid_y = assign_quadrants_absolute_median(xs, ys)
    for i, r in enumerate(kept):
        r["Quadrant"] = quads[i]
        r["Overall Score"] = round((xs[i] + ys[i]) / 2)

    audit = {
        "axis_x": "Therapeutic & Manufacturing Capability",
        "axis_y": "Commercial Reach & Pipeline Strategy",
        "industry_group": "Healthcare",
        "industry_category": "Pharmaceutical",
        "mid_x": mid_x,
        "mid_y": mid_y,
        "method": "absolute median; chart dots+labels kept inside own quadrant",
    }
    details = to_company_detail_rows(kept, QUERY, audit)
    for det in details:
        key = _norm(det.get("Brand") or "")
        src = next((r for r in kept if _norm(r.get("Brand") or "") == key), None)
        if not src:
            continue
        det["Role"] = src.get("Role") or "Brand"
        if src.get("_display_company"):
            det["Company"] = src["_display_company"]
        elif not str(det.get("Company") or "").startswith("("):
            det["Company"] = det.get("Brand") or det.get("Company")

    write_final_xlsx(
        XLSX,
        kept,
        "Companies",
        {"query": QUERY, "xy_scoring": audit},
        detail_rows=details,
    )
    export_expand_quadrant_outputs(
        OUT, details, QUERY, country="global", audit=audit, chart_n=20
    )
    with (OUT / f"{SLUG}_companies.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(
            f,
            fieldnames=["Brand", "Company", "Role", "Quadrant", "X", "Y", "Overall", "Found in"],
        )
        w.writeheader()
        for d in details:
            w.writerow({k: d.get(k, "") for k in w.fieldnames})

    print(f"n={len(kept)} mid_x={mid_x:.1f} mid_y={mid_y:.1f}")
    print("quadrants", Counter(r.get("Quadrant") for r in kept))
    print("terns", next(r for r in kept if "terns" in _norm(r.get("Brand") or "")))
    print("html", OUT / f"{SLUG}_report.html")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
