#!/usr/bin/env python3
"""Re-assign quadrants from existing X/Y scores and rebuild HTML/Excel.

Use after fixing assign_quadrants_half_median so all four cells fill even when
scores are heavily tied (score_floor collapse).
"""
from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx
from vendor_intel.quadrant.rating_map import assign_quadrants_half_median

OUT = ROOT / "output" / "chatgpt_expand"

MARKETS = [
    ("global_flexible_packaging_market_global", "Global Flexible Packaging Market"),
    ("global_wearable_medical_devices_market_global", "Global Wearable Medical Devices Market"),
    ("global_semiconductor_market_global", "Global Semiconductor Market"),
    ("global_liquefied_natural_gas_market_global", "Global Liquefied Natural Gas Market"),
]


def _load(xlsx: Path) -> tuple[list[dict], dict]:
    wb = load_workbook(xlsx, data_only=True)
    landscape: list[dict] = []
    land = "Landscape" if "Landscape" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[land]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows[0]]
    for r in rows[1:]:
        if not r:
            continue
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        if str(d.get("Company") or "").strip():
            landscape.append(d)

    details_by: dict[str, dict] = {}
    if "Company Details" in wb.sheetnames:
        ws = wb["Company Details"]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(h) for h in rows[0]]
        for r in rows[1:]:
            if not r:
                continue
            d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
            key = str(d.get("Brand") or d.get("Company") or "").strip().lower()
            if key:
                details_by[key] = d

    for row in landscape:
        key = str(row.get("Company") or "").strip().lower()
        det = details_by.get(key)
        if not det:
            for dk, dv in details_by.items():
                if dk in key or key in dk:
                    det = dv
                    break
        if not det:
            continue
        brand = str(det.get("Brand") or "").strip()
        if brand:
            row["Company"] = brand
            row["Brand"] = brand
        for col in ("Quadrant", "X", "Y", "Overall", "Role"):
            if det.get(col) not in (None, ""):
                row[col] = det[col]
                if col == "X":
                    row["X Score"] = det[col]
                elif col == "Y":
                    row["Y Score"] = det[col]
                elif col == "Overall":
                    row["Overall Score"] = det[col]
        if det.get("Role"):
            row["Distribution Type"] = det["Role"]
        if det.get("Found in"):
            row["Headquarters"] = det["Found in"]
            row["Found in"] = det["Found in"]
        company_col = str(det.get("Company") or "").strip()
        if company_col and company_col != brand:
            row["_display_company"] = company_col
    return landscape, details_by


def _int(v) -> int:
    try:
        return int(float(str(v or "0").split()[0]))
    except (TypeError, ValueError):
        return 0


def main() -> int:
    only = (sys.argv[1] if len(sys.argv) > 1 else "").strip().lower()
    for slug, query in MARKETS:
        if only and only not in slug and only not in query.lower():
            continue
        folder = OUT / slug
        xlsx = folder / f"{slug}_FINAL.xlsx"
        if not xlsx.exists():
            print(f"skip {slug}: no FINAL")
            continue
        landscape, _ = _load(xlsx)
        xs = [_int(r.get("X Score") or r.get("X")) for r in landscape]
        ys = [_int(r.get("Y Score") or r.get("Y")) for r in landscape]
        quads, mid_x, mid_y = assign_quadrants_half_median(xs, ys)
        for row, q in zip(landscape, quads):
            row["Quadrant"] = q
            if not row.get("X Score"):
                row["X Score"] = str(_int(row.get("X")))
            if not row.get("Y Score"):
                row["Y Score"] = str(_int(row.get("Y")))
            if not row.get("Overall Score"):
                x, y = _int(row.get("X Score")), _int(row.get("Y Score"))
                row["Overall Score"] = str(int(round((x + y) / 2.0)))

        audit_path = folder / "chatgpt_xy_scores_batch_all.json"
        audit: dict = {}
        if audit_path.exists():
            try:
                raw = json.loads(audit_path.read_text(encoding="utf-8"))
                # File may be full xy payload or already an audit dict
                audit = raw.get("xy_scoring") if isinstance(raw.get("xy_scoring"), dict) else raw
                if not isinstance(audit, dict):
                    audit = {}
            except Exception:
                audit = {}
        # Keep market axes / param defs from existing quadrant JSON
        qjson = folder / f"{slug}_quadrant.json"
        if qjson.exists():
            try:
                payload = json.loads(qjson.read_text(encoding="utf-8"))
                crit = payload.get("criteria") or {}
                if crit.get("parameter_definitions"):
                    audit["parameter_definitions"] = crit["parameter_definitions"]
                labels = crit.get("axis_labels") or {}
                if labels.get("x"):
                    audit["axis_x"] = labels["x"]
                if labels.get("y"):
                    audit["axis_y"] = labels["y"]
                if crit.get("x_axis"):
                    audit["x_features"] = crit["x_axis"]
                if crit.get("y_axis"):
                    audit["y_features"] = crit["y_axis"]
                if payload.get("industry_group"):
                    audit["industry_group"] = payload["industry_group"]
                if payload.get("industry_category"):
                    audit["industry_category"] = payload["industry_category"]
            except Exception:
                pass
        audit["quadrant_mid_x"] = mid_x
        audit["quadrant_mid_y"] = mid_y
        audit["quadrant_counts"] = dict(Counter(quads))

        detail_rows = to_company_detail_rows(landscape, query, audit)
        by_land = {
            str(r.get("Company") or r.get("Brand") or "").strip().lower(): r
            for r in landscape
        }
        for det in detail_rows:
            key = str(det.get("Brand") or "").strip().lower()
            src = by_land.get(key)
            if src and src.get("_display_company"):
                det["Company"] = src["_display_company"]

        write_final_xlsx(
            xlsx,
            landscape,
            "Companies",
            {"query": query, "xy_scoring": {k: v for k, v in audit.items() if k != "rows"}},
            detail_rows=detail_rows,
        )
        extras = export_expand_quadrant_outputs(
            folder, detail_rows, query, country="global", audit=audit, chart_n=20
        )
        print(f"{slug}: {dict(Counter(quads))} → {extras.get('html')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
