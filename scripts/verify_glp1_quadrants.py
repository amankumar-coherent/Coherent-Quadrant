#!/usr/bin/env python3
"""Verify GLP-1 Quadrant column vs Coherent Quadrant definitions (no hallucination)."""
from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.quadrant.rating_map import assign_quadrants_absolute_median

FOLDER = (
    ROOT
    / "output"
    / "chatgpt_expand"
    / "global_glp_1_receptor_agonist_market_global"
)
XLSX = FOLDER / "global_glp_1_receptor_agonist_market_global_FINAL.xlsx"
AUDIT = ROOT / "output" / "chatgpt_expand" / "_audit" / "glp1_quadrant_verify.json"

# Image + codebase semantic rules (high/low combinations):
#   Leaders           = High Solution Capability (X) + High Business Strategy (Y)
#   Challengers       = Lower X + High Y
#   Trailblazers      = High X + Lower Y
#   Emerging Players  = Lower X + Lower Y
# High/Lower = vs cohort absolute median (X>=mid_x, Y>=mid_y).


def _f(v) -> float | None:
    if v is None or str(v).strip() == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def main() -> int:
    wb = load_workbook(XLSX, data_only=True)
    dws = wb["Company Details"]
    drows = list(dws.iter_rows(values_only=True))
    dh = [str(h) for h in drows[0]]
    details = []
    for r in drows[1:]:
        if not r or not r[0]:
            continue
        details.append({dh[i]: ("" if r[i] is None else r[i]) for i in range(len(dh))})

    lws = wb["Landscape"]
    lrows = list(lws.iter_rows(values_only=True))
    lh = [str(h) for h in lrows[0]]
    landscape = {}
    for r in lrows[1:]:
        if not r or not r[0]:
            continue
        d = {lh[i]: ("" if r[i] is None else r[i]) for i in range(min(len(lh), len(r)))}
        brand = str(d.get("Company") or d.get("Brand") or "").strip()
        landscape[brand.lower()] = d

    xs, ys, brands, current = [], [], [], []
    missing_xy = []
    for d in details:
        brand = str(d.get("Brand") or "").strip()
        x = _f(d.get("X"))
        y = _f(d.get("Y"))
        q = str(d.get("Quadrant") or "").strip()
        if x is None or y is None:
            missing_xy.append(brand)
            continue
        brands.append(brand)
        xs.append(x)
        ys.append(y)
        current.append(q)

    expected, mid_x, mid_y = assign_quadrants_absolute_median(xs, ys)

    mismatches = []
    sheet_mismatch = []
    for i, brand in enumerate(brands):
        exp = expected[i]
        got = current[i]
        if got != exp:
            mismatches.append(
                {
                    "brand": brand,
                    "x": xs[i],
                    "y": ys[i],
                    "current": got,
                    "expected": exp,
                    "reason": (
                        f"X={xs[i]} {'>=' if xs[i] >= mid_x else '<'} mid_x={mid_x:.1f}; "
                        f"Y={ys[i]} {'>=' if ys[i] >= mid_y else '<'} mid_y={mid_y:.1f} "
                        f"→ {exp}"
                    ),
                }
            )
        land = landscape.get(brand.lower())
        if land:
            lq = str(land.get("Quadrant") or "").strip()
            lx = _f(land.get("X Score"))
            ly = _f(land.get("Y Score"))
            if lq and lq != got:
                sheet_mismatch.append(
                    {"brand": brand, "details_q": got, "landscape_q": lq}
                )
            if lx is not None and abs(lx - xs[i]) > 0.01:
                sheet_mismatch.append(
                    {"brand": brand, "issue": "X mismatch", "details": xs[i], "landscape": lx}
                )
            if ly is not None and abs(ly - ys[i]) > 0.01:
                sheet_mismatch.append(
                    {"brand": brand, "issue": "Y mismatch", "details": ys[i], "landscape": ly}
                )

    # Boundary examples for audit clarity
    by_quad: dict[str, list] = {q: [] for q in ("Leaders", "Challengers", "Trailblazers", "Emerging Players")}
    for i, brand in enumerate(brands):
        by_quad.setdefault(expected[i], []).append(
            {"brand": brand, "x": xs[i], "y": ys[i], "current": current[i]}
        )

    # Sanity: Leaders must be high-high per definition
    def_violations = []
    for i, brand in enumerate(brands):
        q = expected[i]
        hx, hy = xs[i] >= mid_x, ys[i] >= mid_y
        if q == "Leaders" and not (hx and hy):
            def_violations.append(brand)
        if q == "Challengers" and not ((not hx) and hy):
            def_violations.append(brand)
        if q == "Trailblazers" and not (hx and (not hy)):
            def_violations.append(brand)
        if q == "Emerging Players" and not ((not hx) and (not hy)):
            def_violations.append(brand)

    report = {
        "n": len(brands),
        "missing_xy": missing_xy,
        "mid_x": mid_x,
        "mid_y": mid_y,
        "axis_x": "Solution Capability (= Therapeutic & Manufacturing Capability / X)",
        "axis_y": "Business Strategy (= Commercial Reach & Pipeline Strategy / Y)",
        "definitions": {
            "Leaders": "High X + High Y — mature/scalable, balanced product + GTM",
            "Challengers": "Lower X + High Y — broad reach/partnerships, deepening product",
            "Trailblazers": "High X + Lower Y — advanced capability, limited scale/GTM",
            "Emerging Players": "Lower X + Lower Y — newer/niche, limited presence",
        },
        "rule": "absolute median: High = score >= cohort median",
        "current_counts": dict(Counter(current)),
        "expected_counts": dict(Counter(expected)),
        "mismatches": mismatches,
        "mismatch_count": len(mismatches),
        "sheet_mismatch_count": len(sheet_mismatch),
        "sheet_mismatches": sheet_mismatch[:50],
        "definition_logic_violations": def_violations,
        "examples": {
            q: sorted(rows, key=lambda r: -(r["x"] + r["y"]))[:5]
            for q, rows in by_quad.items()
        },
    }

    AUDIT.parent.mkdir(parents=True, exist_ok=True)
    AUDIT.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")

    print(f"n={len(brands)} mid_x={mid_x:.2f} mid_y={mid_y:.2f}")
    print(f"current:  {Counter(current)}")
    print(f"expected: {Counter(expected)}")
    print(f"mismatches vs definition rule: {len(mismatches)}")
    print(f"Landscape vs Details sheet issues: {len(sheet_mismatch)}")
    print(f"logic violations in expected labels: {len(def_violations)}")
    if mismatches[:15]:
        print("\nMISMATCHES (first 15):")
        for m in mismatches[:15]:
            print(f"  {m['brand']}: got={m['current']} expected={m['expected']} | {m['reason']}")
    print(f"\nAudit: {AUDIT}")
    return 0 if not mismatches and not missing_xy else 1


if __name__ == "__main__":
    raise SystemExit(main())
