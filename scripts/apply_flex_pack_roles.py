"""Assign Brand OR Marketer (never both) for Flexible Packaging FINAL + rebuild HTML.



Rule:

  Brand    = primary flexible packaging CONVERTER (pouches/bags/laminates for CPG)

  Marketer = base-film producer (BOPP/BOPET/CPP rolls) OR specialty film/resin brand

             marketer (Aclar, EVAL, Scotchpak, Tyvek, EVOH, polyester films, etc.)

"""

from __future__ import annotations



import json

import re

import sys

from pathlib import Path



ROOT = Path(__file__).resolve().parents[1]

sys.path.insert(0, str(ROOT / "src"))



from openpyxl import load_workbook



from vendor_intel.pipeline.expand_quadrant_score import (

    export_expand_quadrant_outputs,

    to_company_detail_rows,

)

from vendor_intel.pipeline.role_split import classify_packaging_role

from vendor_intel.pipeline.web_expand import write_final_xlsx



SLUG = "global_flexible_packaging_market_global"

QUERY = "Global Flexible Packaging Market"

OUT = ROOT / "output" / "chatgpt_expand" / SLUG



# Finished-pack converters — always Brand

BRAND_FORCE = {

    "amcor",

    "constantia",

    "proampac",

    "printpack",

    "coveris",

    "schur",

    "südpack",

    "sudpack",

    "wipak",

    "goglio",

    "glenroy",

    "epac",

    "bryce",

    "interflex",

    "american packaging",

    "c-p flexible",

    "parkside",

    "paharpur",

    "bilcare",

    "epl",

    "essel",

    "clondalkin",

    "bischof",

    "huhtamaki",

    "mondi",

    "sealed air",

    "sonoco",

    "winpak",

    "transcontinental",

    "tcpl",

    "novolex",

    "pactiv",

    "berry",

    "bemis",

    "ampac",

    "aluflexpack",

    "flexopack",

    "wipf",

    "fabbri",

    "saica flex",

    "oliver healthcare",

    "selig",

    "fujimori",

    "time technoplast",

    "rpc bpi",

    "rpc group",

    "klöckner",

    "klockner",

    "huangshan novel",

    "uflex ltd",

}





def _norm(s: str) -> str:

    s = re.sub(r"\s+", " ", str(s or "").strip().lower())

    return s.replace("ö", "o").replace("ü", "u").replace("ä", "a")





def classify_role(brand: str, company: str = "") -> str:

    blob = _norm(f"{brand} {company}")

    for stem in sorted(BRAND_FORCE, key=len, reverse=True):

        if stem in blob:

            return "Brand"

    return classify_packaging_role(company, brand)





def main() -> None:

    xlsx = OUT / f"{SLUG}_FINAL.xlsx"

    wb = load_workbook(xlsx, data_only=True)

    land_ws = wb["Landscape"] if "Landscape" in wb.sheetnames else wb[wb.sheetnames[0]]

    land_rows = list(land_ws.iter_rows(values_only=True))

    land_hdr = [str(h) for h in land_rows[0]]

    landscape: list[dict] = []

    for r in land_rows[1:]:

        if not r:

            continue

        d = {

            land_hdr[i]: ("" if r[i] is None else r[i])

            for i in range(len(land_hdr))

            if i < len(r)

        }

        if str(d.get("Company") or "").strip():

            landscape.append(d)



    details_by: dict[str, dict] = {}

    if "Company Details" in wb.sheetnames:

        ws = wb["Company Details"]

        rows = list(ws.iter_rows(values_only=True))

        hdr = [str(h) for h in rows[0]]

        for r in rows[1:]:

            if not r:

                continue

            d = {

                hdr[i]: ("" if r[i] is None else str(r[i]))

                for i in range(len(hdr))

                if i < len(r)

            }

            key = _norm(d.get("Brand") or d.get("Company") or "")

            if key:

                details_by[key] = d



    report = []

    brand_n = marketer_n = 0

    for row in landscape:

        name = str(row.get("Company") or "").strip()

        det = details_by.get(_norm(name))

        brand = str((det or {}).get("Brand") or name)

        role = classify_role(brand, name)

        row["Distribution Type"] = role

        row["Role"] = role

        if det:

            for col in ("Quadrant", "X", "Y", "Overall"):

                if det.get(col) not in (None, ""):

                    row[col] = det[col]

                    if col == "X":

                        row["X Score"] = det[col]

                    elif col == "Y":

                        row["Y Score"] = det[col]

                    elif col == "Overall":

                        row["Overall Score"] = det[col]

            if det.get("Found in"):

                row["Headquarters"] = det["Found in"]

        if role == "Brand":

            brand_n += 1

        else:

            marketer_n += 1

        report.append({"brand": brand, "company": name, "role": role})



    audit_path = OUT / "chatgpt_expand_batch_all_audit.json"

    audit: dict = {}

    if audit_path.exists():

        audit = json.loads(audit_path.read_text(encoding="utf-8")).get("xy_scoring") or {}



    detail_rows = to_company_detail_rows(landscape, QUERY, audit)

    by_brand = {_norm(r["brand"]): r["role"] for r in report}

    by_co = {_norm(r["company"]): r["role"] for r in report}

    for d in detail_rows:

        role = by_brand.get(_norm(d.get("Brand") or "")) or by_co.get(

            _norm(d.get("Company") or "")

        )

        if role in ("Brand", "Marketer"):

            d["Role"] = role



    write_final_xlsx(

        xlsx,

        landscape,

        "Companies",

        {"query": QUERY, "xy_scoring": audit, "role_split": True},

        detail_rows=detail_rows,

    )

    extras = export_expand_quadrant_outputs(

        OUT, detail_rows, QUERY, country="global", audit=audit, chart_n=20

    )

    audit_out = ROOT / "output" / "chatgpt_expand" / "_audit" / f"{SLUG}_roles.json"

    audit_out.parent.mkdir(exist_ok=True)

    audit_out.write_text(

        json.dumps(

            {"brand": brand_n, "marketer": marketer_n, "rows": report},

            indent=2,

            ensure_ascii=False,

        ),

        encoding="utf-8",

    )

    print(f"Brand={brand_n} Marketer={marketer_n}")

    print(f"html -> {extras.get('html')}")

    print("Marketers:")

    for r in report:

        if r["role"] == "Marketer":

            print(f"  - {r['brand']}")





if __name__ == "__main__":

    main()

