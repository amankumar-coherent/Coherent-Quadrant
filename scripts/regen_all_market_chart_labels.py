#!/usr/bin/env python3
"""Regenerate all 5 market HTML reports with company-name-only chart labels."""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import export_expand_quadrant_outputs

MARKETS = [
    (
        "global_flexible_packaging_market_global",
        "Global Flexible Packaging Market",
        "Solution Capability",
        "Business Strategy",
    ),
    (
        "global_glp_1_receptor_agonist_market_global",
        "Global GLP-1 Receptor Agonist Market",
        "Therapeutic & Manufacturing Capability",
        "Commercial Reach & Pipeline Strategy",
    ),
    (
        "global_liquefied_natural_gas_market_global",
        "Global Liquefied Natural Gas Market",
        "Solution Capability",
        "Business Strategy",
    ),
    (
        "global_semiconductor_market_global",
        "Global Semiconductor Market",
        "Solution Capability",
        "Business Strategy",
    ),
    (
        "global_wearable_medical_devices_market_global",
        "Global Wearable Medical Devices Market",
        "Solution Capability",
        "Business Strategy",
    ),
]


def _load_details(folder: Path, slug: str) -> list[dict]:
    xlsx = folder / f"{slug}_FINAL.xlsx"
    wb = load_workbook(xlsx, data_only=True)
    if "Company Details" not in wb.sheetnames:
        raise SystemExit(f"missing Company Details: {xlsx}")
    rows = list(wb["Company Details"].iter_rows(values_only=True))
    hdr = [str(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not r or not r[0]:
            continue
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr))}
        out.append(
            {
                "Brand": d.get("Brand") or "",
                "Company": d.get("Company") or d.get("Brand") or "",
                "Role": d.get("Role") or "",
                "Quadrant": d.get("Quadrant") or "",
                "X": d.get("X"),
                "Y": d.get("Y"),
                "Overall": d.get("Overall"),
                "Found in": d.get("Found in") or "",
            }
        )
    return out


def main() -> int:
    base = ROOT / "output" / "chatgpt_expand"
    for slug, query, axis_x, axis_y in MARKETS:
        folder = base / slug
        details = _load_details(folder, slug)
        audit = {
            "axis_x": axis_x,
            "axis_y": axis_y,
            "method": "regenerate HTML — chart labels = company/brand name only",
        }
        export_expand_quadrant_outputs(
            folder, details, query, country="global", audit=audit, chart_n=20
        )
        html = folder / f"{slug}_report.html"
        print(f"OK {slug} n={len(details)} -> {html.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
