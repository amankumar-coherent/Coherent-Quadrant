"""Verify + fix Global Liquefied Natural Gas FINAL: ownership, brands, roles, duplicates.

Web-verified Aug 2026. Updates Landscape Ownership / names / roles, rebuilds
Company Details + quadrant HTML.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx

SLUG = "global_liquefied_natural_gas_market_global"
QUERY = "Global Liquefied Natural Gas Market"
OUT = ROOT / "output" / "chatgpt_expand" / SLUG
XLSX = OUT / f"{SLUG}_FINAL.xlsx"


def _norm(s: str) -> str:
    s = str(s or "").strip().lower()
    s = (
        s.replace("ö", "o")
        .replace("ü", "u")
        .replace("ä", "a")
        .replace("ß", "ss")
        .replace("÷", "o")
        .replace("ń", "n")
    )
    s = re.sub(r"[+/|&,_.:\-]+", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def _stem(s: str) -> str:
    s = _norm(s)
    s = re.sub(r"\([^)]*\)", " ", s)
    s = re.sub(
        r"\b(inc|llc|ltd|limited|plc|gmbh|ag|sa|co|corp|corporation|group|holdings?)\b",
        " ",
        s,
    )
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


# Off-market / defunct / duplicate shells (not active independent LNG players).
# Use exact normalized names (or safe unique substrings) — never bare "lng limited"
# as a substring (would wipe Petronet LNG Limited / Pakistan LNG Limited).
REMOVE_EXACT = {
    "jordan cove lng",
    "rasgas",
    "qatar gas",
    "prelude flng shell operated",
    "prelude flng",
    "lng limited",  # Australian LNGL after Magnolia sale — exact only
    "zeta energy",
    "fox petroleum",
    "indra gas",
    "zodiac energy ltd",
    "zodiac energy",
    "goldeneye lng",
    "stx lng",
}
REMOVE_STEMS: set[str] = set()  # kept for report compatibility; prefer REMOVE_EXACT

# Canonical rename: match key -> new Company display name
RENAME = {
    "apache corporation": "APA Corporation",
    "apache": "APA Corporation",
    "chesapeake energy": "Expand Energy Corporation",
    "mubadala petroleum": "Mubadala Energy",
    "hoegh lng holdings ltd": "Höegh Evi (formerly Höegh LNG)",
    "hoegh lng holdings": "Höegh Evi (formerly Höegh LNG)",
    "hoegh lng": "Höegh Evi (formerly Höegh LNG)",
    "noble group": "Noble Resources",
    "sk e s": "SK Innovation E&S (formerly SK E&S)",
    "sk e&s": "SK Innovation E&S (formerly SK E&S)",
    "polskie lng": "GAZ-SYSTEM (Polskie LNG / Świnoujście)",
    "china gas holdings limited": "China Gas Holdings Limited",
    "china gas holdings": "China Gas Holdings Limited",
    "exxonmobil corporation": "ExxonMobil",
    "exxonmobil": "ExxonMobil",
    "china national offshore oil corporation cnooc": "CNOOC",
    "china national offshore oil corporation": "CNOOC",
    "cnooc": "CNOOC",
    "adnoc": "ADNOC",
    "qatarenergy lng": "QatarEnergy LNG",
    "wintershall dea ag": "Wintershall Dea AG",
    "sempra infrastructure": "Sempra Infrastructure Partners",
    "tellurian inc": "Tellurian Inc.",
    "pavilion energy": "Pavilion Energy",
    "anadarko petroleum corporation": "Anadarko Petroleum Corporation",
    "magnolia lng": "Magnolia LNG",
    "senex energy": "Senex Energy",
    "endesa": "Endesa",
    "kunlun energy": "Kunlun Energy",
    "equinor uk limited": "Equinor (UK) Limited",
    "vng handel vertrieb gmbh": "VNG Handel & Vertrieb GmbH",
    "petrochina international london ltd": "PetroChina International (London) Ltd",
    "ptt lng company limited": "PTT LNG Company Limited",
    "ptt global lng pttgl": "PTT Global LNG (PTTGL)",
    "adnoc gas": "ADNOC Gas",
    "jera co inc": "JERA Co., Inc.",
    "uniper se": "Uniper SE",
    "fortisbc": "FortisBC",
}

# Prefer keeping these when deduping (stem of keep-name / family id)
DEDUP_FAMILY = {
    "exxonmobil": "exxonmobil",
    "cnooc": "cnooc",
    "china national offshore oil": "cnooc",
    "china gas holdings": "china gas holdings limited",
    "qatarenergy lng": "qatarenergy lng",
    "qatar gas": "qatarenergy lng",
    "rasgas": "qatarenergy lng",
    "adnoc gas": "adnoc gas",
    "adnoc": "adnoc",
    "petrochina international": "petrochina",
    "petrochina": "petrochina",
    "apache": "apa corporation",
    "apa corporation": "apa corporation",
    "chesapeake": "expand energy corporation",
    "expand energy": "expand energy corporation",
    "sk e s": "sk innovation e s formerly sk e s",
    "sk innovation e s": "sk innovation e s formerly sk e s",
    "hoegh": "hoegh evi formerly hoegh lng",
    "noble": "noble resources",
    "sempra infrastructure": "sempra infrastructure partners",
    "polskie lng": "gaz system polskie lng swinoujscie",
    "gaz system": "gaz system polskie lng swinoujscie",
}

# Ownership overrides: stem -> (Ownership text, ownership_relation)
# Ownership text must start with acquired by / subsidiary of / merged into / owned by
OWNERSHIP: dict[str, tuple[str, str]] = {
    # Completed third-party acquisitions
    "anadarko petroleum": ("acquired by Occidental Petroleum Corporation", "acquired_by"),
    "pavilion energy": ("acquired by Shell plc", "acquired_by"),
    "tellurian": ("acquired by Woodside Energy Group Ltd", "acquired_by"),
    "noble resources": ("acquired by Vitol", "acquired_by"),
    "senex energy": (
        "acquired by POSCO International (50.1%) and Hancock Energy (49.9%)",
        "acquired_by",
    ),
    "magnolia lng": ("acquired by Glenfarne Group", "acquired_by"),
    "wintershall dea": (
        "acquired by Harbour Energy plc (E&P portfolio excl. Russia, Sep 2024)",
        "acquired_by",
    ),
    # Holding-company reorg / rename — NOT a third-party acquisition.
    # Clear false "acquired by APA" from prior fill; APA is the public parent.
    "apa corporation": ("Public (holding-company reorg of Apache Corporation, Mar 2021)", ""),
    # Mergers / brand consolidations
    "sk innovation e s": ("merged into SK Innovation", "merged_into"),
    "sk e s": ("merged into SK Innovation", "merged_into"),
    "gaz system": (
        "Public (state-owned Polish gas TSO; absorbed Polskie LNG Mar 2021)",
        "",
    ),
    # Subsidiaries / group companies
    "endesa": ("subsidiary of Enel S.p.A. (~70%)", "subsidiary_of"),
    "kunlun energy": ("subsidiary of PetroChina / CNPC", "subsidiary_of"),
    "mubadala energy": ("subsidiary of Mubadala Investment Company", "subsidiary_of"),
    "equinor uk": ("subsidiary of Equinor ASA", "subsidiary_of"),
    "equinor limited": ("subsidiary of Equinor ASA", "subsidiary_of"),
    "vng handel": ("subsidiary of VNG AG", "subsidiary_of"),
    "petrochina international": ("subsidiary of PetroChina", "subsidiary_of"),
    "ptt lng": ("subsidiary of PTT Public Company Limited", "subsidiary_of"),
    "ptt global lng": ("subsidiary of PTT Public Company Limited", "subsidiary_of"),
    "adnoc gas": ("subsidiary of ADNOC", "subsidiary_of"),
    "fortisbc": ("subsidiary of Fortis Inc.", "subsidiary_of"),
    "petrovietnam gas": ("subsidiary of Petrovietnam", "subsidiary_of"),
    "hoegh evi": (
        "owned by Aequitas Limited (Höegh family) 50% and Igneo Infrastructure Partners 50% "
        "(Igneo acquired MSIP stake; close expected H1 2025)",
        "owned_by",
    ),
    "jera": (
        "owned by TEPCO Fuel & Power and Chubu Electric Power (50/50 JV)",
        "owned_by",
    ),
    # Sempra Infrastructure: NOT "acquired by Sempra" — Sempra platform with minority partners;
    # KKR-led controlling stake sale announced Sep 2025, expected close Q2–Q3 2026 (pending).
    "sempra infrastructure": (
        "owned by Sempra (~70%), KKR (~20%), ADIA (~10%); "
        "KKR-led consortium 65% sale announced Sep 2025 (close pending 2026)",
        "owned_by",
    ),
    # State / public control (Company Details shows owned-by when relation set)
    "uniper": ("owned by Federal Republic of Germany (~99.12%)", "owned_by"),
}

# Role overrides — LNG landscape uses Brand / Marketer; pure commodity traders = Marketer
ROLE: dict[str, str] = {
    "vitol": "Marketer",
    "gunvor": "Marketer",
    "trafigura": "Marketer",
    "mercuria": "Marketer",
    "glencore": "Marketer",
}


def _match_key(name: str, table: dict) -> str | None:
    # Strip parentheticals so prior annotations do not match the wrong parent key.
    n = _norm(re.sub(r"\([^)]*\)", " ", str(name or "")))
    n = re.sub(r"\s+", " ", n).strip()
    best = None
    best_len = -1
    for k in table:
        kn = _norm(k)
        if kn in n and len(kn) > best_len:
            best = k
            best_len = len(kn)
    return best


def _family_key(name: str) -> str:
    n = _norm(name)
    for fam, group in sorted(DEDUP_FAMILY.items(), key=lambda x: -len(x[0])):
        if fam in n:
            return group
    return _stem(name)


def main() -> None:
    wb = load_workbook(XLSX, data_only=True)
    land_ws = wb["Landscape"]
    land_rows = list(land_ws.iter_rows(values_only=True))
    hdr = [str(h) for h in land_rows[0]]
    landscape: list[dict] = []
    for r in land_rows[1:]:
        if not r:
            continue
        d = {
            hdr[i]: ("" if r[i] is None else r[i])
            for i in range(len(hdr))
            if i < len(r)
        }
        if str(d.get("Company") or "").strip():
            landscape.append(d)

    report: dict = {
        "removed": [],
        "renamed": [],
        "ownership_fixed": [],
        "role_fixed": [],
        "deduped": [],
        "before": len(landscape),
        "sources_checked": [
            "Shell completes Pavilion Energy acquisition (Apr 2025)",
            "Woodside completes Tellurian acquisition (Oct 2024)",
            "Vitol completes Noble Resources acquisition (Jan 2025)",
            "Anadarko / Occidental (2019)",
            "Apache → APA holding-company reorg (Mar 2021) — not a third-party acquisition",
            "Sempra Infrastructure: Sempra/KKR/ADIA; KKR 65% sale announced Sep 2025, close pending 2026",
            "Wintershall Dea E&P → Harbour Energy (Sep 2024)",
            "SK E&S merged into SK Innovation (Nov 2024)",
            "Senex → POSCO International / Hancock (Apr 2022)",
            "Magnolia LNG → Glenfarne (2020)",
            "RasGas merged into Qatargas (2018) → QatarEnergy LNG",
            "Höegh LNG → Höegh Evi; Aequitas/Igneo ownership",
            "Chesapeake + Southwestern → Expand Energy (Oct 2024)",
            "Polskie LNG merged into GAZ-SYSTEM (Mar 2021)",
            "Endesa ~70% Enel; Kunlun → PetroChina/CNPC; Uniper ~99% Germany",
            "Jordan Cove cancelled by Pembina (2021)",
        ],
    }

    # 1) Remove off-market (exact normalized name match)
    kept: list[dict] = []
    for row in landscape:
        name = str(row.get("Company") or "")
        nk = _norm(name)
        nk_noparen = _norm(re.sub(r"\([^)]*\)", " ", name))
        drop_reason = None
        if nk in REMOVE_EXACT or nk_noparen in REMOVE_EXACT:
            drop_reason = nk if nk in REMOVE_EXACT else nk_noparen
        else:
            for stem in REMOVE_EXACT:
                # Allow unique multi-word stems only when they are the whole name prefix
                if len(stem.split()) >= 2 and (nk == stem or nk_noparen == stem or nk.startswith(stem + " ")):
                    # still block "lng limited" substring false-positives
                    if stem == "lng limited" and nk != "lng limited" and nk_noparen != "lng limited":
                        continue
                    drop_reason = stem
                    break
        if drop_reason:
            report["removed"].append({"brand": name, "reason": f"off_market:{drop_reason}"})
        else:
            kept.append(row)
    landscape = kept

    # 1b) Restore critical LNG players accidentally dropped by prior runs
    ck_path = OUT / "chatgpt_checkpoint_batch_all.json"
    report.setdefault("restored", [])
    if ck_path.exists():
        ck = json.loads(ck_path.read_text(encoding="utf-8"))
        xy_rows = (ck.get("data") or {}).get("xy_rows") or []
        have = {_norm(str(r.get("Company") or "")) for r in landscape}

        def _already(nn: str) -> bool:
            needles = [n for n in ("petronet lng limited", "pakistan lng limited") if n in nn]
            if not needles:
                return True
            for h in have:
                if any(n in h for n in needles):
                    return True
            return False

        for r in xy_rows:
            nm = str(r.get("Company") or "").strip()
            nn = _norm(nm)
            if "petronet lng limited" not in nn and "pakistan lng limited" not in nn:
                continue
            if _already(nn):
                continue
            row = dict(r)
            row["Role"] = str(row.get("Role") or row.get("Distribution Type") or "Brand")
            row["Distribution Type"] = row["Role"]
            landscape.append(row)
            have.add(nn)
            report["restored"].append(nm)

    # 2) Rename + ownership + role
    for row in landscape:
        name = str(row.get("Company") or "").strip()
        rk = _match_key(name, RENAME)
        if rk:
            new_name = RENAME[rk]
            if new_name != name:
                report["renamed"].append({"from": name, "to": new_name})
                row["Company"] = new_name
                name = new_name

        ok = _match_key(name, OWNERSHIP)
        if ok:
            own, rel = OWNERSHIP[ok]
            old = str(row.get("Ownership") or "")
            if old != own:
                report["ownership_fixed"].append(
                    {"brand": name, "from": old, "to": own, "relation": rel}
                )
            row["Ownership"] = own
            row["ownership_relation"] = rel
            row["parent_owner"] = ""
            if rel:
                m = re.match(
                    r"(?i)^\s*(?:acquired by|subsidiary of|merged into|owned by)\s+(.+?)\s*$",
                    own,
                )
                if m:
                    row["parent_owner"] = m.group(1).strip()
            else:
                row.pop("parent_owner", None)
                row["ownership_relation"] = ""

        role_k = _match_key(name, ROLE)
        if role_k:
            new_role = ROLE[role_k]
            old_role = str(row.get("Role") or row.get("Distribution Type") or "")
            if old_role != new_role:
                report["role_fixed"].append(
                    {"brand": name, "from": old_role, "to": new_role}
                )
            row["Role"] = new_role
            row["Distribution Type"] = new_role
        else:
            cur = str(row.get("Distribution Type") or row.get("Role") or "Brand")
            if cur not in ("Brand", "Marketer"):
                cur = "Brand"
            row["Role"] = cur
            row["Distribution Type"] = cur

    # 3) Deduplicate families
    by_fam: dict[str, list[dict]] = {}
    for row in landscape:
        fam = _family_key(str(row.get("Company") or ""))
        by_fam.setdefault(fam, []).append(row)

    final_rows: list[dict] = []
    for fam, rows in by_fam.items():
        if len(rows) == 1:
            final_rows.append(rows[0])
            continue
        prefer = fam
        scored = []
        for r in rows:
            nm = _norm(str(r.get("Company") or ""))
            score = 0
            if prefer and prefer == nm:
                score += 200
            elif prefer and prefer in nm:
                score += 100
            # Prefer shorter canonical brand when both match family
            score -= min(len(nm), 80) // 5
            own = str(r.get("Ownership") or "").lower()
            if "subsidiary of" in own and prefer in nm and "subsidiary" not in prefer:
                score -= 20
            if any(x in own for x in ("acquired", "subsidiary", "owned by", "merged")):
                score += 5
            try:
                score += int(float(r.get("Overall Score") or 0))
            except (TypeError, ValueError):
                pass
            scored.append((score, r))
        scored.sort(key=lambda x: -x[0])
        chosen = scored[0][1]
        for r in rows:
            if r is not chosen:
                report["deduped"].append(
                    {
                        "kept": str(chosen.get("Company")),
                        "dropped": str(r.get("Company")),
                        "family": fam,
                    }
                )
        final_rows.append(chosen)

    landscape = final_rows

    for row in landscape:
        blob = _norm(str(row.get("Company") or ""))
        for stem, role in ROLE.items():
            if stem in blob:
                row["Role"] = role
                row["Distribution Type"] = role

    brand_n = sum(1 for r in landscape if str(r.get("Role") or "") == "Brand")
    marketer_n = sum(1 for r in landscape if str(r.get("Role") or "") == "Marketer")

    audit_path = OUT / "chatgpt_expand_batch_all_audit.json"
    audit: dict = {}
    if audit_path.exists():
        audit = json.loads(audit_path.read_text(encoding="utf-8")).get("xy_scoring") or {}

    detail_rows = to_company_detail_rows(landscape, QUERY, audit)
    by_co = {
        _norm(str(r.get("Company") or "")): str(r.get("Role") or "Brand")
        for r in landscape
    }
    for d in detail_rows:
        role = by_co.get(_norm(d.get("Brand") or "")) or by_co.get(
            _norm(d.get("Company") or "")
        )
        if role in ("Brand", "Marketer"):
            d["Role"] = role

    write_final_xlsx(
        XLSX,
        landscape,
        "Companies",
        {
            "query": QUERY,
            "xy_scoring": audit,
            "role_split": True,
            "ownership_verify": "2026-08-14",
        },
        detail_rows=detail_rows,
    )
    extras = export_expand_quadrant_outputs(
        OUT, detail_rows, QUERY, country="global", audit=audit, chart_n=20
    )

    report["after"] = len(landscape)
    report["brand"] = brand_n
    report["marketer"] = marketer_n
    report["html"] = str(extras.get("html") or "")

    audit_out = (
        ROOT
        / "output"
        / "chatgpt_expand"
        / "_audit"
        / f"{SLUG}_ownership_verify.json"
    )
    audit_out.parent.mkdir(exist_ok=True)
    audit_out.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")

    lines = ["Brand\tCompany\tRole\tFound in"]
    for d in detail_rows:
        lines.append(
            f"{d.get('Brand')}\t{d.get('Company')}\t{d.get('Role')}\t{d.get('Found in')}"
        )
    (
        ROOT
        / "output"
        / "chatgpt_expand"
        / "_audit"
        / f"{SLUG}_verified_rows.tsv"
    ).write_text("\n".join(lines), encoding="utf-8")

    # Markdown summary of annotated Company columns
    acq_lines = [
        "# LNG Brand / Company ownership verification",
        "",
        f"Total companies: **{len(detail_rows)}**",
        "",
        "| Brand | Company | Role |",
        "|---|---|---|",
    ]
    for d in detail_rows:
        co = str(d.get("Company") or "")
        if co.startswith("(") or co != str(d.get("Brand") or ""):
            if co.startswith("(") or "acquir" in co.lower() or "subsidiar" in co.lower() or "merged" in co.lower() or "owned by" in co.lower():
                acq_lines.append(
                    f"| {d.get('Brand')} | {co} | {d.get('Role')} |"
                )
    md_path = (
        ROOT
        / "output"
        / "chatgpt_expand"
        / "_audit"
        / f"{SLUG}_brand_company_ownership.md"
    )
    md_path.write_text("\n".join(acq_lines) + "\n", encoding="utf-8")

    print(
        f"before={report['before']} after={report['after']} "
        f"Brand={brand_n} Marketer={marketer_n}"
    )
    print(
        f"removed={len(report['removed'])} renamed={len(report['renamed'])} "
        f"ownership={len(report['ownership_fixed'])} role={len(report['role_fixed'])} "
        f"deduped={len(report['deduped'])} restored={len(report.get('restored') or [])}"
    )
    print(f"audit -> {audit_out}")
    print(f"html -> {extras.get('html')}")
    for x in report["removed"]:
        print("REMOVED:", x["brand"], x["reason"])
    for x in report.get("restored") or []:
        print("RESTORED:", x)
    for x in report["deduped"]:
        print("DEDUP:", x["dropped"], "->", x["kept"])
    for x in report["ownership_fixed"]:
        print("OWN:", x["brand"])
    for x in report["role_fixed"]:
        print("ROLE:", x["brand"], x["from"], "->", x["to"])


if __name__ == "__main__":
    main()
