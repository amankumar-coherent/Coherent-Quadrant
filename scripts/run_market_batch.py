#!/usr/bin/env python3
"""Run several markets through the expand pipeline SEQUENTIALLY.

Google AI Mode is one browser, one query at a time, and Google rate-limits by
IP — two pipelines running at once is the direct cause of repeated CAPTCHA
throttling. So markets are never run in parallel here, no matter how many are
queued.

Each market is a separate subprocess, so a crash or a hard CAPTCHA block in
one market does not kill the queue: it is recorded and the next market
starts. Every market checkpoints independently, so re-running this script
resumes each one where it stopped rather than starting over.

  .\\.venv\\Scripts\\python.exe scripts\\run_market_batch.py --markets-file queries\\batch.txt
  .\\.venv\\Scripts\\python.exe scripts\\run_market_batch.py -q "Global Polyurethane Market" -q "Global Solar Rooftop Market"
"""
from __future__ import annotations

import argparse
import datetime
import subprocess
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "src") not in sys.path:
    sys.path.insert(0, str(ROOT / "src"))


def _log(msg: str) -> None:
    stamp = datetime.datetime.now().strftime("%H:%M:%S")
    print(f"[batch {stamp}] {msg}", flush=True)


def _company_count(market: str, country: str) -> int:
    """Rows in the finished export, for the summary line."""
    try:
        from openpyxl import load_workbook

        from vendor_intel.pipeline.web_expand import default_output_dir, resolve_final_path

        out_dir = default_output_dir(market, country)
        xlsx = resolve_final_path(market, country, out_dir, None)
        if not Path(xlsx).exists():
            return 0
        wb = load_workbook(xlsx, read_only=True)
        sheet = "Company Details" if "Company Details" in wb.sheetnames else wb.sheetnames[0]
        n = max(0, wb[sheet].max_row - 1)
        wb.close()
        return n
    except Exception:  # noqa: BLE001
        return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--markets-file", help="one market name per line; # comments allowed")
    ap.add_argument("--query", "-q", action="append", default=[], help="repeatable")
    ap.add_argument("--country", "-c", default="global")
    ap.add_argument("--target", "-t", type=int, default=400)
    ap.add_argument(
        "--gap-seconds",
        type=int,
        default=120,
        help="cool-off between markets so the IP is not hammered back-to-back",
    )
    args = ap.parse_args()

    markets: list[str] = list(args.query)
    if args.markets_file:
        for line in Path(args.markets_file).read_text(encoding="utf-8").splitlines():
            name = line.split("#", 1)[0].strip()
            if name:
                markets.append(name)
    markets = list(dict.fromkeys(markets))  # de-dupe, keep order
    if not markets:
        ap.print_help()
        return 2

    _log(f"{len(markets)} market(s) queued, target={args.target} each")
    for i, m in enumerate(markets, 1):
        _log(f"  {i}. {m}")

    results: list[tuple[str, str, int, float]] = []
    for i, market in enumerate(markets, 1):
        _log("")
        _log(f"=== [{i}/{len(markets)}] START {market} ===")
        started = time.monotonic()
        cmd = [
            sys.executable,
            str(ROOT / "scripts" / "run_chatgpt_expand.py"),
            "--query",
            market,
            "--country",
            args.country,
            "--target",
            str(args.target),
        ]
        try:
            proc = subprocess.run(cmd, cwd=str(ROOT))
            status = "ok" if proc.returncode == 0 else f"exit {proc.returncode}"
        except KeyboardInterrupt:
            _log("interrupted by user — stopping the queue (checkpoints are saved)")
            results.append((market, "interrupted", _company_count(market, args.country),
                            time.monotonic() - started))
            break
        except Exception as err:  # noqa: BLE001
            status = f"error: {type(err).__name__}"

        mins = (time.monotonic() - started) / 60.0
        rows = _company_count(market, args.country)
        results.append((market, status, rows, mins))
        _log(f"=== [{i}/{len(markets)}] DONE {market}: {status}, {rows} companies, {mins:.0f} min ===")

        if i < len(markets) and args.gap_seconds > 0:
            _log(f"cooling off {args.gap_seconds}s before the next market…")
            time.sleep(args.gap_seconds)

    _log("")
    _log("=== BATCH SUMMARY ===")
    for market, status, rows, mins in results:
        _log(f"  {status:12} {rows:4} companies  {mins:5.0f} min  {market}")
    failed = [r for r in results if r[1] != "ok"]
    _log(f"{len(results) - len(failed)}/{len(results)} completed cleanly")
    if failed:
        _log("re-run this script to resume the unfinished markets from their checkpoints")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
