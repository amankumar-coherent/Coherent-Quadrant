#!/usr/bin/env python3
"""Verify + fix Brand / Company / Role for Global GLP-1 Receptor Agonist FINAL.

Rules (consumer_brand market):
  Brand   = plain GLP-1 brand owner / developer name
  Company = (acquired by|subsidiary of|merged into Parent) when owned;
            else Company = Brand
  Role    = Brand

Drops: regional arms of majors, pharmacies already gated, off-market (non-GLP-1),
generic-only houses without originator GLP-1 assets, junk placeholders, dupes.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx

SLUG = "global_glp_1_receptor_agonist_market_global"
QUERY = "Global GLP-1 Receptor Agonist Market"
OUT = ROOT / "output" / "chatgpt_expand" / SLUG
XLSX = OUT / f"{SLUG}_FINAL.xlsx"
AUDIT = ROOT / "output" / "chatgpt_expand" / "_audit"

# Canonical keepers → (display Brand, Key Brands Represented, HQ override optional)
KEEPERS: dict[str, tuple[str, str]] = {
    "novo nordisk": (
        "Novo Nordisk",
        "Ozempic; Wegovy; Rybelsus; Victoza; Saxenda; Xultophy",
    ),
    "eli lilly": (
        "Eli Lilly and Company",
        "Mounjaro; Zepbound; Trulicity; Foundayo (orforglipron)",
    ),
    "boehringer ingelheim": (
        "Boehringer Ingelheim",
        "Survodutide (with Zealand Pharma)",
    ),
    "zealand pharma": (
        "Zealand Pharma",
        "Survodutide (with Boehringer); petrelintide",
    ),
    "amgen": ("Amgen", "MariTide (maridebart cafraglutide)"),
    "viking therapeutics": ("Viking Therapeutics", "VK2735"),
    "structure therapeutics": (
        "Structure Therapeutics",
        "GSBR-1290 / aleniglipron",
    ),
    "altimmune": ("Altimmune", "Pemvidutide"),
    "innovent": ("Innovent Biologics", "Mazdutide"),
    "hansoh": ("Hansoh Pharmaceutical Group", "Fulaimei (loxenatide); HS-20094"),
    "jiangsu hansoh": (
        "Hansoh Pharmaceutical Group",
        "Fulaimei (loxenatide); HS-20094",
    ),
    "hengrui": ("Jiangsu Hengrui Medicine", "HRS9531"),
    "cspc": ("CSPC Pharmaceutical Group", "SYH2086 (oral GLP-1; licensed out)"),
    "benemae": ("Shanghai Benemae Pharmaceutical", "GLP-1 pipeline (China)"),
    "roche": ("Roche", "CT-388; CT-996; CT-868 (via Carmot)"),
    "carmot": (
        "Carmot Therapeutics",
        "CT-388; CT-996; CT-868",
    ),
    "hanmi": (
        "Hanmi Pharmaceutical",
        "Efpeglenatide (Efe); HM15275",
    ),
    "astrazeneca": (
        "AstraZeneca",
        "Byetta; Bydureon (discontinued US 2024); pipeline incretins",
    ),
    "sanofi": (
        "Sanofi",
        "Soliqua (lixisenatide+insulin); Adlyxin/Lyxumia discontinued",
    ),
    "pfizer": (
        "Pfizer",
        "Danuglipron (oral GLP-1; program discontinued)",
    ),
    "merck & co": (
        "Merck & Co.",
        "Metabolic / incretin pipeline",
    ),
    "merck sharp": (
        "Merck & Co.",
        "Metabolic / incretin pipeline",
    ),
    "adocia": ("Adocia", "BioChaperone GLP-1 programs"),
    "lg chem": ("LG Chem", "GLP-1 / metabolic pipeline"),
    "tonghua dongbao": ("Tonghua Dongbao", "GLP-1 / diabetes pipeline (China)"),
    "gan & lee": ("Gan & Lee Pharmaceuticals", "GLP-1 / diabetes pipeline (China)"),
    "yuhan": ("Yuhan Corporation", "GLP-1 / metabolic pipeline"),
    "dong-a st": ("Dong-A ST", "GLP-1 / metabolic pipeline"),
    "dong a st": ("Dong-A ST", "GLP-1 / metabolic pipeline"),
}

# Ownership annotations (stem → ownership text, relation, optional year)
OWNERSHIP: dict[str, tuple[str, str, str]] = {
    "carmot": ("acquired by Roche", "acquired_by", "2024"),
    "eli lilly do brasil": ("subsidiary of Eli Lilly and Company", "subsidiary_of", ""),
    "laboratorio chile": ("subsidiary of Teva Pharmaceutical Industries", "subsidiary_of", ""),
    "ratiopharm": ("acquired by Teva Pharmaceutical Industries", "acquired_by", "2010"),
    "hexal": ("subsidiary of Sandoz", "subsidiary_of", ""),
    "mylan": ("merged into Viatris", "merged_into", "2020"),
    "provention": ("acquired by Sanofi", "acquired_by", "2023"),
    "intarcia": ("assets acquired by i2o Therapeutics", "acquired_by", "2023"),
    "sanofi medley": ("subsidiary of Sanofi (EMS acquisition pending 2026)", "subsidiary_of", ""),
}

# Canonical HQ (City, Country) for keepers when fill data is garbage
HQ_FIX: dict[str, str] = {
    "novo nordisk": "Bagsværd, Denmark",
    "eli lilly and company": "Indianapolis, United States",
    "boehringer ingelheim": "Ingelheim am Rhein, Germany",
    "astrazeneca": "Cambridge, United Kingdom",
    "zealand pharma": "Søborg, Denmark",
    "viking therapeutics": "San Diego, United States",
    "cspc pharmaceutical group": "Shijiazhuang, China",
    "innovent biologics": "Suzhou, China",
    "merck & co.": "Rahway, United States",
    "hansoh pharmaceutical group": "Lianyungang, China",
    "jiangsu hengrui medicine": "Lianyungang, China",
    "amgen": "Thousand Oaks, United States",
    "hanmi pharmaceutical": "Seoul, South Korea",
    "shanghai benemae pharmaceutical": "Shanghai, China",
    "pfizer": "New York, United States",
    "roche": "Basel, Switzerland",
    "adocia": "Lyon, France",
    "tonghua dongbao": "Tonghua, China",
    "yuhan corporation": "Seoul, South Korea",
    "carmot therapeutics": "Berkeley, United States",
    "lg chem": "Seoul, South Korea",
    "sanofi": "Paris, France",
    "dong-a st": "Seoul, South Korea",
    "gan & lee pharmaceuticals": "Beijing, China",
    "structure therapeutics": "San Francisco, United States",
    "altimmune": "Gaithersburg, United States",
}

# Explicit drops even if name might fuzzy-match a keeper stem incorrectly
DROP_STEMS: dict[str, str] = {
    "provention": "not GLP-1 (Tzield / T1D immunotherapy; acquired by Sanofi)",
    "mannkind": "not GLP-1 (Afrezza inhaled insulin)",
    "gilead": "not a GLP-1 brand owner",
    "intarcia": "program failed FDA; assets sold to i2o — not active GLP-1 marketer",
    "moksha8": "LATAM specialty partner, not GLP-1 brand owner",
    "kallyope": "gut-brain platform; not commercial GLP-1 brand owner",
    "bristol-myers": "no commercial GLP-1 brand / pipeline in landscape",
    "bristol myers": "no commercial GLP-1 brand / pipeline in landscape",
    "johnson & johnson": "no commercial GLP-1 brand owner role",
    "glaxosmithkline": "Tanzeum discontinued 2017; not active GLP-1 player",
    "novartis": "no originator GLP-1 brand (Sandoz generics separate)",
    "sandoz": "generics arm — not GLP-1 originator",
    "hexal": "Sandoz/Novartis generics subsidiary",
    "ratiopharm": "Teva generics subsidiary — not GLP-1 originator",
    "viatris": "generics (ex-Mylan) — not GLP-1 originator",
    "mylan": "merged into Viatris — not GLP-1 originator",
    "teva": "generics / liraglutide AG only — not originator brand owner",
    "laboratorio chile": "Teva regional generics subsidiary",
    "fresenius": "not GLP-1 brand owner",
    "bausch": "not GLP-1 brand owner",
    "stada": "generics — not GLP-1 originator",
    "hikma": "liraglutide AG only — not originator",
    "merck kgaa": "German Merck — not US Merck GLP-1 pipeline",
    "sumitomo": "no confirmed GLP-1 brand owner role",
    "takeda": "no confirmed commercial GLP-1 brand",
    "crystalgenomics": "not verified GLP-1 brand owner",
    "farmacêutica brasileira": "placeholder / not a real company name",
    "farmaceutica brasileira": "placeholder / not a real company name",
    "laboratório teuto": "generics — not GLP-1 originator",
    "laboratorio teuto": "generics — not GLP-1 originator",
    "prati-donaduzzi": "generics — not GLP-1 originator",
    "prati donaduzzi": "generics — not GLP-1 originator",
    "eurofarma": "generics — not GLP-1 originator",
    "ache laboratorio": "generics — not GLP-1 originator",
    "aché": "generics — not GLP-1 originator",
    "blau": "generics — not GLP-1 originator",
    "cristalia": "generics — not GLP-1 originator",
    "cristália": "generics — not GLP-1 originator",
    "libbs": "generics — not GLP-1 originator",
    "hypera": "not GLP-1 brand owner",
    "ems": "generics — not GLP-1 originator",
    "laboratório farmacêutico ems": "generics — not GLP-1 originator",
    "nordic pharma": "not GLP-1 brand owner",
    "tecnoquimicas": "not GLP-1 brand owner",
    "tecnoquímicas": "not GLP-1 brand owner",
    "laboratorios bago": "not GLP-1 brand owner",
    "laboratorios silanes": "not GLP-1 brand owner",
    "productos cientificos": "not GLP-1 brand owner",
    "productos científicos": "not GLP-1 brand owner",
    "quimica ariston": "not GLP-1 brand owner",
    "química ariston": "not GLP-1 brand owner",
    "julphar": "not verified GLP-1 brand owner",
    "aspen": "not verified GLP-1 brand owner",
    "biocon": "insulin/biosimilars — not GLP-1 originator brand",
    "biomm": "not verified GLP-1 brand owner",
    "servier": "not GLP-1 brand owner",
    "gedeon richter": "not GLP-1 brand owner",
    "aurobindo": "generics — not GLP-1 originator",
    "emcure": "generics — not GLP-1 originator",
    "alkem": "generics — not GLP-1 originator",
    "cipla": "generics — not GLP-1 originator",
    "intas": "generics — not GLP-1 originator",
    "lupin": "generics — not GLP-1 originator",
    "mankind": "generics — not GLP-1 originator",
    "sun pharmaceutical": "generics — not GLP-1 originator",
    "dr reddy": "generics — not GLP-1 originator",
    "glenmark": "generics — not GLP-1 originator",
    "zydus": "generics — not GLP-1 originator",
    "wockhardt": "generics — not GLP-1 originator",
    "usv": "generics — not GLP-1 originator",
    "micro labs": "generics — not GLP-1 originator",
    "fdc": "generics — not GLP-1 originator",
    "pharmascience": "generics — not GLP-1 originator",
    "sanofi medley": "Brazilian generics unit — not GLP-1 brand owner",
}

# Country / regional tokens that mark subsidiary arms of a parent already kept
_REGIONAL_MARKERS = (
    "australia",
    "new zealand",
    "canada",
    "brazil",
    "brasil",
    "china",
    "japan",
    "thailand",
    "india",
    "south africa",
    "middle east",
    "uk",
    "ltd",
    "gmbh",
    "k.k",
    "kk",
    "inc",
    "pty",
    "ulc",
    "lp",
    "pharma ltd",
    "pharmaceuticals lp",
    "do brasil",
    "south africa",
    "us",
    "usa",
)


def _norm(s: str) -> str:
    s = str(s or "").strip().lower()
    s = (
        s.replace("ö", "o")
        .replace("ü", "u")
        .replace("ä", "a")
        .replace("ß", "ss")
        .replace("é", "e")
        .replace("á", "a")
        .replace("í", "i")
        .replace("ó", "o")
        .replace("ú", "u")
        .replace("ç", "c")
        .replace("ã", "a")
        .replace("ê", "e")
    )
    s = re.sub(r"\([^)]*\)", " ", s)
    s = re.sub(r"[^a-z0-9&]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _match(name: str, table: dict) -> str | None:
    n = _norm(name)
    best = None
    best_len = -1
    for k in table:
        kn = _norm(k)
        if not kn:
            continue
        if n == kn or n.startswith(kn + " ") or kn in n:
            if len(kn) > best_len:
                best = k
                best_len = len(kn)
    return best


def _is_regional_of_keeper(name: str, keeper_stem: str) -> bool:
    """True when name is a geo/legal arm of a keeper, not the parent itself."""
    n = _norm(name)
    k = _norm(keeper_stem)
    # Exact / near-exact parent names are keepers
    parent_aliases = {
        "novo nordisk": {"novo nordisk"},
        "eli lilly": {"eli lilly", "eli lilly and company"},
        "boehringer ingelheim": {"boehringer ingelheim"},
        "amgen": {"amgen"},
        "astrazeneca": {"astrazeneca"},
        "sanofi": {"sanofi"},
        "pfizer": {"pfizer"},
        "roche": {"roche"},
        "merck & co": {"merck & co", "merck and co"},
    }
    aliases = parent_aliases.get(k, {k})
    if n in aliases:
        return False
    # Strip common legal suffixes and compare
    plain = re.sub(
        r"\b(inc|ltd|llc|gmbh|ag|sa|plc|co|corp|corporation|limited|company|"
        r"pharmaceuticals?|pharma|group)\b\.?",
        "",
        n,
    )
    plain = re.sub(r"\s+", " ", plain).strip(" ,.")
    if plain in aliases or plain == k:
        # Still regional if geo token remains in original name
        if any(m in n for m in _REGIONAL_MARKERS) and plain in aliases:
            # "Eli Lilly and Company" has "company" stripped → "eli lilly and" — handle
            pass
    # Explicit parent legal forms to KEEP
    keep_exact = {
        "novo nordisk",
        "eli lilly and company",
        "eli lilly",
        "boehringer ingelheim",
        "amgen",
        "astrazeneca",
        "sanofi",
        "pfizer",
        "roche",
        "merck & co",
        "merck and co",
    }
    if n in keep_exact:
        return False
    # BI operating entities / geos — keep one global parent only
    if k == "boehringer ingelheim" and n != "boehringer ingelheim":
        return True
    if k.startswith("eli lilly") or k == "eli lilly":
        if n not in {"eli lilly", "eli lilly and company"}:
            return True
    if k == "novo nordisk" and n != "novo nordisk":
        return True
    if k == "amgen" and n != "amgen":
        return True
    if k == "astrazeneca" and n != "astrazeneca":
        return True
    if k == "sanofi" and n not in {"sanofi"}:
        return True
    if k == "pfizer" and n != "pfizer":
        return True
    if k == "roche" and n != "roche":
        return True
    if "merck" in k and n not in {"merck & co", "merck and co", "merck & co."}:
        # MSD Australia etc.
        if "australia" in n or "south africa" in n or "sharp" in n:
            return True
    return False


def main() -> None:
    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Landscape"]
    rows_raw = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows_raw[0]]

    landscape: list[dict] = []
    removed: list[dict] = []
    ownership_fixed: list[dict] = []
    renamed: list[dict] = []
    brands_filled: list[dict] = []
    seen_keeper: set[str] = set()

    for r in rows_raw[1:]:
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        name = str(d.get("Company") or "").strip()
        if not name:
            continue
        # Company column sometimes holds only "(acquired by …)" — recover Brand
        if name.lower().startswith("(") and "acquired" in name.lower():
            brand_fallback = str(d.get("Brand") or "").strip()
            if brand_fallback and not brand_fallback.lower().startswith("("):
                name = brand_fallback
                d["Company"] = name

        drop = _match(name, DROP_STEMS)
        if drop:
            removed.append({"brand": name, "reason": DROP_STEMS[drop]})
            continue

        keeper = _match(name, KEEPERS)
        if not keeper:
            removed.append({"brand": name, "reason": "not a verified GLP-1 brand owner / developer"})
            continue

        if _is_regional_of_keeper(name, keeper):
            removed.append(
                {
                    "brand": name,
                    "reason": f"regional/legal arm of {KEEPERS[keeper][0]} (parent kept)",
                }
            )
            continue

        display, key_brands = KEEPERS[keeper]
        canon_key = _norm(display)
        if canon_key in seen_keeper:
            removed.append({"brand": name, "reason": f"duplicate of {display}"})
            continue
        seen_keeper.add(canon_key)

        if name != display:
            renamed.append({"from": name, "to": display})
            d["Company"] = display
            name = display

        # Ownership
        own_key = _match(name, OWNERSHIP) or _match(display, OWNERSHIP)
        if own_key:
            own, rel, year = OWNERSHIP[own_key]
            old = str(d.get("Ownership") or "")
            if old != own:
                ownership_fixed.append({"brand": name, "from": old, "to": own})
            d["Ownership"] = own
            d["ownership_relation"] = rel
            d["ownership_year"] = year
            m = re.match(
                r"(?i)^\s*(?:acquired by|subsidiary of|merged into|owned by)\s+(.+?)\s*$",
                own,
            )
            d["parent_owner"] = m.group(1).strip() if m else ""
        else:
            d["Ownership"] = "Independent"
            d["ownership_relation"] = ""
            d["ownership_year"] = ""
            d.pop("parent_owner", None)

        hq_key = _match(name, HQ_FIX) or _match(display, HQ_FIX)
        if hq_key:
            d["Headquarters"] = HQ_FIX[hq_key]

        # Role + distribution
        d["Role"] = "Brand"
        d["Distribution Type"] = "Brand"

        # Key brands
        old_b = str(d.get("Key Brands Represented") or "").strip()
        if old_b.lower() in {"", "not publicly disclosed", "n/a", "na", "-"} or old_b != key_brands:
            brands_filled.append({"brand": name, "brands": key_brands})
            d["Key Brands Represented"] = key_brands

        d["Specialty Focus"] = (
            str(d.get("Specialty Focus") or "").strip()
            or f"GLP-1 / incretin therapies — {key_brands}"
        )
        d["Core Categories"] = "Healthcare / Pharmaceuticals"
        d["Industry Category"] = "Pharmaceuticals"
        d["Retail / E-commerce / Both"] = "No"

        landscape.append(d)

    audit: dict = {}
    ap = OUT / "chatgpt_expand_batch_all_audit.json"
    if ap.exists():
        audit = json.loads(ap.read_text(encoding="utf-8")).get("xy_scoring") or {}

    detail_rows = to_company_detail_rows(landscape, QUERY, audit)

    by_land = {_norm(str(r.get("Company") or "")): r for r in landscape}
    fixed_company: list[dict] = []
    for drow in detail_rows:
        brand = str(drow.get("Brand") or "").strip()
        land = by_land.get(_norm(brand)) or {}
        parent_owner = str(land.get("parent_owner") or "").strip()
        relation = str(land.get("ownership_relation") or "").strip() or "acquired_by"
        year = str(land.get("ownership_year") or "").strip()
        if parent_owner:
            from vendor_intel.quadrant.brand_meta import format_acquired_suffix

            company_col = format_acquired_suffix(
                parent_owner, relation=relation, year=year
            )
        else:
            company_col = brand
        old_co = str(drow.get("Company") or "")
        if old_co != company_col:
            fixed_company.append({"brand": brand, "from": old_co, "to": company_col})
        drow["Brand"] = brand
        drow["Company"] = company_col
        drow["Role"] = "Brand"
        hq = str(land.get("Headquarters") or "").strip()
        if hq:
            drow["Found in"] = hq

    write_final_xlsx(
        XLSX,
        landscape,
        "Companies",
        {
            "query": QUERY,
            "brand_company_verify": "2026-08-14",
            "kept": len(landscape),
            "removed": len(removed),
            "ownership_fixed": len(ownership_fixed),
            "company_col_fixed": len(fixed_company),
            "xy_scoring": audit,
        },
        detail_rows=detail_rows,
    )
    extras = export_expand_quadrant_outputs(
        OUT, detail_rows, QUERY, country="global", audit=audit, chart_n=20
    )

    # Sync CSV
    csv_path = OUT / f"{SLUG}_companies.csv"
    import csv

    with csv_path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(
            f,
            fieldnames=["Brand", "Company", "Role", "Quadrant", "X", "Y", "Overall", "Found in"],
        )
        w.writeheader()
        for drow in detail_rows:
            w.writerow({k: drow.get(k, "") for k in w.fieldnames})

    annotated = [
        {"brand": d.get("Brand"), "company": d.get("Company"), "role": d.get("Role")}
        for d in detail_rows
        if str(d.get("Company") or "").startswith("(")
    ]
    report = {
        "total": len(detail_rows),
        "annotated": len(annotated),
        "removed_n": len(removed),
        "removed": removed,
        "renamed": renamed,
        "ownership_fixed": ownership_fixed,
        "company_col_changed_n": len(fixed_company),
        "brands_filled_n": len(brands_filled),
        "annotated_rows": annotated,
        "html": str(extras.get("html") or ""),
    }
    AUDIT.mkdir(parents=True, exist_ok=True)
    (AUDIT / f"{SLUG}_brand_company_verify.json").write_text(
        json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    lines = ["Brand\tCompany\tRole\tKey Brands"]
    for r in landscape:
        lines.append(
            f"{r.get('Company')}\t"
            f"{next((d.get('Company') for d in detail_rows if _norm(d.get('Brand')) == _norm(r.get('Company'))), '')}\t"
            f"{r.get('Role')}\t{r.get('Key Brands Represented')}"
        )
    (AUDIT / f"{SLUG}_brand_company_after.tsv").write_text(
        "\n".join(lines), encoding="utf-8"
    )

    print(
        f"kept={report['total']} removed={len(removed)} "
        f"annotated={report['annotated']} renamed={len(renamed)} "
        f"ownership_fixed={len(ownership_fixed)}"
    )
    print("--- KEPT ---")
    for d in detail_rows:
        print(f"  {d.get('Brand')} | {d.get('Company')} | {d.get('Role')}")
    print("--- REMOVED (sample) ---")
    for x in removed[:40]:
        print(f"  DROP: {x['brand']} — {x['reason']}")
    if len(removed) > 40:
        print(f"  ... +{len(removed) - 40} more")


if __name__ == "__main__":
    main()
