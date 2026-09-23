#!/usr/bin/env python3
"""Re-score one or more markets with market-specific X/Y axes + 5×5 parameters.

Forces a fresh XY pass (clears prior X/Y), keeps companies/roles/Found in, rebuilds
FINAL Excel + HTML so scorecards show the market parameters used for scoring.

Usage:
  .\\.venv\\Scripts\\python.exe scripts\\rescore_market_axes.py --market "<Market name>"
  .\\.venv\\Scripts\\python.exe scripts\\rescore_market_axes.py --market "<A>" --market "<B>" --crawl
"""
from __future__ import annotations

import argparse
import asyncio
import json
import os
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "src") not in sys.path:
    sys.path.insert(0, str(ROOT / "src"))

if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    score_expand_rows,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx

OUT = ROOT / "output" / "chatgpt_expand"



def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def _load_final(xlsx: Path) -> tuple[list[dict], dict[str, dict]]:
    wb = load_workbook(xlsx, data_only=True)
    landscape: list[dict] = []
    land = "Landscape" if "Landscape" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[land]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows[0]]
    for r in rows[1:]:
        if not r:
            continue
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        if str(d.get("Company") or d.get("Brand") or "").strip():
            landscape.append(d)

    details_by: dict[str, dict] = {}
    if "Company Details" in wb.sheetnames:
        ws = wb["Company Details"]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(h) for h in rows[0]]
        for r in rows[1:]:
            if not r:
                continue
            d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
            key = _norm(d.get("Brand") or d.get("Company") or "")
            if key:
                details_by[key] = d

    for row in landscape:
        key = _norm(row.get("Company") or row.get("Brand") or "")
        det = details_by.get(key)
        if not det:
            for dk, dv in details_by.items():
                if dk in key or key in dk:
                    det = dv
                    break
        if not det:
            continue
        # Prefer Brand as scoring identity when present
        brand = str(det.get("Brand") or "").strip()
        if brand:
            row["Company"] = brand
            row["Brand"] = brand
        if det.get("Role"):
            row["Distribution Type"] = det["Role"]
            row["Role"] = det["Role"]
        if det.get("Found in"):
            row["Headquarters"] = det["Found in"]
            row["Found in"] = det["Found in"]
        # Preserve ownership display for re-apply after score
        company_col = str(det.get("Company") or "").strip()
        if company_col and company_col != brand:
            row["_display_company"] = company_col
            if company_col.startswith("(") or "acquired" in company_col.lower():
                row["Ownership"] = company_col.strip("()")
        # Clear scores so score_expand_rows re-scores every row
        for col in (
            "X Score",
            "Y Score",
            "Overall Score",
            "X",
            "Y",
            "Overall",
            "Quadrant",
        ):
            row[col] = ""
    return landscape, details_by


def _apply_display_company(detail_rows: list[dict], landscape: list[dict]) -> None:
    by_brand = {_norm(r.get("Company") or r.get("Brand") or ""): r for r in landscape}
    for det in detail_rows:
        key = _norm(det.get("Brand") or det.get("Company") or "")
        src = by_brand.get(key)
        if not src:
            continue
        disp = str(src.get("_display_company") or "").strip()
        if disp:
            det["Company"] = disp


async def rescore_one(slug: str, query: str, *, skip_crawl: bool) -> dict:
    folder = OUT / slug
    xlsx = folder / f"{slug}_FINAL.xlsx"
    if not xlsx.exists():
        return {"slug": slug, "skipped": True, "reason": "no FINAL.xlsx"}

    print(f"\n=== {query} ===", flush=True)
    print(f"  file: {xlsx}", flush=True)
    landscape, _ = _load_final(xlsx)
    print(f"  companies: {len(landscape)} (forcing XY rescore)", flush=True)

    if skip_crawl:
        os.environ["EXPAND_XY_SKIP_CRAWL"] = "1"
    os.environ.setdefault("EXPAND_MARKET_AXIS_LLM", "1")

    scored, audit = await score_expand_rows(
        landscape, query, country="global", concurrent=int(os.getenv("EXPAND_XY_CONCURRENT") or "4")
    )
    detail_rows = to_company_detail_rows(scored, query, audit)
    _apply_display_company(detail_rows, scored)

    write_final_xlsx(
        xlsx,
        scored,
        "Companies",
        {
            "query": query,
            "xy_scoring": {k: v for k, v in (audit or {}).items() if k != "rows"},
            "market_axis_rescore": True,
        },
        detail_rows=detail_rows,
    )
    extras = export_expand_quadrant_outputs(
        folder, detail_rows, query, country="global", audit=audit, chart_n=20
    )

    # Persist full xy audit (incl. parameters) for UI/debug
    (folder / "chatgpt_xy_scores_batch_all.json").write_text(
        json.dumps(audit, indent=2, ensure_ascii=False, default=str),
        encoding="utf-8",
    )
    summary = {
        "slug": slug,
        "query": query,
        "n": len(detail_rows),
        "axis_x": audit.get("axis_x"),
        "axis_y": audit.get("axis_y"),
        "x_features": audit.get("x_features"),
        "y_features": audit.get("y_features"),
        "axis_definition_method": audit.get("axis_definition_method"),
        "html": extras.get("html"),
    }
    print(
        f"  axes: {summary['axis_x']} / {summary['axis_y']} "
        f"({summary.get('axis_definition_method')})",
        flush=True,
    )
    print(f"  X params: {summary['x_features']}", flush=True)
    print(f"  Y params: {summary['y_features']}", flush=True)
    print(f"  html: {summary['html']}", flush=True)
    return summary


async def amain(args: argparse.Namespace) -> int:
    try:
        from vendor_intel.placeholders.load_keys import apply_env_overrides

        apply_env_overrides()
    except Exception:
        pass
    os.environ.setdefault("EXPAND_MARKET_AXIS_LLM", "1")
    from vendor_intel.pipeline.web_expand import default_output_dir

    if not args.skip_crawl:
        os.environ["EXPAND_XY_SKIP_CRAWL"] = "0"
    results = []
    for query in args.market:
        slug = default_output_dir(query, args.country).name
        try:
            results.append(await rescore_one(slug, query, skip_crawl=args.skip_crawl))
        except Exception as exc:
            print(f"  FAILED {slug}: {type(exc).__name__}: {exc}", flush=True)
            results.append({"slug": slug, "error": f"{type(exc).__name__}: {exc}"})

    audit_path = OUT / "_audit" / "market_axis_rescore.json"
    audit_path.parent.mkdir(exist_ok=True)
    audit_path.write_text(json.dumps(results, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"\nDone. Audit → {audit_path}", flush=True)
    return 0


def main() -> int:
    p = argparse.ArgumentParser(description="Rescore markets with market-wise X/Y parameters")
    p.add_argument("--market", action="append", required=True,
                   help="market name to rescore (repeat for several markets)")
    p.add_argument("--country", default="global")
    p.add_argument(
        "--skip-crawl",
        action="store_true",
        default=True,
        help="Skip deep crawl (default on) — score from existing KB/model knowledge",
    )
    p.add_argument("--crawl", action="store_true", help="Enable deep crawl during rescore")
    args = p.parse_args()
    if args.crawl:
        args.skip_crawl = False
    return asyncio.run(amain(args))


if __name__ == "__main__":
    raise SystemExit(main())
