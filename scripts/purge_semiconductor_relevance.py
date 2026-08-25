"""Purge audited off-market / duplicate rows from Semiconductor FINAL Excel."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx

OUT = ROOT / "output" / "chatgpt_expand"
SLUG = "global_semiconductor_market_global"
QUERY = "Global Semiconductor Market"

# Exact / substring match keys (lowercase). Matched against Brand|Company.
SEMI_REMOVE = {
    # Wrong vertical
    "mitel networks",
    "weg s.a.",
    "weg s.a",
    "embraer",
    "positivo tecnologia",
    "nortech systems",
    "cascodium",
    "semicab",
    "smar equipamentos industriais",
    "ael sistemas",
    "ciena corporation",
    "siemens ag",
    "raytheon australia",
    "centro de tecnologia da informação renato archer",
    "centro de tecnologia da informacao renato archer",
    "renato archer",
    "instituto de pesquisas eldorado",
    "c-dac (",
    "c-dac",
    "algeria semiconductor",
    "egyptian semiconductor",
    "tunisia semiconductor",
    "moroccan semiconductor",
    "ethiopia semiconductor",
    "south african microelectronics",
    "nanolithography solutions",
    "sahasra electronics",
    # Duplicates / regional / acquired
    "globalfoundries (already listed)",
    "intel (australia)",
    "intel india",
    "intel mexico",
    "alphawave ip (canada)",
    "semtech corporation (canadian",
    "lumentum (canada)",
    "huawei technologies - chile",
    "skyworks solutions - mexico",
    "cypress semiconductor",
    "dialog semiconductor",
    "dsp group",
    "xilinx",
    "mellanox",
    "habana labs",
    "sierra wireless",
}


def _norm(s: str) -> str:
    s = str(s or "").strip().lower()
    s = re.sub(r"\s+", " ", s)
    return s


def _identity(row: dict) -> str:
    return _norm(row.get("Brand") or row.get("Company") or row.get("name") or "")


def _should_remove(ident: str, needles: set[str]) -> str | None:
    for n in needles:
        if not n:
            continue
        if ident == n or n in ident or ident.startswith(n):
            return n
    return None


def _load(xlsx: Path) -> tuple[list[dict], dict[str, dict]]:
    wb = load_workbook(xlsx, data_only=True)
    landscape: list[dict] = []
    land_ws = wb["Landscape"] if "Landscape" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(land_ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows[0]]
    for r in rows[1:]:
        if not r:
            continue
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        if str(d.get("Company") or "").strip():
            landscape.append(d)

    details_by: dict[str, dict] = {}
    if "Company Details" in wb.sheetnames:
        ws = wb["Company Details"]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(h) for h in rows[0]]
        for r in rows[1:]:
            if not r:
                continue
            d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
            key = _norm(d.get("Brand") or d.get("Company") or "")
            if key:
                details_by[key] = d

    for row in landscape:
        key = _norm(row.get("Company") or "")
        det = details_by.get(key)
        if not det:
            for dk, dv in details_by.items():
                if dk in key or key in dk:
                    det = dv
                    break
        if det:
            for col in ("Quadrant", "X", "Y", "Overall", "Role"):
                if det.get(col) not in (None, ""):
                    row[col] = det[col]
                    if col == "X":
                        row["X Score"] = det[col]
                    elif col == "Y":
                        row["Y Score"] = det[col]
                    elif col == "Overall":
                        row["Overall Score"] = det[col]
            if det.get("Role"):
                row["Distribution Type"] = det["Role"]
            if det.get("Found in"):
                row["Headquarters"] = det["Found in"]
    return landscape, details_by


def _purge(landscape: list[dict], needles: set[str]) -> tuple[list[dict], list[tuple[str, str]]]:
    kept: list[dict] = []
    removed: list[tuple[str, str]] = []
    for row in landscape:
        ident = _identity(row)
        hit = _should_remove(ident, needles)
        if hit:
            removed.append((ident, hit))
            continue
        kept.append(row)
    return kept, removed


def main() -> None:
    folder = OUT / SLUG
    xlsx = folder / f"{SLUG}_FINAL.xlsx"
    landscape, _ = _load(xlsx)
    before = len(landscape)
    landscape, removed = _purge(landscape, SEMI_REMOVE)
    after = len(landscape)

    audit_path = folder / "chatgpt_expand_batch_all_audit.json"
    audit: dict = {}
    if audit_path.exists():
        audit = json.loads(audit_path.read_text(encoding="utf-8")).get("xy_scoring") or {}

    detail_rows = to_company_detail_rows(landscape, QUERY, audit)
    write_final_xlsx(
        xlsx,
        landscape,
        "Companies",
        {"query": QUERY, "xy_scoring": audit, "purged": [r[0] for r in removed]},
        detail_rows=detail_rows,
    )
    extras = export_expand_quadrant_outputs(
        folder, detail_rows, QUERY, country="global", audit=audit, chart_n=20
    )

    audit_out = OUT / "_audit" / f"{SLUG}_purge.json"
    audit_out.parent.mkdir(exist_ok=True)
    audit_out.write_text(
        json.dumps(
            {
                "before": before,
                "after": after,
                "removed": [{"name": n, "matched": m} for n, m in removed],
                "html": str(extras.get("html") or ""),
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    print(f"{SLUG}: {before} → {after} (removed {len(removed)})")
    for n, m in removed:
        print(f"  - {n}  [{m}]")
    print(f"  html -> {extras.get('html')}")


if __name__ == "__main__":
    main()
