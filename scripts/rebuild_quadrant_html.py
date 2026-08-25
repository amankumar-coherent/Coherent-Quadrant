"""Rebuild HTML + Company Details Excel with Found in (location) column."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    DETAIL_COLUMNS,
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx

ROOT = Path(r"D:\Coherent-Quadrant\output\chatgpt_expand")
MARKETS = [
    (
        "global_flexible_packaging_market_global",
        "Global Flexible Packaging Market",
    ),
    (
        "global_wearable_medical_devices_market_global",
        "Global Wearable Medical Devices Market",
    ),
]


def _load(xlsx: Path) -> tuple[list[dict], dict]:
    wb = load_workbook(xlsx, data_only=True)
    landscape: list[dict] = []
    land_name = "Landscape" if "Landscape" in wb.sheetnames else wb.sheetnames[-1]
    if "Landscape" in wb.sheetnames or land_name:
        ws = wb["Landscape"] if "Landscape" in wb.sheetnames else wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(h) for h in rows[0]]
        for r in rows[1:]:
            if not r:
                continue
            d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
            if str(d.get("Company") or "").strip():
                landscape.append(d)

    details_by: dict[str, dict] = {}
    if "Company Details" in wb.sheetnames:
        ws = wb["Company Details"]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(h) for h in rows[0]]
        for r in rows[1:]:
            if not r:
                continue
            d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
            key = str(d.get("Brand") or d.get("Company") or "").strip().lower()
            if key:
                details_by[key] = d

    audit = {}
    if "Web Expand Audit" in wb.sheetnames:
        pass
    # Merge scores from Company Details onto landscape for rebuild
    for row in landscape:
        key = str(row.get("Company") or "").strip().lower()
        det = details_by.get(key)
        if not det:
            # try brand match
            for dk, dv in details_by.items():
                if dk in key or key in dk:
                    det = dv
                    break
        if det:
            for col in ("Quadrant", "X", "Y", "Overall", "Role"):
                if det.get(col) not in (None, ""):
                    if col in ("X", "Y", "Overall"):
                        row[f"{col} Score" if col != "Overall" else "Overall Score"] = det[col]
                        row[col] = det[col]
                    else:
                        row[col] = det[col]
            if det.get("Role"):
                row["Distribution Type"] = det["Role"]
    return landscape, details_by


def main() -> None:
    print("DETAIL_COLUMNS", DETAIL_COLUMNS)
    for slug, query in MARKETS:
        folder = ROOT / slug
        xlsx = folder / f"{slug}_FINAL.xlsx"
        if not xlsx.exists():
            print("MISSING", xlsx)
            continue
        landscape, _ = _load(xlsx)
        audit_path = folder / "chatgpt_expand_batch_all_audit.json"
        audit = {}
        if audit_path.exists():
            audit = json.loads(audit_path.read_text(encoding="utf-8")).get("xy_scoring") or {}
        detail_rows = to_company_detail_rows(landscape, query, audit)
        # ensure Found in is location
        for r in detail_rows:
            loc = str(r.get("Found in") or "")
            if loc.isdigit() and len(loc) == 4:
                r["Found in"] = str((r.get("_meta") or {}).get("hq_location") or "")
        write_final_xlsx(
            xlsx,
            landscape,
            "Companies",
            {"query": query, "xy_scoring": audit},
            detail_rows=detail_rows,
        )
        extras = export_expand_quadrant_outputs(
            folder, detail_rows, query, country="global", audit=audit, chart_n=20
        )
        sample = detail_rows[:3]
        print(f"{slug}: rows={len(detail_rows)} Found in samples:")
        for s in sample:
            print(f"  {s.get('Brand')}: {s.get('Found in')!r}")
        print(f"  html -> {extras['html']}")


if __name__ == "__main__":
    main()
