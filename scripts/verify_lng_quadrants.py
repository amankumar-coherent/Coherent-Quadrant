"""Audit LNG FINAL Quadrant vs X/Y using Coherent definitions (no re-score).

X = Solution Capability, Y = Business Strategy
Leaders:            high X + high Y
Challengers:        low X  + high Y  (high Business Strategy, lower Solution Capability)
Trailblazers:       high X + low Y   (high Solution Capability, lower Business Strategy)
Emerging Players:   low X  + low Y
"""
from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx
from vendor_intel.quadrant.rating_map import assign_quadrants_absolute_median

SLUG = "global_liquefied_natural_gas_market_global"
QUERY = "Global Liquefied Natural Gas Market"
OUT = ROOT / "output" / "chatgpt_expand" / SLUG
XLSX = OUT / f"{SLUG}_FINAL.xlsx"
AUDIT = ROOT / "output" / "chatgpt_expand" / "_audit"


def _int(v) -> int:
    try:
        return int(float(str(v or "0").strip() or "0"))
    except Exception:
        return 0


def main() -> None:
    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Landscape"]
    rows_raw = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows_raw[0]]
    landscape: list[dict] = []
    for r in rows_raw[1:]:
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        if str(d.get("Company") or "").strip():
            landscape.append(d)

    xs = [_int(r.get("X Score") or r.get("X")) for r in landscape]
    ys = [_int(r.get("Y Score") or r.get("Y")) for r in landscape]
    expected, mid_x, mid_y = assign_quadrants_absolute_median(xs, ys)

    mismatches: list[dict] = []
    before = Counter()
    after = Counter()
    for i, row in enumerate(landscape):
        old = str(row.get("Quadrant") or "").strip() or "?"
        new = expected[i]
        before[old] += 1
        after[new] += 1
        name = str(row.get("Company") or "")
        if old != new:
            mismatches.append(
                {
                    "company": name,
                    "x": xs[i],
                    "y": ys[i],
                    "from": old,
                    "to": new,
                    "mid_x": mid_x,
                    "mid_y": mid_y,
                }
            )
        row["Quadrant"] = new
        row["X Score"] = str(xs[i])
        row["Y Score"] = str(ys[i])
        overall = _int(row.get("Overall Score") or row.get("Overall"))
        if not overall:
            overall = int(round((xs[i] + ys[i]) / 2.0))
        row["Overall Score"] = str(overall)

    audit = {}
    ap = OUT / "chatgpt_expand_batch_all_audit.json"
    if ap.exists():
        audit = json.loads(ap.read_text(encoding="utf-8")).get("xy_scoring") or {}
    # Prefer full XY audit (axes + parameter definitions) so HTML end-section stays filled
    xy_path = OUT / "chatgpt_xy_scores_batch_all.json"
    if xy_path.exists():
        try:
            xy = json.loads(xy_path.read_text(encoding="utf-8"))
            for k in (
                "industry_group",
                "industry_category",
                "axis_x",
                "axis_y",
                "x_features",
                "y_features",
                "parameter_definitions",
            ):
                if xy.get(k):
                    audit[k] = xy[k]
        except Exception:
            pass
    audit = {
        **audit,
        "quadrant_mid_x": mid_x,
        "quadrant_mid_y": mid_y,
        "quadrant_verify": "2026-08-14_definitions",
        "quadrant_mismatches_fixed": len(mismatches),
    }

    detail_rows = to_company_detail_rows(landscape, QUERY, audit)
    # Ensure detail Quadrant matches Landscape after reassignment
    by = {str(r.get("Company") or "").strip().lower(): r for r in landscape}
    for d in detail_rows:
        brand = str(d.get("Brand") or "").strip()
        land = by.get(brand.lower())
        if not land:
            # brand may differ from landscape Company; try exact Brand match via Company field
            for lr in landscape:
                if str(lr.get("Company") or "").strip().lower() == brand.lower():
                    land = lr
                    break
        if land:
            d["Quadrant"] = str(land.get("Quadrant") or d.get("Quadrant") or "")
            d["X"] = _int(land.get("X Score"))
            d["Y"] = _int(land.get("Y Score"))
            d["Overall"] = _int(land.get("Overall Score"))

    write_final_xlsx(
        XLSX,
        landscape,
        "Companies",
        {
            "query": QUERY,
            "quadrant_verify": "2026-08-14",
            "quadrant_mid_x": mid_x,
            "quadrant_mid_y": mid_y,
            "mismatches_fixed": len(mismatches),
            "xy_scoring": audit,
        },
        detail_rows=detail_rows,
    )
    extras = export_expand_quadrant_outputs(
        OUT, detail_rows, QUERY, country="global", audit=audit, chart_n=20
    )

    # Re-verify Company Details consistency
    wb2 = load_workbook(XLSX, data_only=True)
    det = list(wb2["Company Details"].iter_rows(values_only=True))
    dh = [str(h or "") for h in det[0]]
    still_bad = []
    dxs = [_int(r[dh.index("X")]) for r in det[1:]]
    dys = [_int(r[dh.index("Y")]) for r in det[1:]]
    exp2, mx2, my2 = assign_quadrants_absolute_median(dxs, dys)
    for i, r in enumerate(det[1:]):
        d = {dh[j]: ("" if r[j] is None else str(r[j])) for j in range(min(len(dh), len(r)))}
        if d.get("Quadrant") != exp2[i]:
            still_bad.append(
                {
                    "brand": d.get("Brand"),
                    "x": dxs[i],
                    "y": dys[i],
                    "got": d.get("Quadrant"),
                    "expected": exp2[i],
                }
            )

    report = {
        "definitions": {
            "X": "Solution Capability",
            "Y": "Business Strategy",
            "Leaders": "high X + high Y",
            "Challengers": "low X + high Y (high Business Strategy, lower Solution Capability)",
            "Trailblazers": "high X + low Y (high Solution Capability, lower Business Strategy)",
            "Emerging Players": "low X + low Y",
            "split": "cohort absolute median",
            "mid_x": mid_x,
            "mid_y": mid_y,
        },
        "before_counts": dict(before),
        "after_counts": dict(after),
        "mismatches_fixed": mismatches,
        "still_bad_after_write": still_bad,
        "html": str(extras.get("html") or ""),
        "total": len(detail_rows),
    }
    AUDIT.mkdir(parents=True, exist_ok=True)
    outp = AUDIT / f"{SLUG}_quadrant_verify.json"
    outp.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(
        f"total={len(landscape)} mid_x={mid_x} mid_y={mid_y} "
        f"mismatches={len(mismatches)} still_bad={len(still_bad)}"
    )
    print("before", dict(before))
    print("after", dict(after))
    for m in mismatches[:40]:
        print(
            f"FIX {m['company']}: {m['from']} -> {m['to']} "
            f"(X={m['x']} Y={m['y']})"
        )
    if len(mismatches) > 40:
        print(f"... +{len(mismatches)-40} more")


if __name__ == "__main__":
    main()
