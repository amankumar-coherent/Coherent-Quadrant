"""Fix Company column from Brand + verified ownership for Flexible Packaging FINAL.

Schema:
  Brand   = identity / trading name
  Company = Brand when independent; else (acquired by X) / (subsidiary of X)

Updates Landscape Ownership, rebuilds Company Details via brand_display_fields.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx

SLUG = "global_flexible_packaging_market_global"
QUERY = "Global Flexible Packaging Market"
OUT = ROOT / "output" / "chatgpt_expand" / SLUG
XLSX = OUT / f"{SLUG}_FINAL.xlsx"

# Brand-stem → Ownership text (must start with acquired by|subsidiary of|owned by|merged into)
# Verified Aug 2026
OWNERSHIP_BY_STEM: list[tuple[str, str, str]] = [
    # relation hint stored separately when needed
    # (stem, ownership_text, optional brand_rename)
    ("aluflexpack", "acquired by Constantia Flexibles", ""),
    ("bemis", "acquired by Amcor plc", ""),
    ("berry global", "acquired by Amcor plc", ""),
    ("ampac holdings", "acquired by ProAmpac LLC", ""),
    ("aep industries", "acquired by Berry Global; ultimate parent Amcor plc", ""),
    ("rpc bpi", "acquired by Berry Global; ultimate parent Amcor plc", ""),
    ("rpc group", "acquired by Berry Global; ultimate parent Amcor plc", ""),
    ("nordfolien", "acquired by RPC Group; now part of Amcor via Berry", ""),
    ("innovia", "acquired by CCL Industries", ""),
    ("pactiv evergreen", "acquired by Novolex Holdings, LLC", ""),
    ("selig", "acquired by CC Industries", ""),
    ("walki", "acquired by Oji Holdings Corporation", ""),
    ("alucoat", "acquired by Grupo Alibérico", ""),
    ("al khaleej", "acquired by Taghleef Industries", ""),
    # Treofan Germany → Al Ghurair / Polyopt (Jan 2026) — NOT Jindal anymore
    ("treofan", "acquired by Al Ghurair Group (now Polyopt GmbH)", "Treofan Group (now Polyopt)"),
    # EPL
    (
        "epl limited",
        "acquired by Blackstone (Indorama Ventures ~24.9% minority; Indorama majority merger pending)",
        "",
    ),
    ("essel", "acquired by Blackstone (Indorama Ventures ~24.9% minority; Indorama majority merger pending)", ""),
    # ePac acquired by Butterfly Equity
    ("epac", "acquired by Butterfly Equity", ""),
    # Subsidiaries / group members
    ("cosmo films", "subsidiary of Cosmo First Limited", ""),
    ("wipak", "subsidiary of Wihuri Group", ""),
    ("winpak", "subsidiary of Wihuri Group", ""),
    ("al watania plastics", "subsidiary of Al Watania for Industries", ""),
    ("al ghurair packaging", "subsidiary of Al Ghurair Group", ""),
    ("taghleef", "subsidiary of Al Ghurair Group", ""),
    ("arabian flexible packaging", "subsidiary of Al Ghurair Group", ""),
    ("jindal films", "subsidiary of B.C. Jindal Group", ""),
    ("jindal poly films", "subsidiary of B.C. Jindal Group", ""),
    ("scg packaging", "subsidiary of Siam Cement Group (SCG)", ""),
    ("gascogne flexible", "subsidiary of Gascogne Group", ""),
    ("sudpack iberica", "subsidiary of Südpack Verpackungen GmbH & Co. KG", ""),
    ("südpack ibérica", "subsidiary of Südpack Verpackungen GmbH & Co. KG", ""),
    ("mitsubishi polyester film", "subsidiary of Mitsubishi Chemical Group", ""),
    ("chiripal poly", "subsidiary of Chiripal Group", ""),
    ("toppan speciality films", "subsidiary of Toppan Inc.", ""),
    ("max speciality films", "subsidiary of Toppan Inc.", ""),
    ("nan ya plastics", "subsidiary of Formosa Plastics Group", ""),
    ("polyplex thailand", "subsidiary of Polyplex Corporation Ltd.", ""),
    ("toray advanced film", "subsidiary of Toray Industries", ""),
    ("ester filmtech", "subsidiary of Ester Industries Ltd.", ""),
    ("srf limited packaging", "subsidiary of SRF Limited", ""),
    ("hyosung chemical films", "subsidiary of Hyosung Corporation", ""),
    ("kolon industries films", "subsidiary of Kolon Industries", ""),
    ("asahi kasei packaging", "subsidiary of Asahi Kasei Corporation", ""),
    ("indorama ventures packaging", "subsidiary of Indorama Ventures", ""),
    ("manjushree technopack flexible", "subsidiary of Manjushree Technopack Ltd.", ""),
    ("saica flex", "subsidiary of SAICA Group", ""),
    ("smurfit westrock bag", "subsidiary of Smurfit Westrock", ""),
    ("mpact flexible", "subsidiary of Mpact Limited", ""),
    ("sigma stretch film", "subsidiary of Sigma Plastics Group", ""),
    ("formosa idemitsu", "subsidiary of Formosa Plastics Group / Idemitsu Kosan JV", ""),
    ("far eastern new century films", "subsidiary of Far Eastern New Century", ""),
    ("shinkong synthetic", "subsidiary of Shinkong Synthetic Fibers", ""),
    # Specialty film brand marketers — parent only when Brand is a product-line label
    ("3m scotchpak", "subsidiary of 3M Company", ""),
    ("dupont tyvek", "subsidiary of DuPont de Nemours, Inc.", ""),
    ("soarus", "subsidiary of Mitsubishi Chemical Group (SoarnoL)", ""),
    ("nippon gohsei", "subsidiary of Mitsubishi Chemical Group", ""),
    ("mylar specialty films", "formerly DuPont Teijin Films", ""),
    # PE-owned operating companies (show owner in Company column)
    ("constantia flexibles", "owned by One Rock Capital Partners", ""),
    ("coveris", "owned by Sun Capital Partners", ""),
    ("sealed air", "owned by Clayton, Dubilier & Rice (CD&R)", ""),
    ("novolex", "owned by Apollo Global Management", ""),
    ("clondalkin", "owned by Egeria", ""),
    ("kalle gmbh", "owned by Clayton, Dubilier & Rice (CD&R)", ""),
    ("proampac", "owned by Pritzker Private Capital", ""),
]

INDEPENDENT_PUBLIC = {
    "amcor plc",
    "mondi plc",
    "huhtamaki",
    "sonoco",
    "uflex ltd",
    "polyplex corporation",
    "cosmo first",
    "scientex",
    "flexopack",
    "bilcare",
    "ester industries ltd",
    "tcpl packaging",
    "toppan printing",
    "dai nippon",
    "toyo seikan",
    "toyobo",
    "toray industries",
    "rengo",
    "skc",
    "unitika",
    "futamura",
    "okura industrial",
    "kureha",
    "fujimori",
    "transcontinental",
    "reynolds consumer products",
    "thai film industries",
    "takween",
    "pt trias sentosa",
    "pt argha karya",
    "thong guan",
    "bp plastics",
    "shanghai zijiang",
    "jiangsu shuangxing",
}


def _norm(s: str) -> str:
    import unicodedata

    s = str(s or "").strip().lower()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    s = re.sub(r"[+/|&,_.:\-()]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def match_ownership(name: str) -> tuple[str, str] | None:
    n = _norm(name)
    best = None
    best_len = -1
    rename = ""
    for stem, own, new_brand in OWNERSHIP_BY_STEM:
        if stem in n and len(stem) > best_len:
            best = own
            best_len = len(stem)
            rename = new_brand
    if best:
        return best, rename
    return None


def clear_to_independent(name: str) -> bool:
    n = _norm(name)
    return any(s in n for s in INDEPENDENT_PUBLIC)


def main() -> None:
    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Landscape"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows[0]]
    landscape = []
    for r in rows[1:]:
        d = {hdr[i]: ("" if r[i] is None else str(r[i])) for i in range(len(hdr)) if i < len(r)}
        if d.get("Company"):
            landscape.append(d)

    report = {"ownership_fixed": [], "brand_renamed": [], "cleared_independent": []}

    for row in landscape:
        name = str(row.get("Company") or "").strip()
        old_own = str(row.get("Ownership") or "").strip()
        matched = match_ownership(name)
        if matched:
            own, rename = matched
            if rename and rename != name:
                report["brand_renamed"].append({"from": name, "to": rename})
                row["Company"] = rename
                name = rename
            if old_own != own:
                report["ownership_fixed"].append(
                    {"brand": name, "from": old_own, "to": own}
                )
            row["Ownership"] = own
            # parent_owner for display
            m = re.match(
                r"(?i)^\s*(?:acquired by|subsidiary of|merged into|owned by)\s+(.+)$",
                own,
            )
            row["parent_owner"] = m.group(1).strip() if m else ""
            if own.lower().startswith("subsidiary"):
                row["ownership_relation"] = "subsidiary_of"
            elif own.lower().startswith("owned"):
                row["ownership_relation"] = "owned_by"
            elif own.lower().startswith("merged"):
                row["ownership_relation"] = "merged_into"
            else:
                row["ownership_relation"] = "acquired_by"
        elif clear_to_independent(name):
            # Public independent — Company column should equal Brand (no false acquisition)
            # Keep Ownership as Public / Private text without acquired/subsidiary prefix
            if re.match(
                r"(?i)^\s*(?:acquired by|subsidiary of|merged into|owned by)\b",
                old_own,
            ):
                row["Ownership"] = "Public"
                row["parent_owner"] = ""
                row["ownership_relation"] = ""
                report["cleared_independent"].append(name)
            else:
                # normalize verbose ownership that won't parse as parent
                if "subsidiary" in old_own.lower() and not old_own.lower().startswith(
                    "subsidiary of"
                ):
                    # e.g. "Public / Subsidiary of Wihuri" already handled by match
                    pass
                row["parent_owner"] = ""
                row["ownership_relation"] = ""
        else:
            # Private independent — strip accidental acquired prefixes if Brand is independent
            if re.match(
                r"(?i)^\s*(?:acquired by|subsidiary of|merged into|owned by)\b",
                old_own,
            ):
                # keep if it looks intentional; already handled by match_ownership
                pass
            row.setdefault("parent_owner", "")
            # If Ownership is "Private (Family)" etc., no parent annotation
            if not re.match(
                r"(?i)^\s*(?:acquired by|subsidiary of|merged into|owned by)\b",
                str(row.get("Ownership") or ""),
            ):
                row["parent_owner"] = ""
                row["ownership_relation"] = ""

        # Role lock
        role = str(row.get("Role") or row.get("Distribution Type") or "Brand")
        if role not in ("Brand", "Marketer"):
            role = "Brand"
        row["Role"] = role
        row["Distribution Type"] = role

    audit = {}
    ap = OUT / "chatgpt_expand_batch_all_audit.json"
    if ap.exists():
        audit = json.loads(ap.read_text(encoding="utf-8")).get("xy_scoring") or {}

    details = to_company_detail_rows(landscape, QUERY, audit)
    by = {_norm(r.get("Company") or ""): r.get("Role") for r in landscape}
    for d in details:
        role = by.get(_norm(d.get("Brand") or "")) or by.get(_norm(d.get("Company") or ""))
        if role in ("Brand", "Marketer"):
            d["Role"] = role

    write_final_xlsx(
        XLSX,
        landscape,
        "Companies",
        {
            "query": QUERY,
            "company_column_fix": "2026-08-14",
            "ownership_fixed": len(report["ownership_fixed"]),
        },
        detail_rows=details,
    )
    extras = export_expand_quadrant_outputs(
        OUT, details, QUERY, country="global", audit=audit, chart_n=20
    )

    # Audit Brand|Company pairs
    pairs = [{"Brand": d.get("Brand"), "Company": d.get("Company"), "Role": d.get("Role")} for d in details]
    annotated = [p for p in pairs if any(x in str(p["Company"]).lower() for x in ("acquired", "subsidiary", "owned by", "merged"))]
    report["after"] = len(landscape)
    report["annotated_company_col"] = len(annotated)
    report["pairs_sample_annotated"] = annotated
    report["html"] = str(extras.get("html") or "")

    outp = (
        ROOT
        / "output"
        / "chatgpt_expand"
        / "_audit"
        / f"{SLUG}_company_column_fix.json"
    )
    outp.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")

    print(
        f"DONE n={len(landscape)} ownership_fixed={len(report['ownership_fixed'])} "
        f"renamed={len(report['brand_renamed'])} annotated_company={len(annotated)}"
    )
    for x in report["ownership_fixed"][:40]:
        print(f"  OWN {x['brand']}: {x['from'][:50]} => {x['to'][:60]}")
    for x in report["brand_renamed"]:
        print(f"  RENAME {x['from']} => {x['to']}")
    print("Annotated Company column:")
    for p in annotated:
        print(f"  {p['Brand']} | {p['Company']}")
    print("audit ->", outp)


if __name__ == "__main__":
    main()
