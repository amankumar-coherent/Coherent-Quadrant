#!/usr/bin/env python3
"""
Recalculate X / Y / Overall for wearable medical landscape using REAL web evidence.

- Agent-style: DDGS web search (+ optional site fetch) per company
- DeepSeek scores ONLY from gathered evidence (no invented products/certs/revenue)
- X = Clinical & Device Capability, Y = Market Access & Growth Strategy
- Overall = round((X+Y)/2); quadrants = half-median

Env:
  QUAD_RESCORE_LIMIT=N     score only first N (test)
  QUAD_RESCORE_BATCH=15
  QUAD_RESCORE_CONCURRENT=4
  QUAD_SKIP_FETCH=1        skip homepage fetch
"""
from __future__ import annotations

import asyncio
import json
import os
import re
import sys
import time
from collections import Counter
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from vendor_intel.placeholders.load_keys import apply_env_overrides
from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx
from vendor_intel.quadrant.rating_map import assign_quadrants_half_median

OUT = ROOT / "output" / "chatgpt_expand"
SLUG = "global_wearable_medical_devices_market_global"
QUERY = "Global Wearable Medical Devices Market"
FOLDER = OUT / SLUG
INDUSTRY = "Healthcare / Wearable Medical Devices"

AXIS_X = "Clinical & Device Capability"
AXIS_Y = "Market Access & Growth Strategy"
X_FEATS = [
    "Device Portfolio Breadth",
    "Clinical Performance & Safety",
    "Innovation & R&D Capability",
    "Manufacturing & Quality Systems",
    "Regulatory & Quality Compliance",
]
Y_FEATS = [
    "Geographic & Channel Reach",
    "Brand Trust & Clinical Adoption",
    "Financial Performance",
    "Pipeline & Indication Expansion",
    "Partnerships & Business Expansion",
]
# matrix slot weights
X_W = [0.30, 0.20, 0.20, 0.15, 0.15]
Y_W = [0.25, 0.25, 0.20, 0.15, 0.15]

SYSTEM = """You score a company on the Coherent Quadrant for wearable medical devices.
Use ONLY the EVIDENCE text provided. Do NOT invent products, FDA clearances, revenue,
partnerships, countries, or clinical claims that are not supported by the evidence.

Scoring rules (each feature 1–10):
- 9–10: clear evidence of category leadership / broad validated capability or global scale
- 7–8: strong evidenced capability or solid multi-market commercial presence
- 5–6: moderate / partial evidence
- 3–4: weak or narrow evidence
- 1–2: evidence shows weakness or absence for that feature
If evidence is silent on a feature → score 4–5 and grounding=insufficient for that feature.
When evidence clearly supports strength, DO NOT compress scores to the mid-band — use 8–10.

Return JSON only:
{
  "x_features": [{"feature": "...", "score": 1-10, "grounding": "supported|partial|insufficient", "why": "short cite from evidence"}],
  "y_features": [{"feature": "...", "score": 1-10, "grounding": "supported|partial|insufficient", "why": "short cite from evidence"}],
  "evidence_used": true
}
All five X features and five Y features required."""


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def _domain(website: str) -> str:
    raw = (website or "").strip()
    if not raw:
        return ""
    try:
        host = urlparse(raw if "://" in raw else f"https://{raw}").netloc.lower()
        return host[4:] if host.startswith("www.") else host
    except Exception:
        return ""


def _axis_100(feat_scores: list[dict], features: list[str], weights: list[float]) -> int:
    by = {_norm(f.get("feature") or ""): f for f in feat_scores}
    total = 0.0
    for feat, w in zip(features, weights):
        item = by.get(_norm(feat)) or {}
        try:
            s = float(item.get("score") or 4)
        except (TypeError, ValueError):
            s = 4.0
        s = max(1.0, min(10.0, s))
        # slight upward stretch when supported/partial and mid-high
        g = str(item.get("grounding") or "").lower()
        if g in {"supported", "partial"} and 4 <= s <= 8:
            s = min(10.0, s + 1.0)
        total += w * (s / 10.0)
    return int(round(total * 100))


def _web_search(brand: str, website: str, specialty: str) -> str:
    """Live DDGS search — real snippets only."""
    from ddgs import DDGS

    queries = [
        f"{brand} wearable medical device FDA OR CE OR clinical",
        f"{brand} continuous monitoring OR CGM OR ECG OR hearing aid market",
        f"{brand} headquarters revenue partnership hospital",
    ]
    if specialty:
        queries.append(f"{brand} {specialty[:80]}")
    dom = _domain(website)
    if dom:
        queries.append(f"site:{dom} {brand} product medical")

    chunks: list[str] = []
    seen: set[str] = set()
    try:
        with DDGS() as ddgs:
            for q in queries[:4]:
                try:
                    results = list(ddgs.text(q, max_results=5))
                except Exception:
                    results = []
                for r in results or []:
                    title = str(r.get("title") or "").strip()
                    body = str(r.get("body") or r.get("snippet") or "").strip()
                    href = str(r.get("href") or r.get("link") or "").strip()
                    key = _norm(title + body)[:160]
                    if not key or key in seen:
                        continue
                    seen.add(key)
                    chunks.append(f"- {title}\n  {body}\n  URL: {href}")
                    if len(chunks) >= 14:
                        break
                if len(chunks) >= 14:
                    break
                time.sleep(0.35)
    except Exception as exc:
        chunks.append(f"(web search error: {type(exc).__name__}: {exc})")
    return "\n".join(chunks)


async def _fetch_site(website: str) -> str:
    if not website or (os.getenv("QUAD_SKIP_FETCH") or "").lower() in {"1", "true", "yes"}:
        return ""
    url = website if "://" in website else f"https://{website}"
    try:
        async with httpx.AsyncClient(
            timeout=12.0,
            follow_redirects=True,
            headers={"User-Agent": "Mozilla/5.0 (compatible; CoherentQuadrant/1.0)"},
        ) as client:
            resp = await client.get(url)
            if resp.status_code >= 400:
                return ""
            text = resp.text or ""
            # crude strip tags
            text = re.sub(r"(?is)<script[^>]*>.*?</script>", " ", text)
            text = re.sub(r"(?is)<style[^>]*>.*?</style>", " ", text)
            text = re.sub(r"(?is)<[^>]+>", " ", text)
            text = re.sub(r"\s+", " ", text).strip()
            return text[:4000]
    except Exception:
        return ""


def _deepseek_score(brand: str, role: str, specialty: str, evidence: str) -> dict[str, Any]:
    apply_env_overrides()
    from vendor_intel.placeholders import llm as llm_mod

    # Force DeepSeek path (module-level vars, not only os.environ)
    os.environ["LLM_PROVIDER"] = "deepseek"
    llm_mod.LLM_PROVIDER = "deepseek"
    if not llm_mod.is_configured():
        raise RuntimeError("DeepSeek not configured (DEEPSEEK_API_KEY missing)")

    user = f"""Market: {QUERY}
Brand: {brand}
Role: {role}
Known specialty (may be incomplete): {specialty or 'n/a'}

X axis ({AXIS_X}) features: {X_FEATS}
Y axis ({AXIS_Y}) features: {Y_FEATS}

EVIDENCE (web search + site text — cite only this):
{evidence[:14000]}

Score now. JSON only."""
    parsed = llm_mod.llm_complete_json(
        SYSTEM,
        user,
        max_tokens=1800,
    )
    if not isinstance(parsed, dict):
        raise ValueError(f"no JSON object from DeepSeek for {brand}")
    if parsed.get("status") in {"llm_failed", "placeholder"}:
        raise ValueError(
            f"DeepSeek failed for {brand}: {parsed.get('error') or parsed.get('status')}"
        )
    if not (parsed.get("x_features") and parsed.get("y_features")):
        raise ValueError(f"missing x_features/y_features for {brand}")
    return parsed


def _load_rows() -> list[dict]:
    xlsx = FOLDER / f"{SLUG}_FINAL.xlsx"
    wb = load_workbook(xlsx, data_only=True)
    ws = wb["Landscape"]
    rows_raw = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows_raw[0]]

    details_by: dict[str, dict] = {}
    if "Company Details" in wb.sheetnames:
        dws = wb["Company Details"]
        drows = list(dws.iter_rows(values_only=True))
        dhdr = [str(h) for h in drows[0]]
        for r in drows[1:]:
            if not r:
                continue
            d = {dhdr[i]: ("" if r[i] is None else r[i]) for i in range(len(dhdr)) if i < len(r)}
            key = _norm(d.get("Brand") or "")
            if key:
                details_by[key] = d

    kept: list[dict] = []
    for r in rows_raw[1:]:
        if not r:
            continue
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        brand = str(d.get("Company") or d.get("Brand") or "").strip()
        if not brand:
            continue
        key = _norm(brand)
        det = details_by.get(key)
        if det:
            if det.get("Brand"):
                brand = str(det["Brand"]).strip()
                key = _norm(brand)
                d["Brand"] = brand
            if det.get("Company") and (
                str(det["Company"]).startswith("(") or _norm(det.get("Company")) != key
            ):
                d["_display_company"] = det["Company"]
            if det.get("Found in"):
                d["Headquarters"] = det["Found in"]
                d["Found in"] = det["Found in"]
            if det.get("Role") in {"Brand", "Marketer"}:
                d["Role"] = det["Role"]
                d["Distribution Type"] = det["Role"]
        d["Brand"] = brand
        d["Company"] = brand
        d["Role"] = d.get("Role") if d.get("Role") in {"Brand", "Marketer"} else "Brand"
        d["Distribution Type"] = d["Role"]
        d["Industry Category"] = INDUSTRY
        # clear prior seed/heuristic scores
        for col in ("X Score", "Y Score", "Overall Score", "Quadrant"):
            d[col] = ""
        kept.append(d)
    return kept


def _save(kept: list[dict], *, partial: bool = False, meta: dict | None = None) -> None:
    xs, ys = [], []
    for r in kept:
        try:
            xs.append(float(r.get("X Score") or 50))
        except (TypeError, ValueError):
            xs.append(50.0)
        try:
            ys.append(float(r.get("Y Score") or 50))
        except (TypeError, ValueError):
            ys.append(50.0)
    # only assign quadrants for scored rows; unscored keep blank until end
    scored_idx = [i for i, r in enumerate(kept) if str(r.get("X Score") or "").strip()]
    if len(scored_idx) >= 4:
        sxs = [xs[i] for i in scored_idx]
        sys_ = [ys[i] for i in scored_idx]
        quads, mid_x, mid_y = assign_quadrants_half_median(sxs, sys_)
        for j, i in enumerate(scored_idx):
            kept[i]["Quadrant"] = quads[j]
            kept[i]["Overall Score"] = round((xs[i] + ys[i]) / 2)
    else:
        mid_x = mid_y = 50.0

    audit = {
        "axis_x": AXIS_X,
        "axis_y": AXIS_Y,
        "x_features": X_FEATS,
        "y_features": Y_FEATS,
        "x_feature_weights": X_W,
        "y_feature_weights": Y_W,
        "mid_x": mid_x,
        "mid_y": mid_y,
        "method": "ddgs_web_search + optional site fetch + DeepSeek evidence-only scoring",
        "partial": partial,
        **(meta or {}),
    }
    detail_rows = to_company_detail_rows(kept, QUERY, audit)
    for det in detail_rows:
        key = _norm(det.get("Brand") or "")
        src = next((r for r in kept if _norm(r.get("Brand") or "") == key), None)
        if not src:
            continue
        det["Role"] = src.get("Role") or "Brand"
        if src.get("_display_company"):
            det["Company"] = src["_display_company"]
        elif not str(det.get("Company") or "").startswith("("):
            if not (det.get("Company") and _norm(det.get("Company")) != key):
                det["Company"] = det.get("Brand") or det.get("Company")

    xlsx = FOLDER / f"{SLUG}_FINAL.xlsx"
    write_final_xlsx(
        xlsx,
        kept,
        "Companies",
        {"query": QUERY, "xy_scoring": audit},
        detail_rows=detail_rows,
    )
    if not partial:
        export_expand_quadrant_outputs(
            FOLDER, detail_rows, QUERY, country="global", audit=audit, chart_n=20
        )


async def _score_one(row: dict, sem: asyncio.Semaphore) -> dict[str, Any]:
    brand = str(row.get("Brand") or "")
    async with sem:
        specialty = str(row.get("Specialty Focus") or row.get("Core Categories") or "")
        website = str(row.get("Website") or "")
        role = str(row.get("Role") or "Brand")
        # web search in thread (blocking ddgs)
        search_txt = await asyncio.to_thread(_web_search, brand, website, specialty)
        site_txt = await _fetch_site(website)
        landscape_bits = (
            f"HQ: {row.get('Headquarters') or row.get('Found in') or ''}\n"
            f"Presence: {row.get('Operational Presence') or ''}\n"
            f"Categories: {row.get('Core Categories') or ''}\n"
            f"Specialty: {specialty}\n"
            f"Ownership: {row.get('Ownership') or ''}\n"
            f"Website: {website}\n"
        )
        evidence = (
            f"=== LANDSCAPE FACTS ===\n{landscape_bits}\n"
            f"=== WEB SEARCH ===\n{search_txt or '(no results)'}\n"
            f"=== WEBSITE TEXT ===\n{site_txt or '(not fetched)'}\n"
        )
        # require some real web evidence
        if len(search_txt) < 80 and len(site_txt) < 80:
            return {
                "brand": brand,
                "ok": False,
                "error": "insufficient_web_evidence",
                "x": 45,
                "y": 45,
                "evidence_chars": len(evidence),
            }
        try:
            payload = await asyncio.to_thread(
                _deepseek_score, brand, role, specialty, evidence
            )
            x = _axis_100(payload.get("x_features") or [], X_FEATS, X_W)
            y = _axis_100(payload.get("y_features") or [], Y_FEATS, Y_W)
            return {
                "brand": brand,
                "ok": True,
                "x": x,
                "y": y,
                "overall": round((x + y) / 2),
                "payload": payload,
                "evidence_chars": len(evidence),
            }
        except Exception as exc:
            return {
                "brand": brand,
                "ok": False,
                "error": f"{type(exc).__name__}: {exc}",
                "x": 45,
                "y": 45,
                "evidence_chars": len(evidence),
            }


async def main() -> int:
    apply_env_overrides()
    os.environ["LLM_PROVIDER"] = "deepseek"
    from vendor_intel.placeholders import llm as llm_mod

    llm_mod.LLM_PROVIDER = "deepseek"
    if not llm_mod.is_configured():
        print("ERROR: DeepSeek not configured (set DEEPSEEK_API_KEY)", flush=True)
        return 2

    kept = _load_rows()
    # LIMIT only controls how many to score — NEVER shrink the saved landscape
    limit = int(os.getenv("QUAD_RESCORE_LIMIT") or "0")
    to_score = kept[:limit] if limit > 0 else kept
    if limit > 0:
        print(f"LIMIT {limit} to score (keeping all {len(kept)} rows in FINAL)", flush=True)

    batch_n = int(os.getenv("QUAD_RESCORE_BATCH") or "15")
    concurrent = int(os.getenv("QUAD_RESCORE_CONCURRENT") or "4")
    sem = asyncio.Semaphore(concurrent)

    audit_log: list[dict] = []
    by_brand = {_norm(r.get("Brand") or ""): r for r in kept}

    print(
        f"Evidence rescore {len(to_score)}/{len(kept)} companies "
        f"(batch={batch_n}, concurrent={concurrent}, DeepSeek+DDGS)",
        flush=True,
    )

    for start in range(0, len(to_score), batch_n):
        batch = to_score[start : start + batch_n]
        print(
            f"\n=== Batch {start // batch_n + 1}: "
            f"{start+1}-{start+len(batch)} / {len(to_score)} ===",
            flush=True,
        )
        results = await asyncio.gather(*[_score_one(r, sem) for r in batch])
        for res in results:
            brand = res["brand"]
            row = by_brand.get(_norm(brand))
            if not row:
                continue
            row["X Score"] = res["x"]
            row["Y Score"] = res["y"]
            row["Overall Score"] = res.get("overall") or round((res["x"] + res["y"]) / 2)
            note = "ok" if res.get("ok") else res.get("error")
            row["Summary"] = (
                f"{brand} — evidence-scored X={res['x']} Y={res['y']} "
                f"({note}; evidence_chars={res.get('evidence_chars')})"
            )
            audit_log.append(
                {
                    "brand": brand,
                    "ok": res.get("ok"),
                    "x": res["x"],
                    "y": res["y"],
                    "error": res.get("error"),
                    "evidence_chars": res.get("evidence_chars"),
                    "features": res.get("payload"),
                }
            )
            print(
                f"  {brand}: X={res['x']} Y={res['y']} "
                f"{'OK' if res.get('ok') else 'FALLBACK '+str(res.get('error'))}",
                flush=True,
            )
        _save(kept, partial=True, meta={"scored": len(audit_log)})
        # persist audit progress
        path = OUT / "_audit" / "wearable_xy_evidence_rescore.json"
        path.write_text(
            json.dumps({"scored": len(audit_log), "rows": audit_log}, indent=2, ensure_ascii=False),
            encoding="utf-8",
        )

    # Final export only when every row has X/Y (avoid wiping chart on LIMIT smoke tests)
    all_scored = all(str(r.get("X Score") or "").strip() for r in kept)
    _save(
        kept,
        partial=not all_scored,
        meta={"scored": len(audit_log), "complete": all_scored, "total": len(kept)},
    )

    roles = dict(Counter(r.get("Role") for r in kept))
    quads = dict(Counter(r.get("Quadrant") for r in kept if r.get("Quadrant")))
    ok_n = sum(1 for a in audit_log if a.get("ok"))
    md = OUT / "_audit" / "wearable_xy_evidence_rescore.md"
    lines = [
        "# X/Y/Overall evidence rescore — Wearable Medical Devices",
        "",
        f"Scored **{len(audit_log)}** / {len(kept)} with DDGS web search + DeepSeek (evidence-only).",
        f"OK evidence scores: **{ok_n}** · fallback: **{len(audit_log)-ok_n}**",
        f"Complete landscape scored: **{all_scored}**",
        "",
        f"Axes: **X**={AXIS_X} · **Y**={AXIS_Y} · Overall=(X+Y)/2",
        "",
        f"Roles {roles} · Quadrants {quads}",
        "",
        "## Sample",
        "",
        "| Brand | X | Y | Overall | OK |",
        "|---|---:|---:|---:|---|",
    ]
    for a in sorted(audit_log, key=lambda z: -((z.get("x") or 0) + (z.get("y") or 0)))[:30]:
        o = round(((a.get("x") or 0) + (a.get("y") or 0)) / 2)
        lines.append(
            f"| {a['brand']} | {a.get('x')} | {a.get('y')} | {o} | {a.get('ok')} |"
        )
    md.write_text("\n".join(lines), encoding="utf-8")
    print(f"\nDONE ok={ok_n}/{len(audit_log)} all_scored={all_scored} quads={quads}", flush=True)
    print(f"audit -> {md}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
