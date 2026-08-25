#!/usr/bin/env python3
"""Second-pass brand market verification for Global Semiconductor FINAL.

Removes brands that fail the semiconductor inclusion rule after web check.
Keeps count >= 200 (no adds needed if still above target).
"""
from __future__ import annotations

import csv
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx

OUT = ROOT / "output" / "chatgpt_expand"
SLUG = "global_semiconductor_market_global"
QUERY = "Global Semiconductor Market"
FOLDER = OUT / SLUG
TARGET_MIN = 200

# brand_key -> reason (web-verified Aug 2026)
REMOVE: dict[str, str] = {
    "d-wave systems": "Quantum computing systems OEM — not a commercial semiconductor device/equipment/EDA vendor.",
    "unitec semiconductores": "Brazilian fab project inactive / judicial recovery; not an operating commercial semi vendor.",
    # Acquired shells whose parent Solution Provider is already listed
    "gan systems": "Acquired by Infineon Technologies (2023); Infineon already listed.",
    "autotalks": "Acquired by Qualcomm (2025); Qualcomm already listed.",
    "stats chippac": "Acquired by JCET Group (2015); JCET already listed.",
    "spil (siliconware precision industries)": "Acquired by ASE Technology Holding (2018); ASE already listed.",
    "siliconware precision industries (spil)": "Acquired by ASE Technology Holding (2018); ASE already listed.",
    "silead": "Acquired by GigaDevice (2019); GigaDevice already listed.",
}


def _norm(s: str) -> str:
    s = str(s or "").strip().lower()
    s = re.sub(r"\s*\((?:acquired by|subsidiary of|merged into)[^)]*\)\s*$", "", s)
    return re.sub(r"\s+", " ", s).strip()


def _identity(row: dict) -> str:
    brand = str(row.get("Brand") or "").strip()
    company = str(row.get("Company") or "").strip()
    if brand and not brand.lower().startswith("("):
        return brand
    if company and not company.lower().startswith("("):
        return company
    return brand or company


def _load(xlsx: Path) -> tuple[list[dict], dict]:
    wb = load_workbook(xlsx, data_only=True)
    landscape: list[dict] = []
    ws = wb["Landscape"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows[0]]
    for r in rows[1:]:
        if not r:
            continue
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        if str(d.get("Company") or "").strip():
            landscape.append(d)

    details_by: dict[str, dict] = {}
    if "Company Details" in wb.sheetnames:
        ws = wb["Company Details"]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(h) for h in rows[0]]
        for r in rows[1:]:
            if not r:
                continue
            d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
            key = _norm(d.get("Brand") or d.get("Company") or "")
            if key:
                details_by[key] = d

    for row in landscape:
        key = _norm(row.get("Company") or row.get("Brand") or "")
        det = details_by.get(key)
        if not det:
            for dk, dv in details_by.items():
                if dk in key or key in dk:
                    det = dv
                    break
        if det:
            for col in ("Quadrant", "X", "Y", "Overall", "Role"):
                if det.get(col) not in (None, ""):
                    row[col] = det[col]
                    if col == "X":
                        row["X Score"] = det[col]
                    elif col == "Y":
                        row["Y Score"] = det[col]
                    elif col == "Overall":
                        row["Overall Score"] = det[col]
            if det.get("Found in"):
                row["Headquarters"] = det["Found in"]
                row["Found in"] = det["Found in"]
            brand = str(det.get("Brand") or "").strip()
            if brand:
                row["Brand"] = brand
                # Keep Company display suffix if present on details
                company_col = str(det.get("Company") or "").strip()
                if company_col and company_col != brand:
                    row["_display_company"] = company_col
            if det.get("Role"):
                row["Distribution Type"] = det["Role"]
                row["Role"] = det["Role"]
    return landscape, details_by


def main() -> None:
    xlsx = FOLDER / f"{SLUG}_FINAL.xlsx"
    landscape, _ = _load(xlsx)
    before = len(landscape)

    kept: list[dict] = []
    removed: list[dict] = []
    for row in landscape:
        ident = _identity(row)
        key = _norm(ident)
        reason = REMOVE.get(key)
        if not reason:
            # fuzzy: key contained in REMOVE keys or vice versa
            for rk, rr in REMOVE.items():
                if rk == key or rk in key or key in rk:
                    reason = rr
                    break
        if reason:
            removed.append({"brand": ident, "reason": reason})
            continue
        row["Role"] = "Solution Provider"
        row["Distribution Type"] = "Solution Provider"
        kept.append(row)

    if len(kept) < TARGET_MIN:
        raise SystemExit(f"Would drop below {TARGET_MIN}: {len(kept)} after removing {len(removed)}")

    audit_path = FOLDER / "chatgpt_expand_batch_all_audit.json"
    audit: dict = {}
    if audit_path.exists():
        try:
            audit = json.loads(audit_path.read_text(encoding="utf-8")).get("xy_scoring") or {}
        except Exception:
            pass

    qjson = FOLDER / f"{SLUG}_quadrant.json"
    if qjson.exists():
        try:
            payload = json.loads(qjson.read_text(encoding="utf-8"))
            crit = payload.get("criteria") or {}
            if crit.get("parameter_definitions"):
                audit["parameter_definitions"] = crit["parameter_definitions"]
            labels = crit.get("axis_labels") or {}
            if labels.get("x"):
                audit["axis_x"] = labels["x"]
            if labels.get("y"):
                audit["axis_y"] = labels["y"]
            if crit.get("x_axis"):
                audit["x_features"] = crit["x_axis"]
            if crit.get("y_axis"):
                audit["y_features"] = crit["y_axis"]
        except Exception:
            pass

    detail_rows = to_company_detail_rows(kept, QUERY, audit)
    for det, src in zip(detail_rows, kept):
        det["Role"] = "Solution Provider"
        if src.get("_display_company"):
            det["Company"] = src["_display_company"]
        brand = str(det.get("Brand") or "").strip()
        if brand and not str(det.get("Company") or "").strip():
            det["Company"] = brand

    write_final_xlsx(
        xlsx,
        kept,
        "Companies",
        {
            "query": QUERY,
            "xy_scoring": audit,
            "brand_market_verify": {"removed": removed, "after": len(kept)},
        },
        detail_rows=detail_rows,
    )
    extras = export_expand_quadrant_outputs(
        FOLDER, detail_rows, QUERY, country="global", audit=audit, chart_n=20
    )

    csv_path = FOLDER / f"{SLUG}_companies.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(
            f,
            fieldnames=["Brand", "Company", "Role", "Quadrant", "X", "Y", "Overall", "Found in"],
            extrasaction="ignore",
        )
        w.writeheader()
        for d in detail_rows:
            w.writerow(d)

    md_lines = [
        "# Semiconductor Brand market verification (pass 2)",
        "",
        f"Date: 2026-08-14",
        f"Before: **{before}** → After: **{len(kept)}** (target ≥{TARGET_MIN})",
        "",
        "## Inclusion rule",
        "",
        "Keep Brand if it designs/fabs/packages-tests semiconductors, wafers, fab equipment,",
        "ATE, EDA/IP, substrates/materials, or semiconductor-focused design services.",
        "",
        "## Removed this pass",
        "",
        "| Brand | Why |",
        "|---|---|",
    ]
    for r in sorted(removed, key=lambda x: x["brand"].lower()):
        md_lines.append(f"| {r['brand']} | {r['reason']} |")
    md_lines.extend(
        [
            "",
            "## Verdict",
            "",
            f"All remaining **{len(kept)}** Brands are semiconductor-market Solution Providers",
            "(IDM / fabless / foundry / OSAT / equipment / materials / EDA-IP / design services).",
            "Borderline ecosystem kept: Jenoptik, Rohde & Schwarz, Keysight, Nikon (lithography),",
            "Hamamatsu, Coherent Corp., Mitsubishi Electric / Fuji Electric (power semis), Tata Electronics (OSAT/fab).",
            "",
            f"HTML: `{extras.get('html')}`",
            "",
        ]
    )
    md_path = OUT / "_audit" / "semiconductor_brand_market_verify.md"
    md_path.write_text("\n".join(md_lines), encoding="utf-8")
    (OUT / "_audit" / "semiconductor_brand_market_verify.json").write_text(
        json.dumps(
            {
                "before": before,
                "after": len(kept),
                "removed": removed,
                "html": str(extras.get("html") or ""),
            },
            indent=2,
        ),
        encoding="utf-8",
    )

    print(f"Before {before} -> After {len(kept)}; removed {len(removed)}")
    for r in removed:
        print(f"  - {r['brand']}: {r['reason']}")
    print(f"audit -> {md_path}")


if __name__ == "__main__":
    main()
