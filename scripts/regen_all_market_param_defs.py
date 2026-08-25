#!/usr/bin/env python3
"""Regenerate all 5 market HTML reports with scoring parameter definitions restored."""
from __future__ import annotations

import json
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import export_expand_quadrant_outputs
from vendor_intel.placeholders.load_keys import apply_env_overrides

MARKETS = [
    (
        "global_flexible_packaging_market_global",
        "Global Flexible Packaging Market",
        "Others",
        "Packaging",
    ),
    (
        "global_glp_1_receptor_agonist_market_global",
        "Global GLP-1 Receptor Agonist Market",
        "Healthcare",
        "Pharmaceutical",
    ),
    (
        "global_liquefied_natural_gas_market_global",
        "Global Liquefied Natural Gas Market",
        "Others",
        "Energy",
    ),
    (
        "global_semiconductor_market_global",
        "Global Semiconductor Market",
        "ICT, Automation, Semiconductor",
        "Semiconductors",
    ),
    (
        "global_wearable_medical_devices_market_global",
        "Global Wearable Medical Devices Market",
        "Healthcare",
        "Medical Devices",
    ),
]


def _load_details(folder: Path, slug: str) -> list[dict]:
    xlsx = folder / f"{slug}_FINAL.xlsx"
    wb = load_workbook(xlsx, data_only=True)
    rows = list(wb["Company Details"].iter_rows(values_only=True))
    hdr = [str(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not r or not r[0]:
            continue
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr))}
        out.append(
            {
                "Brand": d.get("Brand") or "",
                "Company": d.get("Company") or d.get("Brand") or "",
                "Role": d.get("Role") or "",
                "Quadrant": d.get("Quadrant") or "",
                "X": d.get("X"),
                "Y": d.get("Y"),
                "Overall": d.get("Overall"),
                "Found in": d.get("Found in") or "",
            }
        )
    return out


def _load_audit(folder: Path) -> dict:
    p = folder / "chatgpt_expand_batch_all_audit.json"
    if not p.exists():
        return {}
    try:
        raw = json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return {}
    if isinstance(raw.get("xy_scoring"), dict):
        return dict(raw["xy_scoring"])
    return dict(raw) if isinstance(raw, dict) else {}


def _ensure_defs(query: str, audit: dict, industry_group: str, industry_category: str) -> dict:
    """Ensure x/y features + parameter_definitions exist (LLM fill if missing)."""
    from vendor_intel.quadrant.criteria_catalog import get_category_features
    from vendor_intel.quadrant.axis_define import explain_market_parameters

    axis_x = str(audit.get("axis_x") or "").strip()
    axis_y = str(audit.get("axis_y") or "").strip()
    x_feats = [str(f).strip() for f in (audit.get("x_features") or []) if str(f).strip()]
    y_feats = [str(f).strip() for f in (audit.get("y_features") or []) if str(f).strip()]
    defs = audit.get("parameter_definitions") if isinstance(audit.get("parameter_definitions"), dict) else {}
    x_defs = defs.get("x") if isinstance(defs.get("x"), dict) else {}
    y_defs = defs.get("y") if isinstance(defs.get("y"), dict) else {}

    if not x_feats or not y_feats or not axis_x or not axis_y:
        try:
            industry = get_category_features(industry_group, industry_category)
        except KeyError:
            industry = {}
        axis_x = axis_x or str(industry.get("axis_x") or "Solution Capability")
        axis_y = axis_y or str(industry.get("axis_y") or "Business Strategy")
        x_feats = x_feats or list(industry.get("x") or [])
        y_feats = y_feats or list(industry.get("y") or [])

    need_defs = (
        not x_defs
        or not y_defs
        or sum(1 for f in x_feats if str(x_defs.get(f) or "").strip()) < len(x_feats)
        or sum(1 for f in y_feats if str(y_defs.get(f) or "").strip()) < len(y_feats)
    )
    if need_defs and x_feats and y_feats:
        print(f"  generating parameter definitions for {query}…", flush=True)
        try:
            defs = explain_market_parameters(
                query,
                x_feats,
                y_feats,
                axis_x=axis_x,
                axis_y=axis_y,
                geography="global",
            )
            x_defs = defs.get("x") if isinstance(defs.get("x"), dict) else x_defs
            y_defs = defs.get("y") if isinstance(defs.get("y"), dict) else y_defs
        except Exception as exc:
            print(f"  definition LLM failed: {exc}", flush=True)

    audit = dict(audit)
    audit["axis_x"] = axis_x
    audit["axis_y"] = axis_y
    audit["x_features"] = x_feats
    audit["y_features"] = y_feats
    audit["parameter_definitions"] = {"x": x_defs or {}, "y": y_defs or {}}
    audit["industry_group"] = audit.get("industry_group") or industry_group
    audit["industry_category"] = audit.get("industry_category") or industry_category
    return audit


def main() -> int:
    apply_env_overrides()
    os.environ.setdefault("LLM_PROVIDER", "deepseek")
    base = ROOT / "output" / "chatgpt_expand"

    for slug, query, group, category in MARKETS:
        folder = base / slug
        details = _load_details(folder, slug)
        audit = _ensure_defs(query, _load_audit(folder), group, category)
        export_expand_quadrant_outputs(
            folder, details, query, country="global", audit=audit, chart_n=20
        )
        html = (folder / f"{slug}_report.html").read_text(encoding="utf-8")
        has_params = "Market Scoring Parameters" in html
        has_defs = "Parameter Definitions" in html
        print(
            f"OK {slug} n={len(details)} "
            f"params={has_params} defs={has_defs} "
            f"xf={len(audit.get('x_features') or [])} "
            f"xdefs={len((audit.get('parameter_definitions') or {}).get('x') or {})}"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
