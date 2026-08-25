#!/usr/bin/env python3
"""Recalculate X / Y / Overall for ALL semiconductor companies from real crawl + DeepSeek.

Clears prior scores, deep-crawls websites, scores features via LLM (DeepSeek),
reassigns half-median quadrants, preserves Brand/Company ownership display.
No hallucinated numeric scores — only pipeline output from evidence.
"""
from __future__ import annotations

import asyncio
import csv
import json
import os
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

from openpyxl import load_workbook

from vendor_intel.placeholders.load_keys import apply_env_overrides
from vendor_intel.pipeline.chatgpt_env import apply_chatgpt_expand_env
from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    score_expand_rows,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx
from vendor_intel.quadrant.rating_map import assign_quadrants_half_median

OUT = ROOT / "output" / "chatgpt_expand"
SLUG = "global_semiconductor_market_global"
QUERY = "Global Semiconductor Market"
FOLDER = OUT / SLUG
INDUSTRY = "ICT, Automation, Semiconductor / Semiconductors"
CKPT = OUT / "_audit" / "semi_xy_rescore_checkpoint.json"
BATCH = int(os.getenv("SEMI_RESCORE_BATCH") or "40")

# Official website fixes (verified) for known bad/redirect domains
WEBSITE_FIXES: dict[str, str] = {
    "hua hong semiconductor": "https://www.huahonggrace.com",
    "ceitec s.a.": "https://www.ceitec.org.br",
    "ceitec sa": "https://www.ceitec.org.br",
    "socionext": "https://www.socionext.com/en/",
    "globalfoundries": "https://gf.com",
    "tokyo electron": "https://www.tel.com",
    "toshiba electronic devices & storage": "https://toshiba.semicon-storage.com",
    "cambridge gan devices": "https://camgandevices.com",
    "win semiconductors": "https://www.winsemiconductorscorp.com",
    "nexchip": "https://www.nexchip.com.cn",
    "sanan optoelectronics": "https://www.sanan-e.com",
    "tokyo ohka kogyo (tok)": "https://www.tok.co.jp/eng",
    "omnivision group": "https://www.omnivision-group.com",
    "intel corporation": "https://www.intel.com",
    "rohm semiconductor": "https://www.rohm.com",
    "melexis": "https://www.melexis.com",
    "silergy": "https://www.silergy.com",
    "alpha and omega semiconductor": "https://www.aosmd.com",
    "pragmatic semiconductor": "https://www.pragmaticsemi.com",
    "xmos": "https://www.xmos.com",
    "vishay intertechnology": "https://www.vishay.com",
    "chipus microelectronics": "https://www.chipus.com.br",
}


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def _plain(name: str) -> str:
    return re.sub(
        r"\s*\((?:acquired by|subsidiary of|merged into)[^)]*\)\s*$",
        "",
        str(name or ""),
        flags=re.I,
    ).strip()


def _load_landscape(xlsx: Path) -> list[dict]:
    wb = load_workbook(xlsx, data_only=True)
    ws = wb["Landscape"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows[0]]
    landscape: list[dict] = []
    for r in rows[1:]:
        if not r:
            continue
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        if str(d.get("Company") or "").strip():
            landscape.append(d)

    details_by: dict[str, dict] = {}
    if "Company Details" in wb.sheetnames:
        ws = wb["Company Details"]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(h) for h in rows[0]]
        for r in rows[1:]:
            if not r:
                continue
            d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
            key = _norm(_plain(d.get("Brand") or d.get("Company") or ""))
            if key:
                details_by[key] = d

    for row in landscape:
        key = _norm(_plain(row.get("Company") or row.get("Brand") or ""))
        det = details_by.get(key)
        if det:
            brand = str(det.get("Brand") or "").strip()
            company_col = str(det.get("Company") or "").strip()
            if brand:
                row["Brand"] = brand
            if company_col and company_col != brand:
                row["_display_company"] = company_col
            if det.get("Found in"):
                row["Headquarters"] = det["Found in"]
                row["Found in"] = det["Found in"]
            if det.get("Role"):
                row["Role"] = det["Role"]
                row["Distribution Type"] = det["Role"]
        # Fix website if known
        fix = WEBSITE_FIXES.get(key)
        if fix:
            row["Website"] = fix
        # Clear scores so score_expand_rows recalculates from evidence
        for col in (
            "X Score",
            "Y Score",
            "Overall Score",
            "X",
            "Y",
            "Overall",
            "Quadrant",
        ):
            row[col] = ""
        row["Role"] = "Solution Provider"
        row["Distribution Type"] = "Solution Provider"
        row["Industry Category"] = INDUSTRY
        if not str(row.get("Brand") or "").strip():
            row["Brand"] = _plain(str(row.get("Company") or ""))
        # Scoring identity = plain brand/company (not ownership suffix)
        row["Company"] = _plain(str(row.get("Brand") or row.get("Company") or ""))
    return landscape


def _save_ckpt(done: dict[str, dict], audit: dict) -> None:
    CKPT.parent.mkdir(exist_ok=True)
    CKPT.write_text(
        json.dumps({"done": done, "audit": audit, "n": len(done)}, indent=2),
        encoding="utf-8",
    )


def _load_ckpt() -> tuple[dict[str, dict], dict]:
    if not CKPT.exists():
        return {}, {}
    try:
        data = json.loads(CKPT.read_text(encoding="utf-8"))
        return dict(data.get("done") or {}), dict(data.get("audit") or {})
    except Exception:
        return {}, {}


async def main() -> int:
    apply_env_overrides()
    apply_chatgpt_expand_env(root=ROOT)
    os.environ["LLM_PROVIDER"] = (os.getenv("LLM_PROVIDER") or "deepseek").strip() or "deepseek"
    os.environ["EXPAND_XY_SKIP_CRAWL"] = "0"
    os.environ["EXPAND_XY_FORCE_CRAWL"] = "1"
    os.environ.setdefault("EXPAND_XY_CONCURRENT", "3")
    os.environ.setdefault("EXPAND_XY_CRAWL_CONCURRENT", "4")
    os.environ.setdefault("QUADRANT_CRAWL_MAX_PAGES", "25")
    os.environ.setdefault("EXPAND_MARKET_AXIS_LLM", "1")

    xlsx = FOLDER / f"{SLUG}_FINAL.xlsx"
    landscape = _load_landscape(xlsx)
    print(f"Loaded {len(landscape)} companies for full X/Y/Overall rescore", flush=True)

    done, audit = _load_ckpt()
    if done:
        print(f"Resuming checkpoint: {len(done)} already scored", flush=True)

    pending: list[dict] = []
    for row in landscape:
        key = _norm(row.get("Company") or "")
        if key in done and done[key].get("X Score") and done[key].get("Y Score"):
            # restore scores from checkpoint
            for col in ("X Score", "Y Score", "Overall Score", "Quadrant", "Industry Category"):
                if done[key].get(col):
                    row[col] = done[key][col]
            if done[key].get("_evidence_snapshot"):
                row["_evidence_snapshot"] = done[key]["_evidence_snapshot"]
        else:
            pending.append(row)

    print(f"Pending to score: {len(pending)}", flush=True)

    # Score in batches for resilience
    for start in range(0, len(pending), BATCH):
        batch = pending[start : start + BATCH]
        print(
            f"\n=== Batch {start // BATCH + 1}: scoring {len(batch)} "
            f"({start + 1}-{start + len(batch)} of {len(pending)}) ===",
            flush=True,
        )
        scored, batch_audit = await score_expand_rows(
            batch,
            QUERY,
            country="global",
            concurrent=int(os.getenv("EXPAND_XY_CONCURRENT") or "3"),
        )
        for k, v in (batch_audit or {}).items():
            if v and not audit.get(k):
                audit[k] = v
            elif k in {
                "parameter_definitions",
                "x_features",
                "y_features",
                "axis_x",
                "axis_y",
                "x_feature_weights",
                "y_feature_weights",
            }:
                audit[k] = v
        for row in scored:
            key = _norm(row.get("Company") or "")
            done[key] = {
                "Company": row.get("Company"),
                "X Score": row.get("X Score"),
                "Y Score": row.get("Y Score"),
                "Overall Score": row.get("Overall Score"),
                "Industry Category": row.get("Industry Category"),
                # snapshot can be large — keep crawl flag only in meta; scores matter
            }
            # copy scores onto matching landscape row
            for lr in landscape:
                if _norm(lr.get("Company") or "") == key:
                    lr["X Score"] = row.get("X Score")
                    lr["Y Score"] = row.get("Y Score")
                    lr["Overall Score"] = row.get("Overall Score")
                    lr["X"] = row.get("X Score")
                    lr["Y"] = row.get("Y Score")
                    lr["Overall"] = row.get("Overall Score")
                    lr["Industry Category"] = row.get("Industry Category") or INDUSTRY
                    break
        _save_ckpt(done, audit)
        print(f"Checkpoint saved: {len(done)}/{len(landscape)}", flush=True)

    # Ensure every row has numeric scores
    missing = [r for r in landscape if not str(r.get("X Score") or "").strip()]
    if missing:
        print(f"WARNING: {len(missing)} still unscored — retrying once", flush=True)
        scored, batch_audit = await score_expand_rows(
            missing, QUERY, country="global", concurrent=2
        )
        for k, v in (batch_audit or {}).items():
            if v:
                audit[k] = v
        for row in scored:
            key = _norm(row.get("Company") or "")
            for lr in landscape:
                if _norm(lr.get("Company") or "") == key:
                    lr["X Score"] = row.get("X Score")
                    lr["Y Score"] = row.get("Y Score")
                    lr["Overall Score"] = row.get("Overall Score")
                    lr["X"] = row.get("X Score")
                    lr["Y"] = row.get("Y Score")
                    lr["Overall"] = row.get("Overall Score")
                    break
            done[key] = {
                "Company": row.get("Company"),
                "X Score": row.get("X Score"),
                "Y Score": row.get("Y Score"),
                "Overall Score": row.get("Overall Score"),
            }
        _save_ckpt(done, audit)

    # Recompute Overall strictly as (X+Y)/2 and assign quadrants
    xs: list[float] = []
    ys: list[float] = []
    for row in landscape:
        try:
            x = float(row.get("X Score") or row.get("X") or 0)
            y = float(row.get("Y Score") or row.get("Y") or 0)
        except (TypeError, ValueError):
            x, y = 0.0, 0.0
        overall = int(round((x + y) / 2.0))
        row["X Score"] = str(int(round(x)))
        row["Y Score"] = str(int(round(y)))
        row["Overall Score"] = str(overall)
        row["X"] = row["X Score"]
        row["Y"] = row["Y Score"]
        row["Overall"] = row["Overall Score"]
        xs.append(x)
        ys.append(y)
        row["Role"] = "Solution Provider"
        row["Distribution Type"] = "Solution Provider"

    quads, mid_x, mid_y = assign_quadrants_half_median(xs, ys)
    for row, q in zip(landscape, quads):
        row["Quadrant"] = q

    audit["score_midpoints"] = {"x": mid_x, "y": mid_y}
    audit["quadrant_mid_x"] = mid_x
    audit["quadrant_mid_y"] = mid_y
    audit["scored"] = len(landscape)
    audit["overall_formula"] = "(X+Y)/2"
    audit["rows"] = [
        {
            "company": r.get("Company"),
            "x": int(str(r.get("X Score") or 0) or 0),
            "y": int(str(r.get("Y Score") or 0) or 0),
            "overall": int(str(r.get("Overall Score") or 0) or 0),
            "quadrant": r.get("Quadrant"),
        }
        for r in landscape
    ]

    # NEVER zip detail_rows with landscape — to_company_detail_rows sorts by Overall.
    detail_rows = to_company_detail_rows(landscape, QUERY, audit)
    by_key = {_norm(_plain(str(r.get("Company") or ""))): r for r in landscape}
    for det in detail_rows:
        key = _norm(_plain(str(det.get("Brand") or det.get("Company") or "")))
        src = by_key.get(key)
        if not src:
            continue
        det["Role"] = "Solution Provider"
        det["Quadrant"] = src.get("Quadrant")
        det["X"] = int(str(src.get("X Score") or 0) or 0)
        det["Y"] = int(str(src.get("Y Score") or 0) or 0)
        det["Overall"] = int(str(src.get("Overall Score") or 0) or 0)
        if src.get("_display_company"):
            det["Company"] = src["_display_company"]
        brand = str(src.get("Brand") or src.get("Company") or "").strip()
        if brand:
            det["Brand"] = brand

    write_final_xlsx(
        xlsx,
        landscape,
        "Companies",
        {"query": QUERY, "xy_scoring": audit, "full_xy_rescore": True},
        detail_rows=detail_rows,
    )
    extras = export_expand_quadrant_outputs(
        FOLDER, detail_rows, QUERY, country="global", audit=audit, chart_n=20
    )

    (FOLDER / "chatgpt_xy_scores_batch_all.json").write_text(
        json.dumps(audit, indent=2),
        encoding="utf-8",
    )
    csv_path = FOLDER / f"{SLUG}_companies.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(
            f,
            fieldnames=["Brand", "Company", "Role", "Quadrant", "X", "Y", "Overall", "Found in"],
            extrasaction="ignore",
        )
        w.writeheader()
        for d in detail_rows:
            w.writerow(d)

    summary = {
        "total": len(landscape),
        "mid_x": mid_x,
        "mid_y": mid_y,
        "x_range": [min(xs), max(xs)],
        "y_range": [min(ys), max(ys)],
        "html": str(extras.get("html") or ""),
        "top10": sorted(
            [
                {
                    "brand": r.get("Brand") or r.get("Company"),
                    "x": int(str(r.get("X Score") or 0) or 0),
                    "y": int(str(r.get("Y Score") or 0) or 0),
                    "overall": int(str(r.get("Overall Score") or 0) or 0),
                    "quadrant": r.get("Quadrant"),
                }
                for r in landscape
            ],
            key=lambda z: -z["overall"],
        )[:10],
    }
    (OUT / "_audit" / "semi_xy_full_rescore.json").write_text(
        json.dumps(summary, indent=2), encoding="utf-8"
    )
    print(json.dumps(summary, indent=2), flush=True)
    print(f"FINAL -> {xlsx}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
