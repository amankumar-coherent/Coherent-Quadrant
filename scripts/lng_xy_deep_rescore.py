#!/usr/bin/env python3
"""
Deep-crawl + DDGS web-search + DeepSeek evidence-only X/Y/Overall rescore
for Global Liquefied Natural Gas Market.

Rules:
  - Score ONLY from crawl + web-search + agent-verified notes (no invented capacity/SPA)
  - When evidence clearly supports global leadership, use high scores (8–10 → 80–95+)
  - Thin/silent evidence → mid-low (3–5), never invent leadership

Axes (Others / Energy catalog):
  X = Asset & Operating Capability
  Y = Portfolio & Commercial Strategy
  Overall = (X+Y)/2
  Quadrant = absolute median (High/Lower vs cohort mid — matches Coherent definitions)

Env:
  QUAD_RESCORE_LIMIT=N
  QUAD_RESCORE_BATCH=5
  QUAD_RESCORE_CONCURRENT=2
  QUAD_CRAWL_CONCURRENT=3
  QUAD_CRAWL_MAX_PAGES=40
  QUAD_SKIP_CRAWL=1
  QUAD_SKIP_DDGS=1
  QUAD_RESCORE_RESUME=1
"""
from __future__ import annotations

import asyncio
import json
import os
import re
import sys
import time
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.config import Settings
from vendor_intel.placeholders.load_keys import apply_env_overrides
from vendor_intel.pipeline.chatgpt_env import apply_chatgpt_expand_env
from vendor_intel.pipeline.expand_quadrant_score import (
    _deep_crawl_rows,
    export_expand_quadrant_outputs,
    kb_from_landscape_row,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx
from vendor_intel.quadrant.company_kb import build_company_kb, kb_text_blob
from vendor_intel.quadrant.rating_map import assign_quadrants_absolute_median

OUT = ROOT / "output" / "chatgpt_expand"
SLUG = "global_liquefied_natural_gas_market_global"
QUERY = "Global Liquefied Natural Gas Market"
FOLDER = OUT / SLUG
INDUSTRY = "Others / Energy"

AXIS_X = "Asset & Operating Capability"
AXIS_Y = "Portfolio & Commercial Strategy"
X_FEATS = [
    "Asset / Project Portfolio Breadth",
    "Operational Efficiency & Throughput",
    "Reliability & Asset Performance",
    "Safety & Regulatory Compliance",
    "Sustainability & Energy Transition",
]
Y_FEATS = [
    "Geographic & Corridor Reach",
    "Offtake / Customer Contracts",
    "Financial Performance",
    "Growth & Project Pipeline",
    "Partnerships & Business Expansion",
]
X_W = [0.25, 0.25, 0.20, 0.15, 0.15]
Y_W = [0.25, 0.25, 0.20, 0.15, 0.15]

SYSTEM = """You score a company on the Coherent Quadrant for the global LNG market
(liquefaction, regas/FSRU, LNG shipping, integrated IOC/NOC LNG, commodity LNG trading).

Use ONLY the EVIDENCE text provided. Do NOT invent mtpa capacity, SPA volumes, fleet size,
terminals, revenues, partners, or countries absent from evidence.

Axes:
- X = Asset & Operating Capability (portfolio breadth, throughput/ops, reliability, HSE/regulatory, transition/CCS)
- Y = Portfolio & Commercial Strategy (geography/corridors, offtake/customers, financials, growth pipeline, partnerships/M&A)

Scoring each feature 1–10 (discriminating — do not give everyone high scores):
- 9–10: ONLY when evidence shows global LNG leadership AND multi-basin / multi-continent scale
  (e.g. top global exporter/marketer, 40+ mtpa operated or marketed, major multi-project portfolio)
- 7–8: strong evidenced multi-region LNG assets OR large contracted offtake / trading franchise
- 5–6: moderate / regional LNG player with partial evidence
- 3–4: weak or narrow evidence (single terminal niche, thin public footprint)
- 1–2: evidence shows weakness or absence
Single-project terminals / regional utilities: typically X 5–7, Y 4–6 — not 9–10.
Commodity traders / sogo shosha marketers: Y may be higher than X when evidence is trading/offtake heavy.
If evidence is silent on a feature → score 3–5 and grounding=insufficient.
When crawl/search CLEARLY supports global leadership, DO NOT compress to mid-band — use 9–10.

Return JSON only:
{
  "x_features": [{"feature": "...", "score": 1-10, "grounding": "supported|partial|insufficient", "why": "short cite from evidence"}],
  "y_features": [{"feature": "...", "score": 1-10, "grounding": "supported|partial|insufficient", "why": "short cite from evidence"}],
  "evidence_used": true,
  "evidence_quality": "rich|moderate|thin"
}
All five X and five Y features required. Never use grounding=model_knowledge."""

# Agent web-verified notes (Aug 2026) — public sources only
AGENT_WEB_EVIDENCE: dict[str, str] = {
    "qatarenergy": (
        "QatarEnergy / QatarEnergy LNG: Qatar LNG capacity ~77 mtpa; North Field East+South "
        "expansion targeting ~126 mtpa by ~2027 (six trains × ~8 mtpa); further North Field West "
        "plans toward ~142 mtpa by 2030. World's largest LNG exporter; long-term Asian/European SPA network."
    ),
    "cheniere": (
        "Cheniere Energy (NYSE:LNG): ~52 mtpa liquefaction in operation across Sabine Pass (~30+ mtpa) "
        "and Corpus Christi (~21+ mtpa); >9 mtpa under construction; >40 mtpa in permitting. Leading US "
        "LNG exporter; cumulative 4,600+ cargos / 315+ Mt loaded. Full-service (gas procurement, "
        "liquefaction, chartering, delivery); offices Houston/London/Singapore/Beijing/Tokyo/Dubai."
    ),
    "sabine pass lng": (
        "Sabine Pass LNG (Cheniere Partners): >30 mtpa liquefaction production capacity; five tanks, "
        "three marine berths; Creole Trail pipeline interconnection — cornerstone of Cheniere platform."
    ),
    "shell": (
        "Shell: industry's largest integrated LNG portfolio (company claim); equity interests incl. "
        "LNG Canada 40% (14 mtpa Phase 1; first cargo ~Jun 2025), Atlantic LNG Trinidad ~14.5 mtpa "
        "(Shell ~47–51%), Peru LNG 20% (~4.5 mtpa), Hazira/Dragon regas access; strategy add up to "
        "~12 mtpa capacity by 2030 via backfill + selective greenfield; publishes annual LNG Outlook."
    ),
    "totalenergies": (
        "TotalEnergies: world's 3rd-largest LNG player claim; ~44 Mt LNG sold in 2025 incl. ~15.1 Mt "
        "equity production; #1 US LNG exporter claim (~19 Mt / ~18% US production 2025). Growth to "
        "~60 mtpa sales by 2030 (+50% equity+LT purchases). Pipeline: NFE/NFS Qatar, Rio Grande LNG, "
        "ECA Mexico, NLNG T7, Mozambique LNG, Marsa/Ruwais; >20 mtpa European regas access."
    ),
    "venture global": (
        "Venture Global (NYSE:VG): US Gulf LNG exporter; Calcasieu Pass COD Apr 2025 (~11.2–12.4 mtpa "
        "peak); Plaquemines producing/commissioning (Phase 1 COD target Q4 2026; Phase 2 mid-2027); "
        "CP2 under construction (first LNG target H2 2027). Company cites 65–81+ mtpa path across "
        "projects toward 2030; 1,000th cargo milestone; high % of 2026 cargos contracted."
    ),
    "calcasieu pass": (
        "Calcasieu Pass LNG (Venture Global): commercial operations from Apr 2025; DOE peak uprate "
        "to 12.4 mtpa; modular mid-scale trains; one-year COD anniversary with no missed cargos claim."
    ),
    "plaquemines": (
        "Plaquemines LNG (Venture Global): first LNG Dec 2024; Phase 1 COD target Q4 2026; Phase 2 "
        "mid-2027; large multi-train mid-scale project (~28 mtpa expected capacity class)."
    ),
    "cp2 lng": (
        "CP2 LNG (Venture Global): DOE export + FERC approvals 2025; Phase 1 SPAs signed; site work "
        "underway; first LNG targeted H2 2027."
    ),
    "woodside": (
        "Woodside Energy: Pluto LNG + Scarborough Energy Project ~8 mtpa (Pluto T2 ~5 + T1 up to 3); "
        "Scarborough ~98% complete, first LNG cargo target Q4 2026. Louisiana LNG foundation ~16.5 mtpa "
        "(permitted up to ~27.6), ~28% complete, first LNG target 2029."
    ),
    "exxonmobil": (
        "ExxonMobil: major LNG equity/operator roles incl. PNG LNG, Golden Pass (30% with QatarEnergy 70%), "
        "Adriatic LNG JV stakes, Rovuma Area 4 interests — multi-basin LNG producer/partner."
    ),
    "bp": (
        "bp: integrated LNG via equity (e.g. Atlantic LNG Trinidad), Tortue/GTA FLNG (operator), "
        "global LNG trading/marketing franchise and European gas portfolio."
    ),
    "chevron": (
        "Chevron: LNG interests incl. Angola LNG, Wheatstone/Gorgon Australia (major equity), "
        "global upstream+LNG marketing capability."
    ),
    "equinor": (
        "Equinor ASA: Norwegian NOC/IOC with European gas/LNG marketing, Hammerfest LNG (Snøhvit), "
        "and Atlantic basin trading presence."
    ),
    "petronas": (
        "Petronas: Malaysia LNG (MLNG) complex operator/partner; equity in LNG Canada (25%); "
        "global LNG marketing arm; Southeast Asia liquefaction leadership."
    ),
    "adnoc": (
        "ADNOC / ADNOC Gas: UAE LNG growth incl. Ruwais LNG project participation; ADNOC Gas listed "
        "processing/marketing platform; expanding LNG export ambitions."
    ),
    "novatek": (
        "Novatek: Yamal LNG operator/lead; Arctic LNG 2 consortium lead — large Russian Arctic "
        "liquefaction portfolio (sanctions/market-access constraints noted in public coverage)."
    ),
    "jera": (
        "JERA: Japan’s largest LNG buyer/importer; equity in Freeport LNG (~21.9% after JAPEX sale); "
        "TEPCO/Chubu 50/50 JV; extensive long-term SPA portfolio."
    ),
    "sempra": (
        "Sempra Infrastructure: Cameron LNG JV partner; Port Arthur LNG development; Energía Costa Azul; "
        "owned ~70% Sempra / ~20% KKR / ~10% ADIA (KKR-led stake sale process reported)."
    ),
    "freeport lng": (
        "Freeport LNG Development: US Gulf export terminal; ownership Freeport LNG Investments ~63.5%, "
        "JERA ~21.9%, Osaka Gas ~10.8%, JAPEX ~3.8%; multi-train export (~15 mtpa class)."
    ),
    "lng canada": (
        "LNG Canada: Kitimat BC ~14 mtpa Phase 1; ownership Shell 40%, Petronas 25%, PetroChina 15%, "
        "Mitsubishi 15%, KOGAS 5%; first cargo mid-2025."
    ),
    "golden pass": (
        "Golden Pass LNG: QatarEnergy 70% / ExxonMobil 30%; large US Gulf liquefaction under construction."
    ),
    "pavilion energy": (
        "Pavilion Energy: Singapore LNG marketer; acquired by Shell plc — now within Shell LNG portfolio."
    ),
    "vitol": (
        "Vitol: leading independent energy trader with large global LNG cargo trading book and "
        "offtake agreements with US exporters (e.g. Venture Global volumes reported)."
    ),
    "trafigura": (
        "Trafigura: major independent commodity trader with significant LNG trading/marketing activity."
    ),
    "gunvor": (
        "Gunvor: independent energy trader active in LNG cargo trading and European gas markets."
    ),
    "mitsubishi": (
        "Mitsubishi Corporation: sogo shosha with LNG equity (e.g. LNG Canada 15%, Brunei LNG 25%) "
        "and large Japanese LNG marketing/trading franchise."
    ),
    "mitsui": (
        "Mitsui & Co.: sogo shosha LNG equity/marketing across multiple liquefaction projects and "
        "long-term Japanese offtake."
    ),
    "marubeni": (
        "Marubeni: Japanese trading house with LNG project equity and marketing (e.g. Peru LNG interests)."
    ),
    "itochu": (
        "Itochu: Japanese trading company with LNG equity stakes and cargo marketing activities."
    ),
    "eni": (
        "Eni: Coral South FLNG (Mozambique Area 4 operator), Damietta/SEGAS interests, global LNG "
        "equity and marketing portfolio."
    ),
    "conocophillips": (
        "ConocoPhillips: Australia Pacific LNG JV partner; global upstream + LNG offtake/marketing."
    ),
    "santos": (
        "Santos: Gladstone LNG lead JV partner; Australian LNG producer with Asian offtake links."
    ),
    "inpex": (
        "Inpex: Ichthys LNG operator/lead JV — major Australian LNG project (~8+ mtpa class)."
    ),
    "golar": (
        "Golar LNG: FLNG pioneer (Cameroon Hilli Episeyo); floating liquefaction / FLNG solutions."
    ),
    "hoegh": (
        "Höegh Evi (ex Höegh LNG): FSRU/FLNG specialist; owned Aequitas (Höegh family) 50% / "
        "Igneo Infrastructure Partners 50%."
    ),
    "excelerate": (
        "Excelerate Energy: FSRU pioneer with multi-country floating regas operations and LNG logistics."
    ),
    "uniper": (
        "Uniper SE: major European gas/LNG importer/marketer; ~99% owned by Federal Republic of Germany."
    ),
    "centrica": (
        "Centrica: UK energy major with LNG offtake/trading and European gas portfolio."
    ),
}


def _norm(s: str) -> str:
    s = str(s or "").strip().lower()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", s).strip()


def _domain(website: str) -> str:
    raw = (website or "").strip()
    if not raw:
        return ""
    try:
        host = urlparse(raw if "://" in raw else f"https://{raw}").netloc.lower()
        return host[4:] if host.startswith("www.") else host
    except Exception:
        return ""


def _axis_100(feat_scores: list[dict], features: list[str], weights: list[float]) -> int:
    by = {_norm(f.get("feature") or ""): f for f in feat_scores}
    total = 0.0
    for feat, w in zip(features, weights):
        item = by.get(_norm(feat)) or {}
        if not item:
            for k, v in by.items():
                if _norm(feat)[:18] in k or k[:18] in _norm(feat):
                    item = v
                    break
        try:
            s = float(item.get("score") or 4)
        except (TypeError, ValueError):
            s = 4.0
        s = max(1.0, min(10.0, s))
        g = str(item.get("grounding") or "").lower()
        if g == "supported" and 6 <= s <= 8:
            s = min(10.0, s + 1.0)
        if g == "insufficient":
            s = min(s, 5.0)
        total += w * (s / 10.0)
    return int(round(total * 100))


def _web_search(brand: str, website: str, specialty: str) -> str:
    if (os.getenv("QUAD_SKIP_DDGS") or "").lower() in {"1", "true", "yes"}:
        return ""
    from ddgs import DDGS

    queries = [
        f"{brand} LNG liquefaction OR regas OR FSRU OR mtpa capacity",
        f"{brand} LNG SPA OR offtake OR cargo OR terminal",
        f"{brand} LNG revenue OR production OR export OR import",
        f"{brand} LNG project FID OR expansion OR partnership",
        f"{brand} LNG shipping OR fleet OR trading",
    ]
    if specialty:
        queries.append(f"{brand} {specialty[:90]}")
    dom = _domain(website)
    if dom:
        queries.append(f"site:{dom} LNG OR liquefaction OR FSRU OR mtpa")

    chunks: list[str] = []
    seen: set[str] = set()
    try:
        with DDGS() as ddgs:
            for q in queries[:6]:
                try:
                    results = list(ddgs.text(q, max_results=6))
                except Exception:
                    results = []
                for r in results or []:
                    title = str(r.get("title") or "").strip()
                    body = str(r.get("body") or r.get("snippet") or "").strip()
                    href = str(r.get("href") or r.get("link") or "").strip()
                    key = _norm(title + body)[:160]
                    if not key or key in seen:
                        continue
                    seen.add(key)
                    chunks.append(f"- {title}\n  {body}\n  URL: {href}")
                    if len(chunks) >= 20:
                        break
                if len(chunks) >= 20:
                    break
                time.sleep(0.2)
    except Exception as exc:
        chunks.append(f"(web search error: {type(exc).__name__}: {exc})")
    return "\n".join(chunks)


def _snap_text(snap: Any) -> str:
    if not isinstance(snap, dict):
        return ""
    parts: list[str] = []
    data = snap.get("data") if isinstance(snap.get("data"), dict) else {}
    for key in ("company", "business", "financials", "relationships", "intel", "location"):
        block = data.get(key)
        if isinstance(block, dict) and block:
            parts.append(f"[{key}]\n{json.dumps(block, ensure_ascii=False)[:4000]}")
        elif isinstance(block, str) and block.strip():
            parts.append(f"[{key}]\n{block[:4000]}")
    page = str(snap.get("page_text") or data.get("page_text") or "").strip()
    if page:
        parts.append(f"[page_text]\n{page[:14000]}")
    snippets = snap.get("snippets") or data.get("snippets") or []
    if isinstance(snippets, list) and snippets:
        parts.append("[snippets]\n" + "\n".join(str(s)[:400] for s in snippets[:12]))
    return "\n\n".join(parts)


def _deepseek_score(brand: str, role: str, specialty: str, evidence: str) -> dict[str, Any]:
    apply_env_overrides()
    from vendor_intel.placeholders import llm as llm_mod

    os.environ["LLM_PROVIDER"] = "deepseek"
    llm_mod.LLM_PROVIDER = "deepseek"
    if not llm_mod.is_configured():
        raise RuntimeError("DeepSeek not configured (DEEPSEEK_API_KEY missing)")

    user = f"""Market: {QUERY}
Brand: {brand}
Role: {role}  (Brand=LNG producer/terminal/FSRU/operator; Marketer=LNG trader/sogo-shosha marketer)
Known specialty (may be incomplete): {specialty or 'n/a'}

X axis ({AXIS_X}) features: {X_FEATS}
Y axis ({AXIS_Y}) features: {Y_FEATS}

EVIDENCE (deep crawl + web search — cite only this; never invent):
{evidence[:24000]}

Score now. JSON only."""
    parsed = llm_mod.llm_complete_json(SYSTEM, user, max_tokens=2400)
    if not isinstance(parsed, dict):
        raise ValueError(f"no JSON object from DeepSeek for {brand}")
    if parsed.get("status") in {"llm_failed", "placeholder"}:
        raise ValueError(
            f"DeepSeek failed for {brand}: {parsed.get('error') or parsed.get('status')}"
        )
    if not (parsed.get("x_features") and parsed.get("y_features")):
        raise ValueError(f"missing x_features/y_features for {brand}")
    return parsed


def _load_rows(*, resume: bool) -> list[dict]:
    xlsx = FOLDER / f"{SLUG}_FINAL.xlsx"
    wb = load_workbook(xlsx, data_only=True)
    ws = wb["Landscape"]
    rows_raw = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows_raw[0]]

    details_by: dict[str, dict] = {}
    if "Company Details" in wb.sheetnames:
        dws = wb["Company Details"]
        drows = list(dws.iter_rows(values_only=True))
        dhdr = [str(h) for h in drows[0]]
        for r in drows[1:]:
            if not r:
                continue
            d = {dhdr[i]: ("" if r[i] is None else r[i]) for i in range(len(dhdr)) if i < len(r)}
            key = _norm(d.get("Brand") or "")
            if key:
                details_by[key] = d

    kept: list[dict] = []
    for r in rows_raw[1:]:
        if not r:
            continue
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        brand = str(d.get("Company") or d.get("Brand") or "").strip()
        if not brand:
            continue
        key = _norm(brand)
        det = details_by.get(key) or {}
        if det.get("Brand"):
            brand = str(det["Brand"]).strip()
            key = _norm(brand)
            d["Brand"] = brand
        if det.get("Company") and (
            str(det["Company"]).startswith("(") or _norm(det.get("Company")) != key
        ):
            d["_display_company"] = det["Company"]
        if det.get("Found in"):
            d["Headquarters"] = det["Found in"]
            d["Found in"] = det["Found in"]
        if det.get("Role") in {"Brand", "Marketer"}:
            d["Role"] = det["Role"]
            d["Distribution Type"] = det["Role"]

        prior_x = str(d.get("X Score") or det.get("X") or "").strip()
        prior_y = str(d.get("Y Score") or det.get("Y") or "").strip()
        prior_q = str(d.get("Quadrant") or det.get("Quadrant") or "").strip()
        prior_o = str(d.get("Overall Score") or det.get("Overall") or "").strip()
        deep_done = "deep-crawl scored" in str(d.get("Summary") or "").lower()

        d["Brand"] = brand
        d["Company"] = brand
        d["Role"] = d.get("Role") if d.get("Role") in {"Brand", "Marketer"} else "Brand"
        d["Distribution Type"] = d["Role"]
        d["Industry Category"] = INDUSTRY

        if resume and prior_x and prior_y and deep_done:
            d["X Score"] = prior_x
            d["Y Score"] = prior_y
            d["Overall Score"] = prior_o or str(
                round((float(prior_x) + float(prior_y)) / 2)
            )
            d["Quadrant"] = prior_q
            d["_already_scored"] = True
        else:
            for col in ("X Score", "Y Score", "Overall Score", "Quadrant"):
                d[col] = ""
            d["_already_scored"] = False
        kept.append(d)
    return kept


def _save(kept: list[dict], *, partial: bool = False, meta: dict | None = None) -> None:
    xs, ys = [], []
    for r in kept:
        try:
            xs.append(float(r.get("X Score") or 50))
        except (TypeError, ValueError):
            xs.append(50.0)
        try:
            ys.append(float(r.get("Y Score") or 50))
        except (TypeError, ValueError):
            ys.append(50.0)
    scored_idx = [i for i, r in enumerate(kept) if str(r.get("X Score") or "").strip()]
    if len(scored_idx) >= 4:
        sxs = [xs[i] for i in scored_idx]
        sys_ = [ys[i] for i in scored_idx]
        quads, mid_x, mid_y = assign_quadrants_absolute_median(sxs, sys_)
        for j, i in enumerate(scored_idx):
            kept[i]["Quadrant"] = quads[j]
            kept[i]["Overall Score"] = round((xs[i] + ys[i]) / 2)
    else:
        mid_x = mid_y = 50.0

    audit = {
        "axis_x": AXIS_X,
        "axis_y": AXIS_Y,
        "x_features": X_FEATS,
        "y_features": Y_FEATS,
        "x_feature_weights": X_W,
        "y_feature_weights": Y_W,
        "mid_x": mid_x,
        "mid_y": mid_y,
        "method": "smart_crawl deep crawl + DDGS + DeepSeek evidence-only (no hallucination)",
        "partial": partial,
        **(meta or {}),
    }
    detail_rows = to_company_detail_rows(kept, QUERY, audit)
    for det in detail_rows:
        key = _norm(det.get("Brand") or "")
        src = next((r for r in kept if _norm(r.get("Brand") or "") == key), None)
        if not src:
            continue
        det["Role"] = src.get("Role") or "Brand"
        if src.get("_display_company"):
            det["Company"] = src["_display_company"]
        elif not str(det.get("Company") or "").startswith("("):
            if not (det.get("Company") and _norm(det.get("Company")) != key):
                det["Company"] = det.get("Brand") or det.get("Company")

    xlsx = FOLDER / f"{SLUG}_FINAL.xlsx"
    write_final_xlsx(
        xlsx,
        kept,
        "Companies",
        {"query": QUERY, "xy_scoring": audit},
        detail_rows=detail_rows,
    )
    if not partial:
        export_expand_quadrant_outputs(
            FOLDER, detail_rows, QUERY, country="global", audit=audit, chart_n=20
        )


async def _kb_evidence(row: dict) -> str:
    from vendor_intel.pipeline.expand_quadrant_score import _scoring_row

    settings = Settings.load()
    score_row = _scoring_row(row)
    try:
        kb = await build_company_kb(
            score_row, market=QUERY, settings=settings, do_search=False
        )
    except Exception:
        kb = kb_from_landscape_row(row)
    blob = kb_text_blob(kb, max_chars=20000)
    snap_extra = _snap_text(row.get("_evidence_snapshot"))
    if snap_extra and snap_extra not in blob:
        blob = f"{blob}\n\n=== SNAPSHOT ===\n{snap_extra}"[:24000]
    return blob


def _agent_evidence(brand: str) -> str:
    # Strip parentheticals so "Mozambique LNG (TotalEnergies)" does not inherit
    # TotalEnergies portfolio evidence.
    n = _norm(re.sub(r"\([^)]*\)", " ", brand))
    best = ""
    best_len = -1
    for stem, text in AGENT_WEB_EVIDENCE.items():
        sn = _norm(stem)
        if not sn:
            continue
        # Prefer whole-token / prefix match over loose substring on short stems
        hit = (
            n == sn
            or n.startswith(sn + " ")
            or n.endswith(" " + sn)
            or f" {sn} " in f" {n} "
        )
        if len(sn) >= 8:
            hit = hit or sn in n
        if hit and len(sn) > best_len:
            best = text
            best_len = len(sn)
    return best


async def _score_one(row: dict, sem: asyncio.Semaphore) -> dict[str, Any]:
    brand = str(row.get("Brand") or "")
    async with sem:
        specialty = str(row.get("Specialty Focus") or row.get("Core Categories") or "")
        website = str(row.get("Website") or "")
        role = str(row.get("Role") or "Brand")
        crawl_txt = await _kb_evidence(row)
        search_txt = await asyncio.to_thread(_web_search, brand, website, specialty)
        agent_txt = _agent_evidence(brand)
        landscape_bits = (
            f"HQ: {row.get('Headquarters') or row.get('Found in') or ''}\n"
            f"Presence: {row.get('Operational Presence') or ''}\n"
            f"Categories: {row.get('Core Categories') or ''}\n"
            f"Specialty: {specialty}\n"
            f"Ownership: {row.get('Ownership') or ''}\n"
            f"Employees: {row.get('Employees') or ''}\n"
            f"Website: {website}\n"
            f"Crawled: {bool(row.get('_evidence_snapshot'))}\n"
        )
        evidence = (
            f"=== LANDSCAPE FACTS ===\n{landscape_bits}\n"
            f"=== AGENT WEB-VERIFIED NOTES ===\n{agent_txt or '(none for this brand)'}\n"
            f"=== DEEP CRAWL / KB ===\n{crawl_txt or '(no crawl text)'}\n"
            f"=== WEB SEARCH ===\n{search_txt or '(no results)'}\n"
        )
        evid_chars = len(crawl_txt) + len(search_txt) + len(agent_txt)
        if evid_chars < 120:
            return {
                "brand": brand,
                "ok": False,
                "error": "insufficient_web_evidence",
                "x": 40,
                "y": 40,
                "evidence_chars": evid_chars,
                "crawled": bool(row.get("_evidence_snapshot")),
            }
        try:
            payload = await asyncio.to_thread(
                _deepseek_score, brand, role, specialty, evidence
            )
            x = _axis_100(payload.get("x_features") or [], X_FEATS, X_W)
            y = _axis_100(payload.get("y_features") or [], Y_FEATS, Y_W)
            return {
                "brand": brand,
                "ok": True,
                "x": x,
                "y": y,
                "overall": round((x + y) / 2),
                "payload": payload,
                "evidence_chars": evid_chars,
                "crawled": bool(row.get("_evidence_snapshot")),
                "evidence_quality": payload.get("evidence_quality"),
            }
        except Exception as exc:
            return {
                "brand": brand,
                "ok": False,
                "error": f"{type(exc).__name__}: {exc}",
                "x": 40,
                "y": 40,
                "evidence_chars": evid_chars,
                "crawled": bool(row.get("_evidence_snapshot")),
            }


async def main() -> int:
    apply_env_overrides()
    apply_chatgpt_expand_env()
    os.environ["LLM_PROVIDER"] = "deepseek"
    os.environ["QUAD_ALLOW_MODEL_KNOWLEDGE"] = "0"
    from vendor_intel.placeholders import llm as llm_mod

    llm_mod.LLM_PROVIDER = "deepseek"
    if not llm_mod.is_configured():
        print("ERROR: DeepSeek not configured", flush=True)
        return 2

    resume = (os.getenv("QUAD_RESCORE_RESUME") or "").lower() in {"1", "true", "yes"}
    kept = _load_rows(resume=resume)
    limit = int(os.getenv("QUAD_RESCORE_LIMIT") or "0")

    batch_n = int(os.getenv("QUAD_RESCORE_BATCH") or "5")
    concurrent = int(os.getenv("QUAD_RESCORE_CONCURRENT") or "2")
    crawl_concurrent = int(os.getenv("QUAD_CRAWL_CONCURRENT") or "3")
    max_pages = int(os.getenv("QUAD_CRAWL_MAX_PAGES") or "40")
    skip_crawl = (os.getenv("QUAD_SKIP_CRAWL") or "").lower() in {"1", "true", "yes"}

    need = [r for r in kept if not r.get("_already_scored")]
    if limit > 0:
        need = need[:limit]
        print(f"LIMIT {limit} (keeping all {len(kept)} rows in FINAL)", flush=True)

    print(
        f"LNG deep-crawl evidence rescore {len(need)}/{len(kept)} "
        f"(batch={batch_n}, score_concurrent={concurrent}, "
        f"crawl_concurrent={crawl_concurrent}, max_pages={max_pages}, "
        f"skip_crawl={skip_crawl}, resume={resume})",
        flush=True,
    )

    settings = Settings.load()
    try:
        object.__setattr__(settings, "quadrant_crawl_max_pages", max_pages)
        object.__setattr__(settings, "quadrant_crawl_mode", "business")
    except Exception:
        pass

    os.environ["EXPAND_XY_CRAWL_CONCURRENT"] = str(crawl_concurrent)

    audit_log: list[dict] = []
    audit_path = OUT / "_audit" / "lng_xy_deep_rescore.json"
    if resume and audit_path.exists():
        try:
            prior = json.loads(audit_path.read_text(encoding="utf-8"))
            audit_log = list(prior.get("rows") or [])
        except Exception:
            pass

    by_brand = {_norm(r.get("Brand") or ""): r for r in kept}
    sem = asyncio.Semaphore(concurrent)

    for start in range(0, len(need), batch_n):
        batch = need[start : start + batch_n]
        print(
            f"\n=== Batch {start // batch_n + 1}: "
            f"{start + 1}-{start + len(batch)} / {len(need)} ===",
            flush=True,
        )
        if not skip_crawl:
            print(f"  deep-crawling {len(batch)} sites…", flush=True)
            stats = await _deep_crawl_rows(batch, country="global", settings=settings)
            print(
                f"  crawl ok={stats.get('ok')} fail={stats.get('fail')} "
                f"skip_no_domain={stats.get('skipped_no_domain')}",
                flush=True,
            )
        results = await asyncio.gather(*[_score_one(r, sem) for r in batch])
        for res in results:
            brand = res["brand"]
            row = by_brand.get(_norm(brand))
            if not row:
                continue
            row["X Score"] = res["x"]
            row["Y Score"] = res["y"]
            row["Overall Score"] = res.get("overall") or round((res["x"] + res["y"]) / 2)
            note = "ok" if res.get("ok") else res.get("error")
            row["Summary"] = (
                f"{brand} — deep-crawl scored X={res['x']} Y={res['y']} "
                f"({note}; evid={res.get('evidence_chars')}; "
                f"crawled={res.get('crawled')}; q={res.get('evidence_quality')})"
            )
            row["_already_scored"] = True
            audit_log.append(
                {
                    "brand": brand,
                    "ok": res.get("ok"),
                    "x": res["x"],
                    "y": res["y"],
                    "error": res.get("error"),
                    "evidence_chars": res.get("evidence_chars"),
                    "crawled": res.get("crawled"),
                    "evidence_quality": res.get("evidence_quality"),
                    "features": res.get("payload"),
                }
            )
            print(
                f"  {brand}: X={res['x']} Y={res['y']} "
                f"{'OK' if res.get('ok') else 'FALLBACK ' + str(res.get('error'))} "
                f"crawl={res.get('crawled')} evid={res.get('evidence_chars')} "
                f"q={res.get('evidence_quality')}",
                flush=True,
            )
        all_scored_so_far = all(str(r.get("X Score") or "").strip() for r in kept)
        _save(
            kept,
            partial=not all_scored_so_far,
            meta={
                "scored": sum(1 for r in kept if str(r.get("X Score") or "").strip())
            },
        )
        audit_path.parent.mkdir(parents=True, exist_ok=True)
        audit_path.write_text(
            json.dumps(
                {
                    "scored": len(audit_log),
                    "rows": audit_log,
                    "method": "deep_crawl_ddgs_deepseek_evidence_only",
                    "axes": {"x": AXIS_X, "y": AXIS_Y},
                },
                indent=2,
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

    all_scored = all(str(r.get("X Score") or "").strip() for r in kept)
    _save(
        kept,
        partial=not all_scored,
        meta={
            "scored": sum(1 for r in kept if str(r.get("X Score") or "").strip()),
            "complete": all_scored,
            "total": len(kept),
        },
    )

    roles = dict(Counter(r.get("Role") for r in kept))
    quads = dict(Counter(r.get("Quadrant") for r in kept if r.get("Quadrant")))
    ok_n = sum(1 for a in audit_log if a.get("ok"))
    crawl_ok = sum(1 for a in audit_log if a.get("crawled"))
    xs = [a.get("x") or 0 for a in audit_log if a.get("ok")]
    ys = [a.get("y") or 0 for a in audit_log if a.get("ok")]
    md = OUT / "_audit" / "lng_xy_deep_rescore.md"
    lines = [
        "# Deep-crawl X/Y/Overall rescore — Global LNG",
        "",
        "Method: **smart_crawl deep site crawl + DDGS web search + DeepSeek evidence-only**.",
        "No invented mtpa/SPA — high scores only when evidence supports.",
        "",
        f"Scored this run: **{len(audit_log)}** · OK **{ok_n}** · crawled **{crawl_ok}**",
        f"Landscape complete: **{all_scored}** "
        f"({sum(1 for r in kept if str(r.get('X Score') or '').strip())}/{len(kept)})",
        "",
        f"Axes: **X**={AXIS_X} · **Y**={AXIS_Y} · Overall=(X+Y)/2",
        f"OK score ranges: X {min(xs) if xs else '-'}–{max(xs) if xs else '-'} · "
        f"Y {min(ys) if ys else '-'}–{max(ys) if ys else '-'}",
        f"Roles {roles} · Quadrants {quads}",
        "",
        "## Top Overall (this run)",
        "",
        "| Brand | X | Y | Overall | Crawled | Quality |",
        "|---|---:|---:|---:|:---:|---|",
    ]
    for a in sorted(
        audit_log, key=lambda z: -((z.get("x") or 0) + (z.get("y") or 0))
    )[:50]:
        o = round(((a.get("x") or 0) + (a.get("y") or 0)) / 2)
        lines.append(
            f"| {a['brand']} | {a.get('x')} | {a.get('y')} | {o} | "
            f"{a.get('crawled')} | {a.get('evidence_quality') or a.get('ok')} |"
        )
    md.write_text("\n".join(lines), encoding="utf-8")
    print(
        f"\nDONE ok={ok_n}/{len(audit_log)} crawled={crawl_ok} "
        f"all_scored={all_scored} quads={quads}",
        flush=True,
    )
    print(f"audit -> {md}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
