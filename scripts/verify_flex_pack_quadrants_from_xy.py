#!/usr/bin/env python3
"""Verify Quadrant from real X/Y only — no invented scores or labels.

Definitions (Coherent Quadrant):
  Leaders:      High Solution Capability (X) + High Business Strategy (Y)
  Challengers:  Lower X + High Y
  Trailblazers: High X + Lower Y
  Emerging:     Lower X + Lower Y

Uses absolute cohort median on evidence-scored X/Y (falls back to half-median
only if a side would be empty). Never invents X/Y.
"""
from __future__ import annotations

import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx
from vendor_intel.quadrant.rating_map import assign_quadrants_absolute_median

SLUG = "global_flexible_packaging_market_global"
QUERY = "Global Flexible Packaging Market"
OUT = ROOT / "output" / "chatgpt_expand" / SLUG
XLSX = OUT / f"{SLUG}_FINAL.xlsx"
AUDIT = (
    ROOT
    / "output"
    / "chatgpt_expand"
    / "_audit"
    / f"{SLUG}_quadrant_verify_post_xy.json"
)
VALID = {"Leaders", "Challengers", "Trailblazers", "Emerging Players"}


def _norm(s: str) -> str:
    s = str(s or "").strip().lower()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", s).strip()


def _f(v, default=0.0) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def main() -> None:
    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Landscape"]
    rows_raw = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows_raw[0]]
    landscape: list[dict] = []
    for r in rows_raw[1:]:
        d = {
            hdr[i]: ("" if r[i] is None else str(r[i]))
            for i in range(len(hdr))
            if i < len(r)
        }
        if d.get("Company"):
            landscape.append(d)

    details_by: dict[str, dict] = {}
    if "Company Details" in wb.sheetnames:
        dws = wb["Company Details"]
        drows = list(dws.iter_rows(values_only=True))
        dhdr = [str(h) for h in drows[0]]
        for r in drows[1:]:
            d = {
                dhdr[i]: ("" if r[i] is None else str(r[i]))
                for i in range(len(dhdr))
                if i < len(r)
            }
            key = _norm(d.get("Brand") or "")
            if key:
                details_by[key] = d

    before: list[dict] = []
    for row in landscape:
        brand = str(row.get("Company") or "").strip()
        det = details_by.get(_norm(brand)) or {}
        if det.get("Brand"):
            brand = str(det["Brand"]).strip()
        row["Brand"] = brand
        row["Company"] = brand
        if det.get("Company") and (
            str(det["Company"]).startswith("(")
            or _norm(det.get("Company")) != _norm(brand)
        ):
            row["_display_company"] = det["Company"]
        if det.get("Role") in ("Brand", "Marketer"):
            row["Role"] = det["Role"]
            row["Distribution Type"] = det["Role"]
        # Prefer evidence scores already on Landscape / Details — do NOT invent
        x = _f(row.get("X Score") or row.get("X") or det.get("X"))
        y = _f(row.get("Y Score") or row.get("Y") or det.get("Y"))
        o = _f(row.get("Overall Score") or row.get("Overall") or det.get("Overall"))
        if not o and (x or y):
            o = round((x + y) / 2)
        old_q = str(row.get("Quadrant") or det.get("Quadrant") or "")
        row["X Score"] = int(round(x))
        row["Y Score"] = int(round(y))
        row["Overall Score"] = int(round(o)) if o else int(round((x + y) / 2))
        row["Found in"] = str(
            row.get("Headquarters") or det.get("Found in") or row.get("Found in") or ""
        )
        before.append({"brand": brand, "q": old_q, "x": x, "y": y})

    xs = [_f(r.get("X Score")) for r in landscape]
    ys = [_f(r.get("Y Score")) for r in landscape]
    # Definition-true placement from real scores
    quads, mid_x, mid_y = assign_quadrants_absolute_median(
        [int(x) for x in xs], [int(y) for y in ys]
    )

    changes = []
    invalid_fixed = []
    for row, q, prev in zip(landscape, quads, before):
        old = prev["q"]
        if old not in VALID and old:
            invalid_fixed.append({"brand": row["Brand"], "from": old, "to": q})
        if old != q:
            changes.append(
                {
                    "brand": row["Brand"],
                    "from": old or "(blank)",
                    "to": q,
                    "x": prev["x"],
                    "y": prev["y"],
                    "role": row.get("Role"),
                }
            )
        row["Quadrant"] = q
        row["X"] = row["X Score"]
        row["Y"] = row["Y Score"]
        row["Overall"] = row["Overall Score"]

    audit: dict = {
        "axis_x": "Packaging Solution Capability",
        "axis_y": "Commercial & Growth Strategy",
        "quadrant_method": "absolute_median_from_evidence_xy",
        "mid_x": mid_x,
        "mid_y": mid_y,
        "quadrant_definitions": {
            "Leaders": "High Solution Capability + High Business Strategy",
            "Challengers": "Lower Solution Capability + High Business Strategy",
            "Trailblazers": "High Solution Capability + Lower Business Strategy",
            "Emerging Players": "Lower Solution Capability + Lower Business Strategy",
        },
        "note": "X/Y unchanged (evidence deep-crawl scores). Quadrant recomputed only.",
    }
    ap = OUT / "chatgpt_expand_batch_all_audit.json"
    if ap.exists():
        try:
            prior = json.loads(ap.read_text(encoding="utf-8")).get("xy_scoring") or {}
            for k in ("x_features", "y_features", "parameter_definitions"):
                if prior.get(k):
                    audit[k] = prior[k]
        except Exception:
            pass

    details = to_company_detail_rows(landscape, QUERY, audit)
    by = {_norm(r.get("Brand") or r.get("Company") or ""): r for r in landscape}
    for d in details:
        src = by.get(_norm(d.get("Brand") or ""))
        if not src:
            continue
        d["Role"] = src.get("Role") or "Brand"
        d["Quadrant"] = src.get("Quadrant")
        d["X"] = src.get("X Score")
        d["Y"] = src.get("Y Score")
        d["Overall"] = src.get("Overall Score")
        d["Found in"] = src.get("Found in") or d.get("Found in")
        if src.get("_display_company"):
            d["Company"] = src["_display_company"]

    write_final_xlsx(
        XLSX,
        landscape,
        "Companies",
        {
            "query": QUERY,
            "quadrant_verify": "2026-08-14-absolute-median",
            "mid_x": mid_x,
            "mid_y": mid_y,
            "changes": len(changes),
        },
        detail_rows=details,
    )
    extras = export_expand_quadrant_outputs(
        OUT, details, QUERY, country="global", audit=audit, chart_n=20
    )

    by_q: dict[str, list] = {q: [] for q in VALID}
    for row in landscape:
        by_q.setdefault(str(row["Quadrant"]), []).append(row)
    top = {}
    for q, items in by_q.items():
        items.sort(key=lambda r: -_f(r.get("Overall Score")))
        top[q] = [
            {
                "brand": r.get("Brand"),
                "x": _f(r.get("X Score")),
                "y": _f(r.get("Y Score")),
                "overall": _f(r.get("Overall Score")),
                "role": r.get("Role"),
            }
            for r in items[:12]
        ]

    report = {
        "n": len(landscape),
        "mid_x": mid_x,
        "mid_y": mid_y,
        "method": "absolute_median_from_evidence_xy",
        "counts_before": dict(Counter(p["q"] for p in before)),
        "counts_after": dict(Counter(quads)),
        "changes": changes,
        "invalid_fixed": invalid_fixed,
        "by_quadrant_top": top,
        "html": str(extras.get("html") or ""),
        "definition_check": {
            "Leaders_rule": f"X>={mid_x:.1f} and Y>={mid_y:.1f}",
            "Challengers_rule": f"X<{mid_x:.1f} and Y>={mid_y:.1f}",
            "Trailblazers_rule": f"X>={mid_x:.1f} and Y<{mid_y:.1f}",
            "Emerging_rule": f"X<{mid_x:.1f} and Y<{mid_y:.1f}",
        },
    }
    AUDIT.parent.mkdir(parents=True, exist_ok=True)
    AUDIT.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")

    print(
        f"DONE n={len(landscape)} mid=({mid_x:.1f},{mid_y:.1f}) "
        f"changes={len(changes)} invalid_fixed={len(invalid_fixed)}"
    )
    print("before", report["counts_before"])
    print("after ", report["counts_after"])
    print("rules", report["definition_check"])
    for q in ("Leaders", "Challengers", "Trailblazers", "Emerging Players"):
        print(f"[{q}]")
        for t in top.get(q, [])[:6]:
            print(
                f"  O={t['overall']:.0f} X={t['x']:.0f} Y={t['y']:.0f} "
                f"[{t['role']}] {t['brand']}"
            )
    # Sanity: every Leaders row must be high-high
    bad = []
    for row, q in zip(landscape, quads):
        x, y = _f(row["X Score"]), _f(row["Y Score"])
        ok = (
            (q == "Leaders" and x >= mid_x and y >= mid_y)
            or (q == "Challengers" and x < mid_x and y >= mid_y)
            or (q == "Trailblazers" and x >= mid_x and y < mid_y)
            or (q == "Emerging Players" and x < mid_x and y < mid_y)
        )
        if not ok:
            bad.append((row["Brand"], q, x, y))
    print("definition_violations", len(bad))
    for b in bad[:10]:
        print("  BAD", b)
    print("audit ->", AUDIT)


if __name__ == "__main__":
    main()
