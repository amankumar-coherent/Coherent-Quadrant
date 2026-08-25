#!/usr/bin/env python3
"""Remove Terns Pharmaceuticals and regenerate GLP-1 FINAL + HTML with fixed chart cells."""
from __future__ import annotations

import csv
import json
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx
from vendor_intel.quadrant.rating_map import assign_quadrants_absolute_median

SLUG = "global_glp_1_receptor_agonist_market_global"
QUERY = "Global GLP-1 Receptor Agonist Market"
OUT = ROOT / "output" / "chatgpt_expand" / SLUG
XLSX = OUT / f"{SLUG}_FINAL.xlsx"
DROP = {"terns pharmaceuticals"}


def _norm(s: str) -> str:
    return " ".join(str(s or "").lower().split())


def main() -> int:
    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Landscape"]
    rows_raw = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows_raw[0]]

    details_by = {}
    if "Company Details" in wb.sheetnames:
        dws = wb["Company Details"]
        drows = list(dws.iter_rows(values_only=True))
        dhdr = [str(h) for h in drows[0]]
        for r in drows[1:]:
            if not r or not r[0]:
                continue
            d = {dhdr[i]: ("" if r[i] is None else r[i]) for i in range(len(dhdr))}
            details_by[_norm(d.get("Brand") or "")] = d

    kept = []
    dropped = []
    for r in rows_raw[1:]:
        if not r or not r[0]:
            continue
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(min(len(hdr), len(r)))}
        brand = str(d.get("Company") or d.get("Brand") or "").strip()
        if _norm(brand) in DROP or any(x in _norm(brand) for x in DROP):
            dropped.append(brand)
            continue
        det = details_by.get(_norm(brand)) or {}
        if det.get("Brand"):
            brand = str(det["Brand"]).strip()
            d["Brand"] = brand
        if det.get("Company") and str(det["Company"]).startswith("("):
            d["_display_company"] = det["Company"]
        if det.get("Found in"):
            d["Headquarters"] = det["Found in"]
            d["Found in"] = det["Found in"]
        if det.get("Role") in {"Brand", "Marketer"}:
            d["Role"] = det["Role"]
            d["Distribution Type"] = det["Role"]
        d["Company"] = brand
        d["Brand"] = brand
        d["X Score"] = det.get("X") if det.get("X") not in (None, "") else d.get("X Score")
        d["Y Score"] = det.get("Y") if det.get("Y") not in (None, "") else d.get("Y Score")
        d["Overall Score"] = (
            det.get("Overall") if det.get("Overall") not in (None, "") else d.get("Overall Score")
        )
        kept.append(d)

    xs = [float(r.get("X Score") or 50) for r in kept]
    ys = [float(r.get("Y Score") or 50) for r in kept]
    quads, mid_x, mid_y = assign_quadrants_absolute_median(xs, ys)
    for i, r in enumerate(kept):
        r["Quadrant"] = quads[i]
        r["Overall Score"] = round((xs[i] + ys[i]) / 2)
        r["X Score"] = int(round(xs[i])) if float(xs[i]).is_integer() else xs[i]
        r["Y Score"] = int(round(ys[i])) if float(ys[i]).is_integer() else ys[i]

    audit = {
        "axis_x": "Therapeutic & Manufacturing Capability",
        "axis_y": "Commercial Reach & Pipeline Strategy",
        "industry_group": "Healthcare",
        "industry_category": "Pharmaceutical",
        "mid_x": mid_x,
        "mid_y": mid_y,
        "method": "absolute median; chart cells cleared of crosshair gap",
        "dropped": dropped,
    }
    details = to_company_detail_rows(kept, QUERY, audit)
    for det in details:
        key = _norm(det.get("Brand") or "")
        src = next((r for r in kept if _norm(r.get("Brand") or "") == key), None)
        if not src:
            continue
        det["Role"] = src.get("Role") or "Brand"
        if src.get("_display_company"):
            det["Company"] = src["_display_company"]
        elif not str(det.get("Company") or "").startswith("("):
            det["Company"] = det.get("Brand") or det.get("Company")

    write_final_xlsx(
        XLSX,
        kept,
        "Companies",
        {"query": QUERY, "xy_scoring": audit, "dropped": dropped},
        detail_rows=details,
    )
    export_expand_quadrant_outputs(
        OUT, details, QUERY, country="global", audit=audit, chart_n=20
    )
    csv_path = OUT / f"{SLUG}_companies.csv"
    with csv_path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(
            f,
            fieldnames=["Brand", "Company", "Role", "Quadrant", "X", "Y", "Overall", "Found in"],
        )
        w.writeheader()
        for d in details:
            w.writerow({k: d.get(k, "") for k in w.fieldnames})

    audit_path = ROOT / "output" / "chatgpt_expand" / "_audit" / "glp1_drop_terns_fix_chart.json"
    audit_path.parent.mkdir(parents=True, exist_ok=True)
    audit_path.write_text(
        json.dumps(
            {
                "dropped": dropped,
                "n": len(kept),
                "mid_x": mid_x,
                "mid_y": mid_y,
                "quadrants": dict(Counter(r.get("Quadrant") for r in kept)),
                "roles": dict(Counter(r.get("Role") for r in kept)),
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    print(f"dropped={dropped} n={len(kept)} mid_x={mid_x:.1f} mid_y={mid_y:.1f}")
    print("quadrants", Counter(r.get("Quadrant") for r in kept))
    print("html", OUT / f"{SLUG}_report.html")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
