"""Verify/fix Quadrant column vs X/Y half-median + Coherent definitions."""
from __future__ import annotations

import csv
import json
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx
from vendor_intel.quadrant.rating_map import assign_quadrants_half_median

OUT = ROOT / "output" / "chatgpt_expand"
SLUG = "global_semiconductor_market_global"
QUERY = "Global Semiconductor Market"
FOLDER = OUT / SLUG


def _load_landscape(xlsx: Path) -> list[dict]:
    wb = load_workbook(xlsx, data_only=True)
    ws = wb["Landscape"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows[0]]
    landscape: list[dict] = []
    for r in rows[1:]:
        if not r:
            continue
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        if str(d.get("Company") or "").strip():
            landscape.append(d)

    details_by: dict[str, dict] = {}
    if "Company Details" in wb.sheetnames:
        ws = wb["Company Details"]
        rows = list(ws.iter_rows(values_only=True))
        hdr = [str(h) for h in rows[0]]
        for r in rows[1:]:
            if not r:
                continue
            d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
            key = str(d.get("Brand") or d.get("Company") or "").strip().lower()
            if key:
                details_by[key] = d

    for row in landscape:
        key = str(row.get("Brand") or row.get("Company") or "").strip().lower()
        # strip ownership suffix for match
        plain = key.split("(")[0].strip()
        det = details_by.get(key) or details_by.get(plain)
        if not det:
            for dk, dv in details_by.items():
                if plain and (plain in dk or dk.startswith(plain)):
                    det = dv
                    break
        if not det:
            continue
        for col in ("Quadrant", "X", "Y", "Overall", "Role"):
            if det.get(col) not in (None, ""):
                row[col] = det[col]
                if col == "X":
                    row["X Score"] = det[col]
                elif col == "Y":
                    row["Y Score"] = det[col]
                elif col == "Overall":
                    row["Overall Score"] = det[col]
        if det.get("Found in"):
            row["Headquarters"] = det["Found in"]
            row["Found in"] = det["Found in"]
        brand = str(det.get("Brand") or "").strip()
        company_col = str(det.get("Company") or "").strip()
        if brand:
            row["Brand"] = brand
        if company_col and company_col != brand:
            row["_display_company"] = company_col
        if det.get("Role"):
            row["Role"] = det["Role"]
            row["Distribution Type"] = det["Role"]
    return landscape


def main() -> None:
    xlsx = FOLDER / f"{SLUG}_FINAL.xlsx"
    landscape = _load_landscape(xlsx)

    brands: list[str] = []
    xs: list[float] = []
    ys: list[float] = []
    current: list[str] = []
    for row in landscape:
        b = str(row.get("Brand") or row.get("Company") or "").strip()
        try:
            x = float(row.get("X Score") if row.get("X Score") not in (None, "") else row.get("X"))
            y = float(row.get("Y Score") if row.get("Y Score") not in (None, "") else row.get("Y"))
        except (TypeError, ValueError):
            x, y = 50.0, 50.0
        brands.append(b)
        xs.append(x)
        ys.append(y)
        current.append(str(row.get("Quadrant") or "").strip())

    expected, mid_x, mid_y = assign_quadrants_half_median(xs, ys)
    fixes: list[dict] = []
    for i, row in enumerate(landscape):
        old = current[i]
        new = expected[i]
        row["Quadrant"] = new
        row["X Score"] = xs[i]
        row["Y Score"] = ys[i]
        row["X"] = xs[i]
        row["Y"] = ys[i]
        row["Overall Score"] = round((xs[i] + ys[i]) / 2.0, 1)
        row["Overall"] = row["Overall Score"]
        row["Role"] = "Solution Provider"
        row["Distribution Type"] = "Solution Provider"
        if old != new:
            fixes.append(
                {
                    "brand": brands[i],
                    "from": old,
                    "to": new,
                    "x": xs[i],
                    "y": ys[i],
                }
            )

    audit_path = FOLDER / "chatgpt_expand_batch_all_audit.json"
    audit: dict = {}
    if audit_path.exists():
        try:
            audit = json.loads(audit_path.read_text(encoding="utf-8")).get("xy_scoring") or {}
        except Exception:
            pass
    audit["score_midpoints"] = {"x": mid_x, "y": mid_y}
    audit["axis_x"] = audit.get("axis_x") or "Solution Capability"
    audit["axis_y"] = audit.get("axis_y") or "Business Strategy"

    qjson = FOLDER / f"{SLUG}_quadrant.json"
    if qjson.exists():
        try:
            payload = json.loads(qjson.read_text(encoding="utf-8"))
            crit = payload.get("criteria") or {}
            if crit.get("parameter_definitions"):
                audit["parameter_definitions"] = crit["parameter_definitions"]
            labels = crit.get("axis_labels") or {}
            if labels.get("x"):
                audit["axis_x"] = labels["x"]
            if labels.get("y"):
                audit["axis_y"] = labels["y"]
            if crit.get("x_axis"):
                audit["x_features"] = crit["x_axis"]
            if crit.get("y_axis"):
                audit["y_features"] = crit["y_axis"]
        except Exception:
            pass

    detail_rows = to_company_detail_rows(landscape, QUERY, audit)
    for det, src in zip(detail_rows, landscape):
        det["Role"] = "Solution Provider"
        det["Quadrant"] = src["Quadrant"]
        det["X"] = src["X"]
        det["Y"] = src["Y"]
        det["Overall"] = src["Overall"]
        if src.get("_display_company"):
            det["Company"] = src["_display_company"]

    write_final_xlsx(
        xlsx,
        landscape,
        "Companies",
        {
            "query": QUERY,
            "xy_scoring": audit,
            "quadrant_verify": {
                "mid_x": mid_x,
                "mid_y": mid_y,
                "fixed": len(fixes),
                "counts": dict(Counter(expected)),
            },
        },
        detail_rows=detail_rows,
    )
    extras = export_expand_quadrant_outputs(
        FOLDER, detail_rows, QUERY, country="global", audit=audit, chart_n=20
    )

    csv_path = FOLDER / f"{SLUG}_companies.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(
            f,
            fieldnames=["Brand", "Company", "Role", "Quadrant", "X", "Y", "Overall", "Found in"],
            extrasaction="ignore",
        )
        w.writeheader()
        for d in detail_rows:
            w.writerow(d)

    md = [
        "# Semiconductor Quadrant column verification",
        "",
        "Mapped to Coherent Quadrant definitions:",
        "",
        "| Quadrant | Solution Capability (X) | Business Strategy (Y) |",
        "|---|---|---|",
        "| **Leaders** | High | High |",
        "| **Challengers** | Lower | High |",
        "| **Trailblazers** | High | Lower |",
        "| **Emerging Players** | Lower | Lower |",
        "",
        "Assignment method: cohort **half-median** (rank-split on X, then on Y within each half).",
        f"Score midpoints: X={mid_x:.1f}, Y={mid_y:.1f}",
        "",
        f"Total brands: **{len(landscape)}**",
        f"Labels corrected: **{len(fixes)}**",
        "",
        "## Counts after fix",
        "",
        "| Quadrant | Count |",
        "|---|---:|",
    ]
    for q, n in Counter(expected).most_common():
        md.append(f"| {q} | {n} |")
    md.extend(["", "## Corrected rows", "", "| Brand | From | To | X | Y |", "|---|---|---|---:|---:|"])
    for frow in sorted(fixes, key=lambda z: z["brand"].lower()):
        md.append(
            f"| {frow['brand']} | {frow['from']} | {frow['to']} | {frow['x']} | {frow['y']} |"
        )
    if not fixes:
        md.append("| — | all matched | — | — | — |")
    md.extend(["", f"HTML: `{extras.get('html')}`", ""])

    audit_md = OUT / "_audit" / "semiconductor_quadrant_verify.md"
    audit_md.write_text("\n".join(md), encoding="utf-8")
    (OUT / "_audit" / "semiconductor_quadrant_verify.json").write_text(
        json.dumps(
            {
                "mid_x": mid_x,
                "mid_y": mid_y,
                "fixed": fixes,
                "counts": dict(Counter(expected)),
                "total": len(landscape),
            },
            indent=2,
        ),
        encoding="utf-8",
    )

    print(f"mid_x={mid_x} mid_y={mid_y}")
    print(f"fixed {len(fixes)} / {len(landscape)}")
    print("counts", dict(Counter(expected)))
    for frow in fixes[:30]:
        print(f"  {frow['brand']}: {frow['from']} -> {frow['to']} (X={frow['x']} Y={frow['y']})")
    if len(fixes) > 30:
        print(f"  ... {len(fixes)-30} more")
    print(f"audit -> {audit_md}")


if __name__ == "__main__":
    main()
