# -*- coding: utf-8 -*-
"""Reassign semiconductor quadrants by absolute median and rebuild chart."""
from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx
from vendor_intel.quadrant.rating_map import assign_quadrants_absolute_median

OUT = ROOT / "output" / "chatgpt_expand"
SLUG = "global_semiconductor_market_global"
QUERY = "Global Semiconductor Market"
FOLDER = OUT / SLUG
XLSX = FOLDER / f"{SLUG}_FINAL.xlsx"


def _norm(s: str) -> str:
    return " ".join(str(s or "").strip().lower().split())


def main() -> int:
    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Landscape"]
    rows_raw = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows_raw[0]]
    landscape: list[dict] = []
    for r in rows_raw[1:]:
        if not r:
            continue
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        if str(d.get("Company") or "").strip():
            landscape.append(d)

    # Prefer Landscape scores; fall back to Company Details by name
    details_by: dict[str, dict] = {}
    if "Company Details" in wb.sheetnames:
        ws = wb["Company Details"]
        rows_raw = list(ws.iter_rows(values_only=True))
        dhdr = [str(h) for h in rows_raw[0]]
        for r in rows_raw[1:]:
            if not r:
                continue
            d = {dhdr[i]: ("" if r[i] is None else r[i]) for i in range(len(dhdr)) if i < len(r)}
            key = _norm(d.get("Brand") or d.get("Company") or "")
            if key:
                details_by[key] = d

    for row in landscape:
        key = _norm(row.get("Company") or row.get("Brand") or "")
        det = details_by.get(key) or {}
        # Landscape X Score / Details X
        x = row.get("X Score")
        y = row.get("Y Score")
        if x in (None, "") and det.get("X") not in (None, ""):
            x = det["X"]
            y = det["Y"]
            row["X Score"] = x
            row["Y Score"] = y
            row["Overall Score"] = det.get("Overall")
        if det.get("Brand"):
            row["Brand"] = det["Brand"]
        if det.get("Role"):
            row["Role"] = det["Role"]
            row["Distribution Type"] = det["Role"]
        if det.get("Company") and det.get("Company") != det.get("Brand"):
            row["_display_company"] = det["Company"]
        if det.get("Found in") and not str(row.get("Headquarters") or "").strip():
            row["Headquarters"] = det["Found in"]
        # Keep HQ from landscape as Found in source of truth when present
        if str(row.get("Headquarters") or "").strip():
            row["Found in"] = row["Headquarters"]
        elif det.get("Found in"):
            row["Found in"] = det["Found in"]
            row["Headquarters"] = det["Found in"]

    xs = [int(round(float(r.get("X Score") or 0))) for r in landscape]
    ys = [int(round(float(r.get("Y Score") or 0))) for r in landscape]
    old_quads = [str(r.get("Quadrant") or "") for r in landscape]
    quads, mid_x, mid_y = assign_quadrants_absolute_median(xs, ys)
    changed = []
    for row, q, old in zip(landscape, quads, old_quads):
        row["Quadrant"] = q
        if old != q:
            changed.append(
                {
                    "brand": row.get("Brand") or row.get("Company"),
                    "from": old,
                    "to": q,
                    "x": row.get("X Score"),
                    "y": row.get("Y Score"),
                }
            )

    # Load axis audit from existing quadrant json if present
    audit: dict = {
        "quadrant_mid_x": mid_x,
        "quadrant_mid_y": mid_y,
        "quadrant_method": "absolute_median",
    }
    qjson = FOLDER / f"{SLUG}_quadrant.json"
    if qjson.exists():
        payload = json.loads(qjson.read_text(encoding="utf-8"))
        crit = payload.get("criteria") or {}
        if crit.get("parameter_definitions"):
            audit["parameter_definitions"] = crit["parameter_definitions"]
        labels = crit.get("axis_labels") or {}
        if labels.get("x"):
            audit["axis_x"] = labels["x"]
        if labels.get("y"):
            audit["axis_y"] = labels["y"]
        if crit.get("x_axis"):
            audit["x_features"] = crit["x_axis"]
        if crit.get("y_axis"):
            audit["y_features"] = crit["y_axis"]
        if payload.get("industry_group"):
            audit["industry_group"] = payload["industry_group"]
        if payload.get("industry_category"):
            audit["industry_category"] = payload["industry_category"]

    detail_rows = to_company_detail_rows(landscape, QUERY, audit)
    # Force Found in / Company display / scores by name (avoid sort zip bugs)
    for det in detail_rows:
        key = _norm(det.get("Brand") or "")
        src = next(
            (r for r in landscape if _norm(r.get("Company") or r.get("Brand") or "") == key),
            None,
        )
        if not src:
            continue
        det["Quadrant"] = src.get("Quadrant") or det.get("Quadrant")
        det["X"] = int(round(float(src.get("X Score") or det.get("X") or 0)))
        det["Y"] = int(round(float(src.get("Y Score") or det.get("Y") or 0)))
        det["Overall"] = int(round((det["X"] + det["Y"]) / 2))
        if src.get("_display_company"):
            det["Company"] = src["_display_company"]
        fi = str(src.get("Found in") or src.get("Headquarters") or "").strip()
        if fi:
            det["Found in"] = fi

    write_final_xlsx(
        XLSX,
        landscape,
        "Companies",
        {
            "query": QUERY,
            "xy_scoring": audit,
            "quadrant_reassign": {
                "method": "absolute_median",
                "mid_x": mid_x,
                "mid_y": mid_y,
                "changed_n": len(changed),
            },
        },
        detail_rows=detail_rows,
    )
    extras = export_expand_quadrant_outputs(
        FOLDER, detail_rows, QUERY, country="global", audit=audit, chart_n=20
    )

    counts = Counter(str(r.get("Quadrant")) for r in detail_rows)
    chart = [r for r in detail_rows if True]  # filled below from json
    qpayload = json.loads((FOLDER / f"{SLUG}_quadrant.json").read_text(encoding="utf-8"))
    on_chart = [
        {
            "brand": b.get("brand"),
            "x": b.get("execution"),
            "y": b.get("innovation"),
            "q": b.get("quadrant"),
            "country": b.get("country"),
        }
        for b in qpayload.get("brands") or []
        if b.get("on_chart")
    ]

    audit_out = {
        "method": "absolute_median",
        "mid_x": mid_x,
        "mid_y": mid_y,
        "counts": dict(counts),
        "changed_n": len(changed),
        "changed_sample": changed[:40],
        "chart": on_chart,
        "html": extras.get("html"),
    }
    out = OUT / "_audit" / "semi_quadrant_absolute_median.json"
    out.write_text(json.dumps(audit_out, indent=2, ensure_ascii=False), encoding="utf-8")

    print(f"mid_x={mid_x} mid_y={mid_y}")
    print(f"counts={dict(counts)}")
    print(f"changed={len(changed)}")
    print("chart:")
    for c in on_chart:
        print(f"  {c['q']}: {c['brand']} ({c['x']}/{c['y']}) [{c.get('country')}]")
    print(f"audit -> {out}")
    print(f"html -> {extras.get('html')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
