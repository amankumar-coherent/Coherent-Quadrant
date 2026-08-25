#!/usr/bin/env python3
"""Fix all 5 markets to use market-specific X/Y axes + parameters (not generic).

Uses catalog baselines + define_market_axes LLM refine, then regenerates HTML
Parameter Definitions. Does NOT rescore company X/Y numbers.
"""
from __future__ import annotations

import json
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.pipeline.expand_quadrant_score import export_expand_quadrant_outputs
from vendor_intel.placeholders.load_keys import apply_env_overrides
from vendor_intel.quadrant.axis_define import define_market_axes
from vendor_intel.quadrant.criteria_catalog import get_category_features

MARKETS = [
    (
        "global_flexible_packaging_market_global",
        "Global Flexible Packaging Market",
        "Others",
        "Packaging",
    ),
    (
        "global_glp_1_receptor_agonist_market_global",
        "Global GLP-1 Receptor Agonist Market",
        "Healthcare",
        "Pharmaceutical",
    ),
    (
        "global_liquefied_natural_gas_market_global",
        "Global Liquefied Natural Gas Market",
        "Others",
        "Energy",
    ),
    (
        "global_semiconductor_market_global",
        "Global Semiconductor Market",
        "ICT, Automation, Semiconductor",
        "Semiconductors",
    ),
    (
        "global_wearable_medical_devices_market_global",
        "Global Wearable Medical Devices Market",
        "Healthcare",
        "Medical Devices",
    ),
]

GENERIC_AXES = {"solution capability", "business strategy"}


def _load_details(folder: Path, slug: str) -> list[dict]:
    xlsx = folder / f"{slug}_FINAL.xlsx"
    wb = load_workbook(xlsx, data_only=True)
    rows = list(wb["Company Details"].iter_rows(values_only=True))
    hdr = [str(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not r or not r[0]:
            continue
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr))}
        out.append(
            {
                "Brand": d.get("Brand") or "",
                "Company": d.get("Company") or d.get("Brand") or "",
                "Role": d.get("Role") or "",
                "Quadrant": d.get("Quadrant") or "",
                "X": d.get("X"),
                "Y": d.get("Y"),
                "Overall": d.get("Overall"),
                "Found in": d.get("Found in") or "",
            }
        )
    return out


def _load_audit(folder: Path) -> dict:
    p = folder / "chatgpt_expand_batch_all_audit.json"
    if not p.exists():
        return {}
    try:
        raw = json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return {}
    if isinstance(raw.get("xy_scoring"), dict):
        return dict(raw["xy_scoring"])
    return dict(raw) if isinstance(raw, dict) else {}


def _save_audit(folder: Path, xy: dict) -> None:
    p = folder / "chatgpt_expand_batch_all_audit.json"
    raw: dict = {}
    if p.exists():
        try:
            raw = json.loads(p.read_text(encoding="utf-8"))
            if not isinstance(raw, dict):
                raw = {}
        except Exception:
            raw = {}
    raw["xy_scoring"] = xy
    p.write_text(json.dumps(raw, indent=2, ensure_ascii=False), encoding="utf-8")


def _is_generic(axis_x: str, axis_y: str, x_feats: list[str]) -> bool:
    ax = (axis_x or "").strip().lower()
    ay = (axis_y or "").strip().lower()
    if ax in GENERIC_AXES or ay in GENERIC_AXES:
        return True
    # Generic packaging/chem-style first features often mean wrong leaf was used
    if x_feats[:1] == ["Product Portfolio Breadth"] and ax in GENERIC_AXES | {""}:
        return True
    return False


def main() -> int:
    apply_env_overrides()
    os.environ["LLM_PROVIDER"] = os.getenv("LLM_PROVIDER") or "deepseek"
    # Allow market-axis LLM refine
    os.environ["EXPAND_MARKET_AXIS_LLM"] = os.getenv("EXPAND_MARKET_AXIS_LLM") or "1"

    base = ROOT / "output" / "chatgpt_expand"
    for slug, query, group, category in MARKETS:
        folder = base / slug
        print(f"\n=== {slug} ===", flush=True)
        catalog = get_category_features(group, category)
        prior = _load_audit(folder)

        industry = {
            "industry_group": group,
            "industry_category": category,
            "axis_x": catalog["axis_x"],
            "axis_y": catalog["axis_y"],
            "x": list(catalog["x"]),
            "y": list(catalog["y"]),
            "parameter_definitions": {"x": {}, "y": {}},
        }

        # Prefer prior market-specific axes if already good; still refresh defs via LLM
        prior_x = str(prior.get("axis_x") or "")
        prior_y = str(prior.get("axis_y") or "")
        prior_xf = [str(f) for f in (prior.get("x_features") or []) if str(f).strip()]
        prior_yf = [str(f) for f in (prior.get("y_features") or []) if str(f).strip()]
        if (
            prior_xf
            and prior_yf
            and not _is_generic(prior_x, prior_y, prior_xf)
            and len(prior_xf) == 5
            and len(prior_yf) == 5
        ):
            industry["axis_x"] = prior_x
            industry["axis_y"] = prior_y
            industry["x"] = prior_xf
            industry["y"] = prior_yf
            print(f"  keep prior market axes: {prior_x} / {prior_y}", flush=True)
        else:
            print(
                f"  catalog baseline: {catalog['axis_x']} / {catalog['axis_y']}",
                flush=True,
            )

        defined = define_market_axes(query, industry, geography="global")
        axis_x = str(defined.get("axis_x") or catalog["axis_x"])
        axis_y = str(defined.get("axis_y") or catalog["axis_y"])
        x_feats = list(defined.get("x") or catalog["x"])
        y_feats = list(defined.get("y") or catalog["y"])
        defs = defined.get("parameter_definitions") or {"x": {}, "y": {}}

        # Never keep generic after refine — fall back to catalog leaf
        if _is_generic(axis_x, axis_y, x_feats):
            print("  LLM returned generic — forcing catalog leaf axes", flush=True)
            axis_x = catalog["axis_x"]
            axis_y = catalog["axis_y"]
            x_feats = list(catalog["x"])
            y_feats = list(catalog["y"])
            # still explain catalog params for this market
            from vendor_intel.quadrant.axis_define import explain_market_parameters

            defs = explain_market_parameters(
                query,
                x_feats,
                y_feats,
                axis_x=axis_x,
                axis_y=axis_y,
                geography="global",
            )

        audit = dict(prior)
        audit.update(
            {
                "axis_x": axis_x,
                "axis_y": axis_y,
                "x_features": x_feats,
                "y_features": y_feats,
                "parameter_definitions": defs,
                "industry_group": group,
                "industry_category": category,
                "axis_definition_method": defined.get("axis_definition_method")
                or "catalog+llm",
                "axis_definition_reason": defined.get("axis_definition_reason") or "",
            }
        )
        _save_audit(folder, audit)

        details = _load_details(folder, slug)
        export_expand_quadrant_outputs(
            folder, details, query, country="global", audit=audit, chart_n=20
        )
        print(f"  FIXED X={axis_x}", flush=True)
        print(f"  FIXED Y={axis_y}", flush=True)
        print(f"  xf[0]={x_feats[0] if x_feats else None}", flush=True)
        print(
            f"  defs x={len((defs.get('x') or {}))} y={len((defs.get('y') or {}))}",
            flush=True,
        )

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
