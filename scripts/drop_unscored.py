#!/usr/bin/env python3
"""Remove rows that carry no X/Y score from a market's outputs.

A company ends up unscored when AI Mode could not answer for it — a quota
block, a CAPTCHA, or a reply that could not be parsed. Those rows render as
em-dashes in the report and sit outside every quadrant, so they are dropped
rather than shipped blank.

Rows are removed from the checkpoint, the scores JSON and the workbook
together, then quadrants are reassigned across what is left: the thresholds
are cohort medians, so removing rows moves them.

    .\\.venv\\Scripts\\python.exe scripts\\drop_unscored.py --all --dry-run
    .\\.venv\\Scripts\\python.exe scripts\\drop_unscored.py --all
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "src") not in sys.path:
    sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook  # noqa: E402

from vendor_intel.quadrant.rating_map import (  # noqa: E402
    assign_quadrants_absolute_median,
)

SHEET = "Company Details"



def _scored(row: dict) -> bool:
    return str(row.get("X") or "").strip() not in ("", "—", "-")


def process(market_dir: Path, *, dry_run: bool) -> int:
    ckpt_path = market_dir / "chatgpt_checkpoint_batch_all.json"
    if not ckpt_path.exists():
        print(f"{market_dir.name}: no checkpoint")
        return 0

    state = json.loads(ckpt_path.read_text(encoding="utf-8"))
    rows = (state.get("data") or {}).get("detail_rows") or []
    keep = [r for r in rows if _scored(r)]
    dropped = [r for r in rows if not _scored(r)]
    if not dropped:
        print(f"{market_dir.name}: nothing unscored")
        return 0

    if dry_run:
        print(f"{market_dir.name}: would drop {len(dropped)} of {len(rows)}")
        for r in dropped[:5]:
            print(f"    - {str(r.get('Company'))[:52]}")
        return len(dropped)

    # Quadrant thresholds are cohort medians, so they shift once rows go.
    quads, _, _ = assign_quadrants_absolute_median(
        [int(r["X"]) for r in keep], [int(r["Y"]) for r in keep]
    )
    for row, quad in zip(keep, quads):
        row["Quadrant"] = quad

    state["data"]["detail_rows"] = keep
    ckpt_path.write_text(json.dumps(state, ensure_ascii=False), encoding="utf-8")

    gone = {str(r.get("Company") or "").strip() for r in dropped}
    _prune_scores(market_dir, gone)
    _prune_workbook(market_dir, keep, gone)
    print(f"{market_dir.name}: dropped {len(dropped)}, kept {len(keep)}")
    return len(dropped)


def _prune_scores(market_dir: Path, gone: set[str]) -> None:
    """Drop the same companies from the evidence JSON, keeping it in step."""
    path = market_dir / "chatgpt_xy_scores_batch_all.json"
    if not path.exists():
        return
    data = json.loads(path.read_text(encoding="utf-8"))
    rows = data.get("rows")
    if isinstance(rows, list):
        data["rows"] = [
            r for r in rows if str(r.get("company") or "").strip() not in gone
        ]
        path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")


def _prune_workbook(market_dir: Path, keep: list[dict], gone: set[str]) -> None:
    """Delete the unscored rows and refresh the quadrant column."""
    matches = list(market_dir.glob("*_FINAL.xlsx"))
    if not matches:
        return
    xlsx = matches[0]
    wb = load_workbook(xlsx)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb[wb.sheetnames[0]]
    header = [c.value for c in ws[1]]
    try:
        ci = header.index("Company") + 1
        qi = header.index("Quadrant") + 1
    except ValueError:
        wb.close()
        return

    # Bottom-up: deleting a row shifts everything below it.
    for r in range(ws.max_row, 1, -1):
        if str(ws.cell(r, ci).value or "").strip() in gone:
            ws.delete_rows(r)

    quad_by_company = {
        str(r.get("Company") or "").strip(): r.get("Quadrant") for r in keep
    }
    for r in range(2, ws.max_row + 1):
        quad = quad_by_company.get(str(ws.cell(r, ci).value or "").strip())
        if quad:
            ws.cell(r, qi).value = quad
    wb.save(xlsx)
    wb.close()


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--market", action="append", default=[],
                    help="market name (repeat for several)")
    ap.add_argument("--country", default="global")
    ap.add_argument("--all", action="store_true",
                    help="every market folder under output/chatgpt_expand/")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    from vendor_intel.pipeline.web_expand import default_output_dir

    if args.all:
        root = default_output_dir("x", args.country).parent
        dirs = sorted(d for d in root.iterdir()
                      if (d / "chatgpt_checkpoint_batch_all.json").exists())
    else:
        dirs = [Path(default_output_dir(m, args.country)) for m in args.market]
    if not dirs:
        ap.print_help()
        return 2

    total = sum(process(d, dry_run=args.dry_run) for d in dirs)
    verb = "would drop" if args.dry_run else "dropped"
    print(f"\n{verb} {total} unscored rows")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
