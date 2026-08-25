#!/usr/bin/env python3
"""Remove uncrawled semiconductor companies; add 75+ web-verified new ones; crawl+score."""
from __future__ import annotations

import asyncio
import json
import os
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

from openpyxl import load_workbook

from vendor_intel.placeholders.load_keys import apply_env_overrides
from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    score_expand_rows,
    to_company_detail_rows,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx

OUT = ROOT / "output" / "chatgpt_expand"
SLUG = "global_semiconductor_market_global"
QUERY = "Global Semiconductor Market"
FOLDER = OUT / SLUG
INDUSTRY = "ICT, Automation, Semiconductor / Semiconductors"

UNCRAWLED = {
    "Inuitive",
    "Key Foundry",
    "LAPIS Technology",
    "Melexis",
    "Rohm Semiconductor",
    "XMOS",
    "Kenya Semiconductor Technologies",
    "LG Innotek",
    "CEITEC S.A.",
    "Chipus Microelectronics",
    "Hua Hong Semiconductor",
    "Intel Corporation",
    "Pragmatic Semiconductor",
    "RIR Power Electronics",
    "SID Microeletrônica",
    "Socionext Inc.",
    "Sony Semiconductor Solutions",
    "SSMC",
    "Vishay Intertechnology",
}


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def _row_from_candidate(c: dict) -> dict:
    name = c["company"]
    hq = c.get("hq") or ""
    return {
        "Company": name,
        "Website": c.get("website") or "",
        "Founded": c.get("founded") or "",
        "Headquarters": hq,
        "Continent / Geography": c.get("continent") or "",
        "Operational Presence": c.get("presence") or "",
        "Ownership": "",
        "Employees": "",
        "Core Categories": c.get("categories") or "",
        "Specialty Focus": c.get("specialty") or "",
        "Key Brands Represented": name,
        "Retail / E-commerce / Both": "No",
        "Distribution Type": "Solution Provider",
        "Contact Person": "",
        "Role": "Solution Provider",
        "Email": "",
        "LinkedIn": "",
        "Office No.": "",
        "Country Code": c.get("country") or "",
        "Region Code": c.get("region") or "",
        "Summary": f"{name} — semiconductor market participant. {c.get('specialty') or ''}".strip(),
        "X Score": "",
        "Y Score": "",
        "Overall Score": "",
        "Quadrant": "",
        "Industry Category": INDUSTRY,
        "Brand": name,
        "Found in": hq,
    }


async def main() -> int:
    apply_env_overrides()
    os.environ["EXPAND_XY_SKIP_CRAWL"] = "0"
    os.environ["EXPAND_XY_FORCE_CRAWL"] = "1"
    os.environ.setdefault("EXPAND_XY_CONCURRENT", "4")
    os.environ.setdefault("QUADRANT_CRAWL_MAX_PAGES", "25")
    os.environ.setdefault("EXPAND_MARKET_AXIS_LLM", "1")

    cands = json.loads(
        (OUT / "_audit" / "semi_new_companies_75plus.json").read_text(encoding="utf-8")
    )
    xlsx = FOLDER / f"{SLUG}_FINAL.xlsx"
    wb = load_workbook(xlsx, data_only=True)
    ws = wb["Landscape"]
    rows_raw = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows_raw[0]]

    kept: list[dict] = []
    removed: list[str] = []
    existing_names = set()
    for r in rows_raw[1:]:
        if not r:
            continue
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        name = str(d.get("Company") or "").strip()
        if not name:
            continue
        if name in UNCRAWLED or _norm(name) in {_norm(x) for x in UNCRAWLED}:
            removed.append(name)
            continue
        existing_names.add(_norm(name))
        # Preserve prior scores from Company Details when present
        kept.append(d)

    # Overlay Details Brand/Role/Found in / scores onto kept
    details_by: dict[str, dict] = {}
    if "Company Details" in wb.sheetnames:
        ws = wb["Company Details"]
        rows_raw = list(ws.iter_rows(values_only=True))
        dhdr = [str(h) for h in rows_raw[0]]
        for r in rows_raw[1:]:
            if not r:
                continue
            d = {dhdr[i]: ("" if r[i] is None else r[i]) for i in range(len(dhdr)) if i < len(r)}
            key = _norm(d.get("Brand") or d.get("Company") or "")
            if key and key not in {_norm(x) for x in UNCRAWLED}:
                details_by[key] = d

    for row in kept:
        key = _norm(row.get("Company") or "")
        det = details_by.get(key)
        if not det:
            continue
        brand = str(det.get("Brand") or "").strip()
        if brand:
            row["Brand"] = brand
            row["Company"] = brand
        if det.get("Role"):
            row["Distribution Type"] = det["Role"]
            row["Role"] = det["Role"]
        if det.get("Found in"):
            row["Headquarters"] = det["Found in"] or row.get("Headquarters")
            row["Found in"] = det["Found in"]
        company_col = str(det.get("Company") or "").strip()
        if company_col and company_col != brand:
            row["_display_company"] = company_col
        if det.get("X") not in (None, ""):
            row["X Score"] = det["X"]
            row["Y Score"] = det["Y"]
            row["Overall Score"] = det["Overall"]
            row["Quadrant"] = det.get("Quadrant") or ""

    new_rows: list[dict] = []
    skipped_dup = []
    for c in cands:
        name = str(c.get("company") or "").strip()
        if not name:
            continue
        if _norm(name) in existing_names:
            skipped_dup.append(name)
            continue
        new_rows.append(_row_from_candidate(c))
        existing_names.add(_norm(name))

    print(f"Removed uncrawled: {len(removed)}", flush=True)
    for n in removed:
        print(f"  - {n}", flush=True)
    print(f"Kept: {len(kept)}", flush=True)
    print(f"Adding new: {len(new_rows)} (skipped dups {len(skipped_dup)})", flush=True)

    # Score only new companies with crawl
    print(f"Crawl+score {len(new_rows)} new companies…", flush=True)
    scored_new, audit_new = await score_expand_rows(
        [dict(r) for r in new_rows],
        QUERY,
        country="global",
        concurrent=int(os.getenv("EXPAND_XY_CONCURRENT") or "4"),
    )
    scored_by = {_norm(r.get("Company") or r.get("Brand") or ""): r for r in scored_new}
    for row in new_rows:
        key = _norm(row.get("Company") or "")
        s = scored_by.get(key)
        if not s:
            continue
        for col in ("X Score", "Y Score", "Overall Score", "Quadrant", "X", "Y", "Overall"):
            if s.get(col) not in (None, ""):
                row[col] = s[col]
        if s.get("_evidence_snapshot"):
            row["_evidence_snapshot"] = s["_evidence_snapshot"]
            # Prefer HQ from crawl backfill when empty
            if not str(row.get("Headquarters") or "").strip() and s.get("Headquarters"):
                row["Headquarters"] = s["Headquarters"]
                row["Found in"] = s["Headquarters"]

    all_rows = kept + new_rows

    # Preserve market axes from prior audit
    full_path = FOLDER / "chatgpt_xy_scores_batch_all.json"
    full = json.loads(full_path.read_text(encoding="utf-8")) if full_path.exists() else {}
    audit = dict(audit_new or {})
    for k in (
        "axis_x",
        "axis_y",
        "x_features",
        "y_features",
        "parameter_definitions",
        "industry_group",
        "industry_category",
        "x_feature_weights",
        "y_feature_weights",
        "question_weights",
        "overall_formula",
        "axis_definition_method",
    ):
        if full.get(k) is not None:
            audit[k] = full[k]

    detail_rows = to_company_detail_rows(all_rows, QUERY, audit)
    for det in detail_rows:
        key = _norm(det.get("Brand") or "")
        src = next((r for r in all_rows if _norm(r.get("Company") or r.get("Brand") or "") == key), None)
        if src and src.get("_display_company"):
            det["Company"] = src["_display_company"]
        # New companies: Brand == Company plain
        if key in {_norm(r.get("Company") or "") for r in new_rows}:
            brand = str(det.get("Brand") or "").strip()
            if brand:
                det["Company"] = brand

    write_final_xlsx(
        xlsx,
        all_rows,
        "Companies",
        {
            "query": QUERY,
            "xy_scoring": {k: v for k, v in audit.items() if k != "rows"},
            "removed_uncrawled": removed,
            "added_companies": [r.get("Company") for r in new_rows],
        },
        detail_rows=detail_rows,
    )
    extras = export_expand_quadrant_outputs(
        FOLDER, detail_rows, QUERY, country="global", audit=audit, chart_n=20
    )

    # Rebuild xy audit rows
    old_rows = [
        r
        for r in (full.get("rows") or [])
        if _norm(r.get("company") or "") not in {_norm(x) for x in UNCRAWLED}
    ]
    by = {_norm(r.get("company") or ""): r for r in old_rows}
    for r in audit_new.get("rows") or []:
        by[_norm(r.get("company") or "")] = r
    full["rows"] = list(by.values())
    full["scored"] = len(full["rows"])
    full["removed_uncrawled"] = removed
    full["added_n"] = len(new_rows)
    full["crawl_new"] = audit_new.get("crawl")
    full_path.write_text(json.dumps(full, indent=2, ensure_ascii=False, default=str), encoding="utf-8")

    summary = {
        "removed": removed,
        "kept": len(kept),
        "added": len(new_rows),
        "added_names": [r.get("Company") for r in new_rows],
        "final_n": len(detail_rows),
        "crawl_new": audit_new.get("crawl"),
        "new_crawled_ok": sum(1 for r in (audit_new.get("rows") or []) if r.get("crawled")),
        "html": extras.get("html"),
    }
    out = OUT / "_audit" / "semi_replace_add75.json"
    out.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    print(
        f"Final companies={len(detail_rows)} kept={len(kept)} added={len(new_rows)} "
        f"new_crawled={summary['new_crawled_ok']}",
        flush=True,
    )
    print(f"audit -> {out}", flush=True)
    print(f"html -> {extras.get('html')}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
