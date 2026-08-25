#!/usr/bin/env python3
"""Last-pass: SSC/DDGS scrape for still-uncrawled semiconductor rows, then rescore."""
from __future__ import annotations

import asyncio
import json
import os
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook

from vendor_intel.placeholders.load_keys import apply_env_overrides
from vendor_intel.pipeline.expand_quadrant_score import (
    export_expand_quadrant_outputs,
    score_expand_rows,
    to_company_detail_rows,
    _domain,
)
from vendor_intel.pipeline.web_expand import write_final_xlsx

OUT = ROOT / "output" / "chatgpt_expand"
SLUG = "global_semiconductor_market_global"
QUERY = "Global Semiconductor Market"
FOLDER = OUT / SLUG


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


async def main() -> int:
    apply_env_overrides()
    os.environ["EXPAND_XY_SKIP_CRAWL"] = "0"
    os.environ["EXPAND_XY_FORCE_CRAWL"] = "1"
    os.environ.setdefault("EXPAND_XY_CONCURRENT", "3")
    os.environ.setdefault("QUADRANT_CRAWL_MAX_PAGES", "20")

    full = json.loads((FOLDER / "chatgpt_xy_scores_batch_all.json").read_text(encoding="utf-8"))
    still = [str(r.get("company") or "").strip() for r in full.get("rows") or [] if not r.get("crawled")]
    still = [n for n in still if n]
    print(f"Still uncrawled: {len(still)}", flush=True)
    if not still:
        return 0

    # Load landscape rows for these companies
    xlsx = FOLDER / f"{SLUG}_FINAL.xlsx"
    wb = load_workbook(xlsx, data_only=True)
    ws = wb["Landscape"]
    rows_raw = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows_raw[0]]
    need: list[dict] = []
    for r in rows_raw[1:]:
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        name = str(d.get("Company") or "").strip()
        if _norm(name) not in {_norm(n) for n in still}:
            continue
        for col in ("X Score", "Y Score", "Overall Score", "X", "Y", "Overall", "Quadrant"):
            d[col] = ""
        need.append(d)

    # Attach role/HQ from details
    if "Company Details" in wb.sheetnames:
        ws = wb["Company Details"]
        rows_raw = list(ws.iter_rows(values_only=True))
        hdr = [str(h) for h in rows_raw[0]]
        details = {}
        for r in rows_raw[1:]:
            d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
            key = _norm(d.get("Brand") or d.get("Company") or "")
            if key:
                details[key] = d
        for row in need:
            det = details.get(_norm(row.get("Company") or ""))
            if not det:
                continue
            brand = str(det.get("Brand") or "").strip()
            if brand:
                row["Company"] = brand
                row["Brand"] = brand
            if det.get("Role"):
                row["Distribution Type"] = det["Role"]
            if det.get("Found in"):
                row["Headquarters"] = det["Found in"]
            company_col = str(det.get("Company") or "").strip()
            if company_col and company_col != brand:
                row["_display_company"] = company_col

    # Pre-enrich via SSC + DDGS (more resilient than smart_crawl alone)
    from vendor_intel.enrichment.smart_enrichment import clear_enrichment_cache, enrich_companies
    from vendor_intel.quadrant.snapshot import build_evidence_snapshot
    from vendor_intel.pipeline.expand_quadrant_score import _scoring_row, _backfill_from_snapshot

    clear_enrichment_cache()
    batch = [
        {"name": str(r.get("Company") or "").strip(), "domain": _domain(str(r.get("Website") or ""))}
        for r in need
        if str(r.get("Company") or "").strip() and _domain(str(r.get("Website") or ""))
    ]
    print(f"SSC/DDGS enrich {len(batch)}…", flush=True)
    enriched = await enrich_companies(
        batch,
        limit=len(batch),
        max_concurrent=3,
        country="global",
        use_ssc=True,
        crawl_mode="business",
        max_pages=20,
    )
    ok = 0
    for row in need:
        name = str(row.get("Company") or "").strip()
        dom = _domain(str(row.get("Website") or ""))
        smart = enriched.get(name) or enriched.get(dom)
        if isinstance(smart, dict) and not smart.get("error"):
            snap = build_evidence_snapshot(smart, _scoring_row(row), _scoring_row(row))
            row["_evidence_snapshot"] = snap
            _backfill_from_snapshot(row, snap)
            ok += 1
    print(f"Attached evidence for {ok}/{len(need)}", flush=True)

    # Skip crawl inside score_expand_rows — we already attached snapshots
    os.environ["EXPAND_XY_SKIP_CRAWL"] = "1"
    scored, audit = await score_expand_rows(need, QUERY, country="global", concurrent=3)

    # Merge into full workbook like the other script
    wb = load_workbook(xlsx, data_only=True)
    ws = wb["Landscape"]
    rows_raw = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows_raw[0]]
    all_rows: list[dict] = []
    for r in rows_raw[1:]:
        d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
        if str(d.get("Company") or "").strip():
            all_rows.append(d)

    details_by: dict[str, dict] = {}
    if "Company Details" in wb.sheetnames:
        ws = wb["Company Details"]
        rows_raw = list(ws.iter_rows(values_only=True))
        hdr = [str(h) for h in rows_raw[0]]
        for r in rows_raw[1:]:
            d = {hdr[i]: ("" if r[i] is None else r[i]) for i in range(len(hdr)) if i < len(r)}
            key = _norm(d.get("Brand") or d.get("Company") or "")
            if key:
                details_by[key] = d

    scored_by = {_norm(r.get("Company") or r.get("Brand") or ""): r for r in scored}
    for row in all_rows:
        key = _norm(row.get("Company") or "")
        det = details_by.get(key)
        if det:
            brand = str(det.get("Brand") or "").strip()
            if brand:
                row["Company"] = brand
                row["Brand"] = brand
                key = _norm(brand)
            if det.get("Role"):
                row["Distribution Type"] = det["Role"]
            if det.get("Found in"):
                row["Headquarters"] = det["Found in"]
            company_col = str(det.get("Company") or "").strip()
            if company_col and company_col != brand:
                row["_display_company"] = company_col
            if det.get("X") not in (None, "") and key not in scored_by:
                row["X Score"] = det["X"]
                row["Y Score"] = det["Y"]
                row["Overall Score"] = det["Overall"]
                row["Quadrant"] = det.get("Quadrant") or ""
        if key in scored_by:
            s = scored_by[key]
            for col in ("X Score", "Y Score", "Overall Score", "Quadrant", "X", "Y", "Overall"):
                if s.get(col) not in (None, ""):
                    row[col] = s[col]
            if s.get("_evidence_snapshot"):
                row["_evidence_snapshot"] = s["_evidence_snapshot"]

    for k in (
        "axis_x",
        "axis_y",
        "x_features",
        "y_features",
        "parameter_definitions",
        "industry_group",
        "industry_category",
        "x_feature_weights",
        "y_feature_weights",
        "question_weights",
        "overall_formula",
        "axis_definition_method",
    ):
        if full.get(k) is not None:
            audit[k] = full[k]

    detail_rows = to_company_detail_rows(all_rows, QUERY, audit)
    for det in detail_rows:
        key = _norm(det.get("Brand") or "")
        src = next((r for r in all_rows if _norm(r.get("Company") or r.get("Brand") or "") == key), None)
        if src and src.get("_display_company"):
            det["Company"] = src["_display_company"]

    write_final_xlsx(
        xlsx,
        all_rows,
        "Companies",
        {"query": QUERY, "xy_scoring": {k: v for k, v in audit.items() if k != "rows"}},
        detail_rows=detail_rows,
    )
    extras = export_expand_quadrant_outputs(
        FOLDER, detail_rows, QUERY, country="global", audit=audit, chart_n=20
    )

    by_old = {_norm(r.get("company") or ""): r for r in (full.get("rows") or [])}
    for r in audit.get("rows") or []:
        # Mark crawled if we attached evidence this pass
        name = _norm(r.get("company") or "")
        src = scored_by.get(name)
        if src and src.get("_evidence_snapshot"):
            r["crawled"] = True
            r["kb_chars"] = len(str((src.get("_evidence_snapshot") or {}).get("kb_text") or src.get("_evidence_snapshot") or ""))
        by_old[name] = r
    full["rows"] = list(by_old.values())
    full["ssc_retry"] = {"tried": len(need), "evidence_ok": ok}
    (FOLDER / "chatgpt_xy_scores_batch_all.json").write_text(
        json.dumps(full, indent=2, ensure_ascii=False, default=str), encoding="utf-8"
    )

    summary = {
        "tried": len(need),
        "evidence_ok": ok,
        "rescored": [
            {
                "company": r.get("company"),
                "x": r.get("x"),
                "y": r.get("y"),
                "overall": r.get("overall"),
                "crawled": r.get("crawled"),
                "kb_chars": r.get("kb_chars"),
            }
            for r in (audit.get("rows") or [])
        ],
        "html": extras.get("html"),
    }
    out = OUT / "_audit" / "semi_ssc_fail_retry.json"
    out.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"evidence_ok={ok}/{len(need)}", flush=True)
    for r in sorted(summary["rescored"], key=lambda x: -int(x.get("overall") or 0))[:15]:
        print(
            f"  {r['company']}: O={r['overall']} crawled={r['crawled']} kb={r['kb_chars']}",
            flush=True,
        )
    print(f"audit -> {out}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
