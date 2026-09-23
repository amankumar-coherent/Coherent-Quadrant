#!/usr/bin/env python3
"""Retry X/Y scoring for the rows a market left UNSCORED.

A company comes back unscored when AI Mode answers UNKNOWN, or when a CAPTCHA
or quota block eats the query. The first is a real judgement; the other two are
accidents of timing, and a later attempt usually succeeds. Either way the row
sits in the report with an em-dash and no quadrant.

This retries ONLY those rows, in batches, and writes the results back into the
checkpoint and the workbook. Anything still unscored afterwards stays unscored
— a blank is honest, an invented number is not.

    .\\.venv\\Scripts\\python.exe scripts\\rescore_unscored.py --market "Global CPaaS Market"
    .\\.venv\\Scripts\\python.exe scripts\\rescore_unscored.py --all
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "src") not in sys.path:
    sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook  # noqa: E402

from vendor_intel.quadrant import ai_mode_scorer  # noqa: E402
from vendor_intel.quadrant.matrix_rollup import normalize_row_score_floor  # noqa: E402
from vendor_intel.quadrant.rating_map import (  # noqa: E402
    assign_quadrants_absolute_median,
)

SHEET = "Company Details"


def _blank(value: object) -> bool:
    return str(value or "").strip() in ("", "—", "-", "0")


def _axis_parameters(market_dir: Path) -> tuple[list[str], list[str]]:
    """The five X and Y parameters this market was scored against.

    Read from the audit file so a retry uses the SAME rubric as the original
    pass — regenerating them would score these rows on a different scale.
    """
    audit = market_dir / "chatgpt_expand_batch_all_audit.json"
    if audit.exists():
        data = json.loads(audit.read_text(encoding="utf-8"))
        xy = data.get("xy_scoring") or data.get("xy") or data.get("axes") or {}
        xs = xy.get("x_features") or xy.get("x") or []
        ys = xy.get("y_features") or xy.get("y") or []
        if xs and ys:
            return list(xs), list(ys)
    raise SystemExit(f"no axis parameters in {audit} — cannot rescore safely")


def rescore(market_dir: Path, batch: int) -> int:
    ckpt_path = market_dir / "chatgpt_checkpoint_batch_all.json"
    if not ckpt_path.exists():
        print(f"{market_dir.name}: no checkpoint")
        return 0
    state = json.loads(ckpt_path.read_text(encoding="utf-8"))
    rows = (state.get("data") or {}).get("detail_rows") or []
    todo = [r for r in rows if _blank(r.get("X")) or _blank(r.get("Y"))]
    if not todo:
        print(f"{market_dir.name}: nothing unscored")
        return 0

    market = state.get("query") or market_dir.name
    x_params, y_params = _axis_parameters(market_dir)
    names = [str(r.get("Company") or "").strip() for r in todo]
    names = [n for n in names if n]
    print(f"{market_dir.name}: retrying {len(names)} unscored companies")

    scores = ai_mode_scorer.score_companies(
        names,
        x_parameters=x_params,
        y_parameters=y_params,
        market=market,
        batch=batch,
    )

    fixed = 0
    for row in todo:
        name = str(row.get("Company") or "").strip()
        x, y = scores.get(name, (None, None))
        if x is None or y is None:
            continue
        # Same floor-65 normalization the pipeline applies, so a retried row
        # is on exactly the same scale as one scored on the first pass.
        nx, ny = normalize_row_score_floor(float(x), float(y))
        row["X"] = int(round(nx))
        row["Y"] = int(round(ny))
        row["Overall"] = int(round((nx + ny) / 2))
        fixed += 1

    if not fixed:
        print(f"{market_dir.name}: none could be scored this time")
        return 0

    # Quadrants shift once new rows join the cohort, so reassign across all
    # scored rows rather than labelling the new ones in isolation.
    scored = [r for r in rows if not _blank(r.get("X")) and not _blank(r.get("Y"))]
    quads, _, _ = assign_quadrants_absolute_median(
        [int(r["X"]) for r in scored], [int(r["Y"]) for r in scored]
    )
    for row, quad in zip(scored, quads):
        row["Quadrant"] = quad

    ckpt_path.write_text(json.dumps(state, ensure_ascii=False), encoding="utf-8")
    _write_workbook(market_dir, rows)
    print(f"{market_dir.name}: scored {fixed} more ({len(scored)} total)")
    return fixed


def _write_workbook(market_dir: Path, rows: list[dict]) -> None:
    """Push X/Y/Overall/Quadrant back into the exported workbook."""
    matches = list(market_dir.glob("*_FINAL.xlsx"))
    if not matches:
        return
    xlsx = matches[0]
    by_company = {str(r.get("Company") or "").strip(): r for r in rows}
    wb = load_workbook(xlsx)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb[wb.sheetnames[0]]
    header = [c.value for c in ws[1]]
    try:
        ci = header.index("Company") + 1
        cols = {k: header.index(k) + 1 for k in ("X", "Y", "Overall", "Quadrant")}
    except ValueError:
        wb.close()
        return
    for r in range(2, ws.max_row + 1):
        src = by_company.get(str(ws.cell(r, ci).value or "").strip())
        if not src:
            continue
        for key, col in cols.items():
            if str(src.get(key) or "").strip():
                ws.cell(r, col).value = src[key]
    wb.save(xlsx)
    wb.close()


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--market", help="market name")
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--batch", type=int, default=ai_mode_scorer.DEFAULT_SCORE_BATCH)
    args = ap.parse_args()

    base = ROOT / "output" / "chatgpt_expand"
    if args.all:
        dirs = sorted(
            p.parent
            for p in base.glob("*/chatgpt_checkpoint_batch_all.json")
            if not any(
                tag in p.parent.name
                for tag in ("_old", "_pre", "_failed", "_blocked", "_soured", "_quota")
            )
        )
    elif args.market:
        from vendor_intel.pipeline.web_expand import default_output_dir

        dirs = [Path(default_output_dir(args.market, "global"))]
    else:
        ap.print_help()
        return 2

    total = sum(rescore(d, args.batch) for d in dirs)
    print(f"\nrescored {total} companies")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
