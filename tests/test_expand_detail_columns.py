"""Company Details columns for ChatGPT expand FINAL Excel."""
from __future__ import annotations

from vendor_intel.pipeline.expand_quadrant_score import (
    DETAIL_COLUMNS,
    build_expand_quadrant_payload,
    to_company_detail_rows,
)


def test_detail_columns_match_user_request():
    assert DETAIL_COLUMNS[:2] == ("Brand", "Company")
    assert "Quadrant" in DETAIL_COLUMNS
    assert "X" in DETAIL_COLUMNS
    assert "Y" in DETAIL_COLUMNS
    assert "Overall" in DETAIL_COLUMNS
    assert "Found in" in DETAIL_COLUMNS
    assert DETAIL_COLUMNS[-1] == "Found in"


def test_acquired_brand_company_column():
    rows = [
        {
            "Company": "Horizon Organic",
            "Ownership": "Acquired by Danone",
            "Founded": "1991",
            "Headquarters": "Boulder, Colorado",
            "Distribution Type": "Brand / Marketer",
            "X Score": "88",
            "Y Score": "84",
            "Overall Score": "86",
            "Quadrant": "Leaders",
        }
    ]
    out = to_company_detail_rows(rows, "Global Organic Milk Market")
    assert len(out) == 1
    assert out[0]["Brand"] == "Horizon Organic"
    assert out[0]["Company"] == "(acquired by Danone)"
    assert out[0]["Role"] == "Brand / Marketer"
    assert out[0]["Quadrant"] == "Leaders"
    assert out[0]["X"] == 88
    assert out[0]["Y"] == 84
    assert out[0]["Overall"] == 86
    assert out[0]["Found in"] == "Boulder, Colorado"
    assert out[0]["Found in"] != "1991"


def test_independent_company_repeats_name():
    rows = [
        {
            "Company": "Amcor",
            "Ownership": "Independent",
            "Founded": "1860",
            "Distribution Type": "Brand / Marketer",
            "X Score": "70",
            "Y Score": "72",
            "Overall Score": "71",
            "Quadrant": "Leaders",
        }
    ]
    out = to_company_detail_rows(rows, "Global Flexible Packaging Market")
    assert out[0]["Brand"] == "Amcor"
    assert out[0]["Company"] == "Amcor"
    assert "(acquired by" not in out[0]["Company"]


def test_tech_solution_provider_parent_plain():
    rows = [
        {
            "Company": "Gemini",
            "Ownership": "Subsidiary of Google",
            "Founded": "2023",
            "Distribution Type": "Solution Provider",
            "X Score": "80",
            "Y Score": "78",
            "Overall Score": "79",
            "Quadrant": "Leaders",
        }
    ]
    out = to_company_detail_rows(
        rows,
        "Global Semiconductor Market",
        {"industry_group": "ICT, Automation, Semiconductor", "industry_category": "Semiconductors"},
    )
    assert out[0]["Brand"] == "Gemini"
    assert out[0]["Company"] == "Google"
    assert out[0]["Role"] == "Solution Provider"


def test_xlsx_company_details_is_first_sheet(tmp_path):
    from pathlib import Path

    from vendor_intel.pipeline.web_expand import read_final_rows, write_final_xlsx

    landscape = [
        {
            "Company": "Horizon Organic",
            "Website": "https://horizon.com",
            "Founded": "1991",
            "Ownership": "Acquired by Danone",
            "Distribution Type": "Brand / Marketer",
            "X Score": "88",
            "Y Score": "84",
            "Overall Score": "86",
            "Quadrant": "Leaders",
        }
    ]
    details = to_company_detail_rows(landscape, "Organic Milk Market")
    xlsx = Path(tmp_path) / "demo_FINAL.xlsx"
    write_final_xlsx(xlsx, landscape, "Companies", {"query": "Organic Milk Market"}, detail_rows=details)
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True)
    assert wb.sheetnames[0] == "Company Details"
    assert "Landscape" in wb.sheetnames
    headers = [c.value for c in next(wb["Company Details"].iter_rows(min_row=1, max_row=1))]
    assert headers == list(DETAIL_COLUMNS)
    data = [c.value for c in next(wb["Company Details"].iter_rows(min_row=2, max_row=2))]
    assert data[0] == "Horizon Organic"
    assert data[1] == "(acquired by Danone)"
    wb.close()
    resumed = read_final_rows(xlsx)
    assert resumed[0]["Company"] == "Horizon Organic"
    assert resumed[0]["Website"] == "https://horizon.com"


def test_sorted_by_overall_and_chart_top_20():
    rows = [
        {
            "Company": f"Co {i}",
            "Ownership": "Independent",
            "Founded": "2000",
            "Headquarters": f"City, Country{(i % 5) + 1}",
            "Distribution Type": "Brand / Marketer",
            "X Score": str(i),
            "Y Score": str(i),
            "Overall Score": str(i),
            "Quadrant": "Leaders" if i >= 50 else "Emerging Players",
        }
        for i in range(1, 26)
    ]
    # Force two quadrants with enough rows
    for i, row in enumerate(rows):
        row["Quadrant"] = "Leaders" if i < 13 else "Emerging Players"
        row["Overall Score"] = str(90 - i)
        row["X Score"] = str(90 - i)
        row["Y Score"] = str(88 - i)
    out = to_company_detail_rows(rows, "Global Flexible Packaging Market")
    assert out[0]["Overall"] >= out[-1]["Overall"]
    payload = build_expand_quadrant_payload(out, "Global Flexible Packaging Market", chart_n=20)
    on_chart = [b for b in payload["brands"] if b["on_chart"]]
    assert len(on_chart) == 20
    assert len(payload["brands"]) == 25
    quads = {b["quadrant"] for b in on_chart}
    assert "Leaders" in quads
    assert "Emerging Players" in quads
    countries = {b.get("country") for b in on_chart}
    assert len(countries) >= 3


def test_country_from_hq():
    from vendor_intel.pipeline.expand_quadrant_score import country_from_hq

    assert country_from_hq("Zurich, Switzerland") == "Switzerland"
    assert country_from_hq("Evansville, IN, USA") == "United States"
    assert country_from_hq("") == "Unknown"
